"""Content agent: drafts Walmart Marketplace listing content for products we decided to buy, for a person to review, edit
and upload BY HAND to Walmart Seller Center. It never calls a Walmart API, never publishes, never sets a price or a UPC.

    python content_agent.py facts-template --approved approved_orders.csv --results "<results.xlsx>"
    python content_agent.py generate --approved approved_orders.csv --results "<results.xlsx>" --facts product_facts.csv [--sku SKU ...] [--allow-listing-facts] [--dry-run]
    python content_agent.py check "<walmart_listings_*.xlsx>"

The language model may only restate facts from product_facts.csv; code checks every claim and every spec-bearing word.
Env (.env): ANTHROPIC_API_KEY, BRAND_NAME (required by `generate` only).
"""
import argparse
import csv
import io
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl
from dotenv import load_dotenv
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from email_agent import excluded_reason, load_exclusions
from order_sheet import (OrderSheetError, approval_expiry, candidates_of, find_candidate, read_approved,
                         read_results)

MODEL = "claude-haiku-4-5"
TEMPERATURE = 0.2
MAX_TOKENS = 1800
# Haiku list prices in USD per million tokens. UNVERIFIED: from memory, check the Anthropic pricing page.
PRICE_PER_MTOK = {"input": Decimal("1.00"), "output": Decimal("5.00")}
COST_WARN = Decimal("0.50")
MAX_CALLS_PER_SKU = 3                       # first try + one invalid-JSON retry + one repair

# ---- Walmart limits. ALL UNVERIFIED: from a third-party summary (lettercounter.org), not Walmart's own pages.
# ---- Confirm each one in Seller Center and correct it here, in one place.
LIMITS = {
    "title_max": 150,                       # UNVERIFIED
    "title_recommended": (50, 75),          # UNVERIFIED (advice, reported as a note, never a block)
    "short_max": 500,                       # UNVERIFIED
    "long_max": 4000,                       # UNVERIFIED
    "long_min_words": 150,                  # UNVERIFIED
    "bullets_min": 3,                       # UNVERIFIED
    "bullets_max": 10,                      # UNVERIFIED
    "bullet_max": 80,                       # UNVERIFIED
    "brand_max": 60,                        # UNVERIFIED
}
TITLE_ALLOWED_CHARS = re.compile(r"[A-Za-z0-9 ,.\-/&()'\"+]")   # anything else in a title is a "special character" (UNVERIFIED reading)
TITLE_SUBJECTIVE = ["premium", "high quality", "perfect", "amazing", "great", "excellent", "ultimate", "superior",
                    "durable", "best", "top quality"]            # UNVERIFIED list of "subjective claims"

# ---- Forbidden terms: whole-word, ignoring case, hyphens, spaces and dots. Edit freely.
FORBIDDEN_LCOM = ["L-Com", "LCom", "L Com", "L-COM"]
FORBIDDEN_PLATFORMS = ["Alibaba", "AliExpress", "Made-in-China"]
FORBIDDEN_COMPETITORS = ["Monoprice", "VCELINK", "Cables To Go", "Amazon Basics", "StarTech", "Tripp Lite"]
FORBIDDEN_PROMO = ["best", "cheap", "cheapest", "#1", "top rated", "free shipping", "guaranteed", "lifetime",
                   "same as", "compatible with L-Com", "OEM", "genuine", "original"]
GENERIC_SUPPLIER_WORDS = {"shenzhen", "dongguan", "ningbo", "guangzhou", "shanghai", "zhejiang", "guangdong", "technology",
                          "technologies", "electronic", "electronics", "industrial", "trading", "company", "limited"}
# Everyday product words a supplier name may contain; a single word from the name is only forbidden if it is not one of these.
COMMON_PRODUCT_WORDS = {"cable", "cables", "network", "networks", "connector", "connectors", "adapter", "adapters", "fiber",
                        "optic", "optical", "communication", "communications", "parts", "wire", "wires", "premier", "precision",
                        "international", "industry", "group", "hardware", "digital", "systems", "machinery", "automation"}
USABLE_STATUSES = ("sample_inspected", "page_read")
ALL_STATUSES = USABLE_STATUSES + ("listing_title", "unverified")
NON_CLAIM_FIELDS = ("selling_price", "upc")      # human-supplied columns, never shown to the model
TEMPLATE_FIELDS = ["product_type", "connector_a", "connector_b", "gender", "category_rating", "shielding", "mounting",
                   "material", "color", "plating", "pack_contents", "dimensions", "certifications", "selling_price", "upc"]
PRICE_PLACEHOLDER, UPC_PLACEHOLDER = "PRICE NEEDED", "UPC NEEDED (GS1)"
BRAND_PLACEHOLDER = "<BRAND NOT SET>"
DRAFT_BANNER = "Draft content for human review. Not published."

HERE = os.path.dirname(os.path.abspath(__file__))
FACTS_PATH = os.path.join(HERE, "product_facts.csv")
REVIEWER_EXCLUSIONS_CSV = os.path.join(HERE, "reviewer_exclusions.csv")
DEFAULT_OUT = os.path.join(HERE, "Walmart Listings")
FACT_COLUMNS = ["sku", "fact_id", "field", "value", "status", "source"]
CENT = Decimal("0.01")
MAX_BULLETS = LIMITS["bullets_max"]


class ContentError(Exception):
    """A refusal, shown as a plain message (nothing is kept)."""


# ---------- spec-token normalisation and scanning ----------

def canon(text: str) -> str:
    """Lower-case text with spec spellings collapsed: 'Cat.6', 'CAT-6', 'cat 6' -> 'cat6'; 'RJ-45' -> 'rj45'; '10 ft' -> '10ft'."""
    t = unicodedata.normalize("NFKC", str(text)).lower()
    t = re.sub(r"(?<![a-z0-9])cat[\s.\-]*(\d[a-z]?)(?![a-z0-9])", r"cat\1", t)
    t = re.sub(r"(?<![a-z0-9])ip[\s\-]?(\d\d)(?![0-9])", r"ip\1", t)
    t = re.sub(r"(?<![a-z0-9])rj[\s\-]?45", "rj45", t)
    t = re.sub(r"(?<![a-z0-9])8p[\s\-]?8c", "8p8c", t)
    t = re.sub(r"(?<![a-z0-9])n[\s\-]type", "ntype", t)
    t = re.sub(r"(?<![a-z0-9])db[\s\-]?9(?![0-9])", "db9", t)
    t = re.sub(r"(?<![a-z0-9])hdmi[\s\-]?(\d)", r"hdmi\1", t)
    t = re.sub(r"(?<![a-z0-9])panel[\s\-]*mount", "panelmount", t)
    t = re.sub(r"(?<![a-z0-9])in[\s\-]line(?![a-z0-9])", "inline", t)
    t = re.sub(r"(\d)\s+(mm|cm|m|ft|in|inch|gbps|mbps|ghz|v|a|ohm|kv|ka)(?![a-z0-9])", r"\1\2", t)
    return t


_WORDS = ("rj45|8p8c|hdmi|dvi|usb|bnc|sma|ntype|db9|shielded|unshielded|gold|nickel|brass|aluminum|aluminium|"
          "waterproof|panelmount|inline|female|male|angled")
TOKEN_RES = [
    re.compile(r"(?<![a-z0-9])cat\d[a-z]?(?![a-z0-9])"),
    re.compile(r"(?<![a-z0-9])ip\d\d(?![0-9])"),
    re.compile(r"(?<![a-z0-9.])\d+(?:\.\d+)?(?:mm|cm|m|ft|in|inch|gbps|mbps|ghz|v|a|ohm|kv|ka)(?![a-z0-9])"),
    re.compile(r"(?<![a-z0-9])[48]k(?![a-z0-9])"),
    re.compile(r"(?<![a-z0-9])hdmi\d(?:\.\d)?(?![a-z0-9])"),
    re.compile(rf"(?<![a-z0-9])(?:{_WORDS})(?![a-z0-9])"),
]
STRICT_CASE = [(re.compile(r"(?<![A-Za-z0-9])(LC|SC|ST|FC)(?![A-Za-z0-9])"), ""),
               (re.compile(r"(?<![A-Za-z0-9])(RoHS|FCC|UL|CE|REACH)(?![A-Za-z0-9])"), "cert:")]
LOOSE_CASE = [(re.compile(r"(?<![a-z0-9])(lc|sc|st|fc)(?![a-z0-9])"), ""),
              (re.compile(r"(?<![a-z0-9])(rohs|fcc|ul|ce|reach)(?![a-z0-9])"), "cert:")]


def spec_tokens(text: str, strict_case: bool) -> set:
    """Spec-bearing tokens in text. Generated text uses strict_case (LC/SC/ST/FC and certifications only count when
    standalone and upper-case); facts are read leniently."""
    c, out = canon(text), set()
    for rx in TOKEN_RES:
        out.update(re.sub(r"\s+", "", m.group(0)) for m in rx.finditer(c))
    out.update("qty:" + m[1] for m in re.finditer(r"pack of (\d+)", c))
    out.update("qty:" + m[1] for m in re.finditer(r"(?<![a-z0-9.])(\d+)\s?(?:pcs|pieces|pc)(?![a-z])", c))
    src, rules = (unicodedata.normalize("NFKC", str(text)), STRICT_CASE) if strict_case else (c, LOOSE_CASE)
    for rx, prefix in rules:
        out.update(prefix + m[1].lower() for m in rx.finditer(src))
    if not strict_case:
        out.update("hdmi" for t in list(out) if t.startswith("hdmi"))     # 'HDMI 2.0' in the facts also supports plain 'HDMI'
    return out


def term_regex(term: str):
    chars = [re.escape(c) for c in re.sub(r"[\s\-_.]+", "", term)]
    return re.compile(r"(?<![a-z0-9])" + r"[\s\-_.]*".join(chars) + r"(?![a-z0-9])", re.IGNORECASE)


def supplier_terms(name: str) -> list:
    name = (name or "").strip()
    if not name or name.lower().startswith("unknown") or name.lower() == "n/a":
        return []
    name = re.sub(r"\s*\([^)]*\)", "", name).strip() or name
    bare = re.sub(r"[,.]?\s*\b(co|ltd|limited|inc|llc|corp|company)\b\.?", "", name, flags=re.IGNORECASE).strip(" ,.")
    terms = {name, bare}
    first = next((w for w in re.split(r"\s+", bare) if len(w) >= 5 and w.lower().strip(",.") not in GENERIC_SUPPLIER_WORDS | COMMON_PRODUCT_WORDS), "")
    if first:
        terms.add(first.strip(",."))
    return [t for t in terms if t]


def forbidden_list(skus, supplier_names) -> list:
    """[(label, compiled regex)] for every forbidden term."""
    terms = [(t, "L-Com") for t in FORBIDDEN_LCOM] + [(t, "platform") for t in FORBIDDEN_PLATFORMS]
    terms += [(t, "competitor") for t in FORBIDDEN_COMPETITORS] + [(t, "promotional") for t in FORBIDDEN_PROMO]
    terms += [(s, "SKU / L-Com part number") for s in skus]
    terms += [(t, "supplier name") for n in supplier_names for t in supplier_terms(n)]
    seen, out = set(), []
    for term, kind in terms:
        key = re.sub(r"[\s\-_.]+", "", term).lower()
        if key and key not in seen:
            seen.add(key)
            out.append((f"{kind} '{term}'", term_regex(term)))
    return out


# ---------- the facts file ----------

def read_facts(path: str) -> list:
    """Fact rows (empty-value template rows skipped), read by header name. Refuses unknown statuses and duplicate ids."""
    try:
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            header = [h.strip() for h in (reader.fieldnames or [])]
    except FileNotFoundError:
        raise ContentError(f"Facts file not found: {path}. Run `facts-template` first, then fill it in.")
    missing = [h for h in FACT_COLUMNS if h not in header]
    if missing:
        raise ContentError(f"{os.path.basename(path)} is missing column(s): {', '.join(missing)}.")
    out, seen = [], set()
    for n, r in enumerate(rows, start=2):
        r = {k.strip(): (v or "").strip() for k, v in r.items() if k}
        if not r["value"]:
            continue
        if not r["sku"] or not r["fact_id"]:
            raise ContentError(f"{os.path.basename(path)} line {n}: sku and fact_id are required.")
        if r["status"] not in ALL_STATUSES:
            raise ContentError(f"{os.path.basename(path)} line {n}: status '{r['status']}' is not one of {', '.join(ALL_STATUSES)}.")
        if (r["sku"].upper(), r["fact_id"]) in seen:
            raise ContentError(f"{os.path.basename(path)} line {n}: duplicate fact_id {r['fact_id']} for {r['sku']}.")
        seen.add((r["sku"].upper(), r["fact_id"]))
        out.append(r)
    return out


def facts_for(rows: list, sku: str) -> list:
    return [r for r in rows if r["sku"].upper() == sku.upper()]


def usable_facts(rows: list, allow_listing: bool) -> list:
    """Facts the model may be shown and claims may use. unverified never; listing_title only with the flag."""
    ok = USABLE_STATUSES + (("listing_title",) if allow_listing else ())
    return [r for r in rows if r["status"] in ok and r["field"] not in NON_CLAIM_FIELDS]


def money_text(value):
    try:
        d = Decimal(str(value).replace("$", "").replace(",", "").strip())
    except InvalidOperation:
        return None
    return d.quantize(CENT, ROUND_HALF_UP) if d.is_finite() and d > 0 else None


def upc_ok(upc: str) -> bool:
    if not re.fullmatch(r"\d{12}", upc or ""):
        return False
    d = [int(c) for c in upc]
    return (10 - (3 * sum(d[0:11:2]) + sum(d[1:11:2])) % 10) % 10 == d[11]


# ---------- validators ----------

def sentences(text: str) -> list:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+|\n+", text or "") if s.strip()]


def nz(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def fields_of(listing: dict) -> list:
    out = [("title", listing.get("title") or ""), ("short_description", listing.get("short_description") or ""),
           ("long_description", listing.get("long_description") or "")]
    return out + [(f"key_feature_{i}", t) for i, t in enumerate(listing.get("key_features") or [], start=1)]


def validate_limits(listing: dict, usable_text: str) -> list:
    f, L = [], LIMITS
    title, short, long_, bullets = (listing.get("title") or ""), (listing.get("short_description") or ""), \
        (listing.get("long_description") or ""), (listing.get("key_features") or [])
    if not title.strip():
        f.append(("limits", "title", "", "title is empty"))
    if len(title) > L["title_max"]:
        f.append(("limits", "title", title, f"title is {len(title)} characters, over the {L['title_max']} limit"))
    bad = sorted({c for c in title if not TITLE_ALLOWED_CHARS.fullmatch(c)})
    if bad:
        f.append(("limits", "title", title, f"title has special characters: {' '.join(bad)}"))
    upper_ok = set(re.findall(r"[A-Z]{3,}", usable_text))
    caps = [w for w in re.findall(r"\b[A-Z]{3,}\b", title) if w not in upper_ok]
    if caps:
        f.append(("limits", "title", title, f"title has ALL CAPS word(s): {', '.join(caps)}"))
    low = nz(title)
    for w in TITLE_SUBJECTIVE:
        if re.search(rf"\b{re.escape(w)}\b", low):
            f.append(("limits", "title", title, f"title has a subjective claim: '{w}'"))
    if not short.strip():
        f.append(("limits", "short_description", "", "short description is empty"))
    if len(short) > L["short_max"]:
        f.append(("limits", "short_description", short, f"short description is {len(short)} characters, over the {L['short_max']} limit"))
    words = len(long_.split())
    if len(long_) > L["long_max"]:
        f.append(("limits", "long_description", long_[:200], f"long description is {len(long_)} characters, over the {L['long_max']} limit"))
    if words < L["long_min_words"]:
        f.append(("limits", "long_description", long_[:200], f"long description has {words} words, under the {L['long_min_words']} minimum"))
    if not L["bullets_min"] <= len(bullets) <= L["bullets_max"]:
        f.append(("limits", "key_features", "", f"{len(bullets)} key features; need {L['bullets_min']}-{L['bullets_max']}"))
    for i, b in enumerate(bullets, start=1):
        if len(b) > L["bullet_max"]:
            f.append(("limits", f"key_feature_{i}", b, f"key feature is {len(b)} characters, over the {L['bullet_max']} limit"))
    for field, text in fields_of(listing):
        if re.search(r"[<>]", text):
            f.append(("limits", field, text[:200], "contains < or >: plain text only, no HTML"))
        if re.search(r"\*\*|__|`|^\s*#|^\s*[-*•]\s|\[[^\]]+\]\([^)]+\)", text, re.MULTILINE):
            f.append(("limits", field, text[:200], "contains markdown markup: plain text only"))
    return f


def validate_evidence(claims, usable: dict, all_facts: dict) -> list:
    f = []
    if not isinstance(claims, list) or not claims:
        return [("evidence", "claims", "", "no claims were given")]
    for c in claims:
        text = c.get("text", "") if isinstance(c, dict) else str(c)
        ids = c.get("fact_ids") if isinstance(c, dict) else None
        if not isinstance(ids, list) or not ids:
            f.append(("evidence", "claims", text, "claim has no fact_ids"))
            continue
        for i in ids:
            if i in usable:
                continue
            if i in all_facts:
                st = all_facts[i]["status"]
                why = ("an unverified fact is never usable" if st == "unverified" else
                       "a listing_title fact needs --allow-listing-facts" if st == "listing_title" else f"status {st} is not usable")
                f.append(("evidence", "claims", text, f"fact {i} cannot be used: {why}"))
            else:
                f.append(("evidence", "claims", text, f"fact_id {i} does not exist for this SKU"))
    return f


def validate_coverage(listing: dict, claims) -> list:
    """Every title, key feature and long-description sentence must sit inside some claim's text."""
    claim_texts = [nz(c.get("text", "")) for c in (claims or []) if isinstance(c, dict)]
    f = []
    items = [("title", listing.get("title") or "")] + [(f"key_feature_{i}", t) for i, t in enumerate(listing.get("key_features") or [], start=1)]
    items += [("long_description", s) for s in sentences(listing.get("long_description") or "")]
    for field, text in items:
        if text.strip() and not any(nz(text) in ct for ct in claim_texts):
            f.append(("evidence", field, text, "not covered by any claim (every title, bullet and description sentence needs a claim with fact_ids)"))
    return f


def validate_spec_tokens(listing: dict, fact_tokens: set) -> list:
    f = []
    for field, text in fields_of(listing):
        for tok in sorted(spec_tokens(text, strict_case=True) - fact_tokens):
            f.append(("spec_token", field, text[:200], f"'{tok}' is not supported by any usable fact"))
    return f


def validate_forbidden(listing: dict, forbidden: list) -> list:
    f = []
    for field, text in fields_of(listing):
        for label, rx in forbidden:
            m = rx.search(text)
            if m:
                f.append(("forbidden", field, text[:200], f"contains forbidden {label} ('{m.group(0)}')"))
    return f


def validate_brand(listing: dict, brand, configured) -> list:
    f = []
    b = (brand or "").strip()
    if not b:
        return [("brand", "brand", "", "brand is empty")]
    if len(b) > LIMITS["brand_max"]:
        f.append(("brand", "brand", b, f"brand is {len(b)} characters, over the {LIMITS['brand_max']} limit"))
    if len(term_regex(b).findall(listing.get("title") or "")) > 1:
        f.append(("brand", "title", listing.get("title", ""), "brand appears more than once in the title"))
    key = re.sub(r"[\s\-_.]+", "", b).lower()
    for t in FORBIDDEN_LCOM + FORBIDDEN_COMPETITORS:
        if key == re.sub(r"[\s\-_.]+", "", t).lower():
            f.append(("brand", "brand", b, f"brand equals the forbidden name '{t}'"))
    if configured and b != configured:
        f.append(("brand", "brand", b, f"brand differs from the configured BRAND_NAME '{configured}'"))
    return f


def title_brand_warning(title: str, brand) -> list:
    """Warning only (never blocks): the title should start with the configured brand, ignoring case. Skipped when no brand is set."""
    if not brand or not (title or "").strip():
        return []
    if re.match(re.escape(brand.strip()) + r"(?![A-Za-z0-9])", title.strip(), re.IGNORECASE):
        return []
    return [f"title does not start with the brand '{brand}': {title.strip()[:80]}"]


def validate_money(listing: dict) -> list:
    f = []
    for field, text in fields_of(listing):
        if re.search(r"\$|\busd\b|\bdollars?\b|\bcents?\b", text, re.IGNORECASE):
            f.append(("price", field, text[:200], "contains a price or cost figure: the model never states money"))
    return f


def validate_price_upc(price_text, upc_text) -> list:
    f = []
    if price_text not in (PRICE_PLACEHOLDER, None) and not (isinstance(price_text, (int, float, Decimal)) and money_text(price_text)):
        f.append(("price", "selling_price", str(price_text), "selling price is not a positive amount"))
    if upc_text not in (UPC_PLACEHOLDER, None) and not upc_ok(str(upc_text)):
        f.append(("upc", "upc", str(upc_text), "UPC must be 12 digits with a valid check digit"))
    return f


def validate_all(listing: dict, ctx: dict, cover: bool = True) -> list:
    """Every validator, in order. ctx: usable, all_facts (id -> fact), forbidden, brand, configured_brand, claims."""
    usable_vals = " ".join(r["value"] for r in ctx["usable"].values())
    tokens = set()
    for r in ctx["usable"].values():
        tokens |= spec_tokens(r["value"], strict_case=False)
    out = validate_limits(listing, usable_vals)
    out += validate_evidence(ctx["claims"], ctx["usable"], ctx["all_facts"])
    if cover:
        out += validate_coverage(listing, ctx["claims"])
    out += validate_spec_tokens(listing, tokens)
    out += validate_forbidden(listing, ctx["forbidden"])
    out += validate_brand(listing, ctx["brand"], ctx["configured_brand"])
    out += validate_money(listing)
    return out


# ---------- the model call ----------

SYSTEM = ("You write Walmart Marketplace listing text. You may state ONLY facts from the list you are given: never add a "
          "specification, a connector, a rating, a certification, a material, a size, a quantity, a compatibility claim or a "
          "benefit that is not in those facts. Never mention prices, costs, other brands, other sellers or stores. Never output "
          "the SKU. Plain text only: no HTML, no markdown, no emoji, no ALL CAPS words in the title. Reply with one JSON object "
          "and nothing else (no code fence).")


def build_prompt(sku: str, brand: str, facts: list) -> str:
    schema = {"title": "...", "short_description": "...", "long_description": "...", "key_features": ["..."],
              "claims": [{"text": "<the exact sentence, bullet or title>", "fact_ids": ["F1"]}]}
    return json.dumps({
        "sku": sku, "brand": brand,
        "usable_facts": [{"fact_id": r["fact_id"], "field": r["field"], "value": r["value"]} for r in facts],
        "limits": {"title_max_chars": LIMITS["title_max"], "title_recommended_chars": list(LIMITS["title_recommended"]),
                   "short_description_max_chars": LIMITS["short_max"], "long_description_max_chars": LIMITS["long_max"],
                   "long_description_min_words": LIMITS["long_min_words"], "key_features": [LIMITS["bullets_min"], LIMITS["bullets_max"]],
                   "key_feature_max_chars": LIMITS["bullet_max"], "brand_max_chars": LIMITS["brand_max"]},
        "rules": ["Every sentence of long_description, every key feature and the title must appear word for word as the "
                  "text of a claim, and each claim lists the fact_ids that support it.",
                  "Use only the facts above. If the facts are too thin to reach the minimum length, write fewer claims "
                  "rather than inventing anything."],
        "reply_format": schema}, ensure_ascii=False, indent=1)


def parse_model_json(text: str):
    """The strict schema, or None. Fences, commentary and extra keys are invalid."""
    try:
        d = json.loads(text)
    except (ValueError, TypeError):
        return None
    keys = {"title", "short_description", "long_description", "key_features", "claims"}
    if not isinstance(d, dict) or set(d) != keys:
        return None
    if not all(isinstance(d[k], str) for k in ("title", "short_description", "long_description")):
        return None
    if not isinstance(d["key_features"], list) or not all(isinstance(x, str) for x in d["key_features"]):
        return None
    if not isinstance(d["claims"], list) or not all(isinstance(c, dict) and isinstance(c.get("text"), str)
                                                    and isinstance(c.get("fact_ids"), list) for c in d["claims"]):
        return None
    return d


BILLING_WORDS = re.compile(r"credit balance|billing|payment|insufficient (?:funds|credit)|api key|authenticat", re.IGNORECASE)


def is_billing_error(exc) -> bool:
    """401/402/403, or a 400 that mentions credit, billing or auth (same rule as sourcing_agent.py, copied)."""
    code = getattr(exc, "status_code", None)
    return code in (401, 402, 403) or (code == 400 and bool(BILLING_WORDS.search(str(getattr(exc, "message", "") or exc))))


def live_model():
    """call(system, messages, max_tokens) -> (text, input_tokens, output_tokens). The key is read from the environment only."""
    key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not key:
        raise ContentError("ANTHROPIC_API_KEY is not set in .env. Nothing was sent. (Use --dry-run to see the prompts.)")
    from anthropic import Anthropic
    client = Anthropic(api_key=key)

    def call(system, messages, max_tokens):
        msg = client.messages.create(model=MODEL, max_tokens=max_tokens, temperature=TEMPERATURE, system=system, messages=messages)
        return "".join(b.text for b in msg.content if getattr(b, "type", "") == "text"), msg.usage.input_tokens, msg.usage.output_tokens
    return call


def call_cost(tokens_in: int, tokens_out: int) -> Decimal:
    return ((Decimal(tokens_in) * PRICE_PER_MTOK["input"] + Decimal(tokens_out) * PRICE_PER_MTOK["output"]) / Decimal(1_000_000))


class Billing(Exception):
    pass


def run_sku(sku, brand, usable, ctx_base, call, log):
    """One SKU: call, parse (one JSON retry), validate, at most one repair. Returns (listing, claims, failures, status)."""
    prompt = build_prompt(sku, brand, usable)
    messages = [{"role": "user", "content": prompt}]

    def ask(msgs):
        try:
            text, tin, tout = call(SYSTEM, msgs, MAX_TOKENS)
        except Exception as e:                                        # noqa: BLE001 - classified below
            if is_billing_error(e):
                raise Billing(str(e)[:200])
            raise
        log.append({"sku": sku, "in": tin, "out": tout, "cost": call_cost(tin, tout)})
        return text

    text = ask(messages)
    data = parse_model_json(text)
    if data is None:
        text = ask(messages + [{"role": "assistant", "content": text}, {"role": "user", "content":
                    "That was not the required JSON. Reply with only the JSON object in the reply_format, no code fence, no commentary."}])
        data = parse_model_json(text)
        if data is None:
            return {}, [], [("json", "model", text[:200], "the model's reply was not valid JSON twice")], "BLOCKED"
    ctx = {**ctx_base, "claims": data["claims"]}
    failures = validate_all(data, ctx)
    if failures:                                                      # one repair attempt, validated again
        listed = "\n".join(f"- [{r}] {fld}: {why} | {txt[:120]}" for r, fld, txt, why in failures[:25])
        text = ask(messages + [{"role": "assistant", "content": text}, {"role": "user", "content":
                    "Code checks found these problems. Fix them using ONLY the usable facts and reply with corrected JSON only:\n" + listed}])
        fixed = parse_model_json(text)
        if fixed is None:
            failures = failures + [("json", "model", text[:200], "the repair reply was not valid JSON")]
        else:
            data = fixed
            failures = validate_all(data, {**ctx_base, "claims": data["claims"]})
    return data, data.get("claims", []), failures, ("READY FOR REVIEW" if not failures else "BLOCKED")


# ---------- selecting the products ----------

def select_lines(approved, results, exclusions, at, sku_filter):
    """[(sku, ap, row, cand)] one per SKU, plus notes for everything skipped. Same safety as the order sheet."""
    by_sku = {str(r["SKU"]).strip().upper(): r for r in results}
    lines, notes, seen = [], [], set()
    for ap in approved:
        sku = ap["sku"]
        if sku_filter and sku.upper() not in sku_filter:
            continue
        if sku.upper() in seen:
            continue
        if at >= approval_expiry(ap):
            notes.append(f"{sku}: approval expired, skipped")
            continue
        row = by_sku.get(sku.upper())
        cand = find_candidate(row, sku, ap["listing_url"]) if row else None
        if cand is None:
            notes.append(f"{sku}: approved listing not found in the results file (matched by SKU + URL), skipped")
            continue
        why = excluded_reason(sku, cand["name"], cand["url"], exclusions)
        if why:
            notes.append(f"{sku}: excluded by the reviewer ({why}), skipped")
            continue
        seen.add(sku.upper())
        lines.append((sku, ap, row, cand))
    return lines, notes


# ---------- facts-template ----------

def facts_template(approved_path, results_path, facts_path=None, at=None, exclusions_path=None) -> tuple:
    facts_path = facts_path or FACTS_PATH
    if os.path.exists(facts_path):
        raise ContentError(f"{facts_path} already exists and was NOT touched. Rename or delete it first if you want a fresh template.")
    at = at or datetime.now(timezone.utc)
    approved, _ = read_approved(approved_path)
    results = read_results(results_path)
    lines, notes = select_lines(approved, results, load_exclusions(exclusions_path or REVIEWER_EXCLUSIONS_CSV), at, None)
    if not lines:
        raise ContentError("No approved, unexpired product to write facts for: " + "; ".join(notes))
    rows = []
    for sku, ap, row, cand in lines:
        n = 0

        def add(field, value, status, source):
            nonlocal n
            n += 1
            rows.append({"sku": sku, "fact_id": f"F{n}", "field": field, "value": value, "status": status, "source": source})
        kw = str(row.get("Keyword") or "").strip()
        if kw:
            add("product_type", kw, "unverified", "results file Keyword (search term); confirm against the sample")
        for part in [p.strip() for p in re.split(r"[;|,]", cand["title"]) if p.strip()]:
            add("spec_line", part, "listing_title", "seller's listing title, not yet checked on the page or the sample")
        note = str(row.get("Ordering Note") or "").strip() if str(row.get("Recommended URL") or "") == cand["url"] else ""
        for part in [p.strip() for p in re.split(r";|\.\s", note) if p.strip()]:
            add("spec_line", part, "unverified", "results file Ordering Note (the search agent's reading, unverified)")
        for fld in TEMPLATE_FIELDS:
            n += 1
            rows.append({"sku": sku, "fact_id": f"F{n}", "field": fld, "value": "", "status": "", "source": ""})
    with open(facts_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FACT_COLUMNS)
        w.writeheader()
        w.writerows(rows)
    return facts_path, len(lines), notes


# ---------- workbook ----------

NAVY = PatternFill("solid", fgColor="FF1F3864")
LISTING_FILL = PatternFill("solid", fgColor="FFFFF2CC")
RED_FILL = PatternFill("solid", fgColor="FFFFC7CE", bgColor="FFFFC7CE")   # Excel reads bgColor for conditional formats
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def arial(**kw):
    return Font(name="Arial", size=kw.pop("size", 10), **kw)


def header_row(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row, col, text)
        c.font, c.fill, c.border = arial(bold=True, color="FFFFFFFF"), NAVY, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 31.5


def body_cell(ws, row, col, value, **kw):
    c = ws.cell(row, col, value)
    c.font, c.border = arial(**kw), BOX
    c.alignment = Alignment(vertical="top", wrap_text=True)
    return c


LISTING_HEADERS = (["SKU", "Status", "Brand", "Title", "Short description", "Long description"]
                   + [f"Key feature {i}" for i in range(1, MAX_BULLETS + 1)]
                   + ["Selling price", "UPC", "Title length", "Short description length", "Long description length",
                      "Long description words"] + [f"Key feature {i} length" for i in range(1, MAX_BULLETS + 1)])
HEADER_ROW = 4
C_TITLE, C_SHORT, C_LONG, C_B1 = 4, 5, 6, 7
C_PRICE, C_UPC = C_B1 + MAX_BULLETS, C_B1 + MAX_BULLETS + 1
C_TLEN = C_UPC + 1


def write_workbook(path, items, run_notes, brand, allow_listing, skus, suppliers):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listings"
    ws["A1"], ws["A1"].font = f"Walmart listing content. {DRAFT_BANNER}", arial(size=14, bold=True)
    ws["A2"], ws["A2"].font = ("Nothing here was sent to Walmart. A person must read every word, fill the yellow-marked gaps and upload by hand. "
                               "Counters turn red when a limit (unverified, see Notes) is broken."), arial()
    header_row(ws, HEADER_ROW, LISTING_HEADERS)
    for i, it in enumerate(items):
        x, L = HEADER_ROW + 1 + i, it["listing"]
        vals = [it["sku"], it["status"], it["brand"], L.get("title", ""), L.get("short_description", ""), L.get("long_description", "")]
        feats = list(L.get("key_features") or [])[:MAX_BULLETS]
        vals += feats + [""] * (MAX_BULLETS - len(feats)) + [it["price"], it["upc"]]
        for col, v in enumerate(vals, start=1):
            body_cell(ws, x, col, v if v != "" else None)
        ws.cell(x, C_UPC).number_format = "@"
        if isinstance(it["price"], float):
            ws.cell(x, C_PRICE).number_format = "\\$#,##0.00"
        for col in (C_PRICE, C_UPC):
            if isinstance(ws.cell(x, col).value, str):
                ws.cell(x, col).fill = LISTING_FILL
        t, s, lg = (get_column_letter(c) for c in (C_TITLE, C_SHORT, C_LONG))
        ws.cell(x, C_TLEN, f'=LEN({t}{x})')
        ws.cell(x, C_TLEN + 1, f'=LEN({s}{x})')
        ws.cell(x, C_TLEN + 2, f'=LEN({lg}{x})')
        ws.cell(x, C_TLEN + 3, f'=IF(TRIM({lg}{x})="",0,LEN(TRIM({lg}{x}))-LEN(SUBSTITUTE(TRIM({lg}{x})," ",""))+1)')
        for b in range(MAX_BULLETS):
            bc = get_column_letter(C_B1 + b)
            ws.cell(x, C_TLEN + 4 + b, f'=IF({bc}{x}="","",LEN({bc}{x}))')
        for col in range(C_TLEN, len(LISTING_HEADERS) + 1):
            body_cell(ws, x, col, ws.cell(x, col).value)
        ws.row_dimensions[x].height = 120
    last = HEADER_ROW + len(items)
    first = HEADER_ROW + 1
    if items:
        def red(col, rule):
            L_ = get_column_letter(col)
            ws.conditional_formatting.add(f"{L_}{first}:{L_}{last}", FormulaRule(formula=[rule.format(c=f"{L_}{first}")], fill=RED_FILL))
        red(C_TLEN, f"AND(ISNUMBER({{c}}),{{c}}>{LIMITS['title_max']})")
        red(C_TLEN + 1, f"AND(ISNUMBER({{c}}),{{c}}>{LIMITS['short_max']})")
        red(C_TLEN + 2, f"AND(ISNUMBER({{c}}),{{c}}>{LIMITS['long_max']})")
        red(C_TLEN + 3, f"AND(ISNUMBER({{c}}),{{c}}<{LIMITS['long_min_words']})")
        for b in range(MAX_BULLETS):
            red(C_TLEN + 4 + b, f"AND(ISNUMBER({{c}}),{{c}}>{LIMITS['bullet_max']})")
    widths = [16, 18, 14, 40, 40, 70] + [30] * MAX_BULLETS + [13, 18] + [11] * (4 + MAX_BULLETS)
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = f"B{HEADER_ROW + 1}"

    wf = wb.create_sheet("Facts used")
    header_row(wf, 1, ["SKU", "Fact ID", "Field", "Value", "Status", "Source", "Claims that relied on it", "Note"])
    r = 2
    for it in items:
        for fact in it["usable"]:
            used = " | ".join(c["text"] for c in it["claims"] if isinstance(c, dict) and fact["fact_id"] in (c.get("fact_ids") or []))
            note = "LISTING-TITLE FACT: not checked on the page or sample; used only because --allow-listing-facts was set" \
                if fact["status"] == "listing_title" else ""
            for col, v in enumerate([it["sku"], fact["fact_id"], fact["field"], fact["value"], fact["status"], fact["source"], used, note], start=1):
                c = body_cell(wf, r, col, v or None)
                if note:
                    c.fill = LISTING_FILL
            r += 1
    for col, w in enumerate([16, 9, 18, 40, 18, 40, 60, 50], start=1):
        wf.column_dimensions[get_column_letter(col)].width = w
    wf.freeze_panes = "A2"

    wr = wb.create_sheet("Review notes")
    header_row(wr, 1, ["SKU", "Rule", "Field", "Text that failed", "Why"])
    r = 2
    for it in items:
        for rule, fld, txt, why in it["failures"]:
            for col, v in enumerate([it["sku"], rule, fld, txt, why], start=1):
                body_cell(wr, r, col, v or None)
            r += 1
    for it in items:
        for why in it.get("warnings", []):
            for col, v in enumerate([it["sku"], "WARNING", "title", "", why], start=1):
                body_cell(wr, r, col, v or None)
            r += 1
    for note in run_notes:
        for col, v in enumerate(["(run)", "run", "", "", note], start=1):
            body_cell(wr, r, col, v or None)
        r += 1
    for col, w in enumerate([16, 13, 20, 60, 70], start=1):
        wr.column_dimensions[get_column_letter(col)].width = w
    wr.freeze_panes = "A2"

    wt = wb.create_sheet("To do before upload")
    header_row(wt, 1, ["SKU", "Still to do (this tool does not do these)", "Done?"])
    todo = ["Product photos, taken from the samples", "Walmart category and item-spec attributes, filled in from the seller account's item spec "
            "(this tool does not invent a category)", "UPC from GS1 (never generated here)", "Final selling price",
            "Brand registration, if Walmart requires it for this brand", "A person reads every word against the sample and the supplier page"]
    r = 2
    for it in items:
        for t in todo:
            for col, v in enumerate([it["sku"], t, "No"], start=1):
                body_cell(wt, r, col, v)
            r += 1
    for col, w in enumerate([16, 90, 10], start=1):
        wt.column_dimensions[get_column_letter(col)].width = w

    wc = wb.create_sheet("Check data")
    header_row(wc, 1, ["Kind", "SKU", "Value", "Extra"])
    rows = [("brand", "", brand, ""), ("allow_listing_facts", "", "yes" if allow_listing else "no", "")]
    rows += [("sku", "", s, "") for s in skus] + [("supplier", "", s, "") for s in suppliers]
    for it in items:
        rows += [("claim", it["sku"], c.get("text", ""), ",".join(c.get("fact_ids") or [])) for c in it["claims"] if isinstance(c, dict)]
    for r, row in enumerate(rows, start=2):
        for col, v in enumerate(row, start=1):
            body_cell(wc, r, col, v or None)
    wc["F1"], wc["F1"].font = "Used by `check` to re-run the validators. Do not edit.", arial(italic=True)
    for col, w in enumerate([20, 16, 80, 30], start=1):
        wc.column_dimensions[get_column_letter(col)].width = w
    wb.save(path)


def write_markdown(path, items):
    out = [f"# Walmart listing content\n\n**{DRAFT_BANNER}**\n"]
    for it in items:
        L = it["listing"]
        out.append(f"\n## {it['sku']} - {it['status']}\n\nBrand: {it['brand']} | Price: {it['price']} | UPC: {it['upc']}\n")
        out.append(f"\n**Title:** {L.get('title', '')}\n\n**Short description:** {L.get('short_description', '')}\n\n"
                   f"**Long description:**\n\n{L.get('long_description', '')}\n\n**Key features:**\n")
        out += [f"- {b}\n" for b in L.get("key_features") or []]
        if it["failures"]:
            out.append("\n**Problems to fix:**\n")
            out += [f"- [{r}] {fld}: {why}\n" for r, fld, _, why in it["failures"]]
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(out))


# ---------- generate ----------

def configured_brand():
    return (os.environ.get("BRAND_NAME") or "").strip() or None


def generate(approved_path, results_path, facts_path=None, skus=None, allow_listing=False, dry_run=False, call_model=None,
             out_dir=None, at=None, exclusions_path=None) -> dict:
    at = at or datetime.now(timezone.utc)
    brand = configured_brand()
    if not brand and not dry_run:
        raise ContentError("BRAND_NAME is not set (environment or .env). Refusing to generate listings without the brand. "
                           "(--dry-run, facts-template and check work without it.)")
    if brand and len(brand) > LIMITS["brand_max"]:
        raise ContentError(f"BRAND_NAME is {len(brand)} characters; the limit is {LIMITS['brand_max']} (unverified).")
    approved, _ = read_approved(approved_path)
    results = read_results(results_path)
    facts = read_facts(facts_path or FACTS_PATH)
    lines, notes = select_lines(approved, results, load_exclusions(exclusions_path or REVIEWER_EXCLUSIONS_CSV), at,
                                {s.upper() for s in skus} if skus else None)
    if not lines:
        raise ContentError("No approved, unexpired product to write for: " + ("; ".join(notes) or "nothing matched --sku"))
    forbidden_skus = sorted({ap["sku"] for ap in approved} | {str(r["SKU"]).strip() for r in results})
    suppliers = sorted({n for _, _, row, _ in lines for n in [str(row.get("Recommended Manufacturer") or "")] + [c["name"] for c in candidates_of(row)] if n})
    forbidden = forbidden_list(forbidden_skus, suppliers)
    shown_brand = brand or BRAND_PLACEHOLDER
    plans = []
    for sku, ap, row, cand in lines:
        mine = facts_for(facts, sku)
        usable = usable_facts(mine, allow_listing)
        plans.append((sku, mine, usable))
    if dry_run:
        for sku, mine, usable in plans:
            print(f"===== DRY RUN {sku}: {len(usable)} usable fact(s), nothing is sent =====\nSYSTEM: {SYSTEM}\n\nUSER:\n{build_prompt(sku, shown_brand, usable)}\n")
            if not usable:
                print(f"(no usable facts for {sku}: a real run would BLOCK it without calling the model)\n")
        for n in notes:
            print("SKIPPED: " + n)
        return {"dry_run": True, "skus": [p[0] for p in plans]}
    calling = [p for p in plans if p[2]]
    est = sum((call_cost(len(build_prompt(s, shown_brand, u)) // 3, MAX_TOKENS) * MAX_CALLS_PER_SKU for s, _, u in calling), Decimal(0))
    if est > COST_WARN:
        print(f"WARNING: worst-case cost for this run is about ${est:.2f} (up to {MAX_CALLS_PER_SKU} calls per SKU), over ${COST_WARN}.")
    call = call_model or live_model()
    log, items, run_notes = [], [], list(notes)
    stopped = None
    for sku, mine, usable in plans:
        facts_by_id = {r["fact_id"]: r for r in mine}
        usable_by_id = {r["fact_id"]: r for r in usable}
        price_row = next((r for r in mine if r["field"] == "selling_price"), None)
        upc_row = next((r for r in mine if r["field"] == "upc"), None)
        price = PRICE_PLACEHOLDER if not price_row else (float(money_text(price_row["value"])) if money_text(price_row["value"]) else price_row["value"])
        upc = UPC_PLACEHOLDER if not upc_row else upc_row["value"].strip()
        extra = validate_price_upc(price, upc)
        item = {"sku": sku, "brand": shown_brand, "usable": usable, "price": price, "upc": upc, "listing": {}, "claims": [], "failures": [],
                "warnings": []}
        if stopped:
            item.update(status="NOT RUN", failures=[("run", "model", "", f"not run: {stopped}")])
        elif not usable:
            item.update(status="BLOCKED", failures=[("evidence", "facts", "", "no usable facts for this SKU (fill product_facts.csv with sample_inspected or page_read facts)")] + extra)
        else:
            ctx = {"usable": usable_by_id, "all_facts": facts_by_id, "forbidden": forbidden, "brand": brand, "configured_brand": brand}
            try:
                listing, claims, failures, status = run_sku(sku, brand, usable, ctx, call, log)
            except Billing as e:
                stopped = f"billing/auth error from the API, batch stopped ({e})"
                run_notes.append(f"{sku}: {stopped}")
                item.update(status="NOT RUN", failures=[("run", "model", "", f"not run: {stopped}")])
                items.append(item)
                continue
            failures = failures + extra
            item.update(listing=listing, claims=claims, failures=failures, status="READY FOR REVIEW" if not failures else "BLOCKED")
            item["warnings"] = title_brand_warning(listing.get("title", ""), brand)
            if listing:
                tl = len(listing.get("title", ""))
                lo, hi = LIMITS["title_recommended"]
                if not lo <= tl <= hi:
                    run_notes.append(f"{sku}: title is {tl} characters; Walmart recommends {lo}-{hi} (unverified advice, not a block)")
        items.append(item)
    total = sum((e["cost"] for e in log), Decimal(0))
    for e in log:
        print(f"[cost] {e['sku']}: {MODEL} in={e['in']} out={e['out']} ${e['cost']:.5f}")
    print(f"[cost] total ${total:.5f} for {len(log)} call(s)")
    os.makedirs(out_dir or DEFAULT_OUT, exist_ok=True)
    stamp = at.strftime("%Y%m%d_%H%M%S")
    xlsx, md = (os.path.join(out_dir or DEFAULT_OUT, f"walmart_listings_{stamp}{ext}") for ext in (".xlsx", ".md"))
    write_workbook(xlsx, items, run_notes, shown_brand, allow_listing, forbidden_skus, suppliers)
    write_markdown(md, items)
    return {"items": items, "path": xlsx, "md": md, "cost": total, "calls": len(log), "notes": run_notes}


# ---------- check ----------

def read_sheet(wb, name, header_name="SKU", header_row=None):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hi = header_row if header_row is not None else next((i for i, r in enumerate(rows) if r and r[0] == header_name), None)
    if hi is None:
        raise ContentError(f"sheet '{name}' has no header row starting with '{header_name}'.")
    head = [str(h) if h is not None else "" for h in rows[hi]]
    return [dict(zip(head, r)) for r in rows[hi + 1:] if r and r[0] not in (None, "")]


def check_workbook(path: str) -> dict:
    """Re-run every validator on a (possibly human-edited) listings file. Returns {sku: [failures]} and warnings."""
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except FileNotFoundError:
        raise ContentError(f"File not found: {path}")
    for need in ("Listings", "Facts used", "Check data"):
        if need not in wb.sheetnames:
            raise ContentError(f"{os.path.basename(path)} has no '{need}' sheet: it was not made by this tool's `generate`.")
    listings = read_sheet(wb, "Listings")
    facts = read_sheet(wb, "Facts used")
    data = read_sheet(wb, "Check data", "Kind")
    warnings = []
    cfg_brand = configured_brand()
    if not cfg_brand:
        warnings.append("BRAND_NAME is not set: skipped the comparison against the configured brand (all other brand checks ran).")
    skus = [d["Value"] for d in data if d["Kind"] == "sku"]
    suppliers = [d["Value"] for d in data if d["Kind"] == "supplier"]
    forbidden = forbidden_list(skus, suppliers)
    result = {}
    for row in listings:
        sku = str(row["SKU"])
        mine = [{"sku": sku, "fact_id": str(f["Fact ID"]), "field": str(f["Field"] or ""), "value": str(f["Value"] or ""),
                 "status": str(f["Status"] or ""), "source": str(f["Source"] or "")} for f in facts if f["SKU"] == sku]
        usable = {r["fact_id"]: r for r in mine}                     # the sheet lists exactly the facts that were usable
        claims = [{"text": str(d["Value"] or ""), "fact_ids": [i for i in str(d["Extra"] or "").split(",") if i]}
                  for d in data if d["Kind"] == "claim" and d["SKU"] == sku]
        listing = {"title": str(row.get("Title") or ""), "short_description": str(row.get("Short description") or ""),
                   "long_description": str(row.get("Long description") or ""),
                   "key_features": [str(row[h]) for h in (f"Key feature {i}" for i in range(1, MAX_BULLETS + 1)) if row.get(h) not in (None, "")]}
        ctx = {"usable": usable, "all_facts": usable, "forbidden": forbidden, "brand": str(row.get("Brand") or ""),
               "configured_brand": cfg_brand, "claims": claims}
        if not (listing["title"] or listing["short_description"] or listing["long_description"] or listing["key_features"]):
            result[sku] = [("run", "listing", "", f"no content to check: status was {row.get('Status')} (nothing was generated)")]
            continue
        failures = validate_all(listing, ctx, cover=False)           # edited text can't be matched to the old claims word for word
        price, upc = row.get("Selling price"), row.get("UPC")
        failures += validate_price_upc(price, None if upc is None else str(upc))
        result[sku] = failures
    return {"results": result, "warnings": warnings}


# ---------- CLI ----------

def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stdin, "reconfigure"):
        sys.stdin.reconfigure(encoding="utf-8")
    load_dotenv(os.path.join(HERE, ".env"))
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    for name in ("facts-template", "generate"):
        s = sub.add_parser(name)
        s.add_argument("--approved", required=True)
        s.add_argument("--results", required=True)
        s.add_argument("--facts", default=None, help="facts file (default product_facts.csv next to the script)")
        if name == "generate":
            s.add_argument("--sku", nargs="+", default=None)
            s.add_argument("--allow-listing-facts", action="store_true")
            s.add_argument("--dry-run", action="store_true")
            s.add_argument("--out", default=DEFAULT_OUT)
    sub.add_parser("check").add_argument("file")
    args = p.parse_args(argv)
    try:
        if args.cmd == "facts-template":
            path, n, notes = facts_template(args.approved, args.results, args.facts)
            print(f"Wrote {path} for {n} product(s). Fill in the empty rows and upgrade statuses to sample_inspected / page_read only after a person checks.")
            for note in notes:
                print("SKIPPED: " + note)
        elif args.cmd == "generate":
            out = generate(args.approved, args.results, args.facts, args.sku, args.allow_listing_facts, args.dry_run, out_dir=args.out)
            if not out.get("dry_run"):
                for it in out["items"]:
                    print(f"{it['sku']}: {it['status']}" + "".join(f"\n    [{r}] {fld}: {why}" for r, fld, _, why in it["failures"]))
                for n in out["notes"]:
                    print("NOTE: " + n)
                print(f"Written: {out['path']}\n         {out['md']}")
        else:
            out = check_workbook(args.file)
            for w in out["warnings"]:
                print("WARNING: " + w)
            bad = 0
            for sku, failures in out["results"].items():
                print(f"{sku}: " + ("OK" if not failures else f"{len(failures)} problem(s)"))
                for r, fld, txt, why in failures:
                    print(f"    [{r}] {fld}: {why}")
                bad += bool(failures)
            if bad:
                sys.exit(1)
    except (ContentError, OrderSheetError) as e:
        sys.exit(f"content_agent: {e}")


if __name__ == "__main__":
    main()
