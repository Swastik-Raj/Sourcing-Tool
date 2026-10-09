"""Decision / approval agent: turns the search agent's recommendations into an approval request for Neeraj, records the
reply a person pastes in, and exports the approvals that are still valid for the (later) order-generation step.

NOTHING HERE ORDERS ANYTHING OR SPENDS MONEY, and there are no LLM, network, Slack or email calls: all parsing and
arithmetic is deterministic. It only reads sourcing_results_*.xlsx, email_state.json and reviewer_exclusions.csv.

    python decision_agent.py request <results.xlsx> [--budget N]             # write the approval request
    python decision_agent.py record  --results <results.xlsx> --request A-000N [--file reply.txt]   # paste Neeraj's reply
    python decision_agent.py status  --results <results.xlsx>                # requested/approved/rejected/expired/revoked
    python decision_agent.py revoke  <SKU> --results <results.xlsx>          # cancel an approval
    python decision_agent.py export  --results <results.xlsx>                # approved_orders.json + .csv (valid only)

In PowerShell quote SKUs that contain & or other special characters:  revoke 'C&P9M' --results ...

Runtime files, none of which belong in git: decision_state.json, approvals_log.jsonl (append-only audit log),
approved_orders.json, approved_orders.csv.
"""
import argparse
import csv
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl

from email_agent import excluded_reason, listing_key, load_exclusions   # same key and exclusion rules as the email agent
from email_agent import statefile   # shared atomic write + lock (re-exported like obs: test_no_network_llm_or_order_code pins this file's imports)
from email_agent import obs   # the shared tracing module (re-exported: test_no_network_llm_or_order_code pins this file's imports)

HERE = os.path.dirname(os.path.abspath(__file__))
STATE_PATH = os.path.join(HERE, "decision_state.json")
LOG_PATH = os.path.join(HERE, "approvals_log.jsonl")
EXPORT_JSON = os.path.join(HERE, "approved_orders.json")
EXPORT_CSV = os.path.join(HERE, "approved_orders.csv")
REPORTS_DIR = os.path.join(HERE, "Reports")
EMAIL_STATE_PATH = os.path.join(HERE, "email_state.json")                  # read only
REVIEWER_EXCLUSIONS_CSV = os.path.join(HERE, "reviewer_exclusions.csv")    # read only

LARGE_TOTAL = Decimal("100")        # a line total at the MOQ above this gets no suggested quantity
APPROVAL_TTL_DAYS = 7               # an approval stops being valid this long after it was recorded
NOT_AN_ORDER = "NOT AN ORDER. Nothing has been purchased."
CENT = Decimal("0.01")

# Columns the search agent's Results sheet must have; read by name, never by position.
REQUIRED_HEADERS = ["SKU", "Product", "Recommendation", "Recommended Manufacturer", "Recommended URL",
                    "Recommended Unit Price", "Ordering Note", "L-Com Unit Price", "Manufacturer 1 Name",
                    "Manufacturer 1 Accuracy", "Manufacturer 1 Unit Price Confidence", "Manufacturer 1 MOQ",
                    "Manufacturer 1 URL", "Manufacturer 1 Match Tier"]


class AgentError(Exception):
    """A problem to show the user as a plain message, not a traceback."""


# Reply-rejection reasons as short codes for tracing (the error text itself quotes the reply, so it is never exported).
REPLY_REASON_CODES = (("approve_needs_qty_and_max", r"approve line needs"), ("unparseable", r"can't be parsed"),
                      ("bad_qty", r"qty '"), ("bad_max_price", r"max unit price '"), ("bad_cap", r"cap '"),
                      ("duplicate_sku", r"appears twice"), ("cap_missing", r"cap is required"),
                      ("cap_repeated", r"cap appears more than once"), ("sku_not_in_request", r"is not in the request"),
                      ("over_cap", r"exceeds the cap"))


class ReplyError(AgentError):
    """The pasted reply was rejected; `errors` says which rule failed. Nothing is stored."""

    def __init__(self, errors):
        self.errors = list(errors)
        super().__init__("Reply rejected - nothing stored:\n" + "\n".join(f"  - {e}" for e in self.errors))


# ---------- money and time ----------

def money(value) -> Decimal:
    """Rounded half-up to cents."""
    return Decimal(str(value)).quantize(CENT, ROUND_HALF_UP)


def fmt_money(d: Decimal) -> str:
    return f"${d:,.2f}"


def fmt_unit(d: Decimal) -> str:
    """Unit prices keep up to 4 decimals ($0.0545) but never fewer than 2 ($1.25)."""
    text = f"{d.quantize(Decimal('0.0001'), ROUND_HALF_UP):f}".rstrip("0")
    whole, _, frac = text.partition(".")
    return f"${int(whole):,}.{frac.ljust(2, '0')}"


def now() -> datetime:
    return datetime.now(timezone.utc).replace(microsecond=0)


def iso(dt: datetime) -> str:
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def parse_iso(text: str) -> datetime:
    return datetime.strptime(text, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)


def to_decimal(value):
    """A sheet cell or typed amount -> finite Decimal, or None."""
    if value is None or value == "":
        return None
    try:
        d = Decimal(str(value).replace(",", "").lstrip("$").strip())
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


# ---------- reading the results file (by header name) ----------

def file_sha256(path: str) -> str:
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def read_results(path: str) -> list:
    """Rows of the Results sheet as {header: value}. Stops with a clear message if an expected header is missing."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise AgentError(f"Results file not found: {path}")
    if "Results" not in wb.sheetnames:
        wb.close()
        raise AgentError(f"{os.path.basename(path)} has no 'Results' sheet - is this a sourcing_results_*.xlsx?")
    rows = list(wb["Results"].iter_rows(values_only=True))
    wb.close()  # Windows keeps the file locked otherwise
    header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else [])]
    missing = [h for h in REQUIRED_HEADERS if h not in header]
    if missing:
        raise AgentError(f"{os.path.basename(path)}: the Results sheet is missing expected column(s): "
                         f"{', '.join(missing)}. The search agent's output format may have changed. Nothing was read.")
    return [dict(zip(header, r)) for r in rows[1:] if r and r[header.index("SKU")]]


def parse_moq(text):
    """First whole number in an MOQ string ("500 pieces", "10,000+", "Min. order: 1,000"), else None."""
    m = re.search(r"\d[\d,]*", str(text or ""))
    n = int(m[0].replace(",", "")) if m else 0
    return n or None


def load_email_state(path: str = None) -> tuple:
    """(state or None, note). Read only. A missing or unrecognised file is noted and ignored."""
    path = path or EMAIL_STATE_PATH
    if not os.path.exists(path):
        return None, "email_state.json not found - seller reply status is unavailable."
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (ValueError, OSError):
        return None, "email_state.json format not recognised (not valid JSON) - ignored."
    if not isinstance(data, dict) or any(not isinstance(v, dict) for k, v in data.items() if not str(k).startswith("_")):
        return None, "email_state.json format not recognised - ignored."
    return data, ""


def reply_status_text(entry: dict, have_state: bool) -> str:
    if not have_state:
        return "unknown (no email state)"
    if entry.get("status"):
        summary = (entry.get("summary") or "").strip()
        return f"{entry['status']}" + (f" - {summary[:160]}" if summary else "")
    return "contacted, no reply recorded" if entry.get("sent") else "not contacted (per email state)"


def build_lines(rows: list, exclusions: list, email_state) -> tuple:
    """(lines, skipped). A line is a row with a recommended seller and a usable unit price, with the seller's
    reviewer-exclusion and email-reply status applied."""
    lines, skipped = [], []
    for r in rows:
        sku = str(r["SKU"]).strip()
        base = {"sku": sku, "product": r.get("Product") or r.get("Keyword") or ""}   # built-in products have no Product name
        maker, url = (r.get("Recommended Manufacturer") or "").strip(), (r.get("Recommended URL") or "").strip()
        if str(r.get("Recommendation") or "").startswith("Error researching"):
            skipped.append({**base, "kind": "error", "why": "the search agent errored on this product"})
            continue
        if not (maker or url):
            skipped.append({**base, "kind": "none", "why": "no recommended seller"})
            continue
        why = excluded_reason(sku, maker, url, exclusions)
        if why:
            skipped.append({**base, "kind": "reviewer", "why": f"excluded by reviewer_exclusions.csv ({why})"})
            continue
        key = listing_key(sku, url, maker)
        entry = (email_state or {}).get(key, {}) if isinstance((email_state or {}).get(key, {}), dict) else {}
        if entry.get("status") == "dead_end":
            skipped.append({**base, "kind": "declined", "why": "seller declined (email state: dead_end)"})
            continue
        unit = to_decimal(r.get("Recommended Unit Price"))
        if unit is None or unit <= 0:
            skipped.append({**base, "kind": "noprice", "why": "no usable unit price"})
            continue
        lines.append(make_line(r, sku, maker, url, key, unit, entry, email_state is not None))
    return lines, skipped


# ---------- the seller's quote (recorded by the email agent; this agent makes no LLM calls) ----------

QUOTE_PRICE_TOLERANCE = Decimal("0.10")   # a quoted unit price more than 10% above the listed price is flagged
MIN_MARGIN_PCT = Decimal("80")            # the search agent's bar: at least this far below the L-Com unit price
QUOTE_CONTENT_KEYS = ("sample_available", "sample_quantity", "sample_unit_price", "sample_total_price", "moq", "bulk_tiers",
                      "branding_possible", "branding_fee", "branding_min_qty", "lead_time_days",
                      "production_lead_time_days", "shipping_cost", "shipping_terms", "needs_from_us", "quote_valid_until")
QUOTE_SNAPSHOT_KEYS = QUOTE_CONTENT_KEYS + ("price_basis", "currency", "evidence", "quote_warnings")
QUOTE_MONEY_KEYS = ("sample_unit_price", "sample_total_price", "branding_fee", "shipping_cost")


def whole(value):
    """A positive whole number from an int or digit string, else None."""
    try:
        n = int(str(value))
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


def entry_quote(entry):
    """The structured quote the email agent recorded for this seller, or None when the entry carries none."""
    if not isinstance(entry, dict):
        return None
    quote = {k: entry.get(k) for k in QUOTE_SNAPSHOT_KEYS}
    quote["evidence"] = quote["evidence"] if isinstance(quote["evidence"], dict) else {}
    quote["quote_warnings"] = list(quote["quote_warnings"] or [])
    return quote if any(quote[k] not in (None, "", [], "unknown") for k in QUOTE_CONTENT_KEYS) else None


def quote_has_prices(quote: dict) -> bool:
    return any(quote.get(k) not in (None, "") for k in QUOTE_MONEY_KEYS) or bool(quote.get("bulk_tiers"))


def bulk_tiers(quote: dict) -> list:
    """[(min_qty, Decimal unit price)] from the quote, lowest quantity first; malformed tiers are skipped."""
    tiers = []
    for t in (quote.get("bulk_tiers") or []):
        qty, price = whole(t.get("min_qty")) if isinstance(t, dict) else None, to_decimal(t.get("unit_price")) if isinstance(t, dict) else None
        if qty and price is not None:
            tiers.append((qty, price))
    return sorted(tiers)


def margin_check(lcom, listing_unit: Decimal, tiers: list) -> dict:
    """The 80% bar on the quoted BULK price at the quantity the quote covers (its lowest tier: the highest bulk price,
    so the conservative one). With no bulk price it uses the listing price and says it is not confirmed. A sample price
    is never used: samples cost more than production."""
    if lcom is None or lcom <= 0:
        return {"basis": "n/a (no L-Com price)", "price": None, "pct": None, "clears": None, "quoted": False}
    if tiers:
        qty, price = tiers[0]
        basis, quoted = f"quoted bulk price {fmt_unit(price)} at >= {qty:,} pcs", True
    else:
        price, basis, quoted = listing_unit, "listing price, not confirmed by the seller", False
    pct = ((1 - price / lcom) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)
    return {"basis": basis, "price": str(price), "pct": str(pct), "quoted": quoted,
            "clears": price * 100 <= lcom * (100 - MIN_MARGIN_PCT)}


def make_line(r: dict, sku: str, maker: str, url: str, key: str, unit: Decimal, entry: dict, have_state: bool) -> dict:
    block = {}
    for n in range(1, 50):  # find the recommended candidate's block by URL, else by name
        if f"Manufacturer {n} URL" not in r:
            break
        if (url and (r[f"Manufacturer {n} URL"] or "").strip() == url) or (
                not url and (r.get(f"Manufacturer {n} Name") or "").strip() == maker):
            block = {k: r.get(f"Manufacturer {n} {k}") for k in ("Accuracy", "Match Tier", "Unit Price Confidence", "MOQ")}
            break
    lcom = to_decimal(r.get("L-Com Unit Price"))
    pct = ((1 - unit / lcom) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP) if lcom and lcom > 0 else None
    moq_text = str(block.get("MOQ") or "").strip()
    stated = entry.get("sample_quantity")
    stated = int(stated) if str(stated).isdigit() and int(stated) > 0 else None
    qty = stated or parse_moq(moq_text)
    source = "seller's stated sample quantity" if stated else ("listing MOQ" if qty else "")
    moq_qty = parse_moq(moq_text)
    moq_total = money(unit * moq_qty) if moq_qty else None
    # A line whose total at the MOQ is over the limit gets no suggested quantity: the MOQ is the seller's minimum for
    # bulk pricing, not a sensible sample order. It needs a seller quote for a smaller quantity first.
    heavy = source == "listing MOQ" and moq_total is not None and moq_total > LARGE_TOTAL

    quote = entry_quote(entry)
    currency = ((quote.get("currency") or "unknown") if quote else None)
    non_usd = bool(quote and quote_has_prices(quote) and currency != "USD")      # excluded from all arithmetic
    waiting = list(quote["needs_from_us"]) if quote and quote.get("needs_from_us") else []
    q_qty = whole(quote.get("sample_quantity")) if quote else None
    q_unit = to_decimal(quote.get("sample_unit_price")) if quote else None
    basis_unclear = bool(quote and q_unit is not None and quote.get("price_basis") != "per_unit")
    price, no_total_reason = unit, ""
    tiers = bulk_tiers(quote) if quote and not non_usd else []
    margin = margin_check(lcom, unit, tiers)
    comparisons = []
    if quote and not non_usd:
        quoted_prices = ([("sample unit price", q_unit)] if q_unit is not None else []) + [
            (f"bulk price at >= {n:,} pcs", p) for n, p in tiers]
        for label, p in quoted_prices:
            if p > unit * (1 + QUOTE_PRICE_TOLERANCE):
                comparisons.append(f"quoted {label} {fmt_unit(p)} vs listed {fmt_unit(unit)} "
                                   f"(+{((p / unit - 1) * 100).quantize(Decimal('0.1'), ROUND_HALF_UP)}%)")
    # A quoted price well above the listing with no bulk price quoted: the listing-based margin can't be trusted.
    margin["listing_may_not_hold"] = bool(comparisons and not margin["quoted"] and margin["pct"] is not None)
    margin_fail = bool(margin["quoted"] and margin["clears"] is False)   # the quoted bulk price misses the 80% bar
    if waiting:
        qty, source, heavy, no_total_reason = None, "", False, "waiting on us"
    elif non_usd:
        qty, source, heavy, no_total_reason = None, "", False, f"quoted in {currency}; needs conversion"
    elif margin_fail:
        qty, source, heavy, no_total_reason = None, "", False, "quote does not clear the margin bar"
    elif quote and q_qty and q_unit is not None and not basis_unclear:
        qty, price, source, heavy = q_qty, q_unit, "seller's quoted sample quantity and unit price", False
    elif heavy:
        qty, source = None, ""
    total = money(price * qty) if qty else None
    if total is None and not heavy and not no_total_reason:
        no_total_reason = "MOQ not stated"

    confidence = str(block.get("Unit Price Confidence") or "not found in candidate columns")
    tier = str(block.get("Match Tier") or "")
    price_source = str(r.get("L-Com Price Source") or "")
    flags = []
    if heavy:
        flags.append("needs_seller_quote_for_smaller_qty")
    elif total is not None and total > LARGE_TOTAL:
        flags.append("large_total")                  # the quoted or stated quantity is large
    if not qty and not heavy and not no_total_reason.startswith(("waiting", "quoted in", "quote does not")):
        flags.append("moq_not_stated")
    if waiting:
        flags.append("waiting_on_us")
    if non_usd:
        flags.append("quote_not_usd")
    if comparisons:
        flags.append("quote_above_listed_price")
    if margin["quoted"] and margin["clears"] is False:
        flags.append("quote_no_longer_clears_margin_bar")
    if basis_unclear and not non_usd:
        flags.append("quote_price_basis_unclear")
    if quote and quote["quote_warnings"]:
        flags.append("quote_has_warnings")
    if not confidence.startswith("stated"):
        flags.append("unit_price_" + re.split(r"[\s(]", confidence)[0])
    if price_source and re.search(r"unverified|unconfirmed", price_source, re.I):
        flags.append("lcom_price_unverified")
    if tier and tier != "Auto-accepted":
        flags.append("manual_review_tier")
    reply = reply_status_text(entry, have_state)
    if reply.startswith("needs_info"):
        flags.append("seller_needs_info")
    if have_state and not entry.get("status"):
        flags.append("no_seller_reply")
    return {
        "sku": sku, "product": r.get("Product") or r.get("Keyword") or "", "seller": maker or "(unnamed)", "url": url, "key": key,
        "accuracy": block.get("Accuracy"), "tier": tier, "unit_price": str(unit), "confidence": confidence,
        "pct_below_lcom": str(pct) if pct is not None else None, "lcom_price": str(lcom) if lcom is not None else None,
        "lcom_price_source": price_source, "moq": moq_text, "qty": qty, "qty_source": source,
        "moq_qty": moq_qty, "moq_total": str(moq_total) if moq_total is not None else None, "moq_heavy": heavy,
        "suggested_unit_price": str(price), "total": str(total) if total is not None else None,
        "no_total_reason": no_total_reason, "note": str(r.get("Ordering Note") or "").strip(),
        "seller_reply": reply, "flags": flags,
        "quote": quote, "quote_currency": currency, "margin": margin, "quote_comparisons": comparisons,
        "waiting_on_us": waiting, "margin_fail": margin_fail,
    }


def skip_summary(skipped: list) -> str:
    if not skipped:
        return "Skipped 0 row(s)."
    labels = {"none": "no recommended seller", "error": "search agent error", "reviewer": "excluded by reviewer",
              "declined": "seller declined", "noprice": "no usable unit price"}
    parts = []
    for kind, label in labels.items():
        skus = [s["sku"] for s in skipped if s["kind"] == kind]
        if skus:
            parts.append(f"{len(skus)} {label} ({', '.join(skus)})")
    return f"Skipped {len(skipped)} row(s): " + "; ".join(parts) + "."


# ---------- state, log ----------

def empty_state() -> dict:
    return {"version": 1, "next_request": 1, "requests": {}, "decisions": []}


def load_state() -> dict:
    if not os.path.exists(STATE_PATH):
        return empty_state()
    try:
        with open(STATE_PATH, encoding="utf-8") as f:
            state = json.load(f)
        assert isinstance(state, dict) and {"next_request", "requests", "decisions"} <= set(state)
        return state
    except (ValueError, AssertionError):
        raise AgentError(f"{STATE_PATH} is not a decision state file - not touching it.")


def save_state(state: dict) -> None:
    statefile.atomic_write(STATE_PATH, json.dumps(state, indent=2, ensure_ascii=False))


def log_event(event: str, **fields) -> None:
    """Append-only audit log: one JSON object per line, never rewritten."""
    with open(LOG_PATH, "a", encoding="utf-8") as f:
        f.write(json.dumps({"at": iso(now()), "event": event, **fields}, ensure_ascii=False) + "\n")


# ---------- the request ----------

def line_section(l: dict) -> str:
    """Which part of the request a line is listed in. A seller waiting on us comes first, then a quote that misses the
    margin bar, then an MOQ too large to suggest."""
    if l["waiting_on_us"]:
        return "waiting"
    if l["margin_fail"]:
        return "margin"
    return "heavy" if l["moq_heavy"] else "normal"


def build_request(results_path: str, budget=None, state: dict = None) -> dict:
    """Everything an approval request says, as data (nothing is written)."""
    rows = read_results(results_path)
    email_state, email_note = load_email_state()
    lines, skipped = build_lines(rows, load_exclusions(REVIEWER_EXCLUSIONS_CSV), email_state)
    subtotal = sum((Decimal(l["total"]) for l in lines if l["total"]), Decimal("0.00"))   # priced lines only
    heavy = [l for l in lines if line_section(l) == "heavy"]
    req = {
        "results_file": os.path.basename(results_path), "results_path": os.path.abspath(results_path),
        "results_sha256": file_sha256(results_path), "budget": str(money(budget)) if budget is not None else None,
        "lines": lines, "skipped": skipped, "subtotal": str(subtotal), "email_note": email_note,
        "lines_without_total": [l["sku"] for l in lines if not l["total"] and line_section(l) == "normal"],
        "no_total_reasons": {l["sku"]: l["no_total_reason"] for l in lines
                             if not l["total"] and line_section(l) == "normal"},
        "waiting": [l["sku"] for l in lines if line_section(l) == "waiting"],
        "margin_failed": [l["sku"] for l in lines if line_section(l) == "margin"],
        "left_out": [l["sku"] for l in heavy],
        "left_out_moq_total": str(sum((Decimal(l["moq_total"]) for l in heavy), Decimal("0.00"))),
    }
    req["orphans"] = orphaned(state, {l["key"] for l in lines}) if state else []
    return req


def qmoney(value, currency) -> str:
    """A quoted amount: dollars for USD, otherwise the number and its currency (never a $ sign for another currency)."""
    d = to_decimal(value)
    if d is None:
        return str(value)
    return fmt_unit(d) if currency == "USD" else f"{value} {currency}"


def quote_block(l: dict) -> list:
    q = l["quote"]
    if not q:
        return ["   Seller quote: no seller quote yet - listing price only"]
    cur = l["quote_currency"] or "n/a"
    out = ["   Seller quote:"]
    sample = []
    if q.get("sample_available") in ("yes", "no"):
        sample.append(f"samples available: {q['sample_available']}")
    if q.get("sample_quantity"):
        sample.append(f"qty {int(q['sample_quantity']):,}")
    if q.get("sample_unit_price") is not None:
        sample.append(f"{qmoney(q['sample_unit_price'], cur)} per unit")
    if q.get("sample_total_price") is not None:
        sample.append(f"{qmoney(q['sample_total_price'], cur)} total")
    if sample:
        out.append("     Sample: " + ", ".join(sample) + (" (price basis unclear)" if q.get("price_basis") == "unclear" else ""))
    tiers = bulk_tiers(q)
    if tiers:
        out.append("     Bulk: " + "; ".join(f">= {n:,} pcs {qmoney(str(p), cur)}" for n, p in tiers))
    more = []
    if q.get("moq"):
        more.append(f"MOQ {int(q['moq']):,}")
    if q.get("branding_possible") in ("yes", "no"):
        brand = f"branding {q['branding_possible']}"
        extra = ([f"fee {qmoney(q['branding_fee'], cur)}"] if q.get("branding_fee") is not None else []) + (
            [f"min qty {int(q['branding_min_qty']):,}"] if q.get("branding_min_qty") else [])
        more.append(brand + (f" ({', '.join(extra)})" if extra else ""))
    lead = ([f"sample {q['lead_time_days']} days"] if q.get("lead_time_days") else []) + (
        [f"production {q['production_lead_time_days']} days"] if q.get("production_lead_time_days") else [])
    if lead:
        more.append("lead time " + ", ".join(lead))
    if more:
        out.append("     " + " | ".join(more))
    ship = ([qmoney(q["shipping_cost"], cur)] if q.get("shipping_cost") is not None else []) + (
        [q["shipping_terms"]] if q.get("shipping_terms") else [])
    tail = ([f"shipping {' '.join(ship)}"] if ship else []) + ([f"valid until {q['quote_valid_until']}"] if q.get("quote_valid_until") else []) + (
        [f"currency {cur}"] if quote_has_prices(q) else [])
    if tail:
        out.append("     " + " | ".join(tail))
    if q["quote_warnings"]:
        out.append("     Quote warnings (fields the email agent could not verify were left blank): " + "; ".join(q["quote_warnings"]))
    if len(out) == 1:
        out.append("     no prices or quantities quoted yet")
    return out


def line_text(i: int, l: dict) -> list:
    unit, qty = Decimal(l["unit_price"]), l["qty"]
    total = Decimal(l["total"]) if l["total"] else None
    if l["waiting_on_us"]:
        qty_part = "Suggested qty: none - waiting on us (not in the subtotal)"
    elif l["margin_fail"]:
        qty_part = "Suggested qty: none - the quote does not clear the margin bar (not in the subtotal)"
    elif l["moq_heavy"]:
        qty_part = (f"MOQ total: {fmt_money(Decimal(l['moq_total']))} | Suggested qty: none - needs a seller quote for a "
                    "smaller quantity first")
    elif total is None:
        qty_part = (f"Suggested qty: none - {l['no_total_reason']} (not in the subtotal)" if l["no_total_reason"].startswith("quoted in")
                    else "Suggested qty: not stated (MOQ missing) - Neeraj to give a qty")
    else:
        at = f" at {fmt_unit(Decimal(l['suggested_unit_price']))}" if l["qty_source"].startswith("seller's quoted") else ""
        qty_part = f"Suggested qty: {qty:,} ({l['qty_source']}){at} | Line total: " + (
            f"LARGE total {fmt_money(total)}  <-- LARGE" if "large_total" in l["flags"] else fmt_money(total))
    pct = f"{l['pct_below_lcom']}% below L-Com ({fmt_money(Decimal(l['lcom_price']))}" + (
        f", {l['lcom_price_source']}" if l["lcom_price_source"] else "") + ")" if l["pct_below_lcom"] else "L-Com comparison n/a"
    m = l["margin"]
    if m["pct"] is None:
        margin = f"   Margin check: {m['basis']}"
    elif m.get("listing_may_not_hold"):                       # the quote is well above the listing and no bulk price came with it
        margin = "   Margin check: bulk price not quoted; the listing price may not hold"
    else:
        margin = (f"   Margin check: {m['pct']}% below L-Com on the {m['basis']} - "
                  + (f"clears the {format(MIN_MARGIN_PCT.normalize(), 'f')}% bar" if m["clears"] else f"DOES NOT clear the {format(MIN_MARGIN_PCT.normalize(), 'f')}% bar"))
    out = [f"{i}) {l['sku']} - {l['product']}",
           f"   Seller: {l['seller']}",
           f"   Listing: {l['url'] or '(no URL)'}",
           f"   Match: {l['accuracy']}% ({l['tier']}) | Unit price: {fmt_unit(unit)} ({l['confidence']}) | {pct}",
           margin,
           f"   MOQ: {l['moq'] or 'not stated'} | {qty_part}"]
    if l["note"]:
        out.append(f"   Ordering note: {l['note']}")
    out.append(f"   Seller reply: {l['seller_reply']}")
    out += quote_block(l)
    if l["waiting_on_us"]:
        out.append(f"   Seller asked us for: {'; '.join(l['waiting_on_us'])}")
    if l["quote_comparisons"]:
        out += [f"   QUOTE ABOVE LISTED PRICE (more than {format((QUOTE_PRICE_TOLERANCE * 100).normalize(), 'f')}%): {c}" for c in l["quote_comparisons"]]
    if "quote_no_longer_clears_margin_bar" in l["flags"]:
        out.append("   QUOTE NO LONGER CLEARS THE MARGIN BAR")
    if "quote_not_usd" in l["flags"]:
        out.append(f"   Quoted in {l['quote_currency']}; needs conversion - the quote is left out of every total and margin check")
    if l["flags"]:
        out.append(f"   Flags: {', '.join(l['flags'])}")
    return out


REPLY_FORMAT = [
    "HOW TO REPLY (one instruction per line, any capitalisation; $ and commas are fine):",
    "  approve <SKU> qty <whole number> max <max unit price>",
    "  reject <SKU>",
    "  cap <the most you will spend in total, including shipping>",
    "Rules: 'cap' is required. Every approve needs its own qty and max unit price - nothing is assumed.",
    "SKUs must be ones listed above, each at most once. The sum of qty x max must not exceed the cap.",
    "A line in the 'needs a seller quote' or 'does not clear the margin bar' section can still be approved: give an explicit qty and max like any other.",
    f"Approvals expire {APPROVAL_TTL_DAYS} days after they are recorded. Replying orders nothing.",
    "EXAMPLE (made-up SKUs, not an approval):",
    "  approve SKU-A qty 100 max 1.30",
    "  approve SKU-B qty 50 max 0.30",
    "  reject SKU-C",
    "  cap 200",
]


def render_text(req: dict) -> str:
    subtotal = Decimal(req["subtotal"])
    waiting = [l for l in req["lines"] if line_section(l) == "waiting"]
    failing = [l for l in req["lines"] if line_section(l) == "margin"]
    heavy = [l for l in req["lines"] if line_section(l) == "heavy"]
    normal = [l for l in req["lines"] if line_section(l) == "normal"]
    reasons = req["no_total_reasons"]
    out = [f"APPROVAL REQUEST {req['id']}   (created {req['created_at']})",
           f"Results file: {req['results_file']}  sha256 {req['results_sha256'][:16]}...",
           f"Products: {len(req['lines'])} | Product subtotal at suggested quantities: {fmt_money(subtotal)} "
           f"(covers {sum(1 for l in req['lines'] if l['total'])} of {len(req['lines'])} lines)"
           + (f"; no qty for {', '.join(f'{s} ({reasons[s]})' for s in req['lines_without_total'])}" if req["lines_without_total"] else "")]
    if req["left_out"]:
        out.append(f"Left out of that subtotal: {len(req['left_out'])} line(s) whose total at the MOQ is over "
                   f"{fmt_money(LARGE_TOTAL)} ({', '.join(req['left_out'])}) - combined MOQ total "
                   f"{fmt_money(Decimal(req['left_out_moq_total']))}. They need a seller quote for a smaller quantity first.")
    if req["margin_failed"]:
        out.append(f"Left out of that subtotal: {len(req['margin_failed'])} line(s) ({', '.join(req['margin_failed'])}) whose "
                   f"quoted bulk price does not clear the {format(MIN_MARGIN_PCT.normalize(), 'f')}% margin bar - not suggested.")
    if req["waiting"]:
        out.append(f"Waiting on us: {len(req['waiting'])} line(s) ({', '.join(req['waiting'])}) - the seller asked us for "
                   "something first; not in the subtotal.")
    out.append("NOT included: shipping, duties and taxes.")
    if req["budget"]:
        over = subtotal - Decimal(req["budget"])
        out.append(f"Budget: {fmt_money(Decimal(req['budget']))} - subtotal (priced lines only) is "
                   + (f"{fmt_money(over)} OVER budget" if over > 0 else f"within budget ({fmt_money(-over)} to spare)"))
    out.append(skip_summary(req["skipped"]))
    if req["email_note"]:
        out.append("Note: " + req["email_note"])
    if req["orphans"]:
        out.append(f"Note: {len(req['orphans'])} earlier approval(s) no longer match the recommended seller and are "
                   f"orphaned (not carried over): {', '.join(req['orphans'])}.")
    out.append("")
    n = 0
    for l in normal:
        n += 1
        out += line_text(n, l) + [""]
    if failing:
        out += [f"QUOTE DOES NOT CLEAR THE MARGIN BAR - NOT SUGGESTED (the quoted bulk price is not at least "
                f"{format(MIN_MARGIN_PCT.normalize(), 'f')}% below L-Com; no quantity is suggested and these are not in the subtotal):", ""]
        for l in failing:
            n += 1
            out += line_text(n, l) + [""]
    if heavy:
        out += [f"NEEDS A SELLER QUOTE FOR A SMALLER QUANTITY FIRST (total at the MOQ is over {fmt_money(LARGE_TOTAL)}; "
                "no quantity is suggested and these are not in the subtotal):", ""]
        for l in heavy:
            n += 1
            out += line_text(n, l) + [""]
    if waiting:
        out += ["WAITING ON US (the seller asked us for something before it can quote; these are not in the subtotal):", ""]
        for l in waiting:
            n += 1
            out += line_text(n, l) + [""]
    out += REPLY_FORMAT
    return "\n".join(out)


def render_markdown(req: dict) -> str:
    text = render_text(req)
    head, _, rest = text.partition("\n\n")
    body, _, fmt = rest.rpartition("HOW TO REPLY")
    return (f"# Approval request {req['id']}\n\n" + "\n".join(f"- {l}" for l in head.splitlines()[1:]) + "\n\n## Lines\n\n"
            + "```\n" + body.rstrip() + "\n```\n\n## How to reply\n\n```\nHOW TO REPLY" + fmt + "\n```\n")


def cmd_request(results_path: str, budget=None) -> dict:
    with statefile.locked(STATE_PATH):
        return _request(results_path, budget)


def _request(results_path: str, budget=None) -> dict:
    state = load_state()
    if budget is not None and (budget <= 0):
        raise AgentError("--budget must be a positive amount.")
    req = build_request(results_path, budget, state)
    req["id"] = f"A-{state['next_request']:04d}"
    obs.tag(f"request:{req['id']}", *(f"sku:{l['sku']}" for l in req["lines"]))
    obs.annotate(request_id=req["id"], lines=len(req["lines"]), subtotal=Decimal(req["subtotal"]), budget=req["budget"],
                 skipped=len(req["skipped"]), lines_left_out=len(req["left_out"]), left_out_moq_total=Decimal(req["left_out_moq_total"]),
                 waiting=len(req["waiting"]), margin_failed=len(req["margin_failed"]), without_total=len(req["lines_without_total"]),
                 orphans=len(req["orphans"]))
    req["created_at"] = iso(now())
    req["reply"] = None
    state["next_request"] += 1
    state["requests"][req["id"]] = req
    text = render_text(req)
    os.makedirs(REPORTS_DIR, exist_ok=True)
    for ext, content in (("txt", text), ("md", render_markdown(req))):
        with open(os.path.join(REPORTS_DIR, f"approval_request_{req['id']}.{ext}"), "w", encoding="utf-8") as f:
            f.write(content + "\n")
    save_state(state)
    log_event("request", request_id=req["id"], results_file=req["results_file"], results_sha256=req["results_sha256"],
              lines=[l["sku"] for l in req["lines"]], subtotal=req["subtotal"], budget=req["budget"])
    return req


# ---------- the reply (the one function to change if the format changes) ----------

def parse_reply(text: str) -> dict:
    """Pasted reply -> {"cap": Decimal, "approvals": [{"sku", "qty", "max"}], "rejections": [sku]}.
    One instruction per line, case-insensitive, `$` and commas allowed:
        approve <SKU> qty <whole number> max <unit price>  |  reject <SKU>  |  cap <amount>
    Raises ReplyError listing every rule that failed. Knows nothing about the request; see check_reply()."""
    errors, approvals, rejections, caps, seen = [], [], [], [], {}
    for n, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue
        tokens = line.split()
        verb = tokens[0].lower()
        if verb == "approve" and len(tokens) == 6 and tokens[2].lower() == "qty" and tokens[4].lower() == "max":
            sku, qty_text, max_text = tokens[1], tokens[3], tokens[5]
            qty_digits = qty_text.replace(",", "")
            qty = int(qty_digits) if re.fullmatch(r"\d+", qty_digits) else 0
            price = to_decimal(max_text)
            if qty <= 0:
                errors.append(f"line {n}: qty '{qty_text}' for {sku} isn't a positive whole number")
            if price is None or price <= 0:
                errors.append(f"line {n}: max unit price '{max_text}' for {sku} isn't a positive amount")
            approvals.append({"sku": sku, "qty": qty, "max": price})
        elif verb == "approve" and len(tokens) >= 2 and not (len(tokens) == 6 and tokens[2].lower() == "qty"):
            errors.append(f"line {n}: can't be parsed - an approve line needs an explicit qty and max unit price "
                          f"('approve <SKU> qty <n> max <price>'): {line!r}")
            sku = tokens[1]
            approvals.append({"sku": sku, "qty": 0, "max": None})
        elif verb == "reject" and len(tokens) == 2:
            sku = tokens[1]
            rejections.append(sku)
        elif verb == "cap" and len(tokens) == 2:
            amount = to_decimal(tokens[1])
            if amount is None or amount <= 0:
                errors.append(f"line {n}: cap '{tokens[1]}' isn't a positive amount")
            caps.append(amount)
            continue
        else:
            errors.append(f"line {n}: can't be parsed: {line!r}")
            continue
        seen.setdefault(sku.upper(), []).append(n)
    errors += [f"SKU {s} appears twice (lines {', '.join(map(str, ns))})" for s, ns in seen.items() if len(ns) > 1]
    if not caps:
        errors.append("cap is required (e.g. 'cap 200') - the most to spend in total, including shipping")
    elif len(caps) > 1:
        errors.append("cap appears more than once")
    if errors:
        raise ReplyError(errors)
    return {"cap": money(caps[0]), "approvals": approvals, "rejections": rejections}


def check_reply(parsed: dict, lines: list) -> dict:
    """Checks a parsed reply against the request's lines and prices it. Raises ReplyError (nothing stored)."""
    by_sku = {l["sku"].upper(): l for l in lines}
    errors = [f"SKU {item if isinstance(item, str) else item['sku']} is not in the request"
              for item in parsed["rejections"] + parsed["approvals"]
              if (item if isinstance(item, str) else item["sku"]).upper() not in by_sku]
    approvals = []
    for a in parsed["approvals"]:
        line = by_sku.get(a["sku"].upper())
        if line is None:
            continue
        line_cap = money(a["qty"] * a["max"])
        listed = Decimal(line["unit_price"])
        warn = (f"max {fmt_unit(a['max'])} is below the listed unit price {fmt_unit(listed)} - this line cannot be "
                f"ordered at the listed price") if a["max"] < listed else ""
        approvals.append({"line": line, "qty": a["qty"], "max": a["max"], "line_cap": line_cap, "warning": warn})
    committed = sum((a["line_cap"] for a in approvals), Decimal("0.00"))
    if not errors and committed > parsed["cap"]:
        errors.append(f"the sum of qty x max ({fmt_money(committed)}) exceeds the cap ({fmt_money(parsed['cap'])})")
    if errors:
        raise ReplyError(errors)
    rejections = [by_sku[s.upper()] for s in parsed["rejections"]]
    decided = {l["sku"] for l in rejections} | {a["line"]["sku"] for a in approvals}
    return {"cap": parsed["cap"], "approvals": approvals, "rejections": rejections, "committed": committed,
            "undecided": [l for l in lines if l["sku"] not in decided]}


def render_echo(req: dict, plan: dict, defaulted: bool = False) -> str:
    out = [f"RECORDING AGAINST REQUEST {req['id']}",
           f"  Created:      {req['created_at']}",
           f"  Results file: {req['results_file']}",
           f"  File hash:    sha256 {req['results_sha256']}",
           "  (picked by default: it is the only open request for this results file)" if defaulted
           else "  (picked with --request)",
           "", f"UNDERSTOOD for request {req['id']}:"]
    for a in plan["approvals"]:
        l = a["line"]
        out.append(f"  APPROVE {l['sku']} ({l['seller']}): qty {a['qty']:,} x max {fmt_unit(a['max'])} = "
                   f"{fmt_money(a['line_cap'])} at most  [listed {fmt_unit(Decimal(l['unit_price']))}]")
        if a["warning"]:
            out.append(f"     WARNING: {a['warning']}.")
    for l in plan["rejections"]:
        out.append(f"  REJECT  {l['sku']} ({l['seller']})")
    out.append(f"  Cap: {fmt_money(plan['cap'])} | Committed at max prices: {fmt_money(plan['committed'])} | "
               f"Left for shipping, duties and taxes: {fmt_money(plan['cap'] - plan['committed'])}")
    if plan["undecided"]:
        out.append(f"  No decision given for: {', '.join(l['sku'] for l in plan['undecided'])} (they stay 'requested')")
    out.append("NOTE: the sender cannot be verified. This records what a person pasted and who they say approved it.")
    out.append("Nothing is ordered by recording this.")
    return "\n".join(out)


def pick_request(state: dict, results_file: str, request_id=None) -> tuple:
    """(request, picked_by_default). A reply is bound to the request it answered: it must be named with --request,
    and only when exactly one request for this results file is still open may that be left out."""
    if request_id:
        m = re.fullmatch(r"A-(\d+)", request_id.strip(), re.IGNORECASE)
        rid = f"A-{int(m[1]):04d}" if m else request_id.strip()
        req = state["requests"].get(rid)
        if req is None:
            raise AgentError(f"No request {rid}. Requests so far: {', '.join(sorted(state['requests'])) or 'none'}.")
        if req["results_file"] != results_file:
            raise AgentError(f"Request {rid} was written for {req['results_file']}, not {results_file}.")
        return req, False
    open_ = sorted((r for r in state["requests"].values() if r["results_file"] == results_file and not r["reply"]),
                   key=lambda r: r["id"])
    if len(open_) == 1:
        return open_[0], True
    if not open_:
        raise AgentError(f"No open request for {results_file} (none was written, or each already has a recorded "
                         "reply). Run 'request' first.")
    raise AgentError(f"{len(open_)} requests are open for {results_file}: "
                     + "; ".join(f"{r['id']} created {r['created_at']}" for r in open_)
                     + ". Say which one Neeraj answered with --request A-000N.")


NO_TERMINAL = ("Can't ask for the request id, channel, approver and confirmation: there is no terminal to type on (stdin is "
               "closed, piped or used up). Nothing was recorded. Run it in a terminal with the reply in a file (--file), or give "
               "--channel, --approver and --confirm-request together.")


def read_console(prompt: str) -> str:
    """input(), only on a real terminal. Never opens the console device: a program started by something else must not
    read the keyboard of whatever terminal launched it."""
    stdin = sys.stdin
    if stdin is None or stdin.closed or not stdin.isatty():
        raise AgentError(NO_TERMINAL)
    try:
        return input(prompt)
    except EOFError:
        raise AgentError(NO_TERMINAL)


ask_user = read_console   # tests replace this


def cmd_record(results_path: str, reply_text: str, request_id=None, gate=None) -> tuple:
    """Parse, echo back, ask which request this answers, the channel, the approver and a final yes, and only then
    store. Returns (plan, stored?). `gate` = (confirm_request, channel, approver) answers the same four questions without
    a terminal (see check_gate); the prompts' own validation then runs unchanged."""
    with statefile.locked(STATE_PATH):
        return _record(results_path, reply_text, request_id, gate)


def check_gate(request_id, confirm_request, channel, approver):
    """All-or-nothing validation of --confirm-request, --channel and --approver. Returns None when none were given,
    else (confirm_request, channel, approver). Typed by a person on every call: no default is ever taken."""
    given = {"--confirm-request": confirm_request, "--channel": channel, "--approver": approver}
    if all(v is None for v in given.values()):
        return None
    missing = [k for k, v in given.items() if v is None]
    if missing:
        raise AgentError(f"{', '.join(missing)} missing: give --confirm-request, --channel and --approver together, or none "
                         "of them to be asked at the keyboard. Nothing was recorded.")
    if request_id is None:
        raise AgentError("--confirm-request needs --request (the request id being answered). Nothing was recorded.")
    for flag, v in given.items():
        if not v.strip() or any(c in v for c in "\r\n\x00"):
            raise AgentError(f"{flag} must be a non-empty single line. Nothing was recorded.")
    if confirm_request.strip().upper() != request_id.strip().upper():
        raise AgentError(f"--confirm-request {confirm_request.strip()!r} is not the same as --request {request_id.strip()!r}. Nothing was recorded.")
    return confirm_request, channel, approver


def _record(results_path: str, reply_text: str, request_id=None, gate=None) -> tuple:
    ask = ask_user
    if gate is not None:
        answers = iter([gate[0], gate[1], gate[2], "yes"])        # the prompts' questions, in their order
        ask = lambda _prompt: next(answers)                       # noqa: E731
    state = load_state()
    req, defaulted = pick_request(state, os.path.basename(results_path), request_id)
    if file_sha256(results_path) != req["results_sha256"]:
        raise AgentError(f"{os.path.basename(results_path)} has changed since request {req['id']} was written "
                         "(hash mismatch). Run 'request' again so the approval matches what Neeraj sees.")
    if req["reply"]:
        raise AgentError(f"Request {req['id']} already has a recorded reply. Run 'request' again for a fresh request "
                         "before recording another reply.")
    obs.annotate(request_id=req["id"], request_lines=len(req["lines"]))
    obs.tag(f"request:{req['id']}", *(f"sku:{l['sku']}" for l in req["lines"]))
    plan = check_reply(parse_reply(reply_text), req["lines"])
    obs.annotate(approvals=len(plan["approvals"]), rejections=len(plan["rejections"]), undecided=len(plan["undecided"]),
                 cap=plan["cap"], committed=plan["committed"])
    print(render_echo(req, plan, defaulted))
    if ask(f"Is this the request Neeraj answered? Type its id ({req['id']}) to confirm: ").strip().upper() != req["id"]:
        print("Request not confirmed. Nothing recorded.")
        return plan, False
    channel = ask("Channel the reply came through (e.g. Slack DM, email): ").strip()
    approver = ask("Who approved (name as given by the sender): ").strip()
    obs.protect(channel, approver)   # typed names: masked wherever they could appear in a trace
    if not channel or not approver:
        print("Channel and approver are both required. Nothing recorded.")
        return plan, False
    if ask("Type 'yes' to record these approvals (nothing will be ordered): ").strip().lower() != "yes":
        print("Nothing recorded.")
        return plan, False
    at = now()
    expires = at + timedelta(days=APPROVAL_TTL_DAYS)
    decisions, n = [], 0
    for kind, items in (("approve", plan["approvals"]), ("reject", plan["rejections"])):
        for item in items:
            n += 1
            line = item["line"] if kind == "approve" else item
            d = {"id": f"{req['id']}-{n}", "request_id": req["id"], "decision": kind, "sku": line["sku"],
                 "key": line["key"], "seller": line["seller"], "url": line["url"], "product": line["product"],
                 "listed_unit_price": line["unit_price"], "flags": line["flags"], "recorded_at": iso(at),
                 "quote": line.get("quote"), "suggested_unit_price": line.get("suggested_unit_price"),   # what Neeraj saw
                 "approver": approver, "channel": channel, "revoked_at": None}
            if kind == "approve":
                d.update(qty=item["qty"], max_unit_price=str(item["max"]), line_cap=str(item["line_cap"]),
                         expires_at=iso(expires))
            decisions.append(d)
    obs.annotate(recorded=True, decisions=len(decisions))
    req["reply"] = {"text": reply_text, "channel": channel, "approver": approver, "recorded_at": iso(at),
                    "cap": str(plan["cap"]), "sender_verified": False, "decision_ids": [d["id"] for d in decisions]}
    state["decisions"] += decisions
    save_state(state)
    log_event("reply_recorded", request_id=req["id"], channel=channel, approver=approver, cap=str(plan["cap"]),
              sender_verified=False, text=reply_text)
    for d in decisions:
        log_event("approval" if d["decision"] == "approve" else "rejection", **{k: v for k, v in d.items() if k != "flags"})
    return plan, True


# ---------- status, revoke, export ----------

def decision_status(d: dict, at: datetime) -> str:
    if d["decision"] == "reject":
        return "rejected"
    if d["revoked_at"]:
        return "revoked"
    return "expired" if parse_iso(d["expires_at"]) <= at else "approved"


def latest_decision(state: dict, key: str):
    found = [d for d in state["decisions"] if d["key"] == key]
    return found[-1] if found else None


def orphaned(state: dict, current_keys: set) -> list:
    """Approvals (not revoked) whose SKU + listing is no longer a current recommended seller: 'SKU (seller)'."""
    return [f"{d['sku']} ({d['seller']})" for d in (state or {}).get("decisions", [])
            if d["decision"] == "approve" and not d["revoked_at"] and d["key"] not in current_keys]


def current_lines(results_path: str) -> tuple:
    rows = read_results(results_path)
    email_state, _ = load_email_state()
    return build_lines(rows, load_exclusions(REVIEWER_EXCLUSIONS_CSV), email_state)


def cmd_status(results_path: str) -> list:
    state = load_state()
    lines, skipped = current_lines(results_path)
    at = now()
    open_keys = {l["key"] for r in state["requests"].values() for l in r["lines"]}   # asked about, no decision yet
    out = []
    for l in lines:
        d = latest_decision(state, l["key"])
        status = decision_status(d, at) if d else ("requested" if l["key"] in open_keys else "not requested")
        detail = ""
        if d and d["decision"] == "approve":
            detail = (f"qty {d['qty']:,} max {fmt_unit(Decimal(d['max_unit_price']))} cap {fmt_money(Decimal(d['line_cap']))} "
                      f"| {d['id']} | approved by {d['approver']} | expires {d['expires_at']}")
        out.append({"sku": l["sku"], "seller": l["seller"], "status": status, "detail": detail})
    return out, orphaned(state, {l["key"] for l in lines}), skipped


def cmd_revoke(sku: str, results_path: str) -> list:
    with statefile.locked(STATE_PATH):
        return _revoke(sku, results_path)


def _revoke(sku: str, results_path: str) -> list:
    state = load_state()
    lines, _ = current_lines(results_path)
    active = [d for d in state["decisions"] if d["sku"].upper() == sku.strip().upper() and d["decision"] == "approve"
              and not d["revoked_at"]]
    if not active:
        known = ", ".join(sorted({d["sku"] for d in state["decisions"] if d["decision"] == "approve" and not d["revoked_at"]})) or "none"
        raise AgentError(f"{sku}: no active approval to revoke. SKUs with a non-revoked approval: {known}.")
    at = iso(now())
    for d in active:
        d["revoked_at"] = at
    save_state(state)
    obs.annotate(revoked=len(active))
    current = {l["sku"].upper(): l["seller"] for l in lines}
    for d in active:
        log_event("revoke", approval_id=d["id"], sku=d["sku"], seller=d["seller"], key=d["key"],
                  current_recommended_seller=current.get(d["sku"].upper()))
    return active


EXPORT_FIELDS = ["sku", "product", "seller", "listing_url", "qty", "max_unit_price", "line_cap", "approval_id",
                 "approved_by", "approved_at", "expires_at", "source_results_file", "flags", "request_id", "channel",
                 "sender_verified", "quoted_unit_price", "quoted_quantity", "quoted_currency", "quoted_bulk_tiers",
                 "quote_evidence", "quote_warnings"]


def export_cell(order: dict, field: str) -> str:
    value = order[field]
    if field == "flags":
        return ";".join(value)
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else ("" if value is None else value)


class _Buffer:
    """What csv.writer writes into, so the whole file can go through statefile.atomic_write in one piece."""
    def __init__(self):
        self.parts = []

    def write(self, s):
        self.parts.append(s)


def cmd_export(results_path: str, out_json: str = None, out_csv: str = None) -> dict:
    """Writes the approvals that are valid now: approved, not expired, not revoked, and still the recommended seller."""
    with statefile.locked(STATE_PATH):
        return _export(results_path, out_json, out_csv)


def _export(results_path: str, out_json: str = None, out_csv: str = None) -> dict:
    state = load_state()
    lines, _ = current_lines(results_path)
    at = now()
    by_key = {l["key"]: l for l in lines}
    orders, dropped = [], {"expired": 0, "revoked": 0, "rejected": 0}
    for key, line in by_key.items():
        d = latest_decision(state, key)
        if d is None:
            continue
        status = decision_status(d, at)
        if status != "approved":
            dropped[status] += 1
            continue
        flags = list(d["flags"])
        if Decimal(line["unit_price"]) > Decimal(d["max_unit_price"]):
            flags.append("listed_price_above_approved_max")
        orders.append({"sku": d["sku"], "product": d["product"], "seller": d["seller"], "listing_url": d["url"],
                       "qty": d["qty"], "max_unit_price": d["max_unit_price"], "line_cap": d["line_cap"],
                       "approval_id": d["id"], "approved_by": d["approver"], "approved_at": d["recorded_at"],
                       "expires_at": d["expires_at"], "source_results_file": os.path.basename(results_path),
                       "flags": flags, "request_id": d["request_id"], "channel": d["channel"], "sender_verified": False,
                       # the seller's quote as Neeraj saw it in the request (None/{}/[] when there was none)
                       "quoted_unit_price": (d.get("quote") or {}).get("sample_unit_price"),
                       "quoted_quantity": (d.get("quote") or {}).get("sample_quantity"),
                       "quoted_currency": (d.get("quote") or {}).get("currency"),
                       "quoted_bulk_tiers": (d.get("quote") or {}).get("bulk_tiers"),
                       "quote_evidence": (d.get("quote") or {}).get("evidence") or {},
                       "quote_warnings": (d.get("quote") or {}).get("quote_warnings") or []})
    orphans = orphaned(state, set(by_key))
    out_json, out_csv = out_json or EXPORT_JSON, out_csv or EXPORT_CSV
    statefile.atomic_write(out_json, json.dumps({"notice": NOT_AN_ORDER, "generated_at": iso(at), "source_results_file": os.path.basename(results_path),
                                                 "valid_approvals": len(orders), "orders": orders}, indent=2, ensure_ascii=False))
    buf = _Buffer()
    w = csv.writer(buf)
    w.writerow([NOT_AN_ORDER])
    w.writerow(EXPORT_FIELDS)
    for o in orders:
        w.writerow([export_cell(o, k) for k in EXPORT_FIELDS])
    statefile.atomic_write(out_csv, "".join(buf.parts), newline="")
    obs.annotate(valid_approvals=len(orders), dropped=dropped, orphans=len(orphans))
    log_event("export", approval_ids=[o["approval_id"] for o in orders], files=[out_json, out_csv],
              source_results_file=os.path.basename(results_path), dropped=dropped, orphaned=len(orphans))
    return {"orders": orders, "dropped": dropped, "orphans": orphans, "json": out_json, "csv": out_csv}


# ---------- CLI ----------

def main(argv=None):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # piped output defaults to cp1252 on Windows
    if hasattr(sys.stdin, "reconfigure"):
        sys.stdin.reconfigure(encoding="utf-8")                  # so does pasted input
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("request")
    s.add_argument("results", help="sourcing_results_*.xlsx")
    s.add_argument("--budget", help="optional budget to compare the subtotal against, e.g. 1500")
    s = sub.add_parser("record")
    s.add_argument("--results", required=True)
    s.add_argument("--request", help="the request id the reply answered, e.g. A-0002 (optional only when exactly one "
                                     "request is open for the results file)")
    s.add_argument("--file", help="UTF-8 text file with the reply (default: read stdin)")
    s.add_argument("--confirm-request", metavar="A-000N", help="with --channel and --approver: answer the prompts without a terminal. "
                   "Must equal --request. All three, or none.")
    s.add_argument("--channel", help="the channel the reply came through, typed by a person (with --confirm-request and --approver)")
    s.add_argument("--approver", help="who approved, typed by a person; never defaulted (with --confirm-request and --channel)")
    for name in ("status", "export"):
        sub.add_parser(name).add_argument("--results", required=True)
    s = sub.add_parser("revoke")
    s.add_argument("sku")
    s.add_argument("--results", required=True)
    args = p.parse_args(argv)
    try:
        with obs.trace_command("decision_agent", args.cmd, workflow=obs.workflow_id(args.results), sku=getattr(args, "sku", None)):
            try:
                run(args)
            except ReplyError as e:
                obs.annotate(reply_rejected=True, reason_codes=sorted({next((c for c, pat in REPLY_REASON_CODES if re.search(pat, err)),
                                                                           "other") for err in e.errors}))
                raise
    except AgentError as e:
        sys.exit(str(e))
    except statefile.StateBusy as e:
        print(f"decision_agent: {e}", file=sys.stderr)
        sys.exit(3)                                              # 3 = state file busy


def run(args) -> None:
    if args.cmd == "request":
        budget = None
        if args.budget is not None:
            budget = to_decimal(args.budget)
            if budget is None:
                raise AgentError(f"--budget '{args.budget}' isn't an amount.")
        req = cmd_request(args.results, budget)
        print(render_text(req))
        print(f"\n(also saved: Reports/approval_request_{req['id']}.txt and .md)")
    elif args.cmd == "record":
        gate = check_gate(args.request, args.confirm_request, args.channel, args.approver)   # before anything is read or written
        text = open(args.file, encoding="utf-8").read() if args.file else sys.stdin.read()
        _, stored = cmd_record(args.results, text, args.request, gate)
        if stored:
            print("Recorded. Nothing has been ordered.")
        elif gate is not None:
            raise AgentError("Nothing was recorded (see above).")
    elif args.cmd == "status":
        rows, orphans, skipped = cmd_status(args.results)
        for r in rows:
            print(f"{r['sku']:<16} {r['status']:<14} {r['seller']}" + (f"  | {r['detail']}" if r["detail"] else ""))
        print(f"{len(rows)} line(s). {skip_summary(skipped)}")
        if orphans:
            print(f"{len(orphans)} earlier approval(s) no longer match the recommended seller and are orphaned: {', '.join(orphans)}.")
    elif args.cmd == "revoke":
        for d in cmd_revoke(args.sku, args.results):
            print(f"Revoked {d['id']}: {d['sku']} ({d['seller']}), qty {d['qty']:,}, cap {fmt_money(Decimal(d['line_cap']))}.")
    elif args.cmd == "export":
        r = cmd_export(args.results)
        print(f"{NOT_AN_ORDER}\nExported {len(r['orders'])} valid approval(s) to {r['json']} and {r['csv']}.")
        d = r["dropped"]
        print(f"Left out: {d['expired']} expired, {d['revoked']} revoked, {d['rejected']} rejected, {len(r['orphans'])} orphaned (seller changed).")


if __name__ == "__main__":
    main()
