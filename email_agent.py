"""Email agent: drafts (and, after confirmation, sends) quotation requests to the sourcing agent's recommended
manufacturers, then checks each seller's quote against the listing the search captured. Reads a sourcing_results_*.xlsx;
never touches sourcing_agent.py.

    python email_agent.py drafts  <results.xlsx>                        # show every draft; read-only, sends nothing
    python email_agent.py send    <results.xlsx> [--yes] [--test-recipient ADDR]
    python email_agent.py mark    <SKU> --results <results.xlsx>        # record a manual inquiry-form submission
    python email_agent.py mark    <SKU> --results <results.xlsx> --followup   # the follow-up questions were sent by hand
    python email_agent.py reply   <SKU> --results <results.xlsx> [--file reply.txt]   # paste reply, summarize, check the quote
    python email_agent.py compare <SKU> --results <results.xlsx>        # check the stored quote again (no AI call)
    python email_agent.py report  <results.xlsx>                        # markdown + Excel report for Neeraj

The quote check and the follow-up questions are code only: one AI call per reply reads the quote, verify_quote() proves
every number against the reply, and check_quote() compares the proven numbers with the listing. Nothing is sent by it.

State (email_state.json) is keyed by SKU + listing URL, so a re-run that recommends a different seller for a SKU
never shows the old seller's reply or sent mark under the new one.

Env (.env): SMTP_HOST, SMTP_PORT (587), SMTP_USER, SMTP_PASSWORD, SMTP_FROM (default SMTP_USER), SENDER_NAME,
ANTHROPIC_API_KEY, and SHIPPING_ADDRESS - one line in double quotes with \\n between lines
    SHIPPING_ADDRESS="Zync Technologies\\n123 Example St\\nPlano, TX 75024"
or a real multi-line value inside the quotes. Real sends are refused while it is unset or still the default.
"""
import argparse
import collections
import csv
import json
import os
import re
import smtplib
import sys
import unicodedata
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from email.message import EmailMessage
from urllib.parse import urlsplit

import openpyxl
from dotenv import load_dotenv
from openpyxl.styles import Font

import observability as obs
import statefile

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(HERE, "email_template.txt")
STATE_PATH = os.path.join(HERE, "email_state.json")  # per SKU + listing: sent / reply / summary
REVIEWER_EXCLUSIONS_CSV = os.path.join(HERE, "reviewer_exclusions.csv")  # shared with the search agent; read only
NEAR_MISSES_CSV = os.path.join(HERE, "near_misses.csv")  # sku, supplier_or_url, concern, question; human-maintained, read only
MODEL = "claude-haiku-4-5"
DEFAULT_SHIPPING = "Zync Technologies\nPlano, TX"
REPLY_STATUSES = ("pricing_provided", "needs_info", "dead_end", "auto_reply_or_spam")


def cfg(name: str, default: str = "") -> str:
    return os.environ.get(name, default)


class SendRefused(Exception):
    """A real send was blocked before anything went out."""


def shipping_address() -> tuple:
    """(address text, is_default). Unset, blank, or still "Zync Technologies, Plano, TX" means the real
    address was never configured."""
    text = cfg("SHIPPING_ADDRESS").replace("\\n", "\n").strip()
    squash = lambda s: re.sub(r"[\s,]+", " ", s).strip().lower()
    return (text or DEFAULT_SHIPPING), (not text or squash(text) == squash(DEFAULT_SHIPPING))


def warnings_for(exclusions_loaded: bool = True) -> list:
    problems = []
    if shipping_address()[1]:
        problems.append("SHIPPING_ADDRESS is not set (or is still the default 'Zync Technologies, Plano, TX'). "
                        "Drafts use the default; real sends are REFUSED until it is changed in .env.")
    if not exclusions_loaded:
        problems.append(f"{os.path.basename(REVIEWER_EXCLUSIONS_CSV)} is missing or empty - known-bad sellers are "
                        "NOT being skipped.")
    return problems


def print_warnings(problems: list) -> None:
    if problems:
        bar = "!" * 72
        print("\n".join([bar, *(f"!!! WARNING: {p}" for p in problems), bar]))


# ---------- input: read the search agent's Results sheet ----------

def listing_key(sku: str, url: str, manufacturer: str) -> str:
    """State key: SKU + listing URL (host and path only), or the manufacturer's name when there is no URL."""
    ref = (url or "").strip()
    if ref:
        parts = urlsplit(ref)
        ref = f"{parts.netloc.lower()}{parts.path.rstrip('/')}"
    else:
        ref = (manufacturer or "").strip().lower()
    return f"{str(sku).strip().upper()}|{ref}"


def load_exclusions(path: str = None) -> list:
    """Rows of reviewer_exclusions.csv (sku, supplier_or_url, reason); [] when missing. Never written."""
    try:
        with open(path or REVIEWER_EXCLUSIONS_CSV, newline="", encoding="utf-8") as f:
            return [r for r in csv.DictReader(f) if (r.get("supplier_or_url") or "").strip()]
    except FileNotFoundError:
        return []


def excluded_reason(sku: str, manufacturer: str, url: str, exclusions: list) -> str:
    """The reviewer's reason when this SKU's seller or listing URL is in the exclusions file, else ""."""
    haystack = f"{manufacturer or ''} {url or ''}".lower()
    for e in exclusions:
        if e["sku"].strip().upper() == str(sku).strip().upper() and e["supplier_or_url"].strip().lower() in haystack:
            return e["reason"].strip()
    return ""


def load_results(path: str, exclusions: list = None) -> tuple:
    """(candidates, skipped). A candidate is a Results row with a recommended seller, ready to draft. Skipped rows
    are listed with the reason: no recommendation, the search agent errored, or a reviewer exclusion."""
    exclusions = load_exclusions() if exclusions is None else exclusions
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Results"].iter_rows(values_only=True))
    wb.close()  # Windows keeps the file locked otherwise
    header, rows = list(rows[0]), rows[1:]
    out, skipped = [], []
    for values in rows:
        r = dict(zip(header, values))
        sku, product = str(r["SKU"]), r.get("Product") or ""
        skip = {"sku": sku, "product": product, "manufacturer": r.get("Recommended Manufacturer") or ""}
        if str(r.get("Recommendation") or "").startswith("Error researching"):
            skipped.append({**skip, "reason": "error researching this product", "kind": "error"})
            continue
        if not (r.get("Recommended Manufacturer") or r.get("Recommended URL")):
            skipped.append({**skip, "reason": "no recommended seller", "kind": "none"})
            continue
        url = r.get("Recommended URL") or ""
        why = excluded_reason(sku, r.get("Recommended Manufacturer"), url, exclusions)
        if why:
            skipped.append({**skip, "reason": f"excluded by reviewer: {why}", "kind": "reviewer"})
            continue
        moq, title, listed, confidence = "", r.get("Recommended Listing Title") or "", "", ""
        for n in range(1, 10):  # per-candidate columns: find the recommended one by URL
            if f"Manufacturer {n} URL" not in r:
                break
            if url and r[f"Manufacturer {n} URL"] == url:
                moq = r.get(f"Manufacturer {n} MOQ") or ""
                title = title or r.get(f"Manufacturer {n} Listing Title") or ""
                listed = str(r.get(f"Manufacturer {n} Listed Price") or "")              # raw listing text, for the quote check
                confidence = str(r.get(f"Manufacturer {n} Unit Price Confidence") or "")
        out.append({
            "sku": sku, "product": product, "description": r.get("Keyword") or "",
            "manufacturer": r.get("Recommended Manufacturer") or "", "email": r.get("Recommended Email") or "",
            "url": url, "unit_price": r.get("Recommended Unit Price"), "moq": moq, "title": title,
            "listed": listed, "confidence": confidence,
            "note": r.get("Ordering Note") or "",
            "key": listing_key(sku, url, r.get("Recommended Manufacturer") or ""),
        })
    return out, skipped


def load_candidates(path: str) -> list:
    return load_results(path)[0]


def load_near_misses(path: str = None):
    """Rows of near_misses.csv (sku, supplier_or_url, concern, question), or None when the file is missing. Never written."""
    try:
        with open(path or NEAR_MISSES_CSV, newline="", encoding="utf-8") as f:
            return [r for r in csv.DictReader(f) if (r.get("sku") or "").strip() and (r.get("supplier_or_url") or "").strip()]
    except FileNotFoundError:
        return None


def near_miss_drafts(results_path: str, recommended: list, exclusions: list, near_rows: list) -> tuple:
    """(drafts, notes): one draft per near_misses.csv row whose candidate is in the results file's candidate columns.
    Opt-in only (drafts --near-miss). The search agent did NOT recommend these, so each draft carries the one concern a
    human must check (printed in the header, never in the text meant for the seller) and one polite question for the
    seller. Reviewer exclusions still apply, and a seller already recommended for that SKU is not drafted twice."""
    wb = openpyxl.load_workbook(results_path, read_only=True, data_only=True)
    rows = list(wb["Results"].iter_rows(values_only=True))
    wb.close()
    header = list(rows[0])
    by_sku = {}
    for values in rows[1:]:
        r = dict(zip(header, values))
        by_sku.setdefault(str(r["SKU"]).strip().upper(), r)
    have = {d["key"] for d in recommended}
    out, notes = [], []
    for near in near_rows:
        sku, wanted = near["sku"].strip(), near["supplier_or_url"].strip()
        r = by_sku.get(sku.upper())
        block = next((n for n in range(1, 10) if r and f"Manufacturer {n} URL" in r and (
            (r[f"Manufacturer {n} URL"] or "").strip() == wanted
            or wanted.lower() in (r.get(f"Manufacturer {n} Name") or "").lower())), None)
        if block is None:
            notes.append(f"near-miss {sku}: not found in the results file ({'no such SKU' if r is None else 'seller or URL not among its candidates'}) - skipped")
            continue
        url, maker = (r[f"Manufacturer {block} URL"] or "").strip(), (r.get(f"Manufacturer {block} Name") or "").strip()
        why = excluded_reason(sku, maker, url, exclusions)
        if why:
            notes.append(f"near-miss {sku} ({maker or url}): excluded by reviewer ({why}) - skipped")
            continue
        key = listing_key(sku, url, maker)
        if key in have:
            notes.append(f"near-miss {sku} ({maker or url}): already a recommended draft - not drafted twice")
            continue
        have.add(key)
        out.append(make_draft({
            "sku": sku, "product": r.get("Product") or "", "description": r.get("Keyword") or "", "manufacturer": maker,
            "email": (r.get(f"Manufacturer {block} Email") or "").strip(), "url": url, "unit_price": None,
            "moq": r.get(f"Manufacturer {block} MOQ") or "", "title": r.get(f"Manufacturer {block} Listing Title") or "",
            "note": "", "key": key, "check": near["concern"].strip(), "questions": [near["question"].strip()]}))
    return out, notes


def format_draft(d: dict) -> str:
    """One draft as printed by `drafts`. The NEEDS HUMAN CHECK and NO EMAIL lines are for us, above the text for the seller."""
    to = d["email"] or "NO EMAIL - inquiry form"
    check = f"NEEDS HUMAN CHECK: {d['check']}\n" if d.get("check") else ""
    label = f"{d['label']}\n" if d["label"] else ""
    return f"{'=' * 70}\n{d['sku']} | {d['manufacturer'] or '(unnamed)'} | To: {to}\n{check}{label}Subject: {d['subject']}\n\n{d['body']}"


def skip_summary(skipped: list) -> str:
    kinds = {"none": "no recommended seller", "error": "search agent error", "reviewer": "excluded by reviewer_exclusions.csv"}
    parts = [f"{sum(s['kind'] == k for s in skipped)} {label}" for k, label in kinds.items() if any(s["kind"] == k for s in skipped)]
    return f"Skipped {len(skipped)} row(s)" + (f": {', '.join(parts)}." if parts else ".")


# ---------- template ----------

# The search agent's Ordering Note is internal ("check before ordering", "target is crimp", the model's own
# caveats). A seller must never see it, so each known note becomes a polite confirmation question and
# anything unrecognised - including all the model's free text after "model:" - is dropped.
NOTE_QUESTIONS = [
    (re.compile(r'title names .+? \("(.+?)"\)'),
     lambda m: f'Your listing title mentions "{m[1]}". Could you confirm the exact variant matches the description above?'),
    (re.compile(r"one listing covers several models \((.+?)\)"),
     lambda m: f"Your listing covers several models ({m[1]}). Could you confirm which one the sample would be?"),
    (re.compile(r"listing offers (.+?) variants"),
     lambda m: f"Your listing offers {m[1]} options. Could you tell us which option to select for this product?"),
    (re.compile(r"price needs an order of ([\d,]+)\+ pieces"),
     lambda m: f"Is the listed price based on an order of {m[1]}+ pieces, and what would it be at sample quantity?"),
    (re.compile(r"price is for the (.+?) tier"),
     lambda m: f"Could you confirm which price tier applies to our sample quantity (the listing shows {m[1]})?"),
    (re.compile(r"wasn't tied to this price"),
     lambda m: "Could you confirm whether the listed price is per piece or per pack?"),
    (re.compile(r"promo price"),
     lambda m: "Could you confirm the standard price for this item?"),
    (re.compile(r"contact type isn't stated"),
     lambda m: "Could you confirm the contact type (crimp, solder cup, PCB or IDC)?"),
    (re.compile(r"several contact types"),
     lambda m: "Your listing mentions more than one contact type. Could you confirm which one applies?"),
    (re.compile(r"contradicts itself on category"),
     lambda m: "Could you confirm the category/rating of this item? The listing seems to mention more than one."),
    (re.compile(r"contradicts itself on gender"),
     lambda m: "Could you confirm the connector gender (male or female)?"),
    (re.compile(r"short cable rather than a plain adapter"),
     lambda m: "Could you confirm this is a standalone adapter/coupler and not a cable with connectors on the ends?"),
]
MAX_QUESTIONS = 4


def confirmation_questions(note: str) -> list:
    """Ordering note -> polite questions for the seller. Free text after "model:" is never used."""
    ours = re.split(r"(?:^|;\s*)model:", note or "", maxsplit=1)[0]
    questions = []
    for segment in ours.split("; "):
        for pattern, ask in NOTE_QUESTIONS:
            m = pattern.search(segment)
            if m:
                q = ask(m)
                if q not in questions:
                    questions.append(q)
                break
    return questions[:MAX_QUESTIONS]


# Names the search agent uses when a listing doesn't name its maker - never greet a seller with them.
PLACEHOLDER_NAME = re.compile(r"unknown|unnamed|generic|not stated|unbranded|seller|^none$", re.IGNORECASE)


QUOTE_QTYS = (200, 500)   # every inquiry asks for these quantities (Neeraj: a minimum of 200 each, and 500 can be ordered)


def qty_list(qtys) -> str:
    return _and(f"{q:,}" for q in qtys)


def salutation(c: dict) -> tuple:
    """(greeting, listing_ref) for a seller: the seller's own listing title when the results file has it, otherwise just
    the URL - never our description - and never a placeholder maker name."""
    name, title = (c["manufacturer"] or "").strip(), (c.get("title") or "").strip()
    if c["url"]:
        listing_ref = f'your listing "{title}" ({c["url"]})' if title else f"your listing ({c['url']})"
    else:
        listing_ref = "your products"
    return ("Hello," if not name or PLACEHOLDER_NAME.search(name) else f"Hello {name} team,"), listing_ref


def render_email(c: dict, template_path: str = TEMPLATE_PATH, sender_name: str = None) -> tuple:
    """Returns (subject, body). Edit email_template.txt to change the wording. `sender_name` is for a caller that does not
    load .env itself (the UI's copy box); the agent's own commands use SENDER_NAME."""
    with open(template_path, encoding="utf-8") as f:
        subject, _, body = f.read().partition("\n")
    greeting, listing_ref = salutation(c)
    asked = len(re.findall(r"^\d+\. ", body, re.M))          # extra questions continue the template's own numbering
    values = {
        "sku": c["sku"], "product": c["description"] or c["product"],
        "greeting": greeting, "listing_ref": listing_ref, "quantities": qty_list(QUOTE_QTYS),
        "note_line": "".join(f"\n{i}. {q}" for i, q in enumerate(
            confirmation_questions(c["note"]) + list(c.get("questions") or []), start=asked + 1)),
        "sender_name": sender_name or cfg("SENDER_NAME", "Swastik Raj"),
        "shipping_address": shipping_address()[0],
    }
    return subject.removeprefix("Subject:").strip().format(**values), body.strip().format(**values) + "\n"


def make_draft(c: dict, sender_name: str = None) -> dict:
    """The body is always clean, ready to paste into a seller's inquiry form. The no-email label is a
    separate field for us: it never travels inside the text a seller receives."""
    subject, body = render_email(c, sender_name=sender_name)
    label = "" if c["email"] else f"NO EMAIL - submit via the inquiry form at {c['url']}"
    key = c.get("key") or listing_key(c["sku"], c["url"], c["manufacturer"])
    return {**c, "subject": subject, "body": body, "label": label, "key": key}


# ---------- state ----------

def load_state() -> dict:
    try:
        with open(STATE_PATH, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


def save_state(state: dict) -> None:
    statefile.atomic_write(STATE_PATH, json.dumps(state, indent=2))     # temp file + rename: a failure keeps the old file


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def entry_for(state: dict, d: dict) -> dict:
    entry = state.setdefault(d["key"], {"sku": d["sku"], "seller": d["manufacturer"] or d["url"]})
    if obs.current_workflow():                     # extra key, only while tracing is on; no existing key changes
        entry["workflow_id"] = obs.current_workflow()
    return entry


# ---------- sending (always behind a confirmation) ----------

def smtp_send(draft: dict) -> None:
    msg = EmailMessage()
    msg["Subject"], msg["To"] = draft["subject"], draft["email"]
    msg["From"] = cfg("SMTP_FROM") or cfg("SMTP_USER")
    msg.set_content(draft["body"])
    with smtplib.SMTP(cfg("SMTP_HOST"), int(cfg("SMTP_PORT", "587"))) as s:
        s.starttls()
        s.login(cfg("SMTP_USER"), cfg("SMTP_PASSWORD"))
        s.send_message(msg)


def as_test(d: dict, test_recipient: str) -> dict:
    """A draft redirected to the test address: subject marked [TEST], body says where it would really have gone."""
    where = d["email"] or f"nobody - no email, inquiry form {d['url']}"
    return {**d, "email": test_recipient, "subject": f"[TEST] {d['subject']}",
            "body": f"[TEST - this would have gone to: {where}]\n\n{d['body']}"}


def duplicate_recipients(todo: list) -> dict:
    """{address: [SKUs]} for addresses that more than one email in this batch is going to."""
    groups = {}
    for d in todo:
        groups.setdefault(d["email"].strip().lower(), []).append(d["sku"])
    return {email: skus for email, skus in groups.items() if len(skus) > 1}


def confirm_and_send(drafts: list, state: dict, yes: bool = False, send=smtp_send, ask=input,
                     test_recipient: str = "", confirm_count=None) -> int:
    """Sends only after an explicit confirmation: 'y' for one email, or the word 'yes' for the whole batch.
    Anything else (Enter, 'n', 'q', typos) sends nothing. Real mode sends only drafts that have an email and
    aren't already sent, and refuses while SHIPPING_ADDRESS is unset or still the default. With test_recipient,
    EVERY draft (even those without an email) goes to that address only, marked [TEST], and is not recorded
    as sent. With confirm_count (the UI's way to confirm without a terminal) the number of emails left AFTER all of the
    checks above must equal it, or nothing is sent; when it matches, nothing is asked."""
    if test_recipient:
        todo = [as_test(d, test_recipient) for d in drafts]
    else:
        if shipping_address()[1]:
            raise SendRefused("SHIPPING_ADDRESS is not set (or is still the default 'Zync Technologies, Plano, TX'). "
                              "Set the real address in .env before sending. Nothing was sent.")
        todo = [d for d in drafts if d["email"] and "sent" not in state.get(d["key"], {})]
    dups = {} if test_recipient else duplicate_recipients(todo)
    for email, skus in dups.items():
        print(f"WARNING: {len(skus)} separate emails ({', '.join(skus)}) are addressed to the same recipient {email}.")
    if confirm_count is not None:
        if len(todo) != confirm_count:
            raise SendRefused(f"--confirm-count is {confirm_count} but {len(todo)} email(s) would be sent. Nothing was sent.")
        yes = True
    if yes:
        print(f"Batch send{' (TEST to ' + test_recipient + ')' if test_recipient else ''}: {len(todo)} email(s).")
        if confirm_count is None and ask("Type 'yes' to send all: ").strip().lower() != "yes":
            return 0
    sent = 0
    for d in todo:
        if test_recipient and d["email"] != test_recipient:
            raise RuntimeError("test mode would send to a real address - aborting")
        if not yes:
            same = f"\nNOTE: same address as {', '.join(s for s in dups[d['email'].lower()] if s != d['sku'])}" \
                if d["email"].lower() in dups else ""
            print(f"\nTo: {d['email']}\nSubject: {d['subject']}{same}\n\n{d['body']}")
            answer = ask(f"Send {d['sku']} to {d['email']}? [y/N/q] ").strip().lower()
            if answer == "q":
                break
            if answer != "y":
                continue
        send(d)
        if not test_recipient:  # a test send must not make the real one look done
            entry = entry_for(state, d)
            entry["sent"], entry["asked_qty"] = now(), list(QUOTE_QTYS)       # the quantities this inquiry named
            save_state(state)
        sent += 1
    return sent


# ---------- reply summarization ----------

# Structured quote fields the summarizer may fill. The model extracts; verify_quote() below checks every number against
# the reply text and nulls what it can't prove. Anything not stated in the reply is None - never guessed.
QUOTE_INT_FIELDS = ("sample_quantity", "moq", "branding_min_qty", "lead_time_days", "production_lead_time_days",
                    "quote_total_quantity")
QUOTE_MONEY_FIELDS = ("sample_unit_price", "sample_total_price", "branding_fee", "shipping_cost", "quote_total_price")
QUOTE_EVIDENCE_FIELDS = QUOTE_INT_FIELDS + QUOTE_MONEY_FIELDS      # each non-null one needs a verbatim evidence quote
QUOTE_KEYS = ("sample_available", "sample_quantity", "sample_unit_price", "sample_total_price", "price_basis", "moq",
              "bulk_tiers", "branding_possible", "branding_fee", "branding_min_qty", "lead_time_days",
              "production_lead_time_days", "shipping_cost", "shipping_terms", "currency", "needs_from_us",
              "quote_valid_until", "quote_total_price", "quote_total_quantity", "evidence", "quote_warnings")
PRICE_BASES = ("per_unit", "total_for_quantity", "unclear")
# Currency counts as USD only if $, US$, USD, 美元 or 美金 is in the evidence; RMB, CNY, the yuan signs, 人民币 or a number
# followed by 元 mean CNY; EUR and the euro sign mean EUR; anything else is "unknown". One place to extend.
# ORDER MATTERS: _currency_codes() blanks each marker it finds before trying the next pattern, so USD (which includes
# 美元 and 美金, both of which contain 元) is matched first and a USD quote is never also read as CNY. A bare 元 only
# counts right after a digit ("22元"), so words like 单元 (unit) are not read as a currency, and 元件 (component) is skipped.
CURRENCY_PATTERNS = (("USD", re.compile(r"US\$|USD|\$|美元|美金", re.IGNORECASE)),
                     ("CNY", re.compile(r"RMB|CNY|¥|￥|人民币|(?<=\d)\s*元(?!件)", re.IGNORECASE)),
                     ("EUR", re.compile(r"EUR|€", re.IGNORECASE)))


def parse_summary(text: str) -> dict:
    data = json.loads(text[text.index("{"): text.rindex("}") + 1])
    if data.get("status") not in REPLY_STATUSES:
        raise ValueError(f"bad status: {data.get('status')!r}")
    extras = {k: data[k] for k in QUOTE_KEYS if k in data}       # the quote fields travel on, still unverified
    return {**extras, "status": data["status"], "summary": str(data.get("summary", ""))}


QUOTE_PROMPT = (
    "A manufacturer replied to our sample-order inquiry or quotation request (possibly several replies, oldest first, in any "
    "language). "
    "Reply with JSON only. Keys: status, summary, and the quote fields below.\n"
    "status: pricing_provided = gave pricing/availability; needs_info = asks us for more; dead_end = a real person "
    "declined or the product is not available; auto_reply_or_spam = out-of-office, automated message or spam with no "
    "real answer.\nsummary: 2-3 sentences, in English.\n"
    "Quote fields - use null for anything the reply does not state. NEVER guess, infer, convert or calculate:\n"
    "  sample_available (yes|no|unknown); sample_quantity (integer); sample_unit_price and sample_total_price (strings, "
    "digits only, either may be null); price_basis (per_unit|total_for_quantity|unclear); moq (integer); "
    "bulk_tiers (list of {min_qty, unit_price, evidence}); branding_possible (yes|no|unknown); branding_fee; "
    "branding_min_qty; lead_time_days (sample) and production_lead_time_days (integers); shipping_cost; "
    "shipping_terms (as stated, e.g. FOB, EXW, DDP); currency (as stated); needs_from_us (list of short strings: what "
    "the seller asks us for); quote_valid_until (as stated); quote_total_price and quote_total_quantity (a total price the "
    "seller states for a stated non-sample quantity, and that quantity).\n"
    "Every non-sample price tied to a quantity goes in bulk_tiers, including a price for the quantity we asked about "
    "(min_qty = that quantity).\n"
    "evidence: an object mapping each non-null price, quantity, fee or lead-time field name to a SHORT VERBATIM quote "
    "copied from the reply that contains that exact number. Each bulk tier carries its own evidence. If you cannot quote "
    "it, set the field to null.\n\n")


def summarize_reply(reply: str) -> dict:
    from anthropic import Anthropic
    obs.instrument_anthropic()                      # record this call (tokens, latency) when tracing is on
    msg = Anthropic(api_key=cfg("ANTHROPIC_API_KEY")).messages.create(
        model=MODEL, max_tokens=1500, messages=[{"role": "user", "content": f"{QUOTE_PROMPT}Reply:\n{reply}"}])
    return parse_summary(msg.content[0].text)


def _currency_codes(text) -> set:
    """Every currency whose marker appears in the text. A marker that matched is blanked before the next pattern runs,
    so 美元 / 美金 (USD) can never also be counted as the 元 of CNY."""
    left, found = str(text), set()
    for code, pattern in CURRENCY_PATTERNS:
        if pattern.search(left):
            found.add(code)
            left = pattern.sub(" ", left)
    return found


def _currency_code(text):
    """The currency a model says it quoted in (a bare "元" is accepted here), else None."""
    said = str(text).strip()
    if said == "元":
        return "CNY"
    codes = _currency_codes(said)
    return next((code for code, _ in CURRENCY_PATTERNS if code in codes), None)


def _norm(text) -> str:
    """Case, width and whitespace-insensitive form used to find a quote inside the reply."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(text))).strip().lower()


def _numbers(text) -> set:
    return {Decimal(n.replace(",", "")) for n in re.findall(r"\d[\d,]*(?:\.\d+)?|\.\d+", unicodedata.normalize("NFKC", str(text)))}


def _whole_number(value):
    """Positive whole number from an int, an integer-valued float or a digit string; else None."""
    if isinstance(value, bool):
        return None
    try:
        d = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None
    return int(d) if d.is_finite() and d == d.to_integral_value() and d >= 1 else None


def _amount(value):
    """Non-negative finite Decimal from a number or a string like "$1,200.50"; else None. Money is never a float."""
    if isinstance(value, bool) or value is None:
        return None
    try:
        d = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except (InvalidOperation, ValueError):
        return None
    return d if d.is_finite() and d >= 0 else None


def _prove(label: str, number, evidence, text_norm: str, warnings: list) -> bool:
    """The evidence must appear in the reply text and the number must appear inside the evidence."""
    if not isinstance(evidence, str) or not evidence.strip():
        warnings.append(f"{label}: no evidence quote given")
    elif _norm(evidence) not in text_norm:
        warnings.append(f"{label}: evidence quote not found in the reply text")
    elif Decimal(number) not in _numbers(evidence):
        warnings.append(f"{label}: the number {number} does not appear in its evidence quote")
    else:
        return True
    return False


def verify_quote(raw: dict, reply_text: str) -> dict:
    """Checks the summarizer's quote fields against the reply and returns the verified quote plus `evidence` and
    `quote_warnings`. A numeric field whose evidence isn't in the reply, or whose number isn't in its evidence, is set
    to None with a warning naming the field and the reason."""
    warnings, evidence, out = [], {}, {k: None for k in QUOTE_KEYS}
    text_norm, given = _norm(reply_text), raw.get("evidence") if isinstance(raw.get("evidence"), dict) else {}
    used = []                                                              # evidence strings that back a price
    for field in QUOTE_INT_FIELDS:
        value = raw.get(field)
        if value is None:
            continue
        n = _whole_number(value)
        if n is None:
            warnings.append(f"{field}: {value!r} is not a positive whole number")
        elif _prove(field, n, given.get(field), text_norm, warnings):
            out[field], evidence[field] = n, given[field]
    for field in QUOTE_MONEY_FIELDS:
        value = raw.get(field)
        if value is None:
            continue
        d = _amount(value)
        if d is None:
            warnings.append(f"{field}: {value!r} is not an amount")
        elif _prove(field, d, given.get(field), text_norm, warnings):
            out[field], evidence[field] = format(d, "f"), given[field]
            used.append(given[field])
    tiers = []
    for i, tier in enumerate(raw.get("bulk_tiers") or []):
        label = f"bulk_tiers[{i}]"
        qty, price = (_whole_number(tier.get("min_qty")), _amount(tier.get("unit_price"))) if isinstance(tier, dict) else (None, None)
        if qty is None or price is None:
            warnings.append(f"{label}: needs a positive whole min_qty and an amount unit_price")
            continue
        ev = tier.get("evidence")
        if _prove(f"{label}.min_qty", qty, ev, text_norm, warnings) and _prove(f"{label}.unit_price", price, ev, text_norm, warnings):
            tiers.append({"min_qty": qty, "unit_price": format(price, "f"), "evidence": ev})
            used.append(ev)
    out["bulk_tiers"] = tiers or None
    for field in ("sample_available", "branding_possible"):
        value = str(raw.get(field) or "unknown").strip().lower()
        if value not in ("yes", "no", "unknown"):
            warnings.append(f"{field}: {value!r} is not yes/no/unknown")
            value = "unknown"
        out[field] = value
    basis = raw.get("price_basis")
    if basis is not None and basis not in PRICE_BASES:
        warnings.append(f"price_basis: {basis!r} is not per_unit/total_for_quantity/unclear")
        basis = None
    if used and basis is None:
        basis = "unclear"                                    # a price without a stated basis is not guessed at
    out["price_basis"] = basis if used else None
    terms = raw.get("shipping_terms")
    if terms:
        if _norm(terms) in text_norm:
            out["shipping_terms"] = str(terms).strip()
        else:
            warnings.append("shipping_terms: not found in the reply text")
    asks = raw.get("needs_from_us")
    out["needs_from_us"] = [str(a).strip() for a in asks if str(a).strip()] if isinstance(asks, list) else None
    out["needs_from_us"] = out["needs_from_us"] or None
    out["quote_valid_until"] = str(raw["quote_valid_until"]).strip() if raw.get("quote_valid_until") else None
    # Currency comes from the evidence, not from the model's say-so.
    if used:
        found = set().union(*(_currency_codes(u) for u in used))
        out["currency"] = next(iter(found)) if len(found) == 1 else "unknown"
        if len(found) > 1:
            warnings.append("currency: the evidence mixes currencies")
        said = raw.get("currency")
        if said and (_currency_code(said) or str(said).strip().upper()) != out["currency"]:
            warnings.append(f"currency: the model said {said!r} but the evidence shows {out['currency']}")
    out["evidence"] = evidence
    out["quote_warnings"] = warnings
    return out


def quote_lines(result: dict) -> list:
    """What was recorded from a reply's quote, for the terminal: the verified fields and any warnings."""
    shown = [f"{k}={result[k]}" for k in QUOTE_KEYS if k not in ("evidence", "quote_warnings", "bulk_tiers", "needs_from_us")
             and result.get(k) not in (None, "unknown")]
    out = ["Quote recorded: " + (", ".join(shown) if shown else "nothing quantified in this reply")]
    if result.get("bulk_tiers"):
        out.append("  bulk tiers: " + "; ".join(f">={t['min_qty']} pcs at {t['unit_price']}" for t in result["bulk_tiers"]))
    if result.get("needs_from_us"):
        out.append("  the seller asks us for: " + "; ".join(result["needs_from_us"]))
    out += [f"  WARNING (field set to null): {w}" for w in result.get("quote_warnings", [])]
    return out


def record_reply(state: dict, draft: dict, text: str, summarize=None) -> dict:
    """Files a reply under this draft's SKU + listing. A second reply is appended to the first, never replacing
    it, and the whole thread is summarized. The summarizer's quote fields are verified against the thread before
    they are stored. Nothing is stored if summarizing fails."""
    entry = state.get(draft["key"], {})
    replies = list(entry.get("replies", []))
    if not replies or replies[-1]["text"].strip() != text.strip():   # pasting the same reply twice adds nothing
        replies.append({"at": now(), "text": text})
    combined = replies[0]["text"] if len(replies) == 1 else "\n\n".join(
        f"--- reply {i} ({r['at']}) ---\n{r['text']}" for i, r in enumerate(replies, start=1))
    with obs.span("summarize_reply", replies=len(replies), reply_chars=len(combined)):
        raw = (summarize or summarize_reply)(combined)
    if raw.get("status") not in REPLY_STATUSES:
        raise ValueError(f"bad status: {raw.get('status')!r}")
    result = {"status": raw["status"], "summary": str(raw.get("summary", "")), **verify_quote(raw, combined)}
    if obs.enabled():                              # what the code-side check did to the model's numbers
        obs.score("verification_nulls", sum(raw.get(k) is not None and result[k] is None for k in QUOTE_EVIDENCE_FIELDS))
        obs.score("quote_warnings", len(result["quote_warnings"]))
        obs.score("currency_detected", result.get("currency") or "none")
    entry = entry_for(state, draft)
    entry.update(replies=replies, reply=combined, **result)
    return result


# ---------- quote check: the verified quote against the listing the search captured (code only, no AI call) ----------

TOLERANCE_PCT, TOLERANCE_MIN = Decimal("0.02"), Decimal("0.01")   # within max(US$0.01, 2%) of the listing counts as equal
BLOCKING, CONFIRM, NOTE = "blocking", "needs confirmation", "note"
COUNT_KEYS = {BLOCKING: "blocking", CONFIRM: "needs_confirmation", NOTE: "note"}
NOT_CHECKED = {"dead_end": "the seller declined", "auto_reply_or_spam": "the reply is automated or spam"}
CENT = Decimal("0.01")

# The results file's "Listed Price" text: tiers in the shapes sourcing_agent.PRICE_TIER reads ("$2.68 (100-999 pieces)",
# "$0.50 (≥50,000 pieces)"), else a range with no quantity ("US$3.82-5.22"), else the search's per-piece figure is used.
LISTING_TIER = re.compile(r"\$\s*(\d+(?:\.\d+)?)\s*\(?\s*(?:(\d[\d,]*)\s*[-–]\s*(\d[\d,]*)|[≥>]=?\s*(\d[\d,]*))\s*(?:pieces|pcs|pc)",
                          re.IGNORECASE)
LISTING_RANGE = re.compile(r"\$\s*(\d+(?:\.\d+)?)\s*[-–]\s*\$?\s*(?![\d.]+\s*%)(\d+(?:\.\d+)?)")
NON_USD_MARK = re.compile(r"¥|￥|RMB|CNY|€|EUR", re.IGNORECASE)

# Rules on the seller's own words (never on anything the AI wrote). A unit is read from a price's verified evidence quote.
PIECE_UNIT = re.compile(r"(?:/|\bper\s+|\ba\s+)\s*(?:pcs?|pieces?|units?|items?|ea)\b|\beach\b|[每/]\s*(?:个|件|只|支|片)|单价", re.IGNORECASE)
PACK_UNIT = re.compile(r"(?:/|\bper\s+|\ba\s+|\beach\s+)\s*(?:packs?|pk|sets?|box(?:es)?|bags?|lots?|rolls?|cartons?|ctns?)\b|"
                       r"[每/]\s*(?:包|套|盒|袋)", re.IGNORECASE)
CHARGES_EXCLUDED = re.compile(
    r"\b(?:EXW|FOB|ex[\s-]?works)\b"
    r"|\b(?:not|without|excluding|excludes?|excluded|excl\.?)\s+(?:\w+\s+){0,3}?(?:shipping|freight|tax(?:es)?|vat|dut(?:y|ies)|delivery)\b"
    r"|\b(?:shipping|freight|tax(?:es)?|vat|dut(?:y|ies)|delivery)\s+(?:cost\s+|fees?\s+|charges?\s+)?(?:is\s+|are\s+)?"
    r"(?:not included|extra|excluded|separate(?:ly)?|additional)"
    r"|不含|不包含|另付|另计|另算|运费自理", re.IGNORECASE)
VAGUE = re.compile(
    r"\b(?:about|around|approx(?:imately|\.)?|roughly|nearly)\s*(?:US\s?\$|USD|\$|RMB)?\s*\d"
    r"|(?<![\w~])~\s*(?:US\s?\$|USD|\$)?\s*\d"
    r"|\b(?:from|starting (?:at|from)|as low as)\s+(?:US\s?\$|USD|\$|RMB)\s*\d"
    r"|\bnegotiable\b|\bdepend(?:s|ing)?\b|\bsubject to\b|\bTBD\b|\bto be (?:confirmed|discussed)\b"
    r"|(?:US\s?\$|USD|\$)\s*\d+(?:\.\d+)?\s*[-–~]\s*(?:US\s?\$|\$)?\s*\d+(?:\.\d+)?(?!\s*(?:pcs|pieces|pc)\b)"
    r"|\d+(?:\.\d+)?\s*[-–~]\s*\d+(?:\.\d+)?\s*(?:USD|US\$|美元|元|RMB)"
    r"|左右|大约|约\s*\d|\d\s*(?:元|美元)?\s*起(?!订)|面议|可议|可谈|视.{0,8}?而定|待定", re.IGNORECASE)
MODEL_TOKEN = re.compile(r"(?<![A-Za-z0-9-])[A-Za-z][A-Za-z0-9]*(?:-[A-Za-z0-9]+)+(?![A-Za-z0-9-])")
CATEGORY = re.compile(r"(?<![A-Za-z0-9])cat\.?\s?-?(5e|6a|6|7a|7|8|5)(?![A-Za-z0-9])", re.IGNORECASE)
# What the inquiry asked, and how an answer is recognised: a verified quote field, or the seller's own words.
ANSWERS = (
    ("your minimum order quantity", ("moq",), re.compile(r"\bMOQ\b|\bmin(?:imum|\.)?\s*order|起订", re.IGNORECASE)),
    ("the lead time", ("lead_time_days", "production_lead_time_days"),
     re.compile(r"lead\s*time|deliver(?:y|ed)?\s+(?:time|in|within)|ready (?:in|within)|ship(?:ped|s)?\s+(?:in|within)|"
                r"\d+\s*(?:-\s*\d+\s*)?(?:working\s+|business\s+)?days|交期|交货|\d+\s*天", re.IGNORECASE)),
    ("the shipping cost to the USA", ("shipping_cost", "shipping_terms"),
     re.compile(r"ship|freight|express|courier|\b(?:DHL|FedEx|UPS|DDP|EMS)\b|运费|快递|物流|包邮", re.IGNORECASE)),
    ("how we would place the order and pay", (),
     re.compile(r"\bpay|paypal|trade assurance|\bT/T\b|\bTT\b|bank|wire|western union|deposit|place (?:an |the |your )?order|"
                r"付款|支付|定金|下单", re.IGNORECASE)),
    ("how long the quotation is valid", (), re.compile(r"\bvalid|validity|expir|有效", re.IGNORECASE)),
)

# The only text a follow-up holds besides the seller's own words: fixed wording with placeholders the code fills.
QUESTIONS = {
    "price_above_listing": "Your listing showed {listing_price} per piece{which} when we checked it on {date}. Your quote is {price} per "
                           "piece for {qty} pieces. Could you confirm whether the listed price still applies, or explain the difference?",
    "price_below_listing": "Your quote of {price} per piece for {qty} pieces is lower than the {listing_price}{which} on your listing "
                           "(checked on {date}). Could you confirm it is for the same item, per single piece?",
    "moq_differs": "Your listing showed a minimum order of {listing_moq} pieces when we checked it on {date}; your reply mentions {moq}. "
                   "Which minimum applies?",
    "moq_above_our_qty": "We are asking about {qtys} pieces, and your minimum order is {moq}. Could you supply {qtys} pieces, and at what "
                         "unit price?",
    "unit_per_pack": "Your price of {price} seems to be per pack. How many pieces are in one pack, and what is the price per single piece?",
    "unit_not_stated": "Is the price of {price} per single piece?",
    "currency_not_usd": "Could you give your prices in US dollars?",
    "currency_not_stated": "Could you confirm that your prices are in US dollars?",
    "charges_excluded": 'You wrote "{snippet}". What would shipping to the USA cost for {qtys} pieces, and are there any other charges?',
    "charges_excluded_plain": "Does your price include shipping to the USA? If not, what would shipping cost for {qtys} pieces, and are "
                              "there any other charges?",
    "qty_not_quoted": "Your quote is for {quoted} pieces. Could you give the unit price for {qtys} pieces?",
    "sample_price_only": "The price you gave ({price}) is for samples. What is the unit price for {qtys} pieces?",
    "no_price": "Could you send the unit price for {qtys} pieces, in US dollars?",
    "product_differs": "Your reply mentions {tokens}, while our inquiry is about {listing_ref}. Is this the same item?",
    "vague_wording": "You wrote {snippets}. Could you give exact figures?",
    "total_missing": "We calculate {totals}, before shipping. Could you confirm {which}?",
    "total_mismatch": "You quoted {price} per piece and a total of {total} for {qty} pieces, but {qty} x {price} = {calc}. Which is correct?",
    "not_answered": "Could you also tell us {missing}?",
}
FOLLOWUP_SUBJECT = "Re: Quotation request - {product} ({sku})"
FOLLOWUP_BODY = ("{greeting}\n\nThank you for your quotation for {listing_ref}. Before we decide, could you please help us with the "
                 "points below?\n\n{questions}\n\nThis is not an order yet. Once these points are clear, we will come back to you.\n\n"
                 "Best regards,\n{sender_name}\nZync Technologies\n")


class FollowupRefused(Exception):
    """The follow-up held a number the code did not supply, so it was not drafted."""


class NothingToCompare(Exception):
    """No reply is recorded for this seller."""


def _and(items) -> str:
    items = list(items)
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def usd(d) -> str:
    """US$ with at least 2 and at most 4 decimals: US$0.0545, US$1.25, US$1,250.00."""
    text = f"{Decimal(d).quantize(Decimal('0.0001'), ROUND_HALF_UP):f}".rstrip("0")
    whole, _, frac = text.partition(".")
    return f"US${int(whole):,}.{frac.ljust(2, '0')}"


def _cents(d: Decimal) -> Decimal:
    return d.quantize(CENT, ROUND_HALF_UP)


def _tolerance(price: Decimal) -> Decimal:
    return max(TOLERANCE_MIN, price * TOLERANCE_PCT)


def _safe(text) -> bool:
    """Seller text may be shown and quoted only if the masking rules would leave it as it is (no email, phone, secret)."""
    return bool(text) and obs.mask_text(text) == text


_CLAUSE_END = re.compile(r"[.,](?!\d)|[;!?\n。；！？，]")      # "1.25" and "10,000" stay whole


def _snippet(text: str, m):
    """The seller's own words around a match: its clause, verbatim, at most about 80 characters; None when unsafe to show."""
    starts = [x.end() for x in _CLAUSE_END.finditer(text, 0, m.start())]
    a = starts[-1] if starts else 0
    e = _CLAUSE_END.search(text, m.end())
    b = e.start() if e else len(text)
    if b - a > 80:
        a, b = max(a, m.start() - 30), min(b, m.end() + 30)
    s = text[a:b].strip()
    return s if _safe(s) else None


def capture_time(results_path: str) -> tuple:
    """(when, how we know) the search read the listings: the time in a sourcing_results_YYYYMMDD_HHMMSS name (when the
    search started; the file is saved after each product), else the file's last-modified time."""
    stamp = re.search(r"sourcing_results_(\d{8}_\d{6})", os.path.basename(results_path))
    if stamp:
        return datetime.strptime(stamp[1], "%Y%m%d_%H%M%S"), "the results file's name (when the search started)"
    return datetime.fromtimestamp(os.path.getmtime(results_path)), "the results file's last-modified time (its name has no search time)"


def listing_terms(c: dict, results_path: str) -> dict:
    """What the results file says about this listing: its price tiers or range, the search's per-piece figure and its
    confidence, its MOQ, and when the search read it. Nothing here comes from the seller's reply."""
    raw = str(c.get("listed") or "")
    num = lambda t: int(t.replace(",", ""))
    tiers = sorted(((num(lo or ge), num(hi) if hi else None, Decimal(p)) for p, lo, hi, ge in LISTING_TIER.findall(raw)),
                   key=lambda t: t[0])
    rng = None if tiers else LISTING_RANGE.search(raw)
    moq = re.search(r"\d[\d,]*", str(c.get("moq") or ""))
    when, source = capture_time(results_path)
    return {"url": c.get("url") or "", "title": c.get("title") or "", "listed": raw, "unit_price": _amount(c.get("unit_price")),
            "confidence": str(c.get("confidence") or ""), "tiers": tiers, "range": (Decimal(rng[1]), Decimal(rng[2])) if rng else None,
            "moq": (num(moq[0]) or None) if moq else None, "captured": when, "captured_from": source,
            "results_file": os.path.basename(results_path)}


def listing_price_at(listing: dict, q: int) -> tuple:
    """(price, which) the listing shows for q pieces: the tier holding q, its (low, high) range, or the search's per-piece
    figure. (None, why) when the listing has tiers and none holds q, or has no usable price."""
    if listing["tiers"]:
        for lo, hi, p in listing["tiers"]:
            if lo <= q and (hi is None or q <= hi):
                return p, f" for {lo:,}-{hi:,} pieces" if hi else f" for {lo:,}+ pieces"
        return None, "its tiers: " + "; ".join(f"{usd(p)} for {lo:,}" + (f"-{hi:,}" if hi else "+") + " pieces" for lo, hi, p in listing["tiers"])
    if listing["range"]:
        return listing["range"], ""
    if listing["unit_price"]:
        return listing["unit_price"], ""
    return None, "the results file has no usable price for it"


def quoted_tiers(quote: dict) -> list:
    """[(min_qty, unit price, evidence)] of the verified bulk tiers, lowest quantity first."""
    out = []
    for t in quote.get("bulk_tiers") or []:
        if isinstance(t, dict) and _whole_number(t.get("min_qty")) and _amount(t.get("unit_price")) is not None:
            out.append((_whole_number(t["min_qty"]), _amount(t["unit_price"]), str(t.get("evidence") or "")))
    return sorted(out, key=lambda t: t[0])


def quoted_at(quote: dict, q: int):
    """The verified tier that applies at q pieces (the largest minimum at or below q), or None."""
    met = [t for t in quoted_tiers(quote) if t[0] <= q]
    return met[-1] if met else None


def _their_words(reply: str, d: dict) -> str:
    """The reply without quoted lines ('>') or lines copied from our own inquiry (a pasted thread often repeats it), so the
    wording checks read only what the seller wrote."""
    ours = {line.strip() for line in make_draft(d)["body"].splitlines() if len(line.strip()) > 15}
    return "\n".join(line for line in reply.splitlines() if not line.lstrip().startswith(">") and line.strip() not in ours)


def _flag(code: str, severity: str, text: str, why: str, question: str = None, **values) -> dict:
    """One finding. `text` and `why` are for us. A question for the seller is a QUESTIONS template filled from `values`: the
    only way a number reaches a follow-up (build_followup checks that)."""
    return {"code": code, "severity": severity, "text": text, "why": why,
            "question": QUESTIONS[question].format(**values) if question else None, "_inserted": [str(v) for v in values.values()]}


def _per_pack(evidence: str) -> bool:
    """The seller's words price a pack, set or box rather than a piece."""
    return bool(PACK_UNIT.search(evidence) and not PIECE_UNIT.search(evidence))


def _per_piece_at(c: dict, q: int):
    """The quoted tier that applies at q, unless its words price a pack: a pack price is never compared or multiplied."""
    at = c["at"][q]
    return at if at and not _per_pack(at[2]) else None


def _price_text(c: dict, p: Decimal) -> str:
    cur = c["q"].get("currency")
    return usd(p) if c["usd"] else format(p, "f") + (f" {cur}" if cur and cur != "unknown" else "")


def check_currency(c: dict) -> list:
    if not c["priced"] or c["usd"]:
        return []
    cur = c["q"].get("currency") or "unknown"
    if cur == "unknown":
        return [_flag("currency_not_stated", CONFIRM, "The quoted prices carry no single currency (no $, USD, RMB...), so they are not "
                      "compared with the listing.", "A price without a currency cannot be compared.", "currency_not_stated")]
    return [_flag("currency_not_usd", BLOCKING, f"The quote is in {cur}. Nothing is converted, so its prices are not compared with the "
                  "listing.", "The listing is in US dollars; a conversion would be a guess.", "currency_not_usd")]


def check_price(c: dict) -> list:
    if not c["usd"]:
        return []
    listing, out = c["listing"], []
    for q in c["asked"]:
        if not _per_piece_at(c, q):
            continue
        price = c["at"][q][1]
        listed, which = listing_price_at(listing, q)
        if listed is None:
            out.append(_flag("listing_price_not_firm", NOTE, f"The listing shows no price for {q:,} pieces ({which}), so the quoted "
                             f"{usd(price)} cannot be compared.", "The listing gives nothing to hold the seller to at this quantity."))
            continue
        lo, hi = listed if isinstance(listed, tuple) else (listed, listed)
        span = f"the range {usd(lo)}-{usd(hi)}" if lo != hi else usd(lo)
        if lo - _tolerance(lo) <= price <= hi + _tolerance(hi):
            gap = f" (difference {usd(abs(price - lo))})" if lo == hi else " (inside it)"
            out.append(_flag("price_matches_listing", NOTE, f"Quoted {usd(price)} per piece for {q:,} pieces matches the listing's {span}"
                             f"{which}{gap}, seen {c['seen']}.", "The quote agrees with the listing."))
        elif price > hi:
            pct = ((price / hi - 1) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)
            out.append(_flag("price_above_listing", BLOCKING, f"Quoted {usd(price)} per piece for {q:,} pieces is {usd(price - hi)} "
                             f"({pct}%) above the listing's {span}{which}, seen {c['seen']}.", "We would pay more than the listing showed.",
                             "price_above_listing", listing_price=usd(hi), which=which or (f" (the top of its range, {usd(lo)}-{usd(hi)})" if lo != hi else ""),
                             date=c["date"], price=usd(price), qty=f"{q:,}"))
        else:
            pct = ((1 - price / lo) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)
            out.append(_flag("price_below_listing", CONFIRM, f"Quoted {usd(price)} per piece for {q:,} pieces is {usd(lo - price)} "
                             f"({pct}%) below the listing's {span}{which}, seen {c['seen']}.",
                             "A cheaper quote may be for a different item, unit or quantity.",
                             "price_below_listing", price=usd(price), qty=f"{q:,}", listing_price=usd(lo),
                             which=which or (f" (the bottom of its range, {usd(lo)}-{usd(hi)})" if lo != hi else ""), date=c["date"]))
    conf = listing["confidence"]
    if out and conf and not conf.lower().startswith("stated"):
        out.append(_flag("listing_price_not_firm", NOTE, f"The search did not find the listing's price stated plainly (it is "
                         f"'{conf}'), so this comparison is less certain.", "The listing price itself may be wrong."))
    if out and NON_USD_MARK.search(listing["listed"]):
        out.append(_flag("listing_price_not_firm", NOTE, f"The listing's own price text is not in dollars ('{listing['listed']}'); the "
                         "search's US$ figure was used.", "The listing price itself may be a conversion."))
    return out


def check_moq(c: dict) -> list:
    q_moq, l_moq, top, out = _whole_number(c["q"].get("moq")), c["listing"]["moq"], max(c["asked"]), []
    if q_moq and l_moq and q_moq != l_moq:
        out.append(_flag("moq_differs", CONFIRM, f"The reply's minimum order is {q_moq:,}; the listing showed {l_moq:,}, seen {c['seen']}.",
                         "The minimum we must buy may not be what the listing showed.", "moq_differs",
                         listing_moq=f"{l_moq:,}", date=c["date"], moq=f"{q_moq:,}"))
    if q_moq and q_moq > top:
        out.append(_flag("moq_above_our_qty", BLOCKING, f"The reply's minimum order ({q_moq:,}) is above every quantity we asked about "
                         f"({qty_list(c['asked'])}).", "We could not buy the quantity we asked about.", "moq_above_our_qty",
                         qtys=qty_list(c["asked"]), moq=f"{q_moq:,}"))
    elif not q_moq and l_moq and l_moq > top:
        out.append(_flag("moq_above_our_qty", NOTE, f"The listing's minimum order ({l_moq:,}, seen {c['seen']}) is above every quantity "
                         f"we asked about ({qty_list(c['asked'])}); the reply states no minimum.", "The seller may not sell our quantity."))
    return out


def check_unit(c: dict) -> list:
    applicable = []
    for q in c["asked"]:
        if c["at"][q] and c["at"][q] not in applicable:
            applicable.append(c["at"][q])
    priced = [(ev, p) for _, p, ev in applicable]
    sample = _amount(c["q"].get("sample_unit_price"))
    if not priced and sample is not None:
        evidence = c["q"].get("evidence") if isinstance(c["q"].get("evidence"), dict) else {}
        priced = [(str(evidence.get("sample_unit_price") or ""), sample)]
    shown = lambda ev: f': "{ev}"' if _safe(ev) else ""
    for ev, p in priced:
        if _per_pack(ev):
            return [_flag("unit_per_pack", BLOCKING, f"The price {_price_text(c, p)} reads as a price per pack{shown(ev)}. It is not "
                          "compared with the listing's per-piece price or multiplied into a total.",
                          "A pack price set against a per-piece listing price is wrong by the pack size.", "unit_per_pack",
                          price=_price_text(c, p))]
    for ev, p in priced:
        if not PIECE_UNIT.search(ev):
            return [_flag("unit_not_stated", CONFIRM, f"The reply does not say whether {_price_text(c, p)} is per piece{shown(ev)}.",
                          "A pack price can look like a cheap piece price.", "unit_not_stated", price=_price_text(c, p))]
    return []


def check_charges(c: dict) -> list:
    m = CHARGES_EXCLUDED.search(c["text"])
    if not m:
        return []
    snip, qtys = _snippet(c["text"], m), qty_list(c["asked"])
    ask = {"question": "charges_excluded", "snippet": snip, "qtys": qtys} if snip else {"question": "charges_excluded_plain", "qtys": qtys}
    return [_flag("charges_excluded", CONFIRM, f"The price leaves out shipping or other charges: \"{snip or m[0]}\".",
                  "What we would really pay is more than the unit price.", **ask)]


def check_quantities(c: dict) -> list:
    missing = [q for q in c["asked"] if not c["at"][q]]
    if not missing:
        return []
    tiers, sample = quoted_tiers(c["q"]), _amount(c["q"].get("sample_unit_price"))
    if tiers:
        quoted = qty_list(t[0] for t in tiers)
        return [_flag("qty_not_quoted", CONFIRM, f"No quoted price applies at {qty_list(missing)} pieces; the quote starts at {quoted} "
                      "pieces.", "We cannot compare or order our quantity without its price.", "qty_not_quoted",
                      quoted=quoted, qtys=qty_list(missing))]
    if sample is not None:
        return [_flag("sample_price_only", CONFIRM, f"Only a sample price ({_price_text(c, sample)}) was quoted. It is not compared "
                      "with the listing's price.", "A sample price is not the price for our quantity.", "sample_price_only",
                      price=_price_text(c, sample), qtys=qty_list(missing))]
    return [_flag("no_price", CONFIRM, f"The reply quotes no unit price for {qty_list(missing)} pieces.",
                  "There is nothing to compare with the listing yet.", "no_price", qtys=qty_list(missing))]


def check_product(c: dict) -> list:
    d, listing = c["d"], c["listing"]
    ours = f"{listing['title']} {d.get('description') or ''} {d.get('product') or ''} {d.get('sku') or ''} {listing['url']}"
    flat = re.sub(r"[^a-z0-9]", "", ours.lower())
    found = []
    for m in MODEL_TOKEN.finditer(c["text"]):
        token = m[0]
        if (re.search(r"\d", token) and not token.lower().startswith("cat") and re.sub(r"[^a-z0-9]", "", token.lower()) not in flat
                and _safe(token) and token not in found):
            found.append(token)
    our_cats = {g.lower() for g in CATEGORY.findall(ours)}
    if our_cats:
        found += [m[0] for m in CATEGORY.finditer(c["text"]) if m[1].lower() not in our_cats and _safe(m[0])]
    found = list(dict.fromkeys(found))[:3]
    if not found:
        return []
    tokens = _and(f'"{t}"' for t in found)
    return [_flag("product_differs", CONFIRM, f"The reply names {tokens}, which the listing title and our description do not.",
                  "The quote may be for a different model or specification.", "product_differs",
                  tokens=tokens, listing_ref=salutation(d)[1])]


def check_vague(c: dict) -> list:
    found = []
    for m in VAGUE.finditer(c["text"]):
        s = _snippet(c["text"], m)
        if s and s not in found:
            found.append(s)
    if not found:
        return []
    snippets = _and(f'"{s}"' for s in found[:3])
    return [_flag("vague_wording", CONFIRM, f"Unclear wording: {snippets}.", "Approximate or open terms cannot be held to.",
                  "vague_wording", snippets=snippets)]


def check_total(c: dict) -> list:
    if not c["usd"]:
        return []
    q, out, stated_for = c["q"], [], None
    t_qty, t_price = _whole_number(q.get("quote_total_quantity")), _amount(q.get("quote_total_price"))
    t_tier = quoted_at(q, t_qty) if t_qty else None
    if t_tier and t_price is not None and not _per_pack(t_tier[2]):
        stated_for, unit = t_qty, t_tier[1]
        calc = _cents(t_qty * unit)
        if calc != _cents(t_price):
            out.append(_flag("total_mismatch", BLOCKING, f"The stated total {usd(t_price)} for {t_qty:,} pieces is not {t_qty:,} x "
                             f"{usd(unit)} = {usd(calc)}.", "One of the two numbers is wrong, and we would pay one of them.",
                             "total_mismatch", price=usd(unit), total=usd(t_price), qty=f"{t_qty:,}", calc=usd(calc)))
    parts = [f"{usd(_cents(n * c['at'][n][1]))} for {n:,} pieces at {usd(c['at'][n][1])} per piece"
             for n in c["asked"] if _per_piece_at(c, n)]
    if parts and t_price is None:                                   # a stated total is checked above; none at all is asked for
        out.append(_flag("total_missing", CONFIRM, f"The reply states no total; the code calculates {_and(parts)}.",
                         "A total the seller confirms avoids a different amount at checkout.", "total_missing",
                         totals=_and(parts), which="this total" if len(parts) == 1 else "these totals"))
    return out


def check_unanswered(c: dict) -> list:
    missing = [what for what, fields, words in ANSWERS
               if not any(c["q"].get(f) not in (None, "") for f in fields) and not words.search(c["text"])]
    if not missing:
        return []
    return [_flag("not_answered", CONFIRM, f"The reply does not say {_and(missing)}.", "We asked; without it the quote is incomplete.",
                  "not_answered", missing=_and(missing))]


def check_for_us(c: dict) -> list:
    out = []
    if c["q"].get("needs_from_us"):
        out.append(_flag("seller_asks_us", NOTE, "The seller asked us for: " + "; ".join(map(str, c["q"]["needs_from_us"])) + ". Answer "
                         "this yourself; it is not in the follow-up.", "The seller may wait for our answer before quoting."))
    if c["q"].get("quote_warnings"):
        out.append(_flag("ai_value_dropped", NOTE, f"{len(c['q']['quote_warnings'])} value(s) the AI read could not be proven against "
                         "the reply and were left out: " + "; ".join(map(str, c["q"]["quote_warnings"])),
                         "Only numbers found in the seller's own words are used."))
    return out


CHECKS = (check_currency, check_price, check_moq, check_unit, check_charges, check_quantities, check_product, check_vague,
          check_total, check_unanswered, check_for_us)


def build_followup(d: dict, flags: list, seen: str):
    """The seller-facing questions, or None when there is nothing to ask. Every number in the subject and body must be one the
    code put there (a filled-in value, a list number, or inside our own title, URL, SKU or the seller's quoted words);
    otherwise FollowupRefused and nothing is drafted."""
    asks = [f for f in flags if f.get("question") and f["severity"] in (BLOCKING, CONFIRM)]
    if not asks:
        return None
    greeting, listing_ref = salutation(d)
    frame = {"product": d["description"] or d["product"], "sku": d["sku"], "greeting": greeting, "listing_ref": listing_ref,
             "sender_name": cfg("SENDER_NAME", "Swastik Raj")}
    questions = "\n".join(f"{i}. {f['question']}" for i, f in enumerate(asks, start=1))
    subject, body = FOLLOWUP_SUBJECT.format(**frame), FOLLOWUP_BODY.format(**frame, questions=questions)
    allowed = {Decimal(i) for i in range(1, len(asks) + 1)}.union(
        *(_numbers(v) for v in [*frame.values(), *(x for f in asks for x in f["_inserted"])]))
    stray = _numbers(f"{subject}\n{body}") - allowed
    if stray:
        raise FollowupRefused("The follow-up was not drafted: it held number(s) the code did not supply ("
                              + ", ".join(sorted(format(n, "f") for n in stray)) + ").")
    for_us = [f["text"] for f in flags if f["code"] == "seller_asks_us"]
    for_us.append(f"The listing price was seen by the search on {seen}. Open the listing to check it has not changed before you "
                  "send this.")
    return {"subject": subject, "body": body, "for_us": for_us, "at": now(), "sent": None}


def check_quote(entry: dict, d: dict, listing: dict, asked: list, asked_from: str) -> tuple:
    """(comparison, follow-up or None) for a recorded reply. Reads only the verified quote stored on the entry (never the AI's
    raw output), the reply's own words and the results file. Flags inform; they never decide and never send anything."""
    when = listing["captured"]
    comparison = {
        "at": now(), "replies": len(entry.get("replies") or []) or 1, "results_file": listing["results_file"],
        "listing": {"url": listing["url"], "title": listing["title"], "listed": listing["listed"],
                    "unit_price": format(listing["unit_price"], "f") if listing["unit_price"] is not None else None,
                    "confidence": listing["confidence"], "moq": listing["moq"],
                    "tiers": [{"min": lo, "max": hi, "unit_price": format(p, "f")} for lo, hi, p in listing["tiers"]],
                    "range": [format(x, "f") for x in listing["range"]] if listing["range"] else None,
                    "captured": f"{when:%Y-%m-%d %H:%M}", "captured_from": listing["captured_from"]},
        "asked_qty": list(asked), "asked_from": asked_from,
        "tolerance": {"pct": format((TOLERANCE_PCT * 100).normalize(), "f"), "min_usd": format(TOLERANCE_MIN, "f")},
        "not_checked": None, "flags": [], "counts": {k: 0 for k in COUNT_KEYS.values()}}
    why = NOT_CHECKED.get(entry.get("status"))
    if why:
        comparison["not_checked"] = f"Not compared: {why}."
        return comparison, None
    q = entry
    c = {"q": q, "d": d, "listing": listing, "asked": list(asked), "text": _their_words(entry.get("reply") or "", d),
         "at": {n: quoted_at(q, n) for n in asked}, "usd": q.get("currency") == "USD",
         "priced": bool(quoted_tiers(q)) or any(_amount(q.get(k)) is not None for k in ("sample_unit_price", "sample_total_price", "quote_total_price")),
         "date": f"{when.day} {when:%B %Y}", "seen": f"{when:%Y-%m-%d %H:%M}"}
    flags = sorted((f for check in CHECKS for f in check(c)), key=lambda f: (BLOCKING, CONFIRM, NOTE).index(f["severity"]))
    try:
        followup = build_followup(d, flags, c["seen"])
    except FollowupRefused as e:
        followup = None
        flags.append(_flag("followup_not_drafted", NOTE, str(e), "Only numbers the code supplied may reach a seller."))
    for f in flags:
        f.pop("_inserted")
    comparison["flags"] = flags
    comparison["counts"] = {key: sum(f["severity"] == sev for f in flags) for sev, key in COUNT_KEYS.items()}
    return comparison, followup


def compare_entry(state: dict, d: dict, results_path: str) -> dict:
    """Checks the reply recorded for this draft's listing against the results file, and stores the comparison and the
    follow-up on its entry. A follow-up already marked sent is kept in followup_history when a different one replaces it."""
    entry = state.get(d["key"])
    if not isinstance(entry, dict) or not entry.get("reply"):
        raise NothingToCompare(f"{d['sku']}: no reply is recorded for this seller, so there is nothing to compare. Nothing stored.")
    recorded = [n for n in map(_whole_number, entry.get("asked_qty") or []) if n]
    asked_from = ("recorded when this seller was contacted" if recorded
                  else "not recorded when this seller was contacted, so the tool's current quantities")
    comparison, followup = check_quote(entry, d, listing_terms(d, results_path), recorded or list(QUOTE_QTYS), asked_from)
    old = entry.get("followup")
    if isinstance(old, dict) and old.get("sent"):
        if followup and (followup["subject"], followup["body"]) == (old.get("subject"), old.get("body")):
            followup = old                                  # the same questions, already sent: they stay sent
        else:
            entry.setdefault("followup_history", []).append(old)
    entry_for(state, d).update(comparison=comparison, followup=followup)
    return comparison


def check_summary(entry: dict) -> str:
    """One line for the report: the counts and when the listing was seen."""
    comp = entry.get("comparison") if isinstance(entry, dict) else None
    if not isinstance(comp, dict):
        return "not checked yet" if isinstance(entry, dict) and entry.get("reply") else ""
    if comp.get("not_checked"):
        return comp["not_checked"]
    n = comp.get("counts") or {}
    return (f"{n.get('blocking', 0)} blocking, {n.get('needs_confirmation', 0)} to confirm, {n.get('note', 0)} note(s) "
            f"(listing seen {comp.get('listing', {}).get('captured', '?')})")


def check_lines(entry: dict) -> list:
    """The stored comparison and follow-up, for the terminal."""
    comp, fu = entry.get("comparison"), entry.get("followup")
    if not isinstance(comp, dict):
        return []
    listing = comp["listing"]
    out = [f"Quote check against the listing as seen by the search on {listing['captured']} (from {listing['captured_from']}; "
           "the listing page was not re-read):"]
    if comp.get("not_checked"):
        return out + [f"  {comp['not_checked']}"]
    out.append(f"  {check_summary(entry)}")
    out += [f"  {f['severity'].upper()}: {f['text']}" for f in comp["flags"]]
    if not fu:
        return out + ["No follow-up: nothing to ask the seller."]
    out += ["", f"FOLLOW-UP DRAFT - not sent. Copy it into the seller's website form, then run: mark {entry.get('sku', '<SKU>')} "
            "--results <file> --followup. Or send it yourself.", f"Subject: {fu['subject']}", "", fu["body"].rstrip(), ""]
    return out + [f"For you, not for the seller: {x}" for x in fu.get("for_us") or []]


# ---------- report ----------

def contact_method(d: dict, entry: dict) -> str:
    if d["email"]:
        return "email sent" if "sent" in entry else "email drafted, not sent"
    return "inquiry form submitted" if "sent" in entry else "inquiry form not yet submitted"


def build_report(drafts: list, state: dict, skipped: list = ()) -> list:
    rows = []
    for d in drafts:
        s = state.get(d["key"], {})
        channel = "email" if d["email"] else "inquiry form"
        sent = f"yes ({channel}, {s['sent']})" if "sent" in s else f"no ({channel})"
        rows.append([d["sku"], d["product"], d["manufacturer"], d["email"] or d["url"], contact_method(d, s), sent,
                     "yes" if "reply" in s else "no", s.get("status", ""), s.get("summary", ""), check_summary(s)])
    for s in skipped:  # rows with no recommended seller, an error, or a reviewer exclusion: nothing was sent
        rows.append([s["sku"], s["product"], s["manufacturer"], "", "not contacted", "no", "no", "", s["reason"], ""])
    return rows


def orphaned(state: dict, drafts: list) -> list:
    """State records whose SKU + listing is no longer a recommended seller (the pick changed on a re-run)."""
    current = {d["key"] for d in drafts}
    return [f"{v.get('sku', k.split('|')[0])} ({v.get('seller', k)})" for k, v in state.items()
            if not k.startswith("_") and k not in current and isinstance(v, dict)]


REPORT_HEADER = ["SKU", "Product", "Manufacturer", "Contact", "Contact method", "Sent", "Reply received",
                 "Reply status", "Summary", "Quote check"]


def write_report(rows: list, stamp: str, notes: list = ()) -> tuple:
    os.makedirs(os.path.join(HERE, "Reports"), exist_ok=True)
    md_path = os.path.join(HERE, "Reports", f"email_report_{stamp}.md")
    xlsx_path = os.path.join(HERE, "Reports", f"email_report_{stamp}.xlsx")
    cell = lambda v: str(v).replace("|", "/").replace("\n", " ")
    contacted = sum(r[4] in ("email sent", "inquiry form submitted") for r in rows)
    lines = ["# Email outreach report", "", f"Generated {now()} | {len(rows)} product(s) | {contacted} contacted | "
             f"{sum(r[6] == 'yes' for r in rows)} replies", "", *[f"- {n}" for n in notes], *([""] if notes else []),
             "| " + " | ".join(REPORT_HEADER) + " |", "|" + "---|" * len(REPORT_HEADER)]
    lines += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows]
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Email Outreach"
    ws.append(REPORT_HEADER)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    wb.save(xlsx_path)
    return md_path, xlsx_path


# ---------- CLI ----------

def find_draft(drafts: list, skipped: list, sku: str, results: str) -> dict:
    """The draft for `sku` in this results file, or a clear error (nothing is stored on error)."""
    for d in drafts:
        if d["sku"].upper() == sku.strip().upper():
            return d
    for s in skipped:
        if s["sku"].upper() == sku.strip().upper():
            sys.exit(f"{sku}: row found in {results} but it has no recommended seller to attach this to "
                     f"({s['reason']}). Nothing stored.")
    known = ", ".join(d["sku"] for d in drafts) or "none"
    sys.exit(f"{sku}: not a SKU in {results}. SKUs with a recommended seller: {known}. Nothing stored.")


def busy_exit(e) -> int:
    """Exit code 3 = the state file is busy (another command holds its lock). The message goes to stderr."""
    print(f"email_agent: {e}", file=sys.stderr)
    return 3


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # piped output defaults to cp1252 on Windows
    if hasattr(sys.stdin, "reconfigure"):
        sys.stdin.reconfigure(encoding="utf-8")                 # so does piped/pasted input: Chinese replies broke
    load_dotenv(os.path.join(HERE, ".env"))
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    for name in ("drafts", "send", "report"):
        s = sub.add_parser(name)
        s.add_argument("results", help="sourcing_results_*.xlsx")
        if name == "drafts":
            s.add_argument("--near-miss", action="store_true",
                           help="also show drafts for the near-miss candidates listed in near_misses.csv (not recommended by the "
                                "search agent; each is labelled NEEDS HUMAN CHECK). Read only; send and report never include them.")
        if name == "send":
            s.add_argument("--yes", action="store_true", help="one confirmation for the whole batch")
            s.add_argument("--test-recipient", default="", metavar="ADDRESS",
                           help="send every email ONLY to this address, subject prefixed [TEST], nothing marked as sent")
            s.add_argument("--confirm-count", type=int, metavar="N",
                           help="send without any prompt, but only if exactly N emails would be sent (after every check); "
                                "otherwise nothing is sent")
    for name in ("mark", "reply", "compare"):
        s = sub.add_parser(name)
        s.add_argument("sku")
        s.add_argument("--results", required=True, help="sourcing_results_*.xlsx the SKU must be in")
        if name == "reply":
            s.add_argument("--file", help="UTF-8 text file with the reply (default: read stdin)")
        if name == "mark":
            s.add_argument("--followup", action="store_true",
                           help="record that the drafted follow-up questions were sent by hand (website form or your own email)")
    args = p.parse_args()
    if getattr(args, "confirm_count", None) is not None and args.confirm_count < 0:
        p.error("--confirm-count must be 0 or more")
    try:
        with obs.trace_command("email_agent", args.cmd, workflow=obs.workflow_id(args.results), sku=getattr(args, "sku", None)):
            run(args)
    except statefile.StateBusy as e:
        sys.exit(busy_exit(e))


def run(args):
    exclusions = load_exclusions()
    results = args.results
    drafts, skipped = load_results(results, exclusions)
    drafts = [make_draft(c) for c in drafts]
    if args.cmd in ("drafts", "send", "report"):
        print_warnings(warnings_for(bool(exclusions)))
        print(skip_summary(skipped))
        obs.annotate(drafts=len(drafts), drafts_with_email=sum(bool(d["email"]) for d in drafts), skipped=len(skipped),
                     skipped_by_kind=collections.Counter(s["kind"] for s in skipped))

    if args.cmd == "drafts":  # read-only: never loads or writes email_state.json
        for d in drafts:
            print(format_draft(d))
        print(f"{len(drafts)} draft(s); {sum(bool(d['email']) for d in drafts)} with an email address. Nothing sent.")
        if args.near_miss:
            near_rows = load_near_misses()
            if near_rows is None:
                print_warnings([f"{os.path.basename(NEAR_MISSES_CSV)} is missing - no near-miss drafts were added."])
            else:
                extra, notes = near_miss_drafts(results, drafts, exclusions, near_rows)
                print(f"\n{'#' * 70}\nNEAR-MISS DRAFTS: the search agent did NOT recommend these. Each needs a human check "
                      f"before anything is sent or submitted.\n{'#' * 70}")
                for d in extra:
                    print(format_draft(d))
                for note in notes:
                    print(f"NOTE: {note}")
                print(f"{len(extra)} near-miss draft(s), {sum(bool(d['email']) for d in extra)} with an email address. Nothing sent.")
    elif args.cmd == "send":
        missing = [k for k in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD") if not cfg(k)]
        if missing:
            obs.annotate(blocked_sends=len(drafts))
            sys.exit(f"Set {', '.join(missing)} in .env first (nothing sent).")
        if args.test_recipient:
            real = {d["email"].lower() for d in drafts if d["email"]}
            if "@" not in args.test_recipient or args.test_recipient.lower() in real:
                sys.exit("--test-recipient must be a test address, not a manufacturer's.")
        extra = {} if args.confirm_count is None else {"confirm_count": args.confirm_count}
        try:
            with statefile.locked(STATE_PATH):
                sent = confirm_and_send(drafts, load_state(), args.yes, test_recipient=args.test_recipient, **extra)
            print(f"Sent {sent} email(s).")
            obs.annotate(sent=sent, test_mode=bool(args.test_recipient))
        except SendRefused as e:
            obs.annotate(blocked_sends=len(drafts))
            sys.exit(f"Refusing to send: {e}")
        except EOFError:     # no terminal to ask on (closed stdin): nothing was sent
            sys.exit("Nothing sent: there is no terminal to confirm on. Run it in a terminal, or pass --confirm-count N "
                     "with the number of emails you expect to send.")
    elif args.cmd == "mark":
        d = find_draft(drafts, skipped, args.sku, results)
        obs.annotate(sku=d["sku"], listing_host=obs.url_host(d["url"]))
        with statefile.locked(STATE_PATH):
            state = load_state()
            if args.followup:
                followup = (state.get(d["key"]) or {}).get("followup")
                if not isinstance(followup, dict):
                    sys.exit(f"{d['sku']}: no follow-up has been drafted for this seller. Nothing stored.")
                followup["sent"] = now()
            else:
                entry = entry_for(state, d)
                entry["sent"], entry["asked_qty"] = now(), list(QUOTE_QTYS)       # the quantities this inquiry named
            save_state(state)
        print(f"{d['sku']} ({d['manufacturer'] or d['url']}): " + ("follow-up questions marked as sent by hand." if args.followup
                                                                   else "marked as submitted via the inquiry form."))
    elif args.cmd in ("reply", "compare"):
        d = find_draft(drafts, skipped, args.sku, results)
        obs.annotate(sku=d["sku"], listing_host=obs.url_host(d["url"]))
        text = (open(args.file, encoding="utf-8").read() if args.file else sys.stdin.read()) if args.cmd == "reply" else None
        with statefile.locked(STATE_PATH):
            state = load_state()
            result = record_reply(state, d, text) if args.cmd == "reply" else None
            try:
                compare_entry(state, d, results)                 # code only: no AI call, nothing is sent
            except NothingToCompare as e:
                sys.exit(str(e))
            save_state(state)
        if result:
            print(f"{d['sku']} ({d['manufacturer'] or d['url']}): {result['status']}\n{result['summary']}")
            print("\n".join(quote_lines(result)))
        print("\n".join(check_lines(state[d["key"]])))
    elif args.cmd == "report":
        state = load_state()
        lost = orphaned(state, drafts)
        notes = [skip_summary(skipped), *warnings_for(bool(exclusions))]
        if lost:
            notes.append(f"{len(lost)} earlier record(s) belong to a seller that is no longer the recommended one "
                         f"for that SKU and are not shown: {', '.join(lost)}.")
        md, xlsx = write_report(build_report(drafts, state, skipped), datetime.now().strftime("%Y%m%d_%H%M%S"), notes)
        print(f"Wrote {md}\nWrote {xlsx}")
        if lost:
            print(notes[-1])


if __name__ == "__main__":
    main()
