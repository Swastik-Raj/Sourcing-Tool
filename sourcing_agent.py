"""
Competitive sourcing research POC.

For each product, connects directly to the Nimble MCP server (as an MCP client - not
Anthropic's server-side MCP connector) and exposes ONLY its search and extract tools to
Claude. Claude is given a hard budget of tool calls per product, must return a structured
(schema-enforced) result with up to 3 ranked candidates, and runs on Haiku for cost.

Setup:
    pip install "anthropic[mcp]" pydantic openpyxl python-dotenv
    .env file in this folder with:
        ANTHROPIC_API_KEY=...   (your Anthropic API key)
        NIMBLE_API_KEY=...      (your Nimble API key, used as the MCP bearer token)

Run:
    python sourcing_agent.py                                    (5 built-in sample products)
    python sourcing_agent.py --input products.xlsx               (SKU/Product/Keyword columns)
    python sourcing_agent.py --input products.xlsx --output results.xlsx
"""

import argparse
import asyncio
import json
import os
import re
import sys
from datetime import datetime
from typing import Literal, Optional

import httpx2
import openpyxl
from dotenv import load_dotenv
from openpyxl.styles import Font, PatternFill
from mcp import ClientSession
from mcp.client.streamable_http import streamable_http_client
from pydantic import BaseModel, ValidationError

from anthropic import AsyncAnthropic
from anthropic.lib.tools import beta_async_tool

MODEL = "claude-haiku-4-5"
# Haiku 4.5 list pricing, used only to print an estimate - not billed exactly.
HAIKU_INPUT_PER_MTOK = 1.00
HAIKU_OUTPUT_PER_MTOK = 5.00

NIMBLE_MCP_URL = "https://mcp.nimbleway.com/mcp"

# 1688.com is deliberately excluded: its search results sit behind a Taobao/Alibaba
# login wall regardless of driver tier (confirmed with both vx8 and vx10 - neither
# reaches real content), so it's not reachable via extract at all.
SOURCING_SITES = ["aliexpress.com", "alibaba.com", "made-in-china.com"]

# Nimble's general web search barely indexes individual listings on these sites, and its
# "shopping" focus mode is hard-restricted to a fixed roster of Western retailers (Amazon,
# Best Buy, etc. - confirmed via the API rejecting any other subagent name) that doesn't
# cover any of these. Each site's own on-site search page, fetched directly with a
# JS-rendering driver, reliably returns real listings - so that's the primary discovery
# method instead. {query} is filled in per-search with URL-encoded search terms.
SITE_SEARCH_URLS = {
    "aliexpress.com": "https://www.aliexpress.com/wholesale?SearchText={query}",
    "alibaba.com": "https://www.alibaba.com/trade/search?SearchText={query}",
    "made-in-china.com": "https://www.made-in-china.com/products-search/hot-china-products/{query}.html",
}

MAX_TOOL_CALLS_PER_PRODUCT = 8  # searches + extracts combined, hard cap
# ^ first guess for finding/verifying 3 distinct candidates instead of 1 - tune based on
# actual cost/results once tested on a real batch.
MAX_ITERATIONS_PER_PRODUCT = MAX_TOOL_CALLS_PER_PRODUCT + 3  # backstop against loops
MAX_TOOL_RESULT_CHARS = 2500  # truncate every tool result before it goes back to Claude
MAX_TOKENS_PER_TURN = 3000

# Sourcing only makes sense with real margin under L-Com's price, and only for candidates
# that aren't rubric-rejected. Both are business knobs - tune with the sourcing team.
MIN_MARGIN_PCT = 80  # 30% got eaten by shipping, storage and import taxes
# A unit price this many times cheaper than the next-cheapest candidate for the same
# product is treated as an extraction error (e.g. price divided by MOQ), not a bargain.
IMPLAUSIBLE_PRICE_RATIO = 10
IMPLAUSIBLE_NOTE = (
    "unit price implausible relative to other candidates for this product - "
    "likely extraction error, verify manually before ordering"
)
MIN_RECOMMEND_ACCURACY = 80

PRODUCTS = [
    {"sku": "ECF504-SC6", "description": "Cat6 RJ45 Coupler Shielded (8x8) Panel Mount Style"},
    {"sku": "TDG1026KS-C6", "description": "Cat6 Coupler - Shielded RJ45 (8x8) Keystone Feed-thru"},
    {"sku": "HDFF", "description": "HDMI Panel Mount Adapter, Female to Female"},
    {"sku": "CAPUSB-A", "description": "USB Protective Cover for Type A Jacks, ABS material"},
    {"sku": "VIC00001", "description": "DVI 24+5 female to female coupler, bulkhead mount"},
]

ATTRIBUTES = ["product_type", "category_spec", "shielding_material", "mount_form", "gender_pins"]
ATTRIBUTE_LABELS = {
    "product_type": "Product type",
    "category_spec": "Category/spec",
    "shielding_material": "Shielding/material",
    "mount_form": "Mount/form factor",
    "gender_pins": "Gender/pins",
}

RUBRIC_INSTRUCTIONS = """\
Score the candidate with this exact rubric (the same method used in our manual sourcing
audit). Five attributes, weighted 20 points each, for a 100-point total:

1. product_type (0-20) - is the listed item fundamentally the same kind of product?
2. category_spec (0-20) - category/performance rating match (e.g. Cat5e vs Cat6 vs Cat6a,
   or the equivalent spec/performance class for this product type).
3. shielding_material (0-20) - shielding and/or material match.
4. mount_form (0-20) - mount/form factor match (e.g. panel mount, keystone feed-thru,
   bulkhead mount).
5. gender_pins (0-20) - gender/pin configuration match (e.g. female-to-female), or the
   equivalent connector-orientation attribute for non-connector products. If a product
   type genuinely has no equivalent attribute, score it 0 - nothing was confirmed.

CRITICAL RULE: award 0 points for any attribute the candidate listing's own text does not
explicitly state. Do not infer or assume a match from category context, product photos, or
"typical" industry defaults - only credit what the listing text actually says. No benefit
of the doubt. match_percent must equal the sum of the five scores.
"""


class AttributeBreakdown(BaseModel):
    product_type: int
    category_spec: int
    shielding_material: int
    mount_form: int
    gender_pins: int


class Candidate(BaseModel):
    manufacturer: Optional[str] = None
    distributor_site: Optional[str] = None
    listing_title: Optional[str] = None
    price: Optional[str] = None  # raw listing price string, display only
    # Required and non-nullable on purpose: every nullable field is an anyOf union, and
    # strict output mode rejects the schema ("Schema is too complex") past ~11 of them.
    # 0 / "" mean "unknown".
    price_total: float  # numeric USD amount that `price` covers
    quantity_covered: int  # how many units price_total buys
    unit_price_confidence: Literal["stated", "inferred", "ambiguous"]
    unit_price_note: str
    moq: Optional[str] = None
    email: Optional[str] = None
    match_percent: Optional[int] = None
    attribute_breakdown: Optional[AttributeBreakdown] = None
    url: Optional[str] = None


class SourcingResult(BaseModel):
    product: str
    no_match: bool
    no_match_reason: Optional[str] = None
    candidates: list[Candidate] = []
    note: Optional[str] = None


def build_prompt(product: dict) -> str:
    sites = ", ".join(SOURCING_SITES)
    search_urls = "\n".join(f"  - {site}: {url}" for site, url in SITE_SEARCH_URLS.items())
    return f"""\
You are a competitive sourcing research analyst. Find this product for sale on Chinese
sourcing/manufacturing sites: {sites}.

Target product:
  SKU: {product['sku']}
  Description: "{product['description']}"

The SKU above is our own internal part number - it is not a manufacturer part number and
will never appear on a competitor's listing. You are looking for a DIFFERENT manufacturer's
equivalent product, judged only by the five rubric attributes below, never by whether the
listing happens to mention this SKU or any particular brand/part number.

The search tool's general web mode barely indexes individual listings on these sites, and
its "shopping" focus mode only covers a fixed set of Western retailers (Amazon, Best Buy,
etc.) - neither will find real listings here. Instead, use each site's own on-site search
page directly as your primary discovery method:
{search_urls}
Build the URL by substituting {{query}} with your search terms (spaces as "+" for the
aliexpress/alibaba forms, "_" for the made-in-china path form), then call extract on it
with driver="vx8" (these are JS-heavy pages - the default driver will return nothing
useful).

Score primarily from the listing titles/prices on that search page itself - they are
usually specific enough to score against the rubric (e.g. a title stating "Female to
Female Panel Mount" directly confirms two attributes). Individual product detail pages on
these sites are unreliable to extract directly (heavier anti-bot throttling than search
pages) and often come back empty even after a wait - treat one extract attempt on a
promising listing URL as an optional bonus for more detail, never as required, and never
let an empty/incomplete result there block you from scoring off the search listing text
you already have.

You have a HARD BUDGET of {MAX_TOOL_CALLS_PER_PRODUCT} tool calls total for this product -
more than before because you're now finding and verifying up to 3 candidates instead of 1.
Spend it deliberately: a handful of site-search extracts across different sites/queries to
surface multiple distinct candidates, plus the occasional individual listing page only if a
title is genuinely ambiguous on a rubric attribute. Once you've used your budget, or you
already have enough to score confidently, STOP calling tools immediately and give your
final answer. Do not narrate your process in free text at any point - every response you
give must be the final structured result, nothing else.

{RUBRIC_INSTRUCTIONS}

PRICING - for every candidate, work out what quantity the listed price actually buys:
- price: the raw price text exactly as shown on the listing.
- price_total: that price as a number in USD. If a range is shown, use the HIGHER end.
  0 if no price is shown.
- quantity_covered: how many units price_total buys (e.g. "$0.33 / piece" -> 1;
  "$19.75 ... (10PCS)" or "pack of 10" -> 10). 0 if unknown. The MOQ / minimum order is
  NOT quantity_covered: "$1.27-1.58, Min. order 500 pieces" is a per-piece price (-> 1).
- unit_price_note: one short line on how you read the quantity ("" if obvious).
- unit_price_confidence:
  "stated"    - the listing explicitly says per piece/unit, or explicitly states the pack size.
  "inferred"  - not explicit, but the listing text makes it clear (e.g. title says "10PCS").
  "ambiguous" - you cannot tell whether the price is per unit or for a pack/MOQ batch
                (e.g. "$19.75, MOQ 10 pcs" with nothing saying which), or the price is not
                in USD. DO NOT GUESS - mark it ambiguous and say why in unit_price_note.
Never divide the price yourself; just report price_total and quantity_covered.

Return up to 3 candidates in `candidates`, ranked best first, balancing match quality
against price (a slightly lower match % at a much lower price can rank above a perfect
match at a high price) - report both numbers per candidate, never silently substitute one
for the other. Each candidate must be a DIFFERENT manufacturer/supplier - never list 2-3
listings from the same company just to fill the slots. If you only find 1 or 2 genuinely
distinct, plausible candidates, return only that many - do not force a weak extra pick just
to reach 3.

If you find no plausible candidate at all, or every candidate you found scores very low,
set no_match to true, leave candidates empty, and briefly say why in no_match_reason - do
not guess or force a weak candidate into looking like a real match.
"""


def truncate_text(text: str, limit: int = MAX_TOOL_RESULT_CHARS) -> str:
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)  # strip markdown images
    text = re.sub(r"(\]\(https?://[^)?\s]+)\?[^)]*\)", r"\1)", text)  # strip link tracking params
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > limit:
        return text[:limit] + " ...[truncated]"
    return text


class ToolBudget:
    def __init__(self, limit: int):
        self.limit = limit
        self.used = 0

    def take(self) -> bool:
        if self.used >= self.limit:
            return False
        self.used += 1
        return True


def pick_tool(tools, exact_names, contains, excludes=("agent", "template", "crawl", "map")):
    for t in tools:
        if t.name.lower() in exact_names:
            return t
    for t in tools:
        name = t.name.lower()
        if contains in name and not any(x in name for x in excludes):
            return t
    return None


def make_bounded_tool(tool, session: ClientSession, budget: ToolBudget):
    tool_name = tool.name
    input_schema = getattr(tool, "input_schema", None) or getattr(tool, "inputSchema", None)
    description = (tool.description or "")[:600]

    async def call(**kwargs):
        if not budget.take():
            return (
                f"Tool call budget exhausted ({budget.limit} calls used for this product). "
                "Stop searching and return your best final answer now."
            )
        print(f"    -> {tool_name}({kwargs})")
        result = await session.call_tool(name=tool_name, arguments=kwargs)
        text = " ".join(
            block.text for block in result.content if getattr(block, "type", None) == "text"
        )
        print(f"    <- {truncate_text(text, limit=300)}")
        return truncate_text(text)

    return beta_async_tool(call, name=tool_name, description=description, input_schema=input_schema)


def strict_json_schema(model: type[BaseModel]) -> dict:
    """model_json_schema() doesn't set additionalProperties, but Claude's strict
    output mode requires it explicitly false on every object (including $defs)."""
    schema = model.model_json_schema()

    def lock_down(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                node.setdefault("additionalProperties", False)
            for value in node.values():
                lock_down(value)
        elif isinstance(node, list):
            for value in node:
                lock_down(value)

    lock_down(schema)
    return schema


def usage_tokens(message) -> tuple:
    usage = getattr(message, "usage", None)
    if usage is None:
        return 0, 0
    return usage.input_tokens or 0, usage.output_tokens or 0


async def research_once(client: AsyncAnthropic, session: ClientSession, tools, product: dict):
    budget = ToolBudget(MAX_TOOL_CALLS_PER_PRODUCT)
    runner_tools = [make_bounded_tool(t, session, budget) for t in tools]

    runner = client.beta.messages.tool_runner(
        model=MODEL,
        max_tokens=MAX_TOKENS_PER_TURN,
        max_iterations=MAX_ITERATIONS_PER_PRODUCT,
        tools=runner_tools,
        output_config={
            "format": {"type": "json_schema", "schema": strict_json_schema(SourcingResult)}
        },
        messages=[{"role": "user", "content": build_prompt(product)}],
    )

    total_in, total_out = 0, 0
    last_message = None
    async for message in runner:
        last_message = message
        i, o = usage_tokens(message)
        total_in += i
        total_out += o

    parsed = None
    if last_message is not None:
        text = " ".join(block.text for block in last_message.content if block.type == "text")
        try:
            parsed = SourcingResult.model_validate_json(text)
        except ValidationError:
            parsed = None
    return parsed, total_in, total_out, budget.used


def is_garbled(result: Optional[SourcingResult]) -> bool:
    """Haiku occasionally emits a whole response with spaces sprayed mid-word
    ("al ib ab a.c om", "$2 .1 0"). URLs and domains never contain whitespace, so any
    in those fields means the response text can't be trusted."""
    return result is not None and any(
        re.search(r"\s", field or "")
        for c in result.candidates
        for field in (c.url, c.distributor_site)
    )


async def research_product(client: AsyncAnthropic, session: ClientSession, tools, product: dict):
    """research_once, retried once if the response comes back garbled. A second garbled
    response is dropped (reported as no structured result) rather than written out."""
    total_in = total_out = calls = 0
    for attempt in (1, 2):
        parsed, tin, tout, used = await research_once(client, session, tools, product)
        total_in, total_out, calls = total_in + tin, total_out + tout, calls + used
        if not is_garbled(parsed):
            return parsed, total_in, total_out, calls
        print(f"  Response text came back garbled (attempt {attempt}) - "
              + ("retrying." if attempt == 1 else "dropping it."))
    return None, total_in, total_out, calls


def candidate_score(candidate: Candidate) -> tuple:
    """Returns (computed_match_percent, {attr: score}) - recomputes from the breakdown
    when present so a miscounted match_percent from the model can't slip through."""
    breakdown = candidate.attribute_breakdown
    if breakdown is not None:
        scores = breakdown.model_dump()
        return sum(max(0, min(20, v)) for v in scores.values()), scores
    return candidate.match_percent, {}


def match_tier(match_percent: Optional[int]) -> str:
    if match_percent is None:
        return ""
    if match_percent >= 90:
        return "Auto-accepted"
    if match_percent >= 80:
        return "Flagged for manual review"
    return "Rejected"


def candidate_comment(scores: dict) -> str:
    if not scores:
        return ""
    zero_attrs = [ATTRIBUTE_LABELS[attr] for attr in ATTRIBUTES if not scores.get(attr)]
    if zero_attrs:
        return f"0 pts on: {', '.join(zero_attrs)} (not stated on listing)"
    return "All five rubric attributes matched something explicit on the listing."


def unit_price(candidate: Candidate) -> Optional[float]:
    """True single-unit cost, or None when it can't be determined confidently."""
    c = candidate
    if c.unit_price_confidence not in ("stated", "inferred") or not c.price_total or not c.quantity_covered:
        return None
    return c.price_total / c.quantity_covered


def price_vs_lcom(unit: Optional[float], lcom_price: Optional[float]) -> Optional[tuple]:
    """(dollars cheaper per unit, percent cheaper) vs L-Com. Negative = more expensive."""
    if unit is None or not lcom_price:
        return None
    diff = lcom_price - unit
    return diff, diff / lcom_price * 100


def format_vs_lcom(vs: Optional[tuple]) -> str:
    if vs is None:
        return "n/a"
    diff, pct = vs
    return f"${abs(diff):.2f} ({abs(pct):.0f}%) {'cheaper' if diff >= 0 else 'more expensive'}"


def format_unit(unit: Optional[float]) -> str:
    if unit is None:
        return "UNDETERMINED"
    return f"${unit:.2f}" if unit >= 0.1 else f"${unit:.4f}"  # don't round a $0.0032 bug to $0.00


def implausible_units(result: Optional[SourcingResult]) -> set:
    """Indices of candidates whose unit price is > IMPLAUSIBLE_PRICE_RATIO x cheaper than
    the next-cheapest candidate for the same product."""
    units = [unit_price(c) for c in (result.candidates if result else [])]
    suspect = set()
    for i, u in enumerate(units):
        others = [x for j, x in enumerate(units) if j != i and x is not None]
        if u is not None and others and u * IMPLAUSIBLE_PRICE_RATIO < min(others):
            suspect.add(i)
    return suspect


def checked_unit(result: SourcingResult, i: int) -> Optional[float]:
    """unit_price() for comparisons against L-Com: None when implausible."""
    return None if i in implausible_units(result) else unit_price(result.candidates[i])


def confidence_label(result: SourcingResult, i: int) -> str:
    c = result.candidates[i]
    if i in implausible_units(result):
        return f"{c.unit_price_confidence} (IMPLAUSIBLE)"
    return (c.unit_price_confidence or "?") + (" (FLAGGED)" if unit_price(c) is None else "")


def recommend(result: Optional[SourcingResult], lcom_price: Optional[float]) -> tuple:
    """Returns (index of recommended candidate or None, one-line reason)."""
    if result is None or result.no_match or not result.candidates:
        return None, "No candidates found - nothing to source."
    if not lcom_price:
        return None, "No L-Com reference price for this product - cannot judge margin, not recommending."
    qualifying = []  # (index, accuracy, margin %, price is stated)
    for i, c in enumerate(result.candidates):
        accuracy, _ = candidate_score(c)
        vs = price_vs_lcom(checked_unit(result, i), lcom_price)
        if vs is None or accuracy is None or accuracy < MIN_RECOMMEND_ACCURACY or vs[1] < MIN_MARGIN_PCT:
            continue
        qualifying.append((i, accuracy, vs[1], c.unit_price_confidence == "stated"))
    if not qualifying:
        suspect = implausible_units(result)
        return None, (
            f"No candidate offers sufficient margin below L-Com's ${lcom_price:.2f} "
            f"(needs a confirmed unit price >= {MIN_MARGIN_PCT}% cheaper at >= {MIN_RECOMMEND_ACCURACY}% "
            "accuracy) - do not source from any of these."
            + "".join(f" Candidate {i + 1}: {IMPLAUSIBLE_NOTE}." for i in sorted(suspect))
        )
    # ponytail: accuracy x margin% is a crude trade-off; swap for a landed-cost model (shipping, duty) if needed
    best = max(qualifying, key=lambda q: q[1] * q[2])
    # A stated price beats an inferred one at equal or better accuracy.
    stated = [q for q in qualifying if q[3] and q[1] >= best[1]]
    if not best[3] and stated:
        best = max(stated, key=lambda q: q[1] * q[2])
    best = best[0]
    c = result.candidates[best]
    accuracy, _ = candidate_score(c)
    unit = unit_price(c)
    return best, (
        f"Buy from Candidate {best + 1} ({c.manufacturer or 'unnamed supplier'}): {accuracy}% accuracy at "
        f"{format_unit(unit)}/unit, {format_vs_lcom(price_vs_lcom(unit, lcom_price))} than L-Com."
    )


def print_product_result(product: dict, result: Optional[SourcingResult], calls_used: int):
    print("-" * 72)
    print(f"{product['sku']} - \"{product['description']}\"  ({calls_used} tool calls used)")
    print()

    if result is None:
        print("  Could not get a structured result for this product (empty, refused or garbled response).")
        return

    if result.no_match or not result.candidates:
        reason = result.no_match_reason or "no plausible candidate found"
        print(f"  No match found - {reason}")
        return

    for i, candidate in enumerate(result.candidates, start=1):
        computed, scores = candidate_score(candidate)
        vs = price_vs_lcom(checked_unit(result, i - 1), product.get("lcom_price"))
        manufacturer = candidate.manufacturer or "Not stated on listing"
        source = candidate.distributor_site or "Unknown site"
        price = candidate.price or "Not stated"

        print(f"  #{i}  {manufacturer}  |  via {source}  |  Tier: {match_tier(computed)}")
        print(f"      Listing:  {candidate.listing_title or '?'}")
        print(f"      Price:    {price}   MOQ: {candidate.moq or 'Not stated'}")
        print(f"      Unit:     {format_unit(unit_price(candidate))} ({confidence_label(result, i - 1)})  "
              f"vs L-Com: {format_vs_lcom(vs)}")
        print(f"      Accuracy: {computed}%")
        print(f"      URL:      {candidate.url or 'Not stated'}")
        if scores:
            line = " | ".join(
                f"{ATTRIBUTE_LABELS[attr]} {max(0, min(20, scores.get(attr, 0)))}" for attr in ATTRIBUTES
            )
            print(f"      Why:      {line}  ->  {candidate_comment(scores)}")
        print()

    if result.note:
        print(f"  Note: {result.note}")
    print(f"  Recommendation: {recommend(result, product.get('lcom_price'))[1]}")


def format_result_markdown(
    product: dict, result: Optional[SourcingResult], calls_used: int, tin: int, tout: int, cost: float
) -> str:
    lines = [f"## {product['sku']} - {product['description']}", ""]

    if result is None:
        lines.append("**No structured result** (empty, refused or garbled response).")
    elif result.no_match or not result.candidates:
        reason = result.no_match_reason or "no plausible candidate found"
        lines.append(f"**No match found** - {reason}")
    else:
        lcom = product.get("lcom_price")
        lcom_text = f"${lcom:.2f}" if lcom else "n/a"
        rec_idx, rec_reason = recommend(result, lcom)
        lines += [
            f"**Recommendation:** {rec_reason}",
            "",
            f"| Candidate | Accuracy | Unit price | Unit price confidence | vs. L-Com ({lcom_text}) | Recommended |",
            "|---|---|---|---|---|---|",
        ]
        for i, c in enumerate(result.candidates):
            accuracy, _ = candidate_score(c)
            cells = [
                f"{i + 1}. {c.manufacturer or 'Not stated'}",
                f"{accuracy}%",
                format_unit(unit_price(c)),
                confidence_label(result, i),
                format_vs_lcom(price_vs_lcom(checked_unit(result, i), lcom)),
                "YES" if i == rec_idx else "no",
            ]
            if i == rec_idx:
                cells = [f"**{x}**" for x in cells]
            lines.append("| " + " | ".join(cells) + " |")
        lines += [f"| L-Com (benchmark) | - | {lcom_text} | - | - | - |", ""]

        for i, candidate in enumerate(result.candidates, start=1):
            computed, scores = candidate_score(candidate)
            lines += [
                f"### Candidate {i}: {computed}% - {match_tier(computed)}",
                "",
                "| Field | Value |",
                "|---|---|",
                f"| Manufacturer / Producer | {candidate.manufacturer or 'Not stated on listing'} |",
                f"| Supplier / Site | {candidate.distributor_site or 'Unknown site'} |",
                f"| Listing | {candidate.listing_title or '?'} |",
                f"| Accuracy | {computed}% |",
                f"| Listed price | {candidate.price or 'Not stated'} |",
                f"| Qty the price covers | {candidate.quantity_covered or '?'} |",
                f"| Unit price | {format_unit(unit_price(candidate))} ({confidence_label(result, i - 1)}) |",
                f"| Unit price note | {candidate.unit_price_note or '-'} |",
                f"| vs. L-Com price | {format_vs_lcom(price_vs_lcom(checked_unit(result, i - 1), lcom))} |",
                f"| MOQ | {candidate.moq or 'Not stated'} |",
                f"| Email | {candidate.email or 'Not stated'} |",
                f"| URL | {candidate.url or 'Not stated'} |",
                "",
            ]
            if scores:
                lines.append("| Rubric attribute | Score (of 20) |")
                lines.append("|---|---|")
                for attr in ATTRIBUTES:
                    lines.append(f"| {ATTRIBUTE_LABELS[attr]} | {max(0, min(20, scores.get(attr, 0)))} |")
                lines += ["", candidate_comment(scores), ""]
        if result.note:
            lines.append(f"*{result.note}*")

    lines += ["", f"*Tool calls used: {calls_used} | Tokens: in={tin} out={tout} | Cost: ~${cost:.4f}*", ""]
    return "\n".join(lines)


def ambiguous_price_flags(rows: list) -> list:
    """Markdown bullets for every candidate whose unit price couldn't be determined or looks implausible."""
    flags = []
    for product, result in rows:
        suspect = implausible_units(result)
        for i, c in enumerate(result.candidates if result else [], start=1):
            if i - 1 in suspect:
                reason = f"{format_unit(unit_price(c))}/unit: {IMPLAUSIBLE_NOTE} ({c.unit_price_note or 'no note'})"
            elif unit_price(c) is None:
                reason = c.unit_price_note or c.unit_price_confidence or "no unit price given"
            else:
                continue
            flags.append(
                f"- **{product['sku']}** candidate {i} ({c.manufacturer or 'unnamed'}): "
                f"listed \"{c.price or 'no price'}\" - {reason}"
            )
    return flags


def write_report(
    sections: list, grand_in: int, grand_out: int, grand_cost: float, num_products: int, path: str,
    flags: list,
) -> None:
    header = [
        "# Competitive Sourcing Research Report",
        "",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"Model: {MODEL} | Sites: {', '.join(SOURCING_SITES)} | Products: {num_products}",
        f"Recommendation rule: >= {MIN_RECOMMEND_ACCURACY}% accuracy and a confirmed unit price "
        f">= {MIN_MARGIN_PCT}% below L-Com's price; best accuracy x margin wins.",
        "",
        "## Unit prices that could not be determined confidently or look implausible",
        "",
        *(flags or ["None."]),
        "",
        "---",
        "",
    ]
    footer = [
        "---",
        "",
        f"**Total tokens:** in={grand_in} out={grand_out} | **Estimated total cost:** ~${grand_cost:.4f}",
    ]
    content = "\n".join(header + sections + footer) + "\n"
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


PRODUCT_XLSX_FIELDS = ["SKU", "Product", "Keyword", "L-Com Unit Price", "Recommendation"]
RECOMMENDED_XLSX_FIELDS = [
    "Recommended Manufacturer", "Recommended Email", "Recommended URL", "Recommended Unit Price",
]
CANDIDATE_XLSX_FIELDS = [
    "Name", "Accuracy", "Listed Price", "Unit Price", "Unit Price Confidence", "vs. L-Com Price",
    "MOQ", "URL", "Email", "Match Tier", "Comment",
]
GREEN = PatternFill("solid", fgColor="C6EFCE")


def write_excel_report(rows: list, path: str) -> None:
    """rows: list of (product: dict, result: Optional[SourcingResult]) tuples.
    Sheet "Results": one row per product. Sheet "Comparison": candidates vs L-Com side by side.
    The green fill marks who to buy from: the Recommended columns on Results, the row on
    Comparison. Nothing is green when no candidate qualifies."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    comp = wb.create_sheet("Comparison")

    header = PRODUCT_XLSX_FIELDS + RECOMMENDED_XLSX_FIELDS
    for n in (1, 2, 3):
        header += [f"Manufacturer {n} {field}" for field in CANDIDATE_XLSX_FIELDS]
    ws.append(header)
    comp.append(["SKU", "Candidate", "Manufacturer", "Accuracy", "Unit Price", "Unit Price Confidence",
                 "vs. L-Com Price", "Recommended"])

    for product, result in rows:
        lcom = product.get("lcom_price")
        rec_idx, rec_reason = recommend(result, lcom)
        row = [product["sku"], product.get("product_name", ""), product["description"], lcom, rec_reason]
        candidates = result.candidates if result else []
        if rec_idx is None:
            row += [""] * len(RECOMMENDED_XLSX_FIELDS)
        else:
            rec = candidates[rec_idx]
            row += [rec.manufacturer or "", rec.email or "", rec.url or "", round(unit_price(rec), 4)]
        for i in range(3):
            if i < len(candidates):
                c = candidates[i]
                computed, scores = candidate_score(c)
                unit = unit_price(c)
                unit_cell = round(unit, 4) if unit is not None else "UNDETERMINED"
                confidence = confidence_label(result, i)
                vs = format_vs_lcom(price_vs_lcom(checked_unit(result, i), lcom))
                row += [
                    c.manufacturer or "",
                    computed,
                    c.price or "",
                    unit_cell,
                    confidence,
                    vs,
                    c.moq or "",
                    c.url or "",
                    c.email or "",
                    match_tier(computed),
                    candidate_comment(scores),
                ]
                comp.append([product["sku"], i + 1, c.manufacturer or "", computed, unit_cell,
                             confidence, vs, "YES" if i == rec_idx else "no"])
                if i == rec_idx:
                    for cell in comp[comp.max_row]:
                        cell.fill = GREEN
            else:
                row += [""] * len(CANDIDATE_XLSX_FIELDS)
        comp.append([product["sku"], "L-Com", "L-Com (benchmark)", None, lcom])
        comp.append([])
        ws.append(row)
        if rec_idx is not None:
            start = len(PRODUCT_XLSX_FIELDS)
            for cell in ws[ws.max_row][start : start + len(RECOMMENDED_XLSX_FIELDS)]:
                cell.fill = GREEN

    for sheet in (ws, comp):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col_cells in sheet.columns:
            width = max((len(str(cell.value)) for cell in col_cells if cell.value), default=10)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 60)

    wb.save(path)


def read_products(path: str) -> list:
    """Reads SKU / Product / Keyword columns from an .xlsx (header row detected by
    scanning for a cell literally containing "SKU"), matching the manual spreadsheet's
    input layout. Description is taken from Keyword, falling back to Product."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    headers, header_row = {}, None
    for row in ws.iter_rows(min_row=1, max_row=5):
        cells = {str(c.value).strip().lower(): c.column for c in row if c.value}
        if "sku" in cells:
            headers, header_row = cells, row[0].row
            break
    if header_row is None:
        raise ValueError(f"Could not find a 'SKU' header row in the first 5 rows of {path}")

    sku_col = headers["sku"]
    keyword_col = headers.get("keyword")
    product_col = headers.get("product")
    lcom_col = headers.get("lcom sale price")

    products = []
    for row in ws.iter_rows(min_row=header_row + 1):
        sku = row[sku_col - 1].value
        if not sku:
            continue
        keyword = row[keyword_col - 1].value if keyword_col else None
        product_name = row[product_col - 1].value if product_col else None
        lcom_price = row[lcom_col - 1].value if lcom_col else None
        # L-Com prices packs as one SKU ("..., Package/10") - compare candidates per unit.
        pack = re.search(r"package\s*/\s*(\d+)", f"{keyword} {product_name}", re.IGNORECASE)
        if pack and isinstance(lcom_price, (int, float)):
            lcom_price /= int(pack.group(1))
        products.append({
            "sku": str(sku).strip(),
            "description": str(keyword or product_name or "").strip(),
            "product_name": str(product_name or "").strip(),
            "lcom_price": float(lcom_price) if isinstance(lcom_price, (int, float)) else None,
        })
    return products


async def connect_tools(nimble_api_key: str):
    http_client = httpx2.AsyncClient(headers={"Authorization": f"Bearer {nimble_api_key}"})
    return http_client


def parse_args():
    parser = argparse.ArgumentParser(description="Competitive sourcing research POC")
    parser.add_argument(
        "--input", help="Path to an .xlsx with SKU/Product/Keyword columns. Defaults to 5 built-in sample products."
    )
    parser.add_argument("--output", help="Path for the .xlsx results file. Defaults to a timestamped filename.")
    parser.add_argument("--limit", type=int, help="Only process the first N products (for a quick/cheap test run).")
    parser.add_argument(
        "--inspect-tools", action="store_true", help="Print Nimble tool schemas and exit (no Anthropic tokens spent)."
    )
    return parser.parse_args()


async def main_async():
    args = parse_args()
    load_dotenv()  # .env next to where you run it; real env vars still win

    anthropic_key = os.environ.get("ANTHROPIC_API_KEY")
    nimble_key = os.environ.get("NIMBLE_API_KEY")
    missing = [n for n, v in [("ANTHROPIC_API_KEY", anthropic_key), ("NIMBLE_API_KEY", nimble_key)] if not v]
    if missing:
        print(f"Missing required environment variable(s): {', '.join(missing)}")
        print("Add them to a .env file in this folder (or set them in the shell), e.g.:")
        print("  ANTHROPIC_API_KEY=sk-ant-...")
        print("  NIMBLE_API_KEY=...")
        sys.exit(1)

    products = read_products(args.input) if args.input else PRODUCTS
    if args.limit:
        products = products[: args.limit]
    if not products:
        print(f"No products found in {args.input} - check it has a 'SKU' header column.")
        sys.exit(1)

    client = AsyncAnthropic(api_key=anthropic_key)

    print("Competitive Sourcing Research POC")
    print(f"Model: {MODEL}  |  Products: {len(products)}  |  Sites: {', '.join(SOURCING_SITES)}")
    print(f"Tool call budget: {MAX_TOOL_CALLS_PER_PRODUCT} per product")

    http_client = httpx2.AsyncClient(
        headers={"Authorization": f"Bearer {nimble_key}"}, timeout=60.0
    )
    async with http_client:
        async with streamable_http_client(NIMBLE_MCP_URL, http_client=http_client) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                tools_result = await session.list_tools()

                search_tool = pick_tool(tools_result.tools, {"nimble_search", "search"}, "search")
                extract_tool = pick_tool(tools_result.tools, {"nimble_extract", "extract"}, "extract")
                if search_tool is None or extract_tool is None:
                    available = ", ".join(t.name for t in tools_result.tools)
                    print(f"Could not find search/extract tools on the Nimble MCP server.")
                    print(f"Available tools: {available}")
                    sys.exit(1)

                tools = [search_tool, extract_tool]
                print(f"Using Nimble tools: {search_tool.name}, {extract_tool.name}")
                for t in (search_tool, extract_tool):
                    schema = getattr(t, "input_schema", None) or getattr(t, "inputSchema", None) or {}
                    print(f"  {t.name} params: {list(schema.get('properties', {}).keys())}")

                if args.inspect_tools:
                    for t in (search_tool, extract_tool):
                        schema = getattr(t, "input_schema", None) or getattr(t, "inputSchema", None) or {}
                        print("-" * 72)
                        print(f"{t.name} full schema:")
                        print(json.dumps(schema, indent=2))
                    print("-" * 72)
                    print("--inspect-tools: exiting before spending any Anthropic tokens.")
                    return

                grand_in, grand_out, grand_cost = 0, 0, 0.0
                report_sections = []
                excel_rows = []

                for idx, product in enumerate(products, start=1):
                    print(f"\nProduct {idx} of {len(products)}: {product['sku']} ...")
                    try:
                        result, tin, tout, calls_used = await research_product(
                            client, session, tools, product
                        )
                        print_product_result(product, result, calls_used)
                        cost = (tin / 1_000_000 * HAIKU_INPUT_PER_MTOK) + (
                            tout / 1_000_000 * HAIKU_OUTPUT_PER_MTOK
                        )
                        grand_in += tin
                        grand_out += tout
                        grand_cost += cost
                        print(f"  Tokens:     in={tin} out={tout}  (~${cost:.4f})")
                        report_sections.append(format_result_markdown(product, result, calls_used, tin, tout, cost))
                        excel_rows.append((product, result))
                    except Exception as exc:  # noqa: BLE001 - POC: keep going on any per-product failure
                        print("-" * 72)
                        print(f"{product['sku']} - \"{product['description']}\"")
                        print(f"  Error researching this product: {exc}")
                        report_sections.append(
                            f"## {product['sku']} - {product['description']}\n\n"
                            f"**Error researching this product:** {exc}\n"
                        )
                        excel_rows.append((product, None))

                print("-" * 72)
                print(
                    f"Total tokens: in={grand_in} out={grand_out}  |  "
                    f"Estimated total cost: ~${grand_cost:.4f}"
                )

                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                report_path = f"sourcing_report_{timestamp}.md"
                flags = ambiguous_price_flags(excel_rows)
                write_report(report_sections, grand_in, grand_out, grand_cost, len(products), report_path, flags)
                if flags:
                    print("Unit prices that could not be determined confidently:")
                    print("\n".join(flags))
                print(f"Report written to: {os.path.abspath(report_path)}")

                excel_path = args.output or f"sourcing_results_{timestamp}.xlsx"
                write_excel_report(excel_rows, excel_path)
                print(f"Excel results written to: {os.path.abspath(excel_path)}")


def main():
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
