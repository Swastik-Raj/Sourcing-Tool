"""Offline, table-driven checks for decision_agent.py. No API, no network, no LLM, nothing is ordered.
Run: python test_decision_agent.py"""
import contextlib
import csv
import io
import json
import os
import re
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal as D

import openpyxl

os.environ["OBS_ENABLED"] = "0"   # tests are offline: never trace, even if .env switches tracing on
import decision_agent as da
from email_agent import listing_key
from sourcing_agent import (CANDIDATE_XLSX_FIELDS, MAX_CANDIDATES, ORDERING_NOTE_FIELD, PRODUCT_XLSX_FIELDS,
                            RECOMMENDED_XLSX_FIELDS)

HERE = os.path.dirname(os.path.abspath(__file__))
HEADER = (PRODUCT_XLSX_FIELDS + RECOMMENDED_XLSX_FIELDS + [ORDERING_NOTE_FIELD]
          + [f"Manufacturer {n} {f}" for n in range(1, MAX_CANDIDATES + 1) for f in CANDIDATE_XLSX_FIELDS]
          + ["Recommended Listing Title"] + [f"Manufacturer {n} Listing Title" for n in range(1, MAX_CANDIDATES + 1)])
T0 = datetime(2026, 10, 5, 12, 0, 0, tzinfo=timezone.utc)


def res_row(sku, maker="Acme", url=None, unit=1.25, lcom=28.19, moq="100", conf="stated", acc=100, tier="Auto-accepted",
            note="", rec="Buy from Candidate 1", lcom_source="Srijan, live check"):
    """A Results row as the search agent writes it (recommended columns plus candidate block 1)."""
    url = url if url is not None else f"http://shop/{sku}"
    r = dict.fromkeys(HEADER)
    r.update({"SKU": sku, "Product": f"{sku} product", "Keyword": f"{sku} keyword", "L-Com Unit Price": lcom,
              "L-Com Price Source": lcom_source, "Recommendation": rec, "Recommended Manufacturer": maker,
              "Recommended URL": url, "Recommended Unit Price": unit, "Ordering Note": note,
              "Manufacturer 1 Name": maker, "Manufacturer 1 Accuracy": acc, "Manufacturer 1 Unit Price Confidence": conf,
              "Manufacturer 1 MOQ": moq, "Manufacturer 1 URL": url, "Manufacturer 1 Match Tier": tier})
    return r


def make_results(path, rows, drop=(), reverse=False):
    header = [h for h in HEADER if h not in drop]
    if reverse:
        header = header[::-1]                       # columns in another order: must still be read by name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])
    wb.save(path)
    wb.close()


def key_for(sku, maker="Acme", url=None):
    return listing_key(sku, url if url is not None else f"http://shop/{sku}", maker)


class Clock:
    t = T0


@contextlib.contextmanager
def sandbox(exclusions=(), email_state=None):
    """Temp copies of every runtime path, a fixed clock, scripted prompts. Reviewer exclusions and email state are
    real files so the read-only rule can be checked byte for byte."""
    names = ("STATE_PATH", "LOG_PATH", "EXPORT_JSON", "EXPORT_CSV", "REPORTS_DIR", "EMAIL_STATE_PATH",
             "REVIEWER_EXCLUSIONS_CSV", "now", "ask_user")
    old = {n: getattr(da, n) for n in names}
    tmp = tempfile.mkdtemp()
    for n, f in (("STATE_PATH", "decision_state.json"), ("LOG_PATH", "approvals_log.jsonl"),
                 ("EXPORT_JSON", "approved_orders.json"), ("EXPORT_CSV", "approved_orders.csv"),
                 ("REPORTS_DIR", "Reports"), ("EMAIL_STATE_PATH", "email_state.json"),
                 ("REVIEWER_EXCLUSIONS_CSV", "reviewer_exclusions.csv")):
        setattr(da, n, os.path.join(tmp, f))
    with open(da.REVIEWER_EXCLUSIONS_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["sku", "supplier_or_url", "reason"])
        w.writerows(exclusions)
    if email_state is not None:
        open(da.EMAIL_STATE_PATH, "w", encoding="utf-8").write(email_state if isinstance(email_state, str) else json.dumps(email_state))
    Clock.t = T0
    da.now = lambda: Clock.t
    da.ask_user = lambda prompt: (_ for _ in ()).throw(AssertionError("unexpected prompt: " + prompt))
    try:
        yield tmp
    finally:
        for n, v in old.items():
            setattr(da, n, v)


def answers(*given):
    it = iter(given)
    da.ask_user = lambda prompt: next(it)


def run_cli(argv, stdin_bytes=b""):
    out = io.TextIOWrapper(io.BytesIO(), encoding="utf-8", newline="")
    old = (sys.stdout, sys.stdin)
    sys.stdout, sys.stdin = out, io.TextIOWrapper(io.BytesIO(stdin_bytes), encoding="cp1252", newline="")  # a Windows pipe
    try:
        da.main(argv)
    finally:
        sys.stdout, sys.stdin = old
    out.flush()
    return out.buffer.getvalue().decode("utf-8")


def cli_fails(argv, stdin_bytes=b""):
    try:
        run_cli(argv, stdin_bytes)
    except SystemExit as e:
        return str(e.code)
    raise AssertionError("expected the command to stop with a message")


def state():
    return json.load(open(da.STATE_PATH, encoding="utf-8"))


def log():
    return [json.loads(l) for l in open(da.LOG_PATH, encoding="utf-8").read().splitlines()] if os.path.exists(da.LOG_PATH) else []


def standard_rows():
    return [res_row("VIC00001", "Eternalstar", unit=1.25, moq="40"),       # $50.00 at the MOQ: under the $100 limit
            res_row("FOA-020C", "FiberMania", unit=0.25, lcom=72.99, moq="10 Pieces"),
            res_row("BIG1", "Big Co", unit=11.00, moq="500"),
            res_row("HDFF", "Farsince", unit=0.48, moq=None)]


def request_for(rows, tmp, name="r.xlsx", budget=None):
    path = os.path.join(tmp, name)
    make_results(path, rows)
    return path, da.cmd_request(path, budget)


def record(path, text, channel="Slack DM", approver="Neeraj", confirm="yes", request_id=None, typed_id="A-0001"):
    """cmd_record with scripted prompts: confirm the request id, then channel, approver, and the final yes."""
    answers(typed_id, channel, approver, confirm)
    with contextlib.redirect_stdout(io.StringIO()) as out:
        plan, stored = da.cmd_record(path, text, request_id)
    return plan, stored, out.getvalue()


# ---------- money ----------

def test_money_rounds_half_up_to_cents():
    table = [(1.005, "1.01"), (2.675, "2.68"), (0.125, "0.13"), (0.1 + 0.2, "0.30"), (1.004999, "1.00"), (5, "5.00"),
             (D("3") * D("0.335"), "1.01"), (D("100") * D("0.0545"), "5.45"), (D("7") * D("0.0545"), "0.38"),
             (D("1234.565"), "1234.57"), (-1.005, "-1.01")]
    for value, want in table:
        assert str(da.money(value)) == want, (value, da.money(value))
    for d, want in [("0.0545", "$0.0545"), ("1.25", "$1.25"), ("1234.5", "$1,234.50"), ("0.3", "$0.30"), ("2", "$2.00"),
                    ("0.12345", "$0.1235")]:
        assert da.fmt_unit(D(d)) == want, d
    assert da.fmt_money(D("5500")) == "$5,500.00" and da.fmt_money(D("0.5")) == "$0.50"


def test_cap_arithmetic_table():
    lines = [{"sku": "A", "unit_price": "1.25", "seller": "s", "key": "A|x", "flags": [], "product": "p", "url": "u"},
             {"sku": "B", "unit_price": "0.30", "seller": "s", "key": "B|x", "flags": [], "product": "p", "url": "u"}]
    # (approve lines, cap, expected committed or the failure text)
    table = [
        ("approve A qty 100 max 1.30\napprove B qty 50 max 0.30", "145", "145.00"),    # exactly the cap is allowed
        ("approve A qty 100 max 1.30\napprove B qty 50 max 0.30", "144.99", "exceeds the cap"),    # one cent over is not
        ("approve A qty 3 max 0.335", "1.01", "1.01"),                                  # 1.005 rounds half up to 1.01
        ("approve A qty 3 max 0.335", "1.00", "exceeds the cap"),
        ("approve A qty 7 max 0.0545", "0.38", "0.38"),                                 # 0.3815 -> 0.38
        ("approve A qty 1,000 max $1.30", "$1,300", "1300.00"),                         # $ and commas
        ("approve A qty 1 max 0.004", "0.01", "0.00"),                                  # a sub-cent line rounds to nothing
        ("approve A qty 100 max 1.30\nreject B", "129.99", "exceeds the cap"),
    ]
    for text, cap, want in table:
        parsed = da.parse_reply(f"{text}\ncap {cap}")
        try:
            plan = da.check_reply(parsed, lines)
            assert want[0].isdigit() and str(plan["committed"]) == want, (text, cap, plan["committed"])
        except da.ReplyError as e:
            assert want in str(e), (text, cap, str(e))


# ---------- the reply parser ----------

VALID_REPLIES = [
    ("approve VIC00001 qty 100 max 1.30\napprove FOA-020C qty 50 max 0.30\nreject HDFF\ncap 200",
     {"cap": "200.00", "approve": [("VIC00001", 100, "1.30"), ("FOA-020C", 50, "0.30")], "reject": ["HDFF"]}),
    ("APPROVE vic00001 QTY 1,000 MAX $1.30\nCap $1,200.50",
     {"cap": "1200.50", "approve": [("vic00001", 1000, "1.30")], "reject": []}),
    ("\n\n   approve A qty 5 max 0.1234   \n\ncap 1\n", {"cap": "1.00", "approve": [("A", 5, "0.1234")], "reject": []}),
    ("reject A\nreject B\ncap 10", {"cap": "10.00", "approve": [], "reject": ["A", "B"]}),
    ("cap 200", {"cap": "200.00", "approve": [], "reject": []}),
    ("approve A qty 2 max $.5\ncap 5", {"cap": "5.00", "approve": [("A", 2, "0.5")], "reject": []}),
    ("Approve A Qty 7 Max 1\r\nREJECT B\r\nCAP 9", {"cap": "9.00", "approve": [("A", 7, "1")], "reject": ["B"]}),
]
INVALID_REPLIES = [  # (reply, text that must appear in the rejection)
    ("approve A qty 5 max 1", "cap is required"),
    ("", "cap is required"),
    ("buy A\ncap 5", "can't be parsed"),
    ("approve A\ncap 5", "explicit qty and max unit price"),
    ("approve A qty 5\ncap 5", "explicit qty and max unit price"),
    ("approve A max 1\ncap 5", "explicit qty and max unit price"),
    ("approve A qty 5 max 1 now\ncap 5", "can't be parsed"),
    ("approve A qty 0 max 1\ncap 5", "isn't a positive whole number"),
    ("approve A qty -3 max 1\ncap 5", "isn't a positive whole number"),
    ("approve A qty 2.5 max 1\ncap 5", "isn't a positive whole number"),
    ("approve A qty ten max 1\ncap 5", "isn't a positive whole number"),
    ("approve A qty 5 max 0\ncap 5", "isn't a positive amount"),
    ("approve A qty 5 max -1\ncap 5", "isn't a positive amount"),
    ("approve A qty 5 max abc\ncap 5", "isn't a positive amount"),
    ("approve A qty 5 max nan\ncap 5", "isn't a positive amount"),
    ("approve A qty 5 max 1\napprove a qty 5 max 1\ncap 50", "appears twice"),
    ("approve A qty 5 max 1\nreject A\ncap 50", "appears twice"),
    ("reject A\nreject A\ncap 50", "appears twice"),
    ("cap 5\ncap 6", "cap appears more than once"),
    ("cap abc", "cap 'abc' isn't a positive amount"),
    ("cap 0", "cap '0' isn't a positive amount"),
    ("cap", "can't be parsed"),
    ("rejectA\ncap 5", "can't be parsed"),
    ("批准 A qty 5 max 1\ncap 5", "批准"),                 # a non-ASCII line is reported, not mangled
]


def test_reply_parser_valid_table():
    for text, want in VALID_REPLIES:
        got = da.parse_reply(text)
        assert str(got["cap"]) == want["cap"], (text, got)
        assert [(a["sku"], a["qty"], str(a["max"])) for a in got["approvals"]] == want["approve"], text
        assert got["rejections"] == want["reject"], text


def test_reply_parser_invalid_table():
    for text, why in INVALID_REPLIES:
        try:
            da.parse_reply(text)
            raise AssertionError(f"accepted: {text!r}")
        except da.ReplyError as e:
            assert why in str(e) and "nothing stored" in str(e), (text, str(e))
    # every failure is listed, not just the first
    try:
        da.parse_reply("buy A\napprove B qty 0 max 1\nreject C\nreject c")
    except da.ReplyError as e:
        assert len(e.errors) == 4, e.errors      # unparseable, qty, duplicate SKU, no cap


def test_check_reply_against_the_request():
    with sandbox() as tmp:
        path, req = request_for(standard_rows(), tmp)
        lines = req["lines"]
        for text, why in [("approve NOPE qty 5 max 1\ncap 50", "SKU NOPE is not in the request"),
                          ("reject NOPE\ncap 50", "SKU NOPE is not in the request"),
                          ("approve VIC00001 qty 100 max 1.30\ncap 129.99", "exceeds the cap")]:
            try:
                da.check_reply(da.parse_reply(text), lines)
                raise AssertionError("accepted " + text)
            except da.ReplyError as e:
                assert why in str(e)
        plan = da.check_reply(da.parse_reply("approve vic00001 qty 100 max 1.10\nreject big1\ncap 200"), lines)
        a = plan["approvals"][0]
        assert a["line"]["sku"] == "VIC00001"                     # case-insensitive match, canonical spelling kept
        assert "below the listed unit price" in a["warning"] and "$1.10" in a["warning"] and "$1.25" in a["warning"]
        assert [l["sku"] for l in plan["rejections"]] == ["BIG1"] and {l["sku"] for l in plan["undecided"]} == {"FOA-020C", "HDFF"}
        ok = da.check_reply(da.parse_reply("approve VIC00001 qty 100 max 1.25\ncap 200"), lines)
        assert ok["approvals"][0]["warning"] == ""                # max equal to the listed price is fine


# ---------- the request ----------

def test_request_contents_and_quantities():
    state_json = {key_for("FOA-020C", "FiberMania"): {"status": "pricing_provided", "summary": "sample USD 0.30", "sample_quantity": 25},
                  key_for("VIC00001", "Eternalstar"): {"sent": "2026-10-04 10:00"}}
    with sandbox(email_state=state_json) as tmp:
        rows = standard_rows() + [res_row("PROMO1", "Promo Co", unit=1.09, moq="5", conf="promo (regular price unknown)", tier="Flagged for manual review", note="promo price - regular price not shown", lcom_source="unverified")]
        path, req = request_for(rows, tmp, budget=D("1000"))
        by = {l["sku"]: l for l in req["lines"]}
        assert req["id"] == "A-0001" and req["results_file"] == "r.xlsx" and req["results_sha256"] == da.file_sha256(path)
        v = by["VIC00001"]
        assert (v["qty"], v["qty_source"], v["total"]) == (40, "listing MOQ", "50.00")
        assert v["accuracy"] == 100 and v["tier"] == "Auto-accepted" and v["confidence"] == "stated"
        assert v["pct_below_lcom"] == "95.6" and v["seller_reply"] == "contacted, no reply recorded"
        f = by["FOA-020C"]                                         # the seller's stated sample qty beats the listing MOQ
        assert (f["qty"], f["qty_source"], f["total"]) == (25, "seller's stated sample quantity", "6.25")
        assert f["seller_reply"].startswith("pricing_provided - sample USD 0.30")
        b = by["BIG1"]                                              # MOQ total $5,500 > $100: no suggested quantity
        assert (b["qty"], b["total"], b["moq_qty"], b["moq_total"], b["moq_heavy"]) == (None, None, 500, "5500.00", True)
        assert b["flags"][0] == "needs_seller_quote_for_smaller_qty" and "moq_not_stated" not in b["flags"]
        h = by["HDFF"]
        assert h["qty"] is None and h["total"] is None and "moq_not_stated" in h["flags"]
        p = by["PROMO1"]
        assert "unit_price_promo" in p["flags"] and "manual_review_tier" in p["flags"] and "lcom_price_unverified" in p["flags"]
        assert p["note"] == "promo price - regular price not shown"
        assert D(req["subtotal"]) == D("50.00") + D("6.25") + D("5.45")           # BIG1 is left out of the subtotal
        assert req["left_out"] == ["BIG1"] and req["left_out_moq_total"] == "5500.00" and req["lines_without_total"] == ["HDFF"]
        text = da.render_text(req)
        for needle in ["APPROVAL REQUEST A-0001", "Products: 5", "$61.70 (covers 3 of 5 lines)", "NOT included: shipping, duties and taxes.",
                       "Left out of that subtotal: 1 line(s) whose total at the MOQ is over $100.00 (BIG1) - combined MOQ total $5,500.00",
                       "no qty for HDFF (MOQ not stated)", "Budget: $1,000.00", "$938.30 to spare",
                       "Seller: Eternalstar", "Listing: http://shop/VIC00001",
                       "Match: 100% (Auto-accepted) | Unit price: $1.25 (stated) | 95.6% below L-Com ($28.19, Srijan, live check)",
                       "Ordering note: promo price", "Seller reply: pricing_provided", "HOW TO REPLY", "EXAMPLE",
                       "approve SKU-A qty 100 max 1.30", "cap 200", "Replying orders nothing"]:       # (BIG1 is covered below)
            assert needle in text, needle
        assert text.rstrip().endswith("cap 200")                       # the reply format and example come last
        files = sorted(os.listdir(da.REPORTS_DIR))
        assert files == ["approval_request_A-0001.md", "approval_request_A-0001.txt"]
        md = open(os.path.join(da.REPORTS_DIR, files[0]), encoding="utf-8").read()
        assert md.startswith("# Approval request A-0001") and "HOW TO REPLY" in md and "NEEDS A SELLER QUOTE" in md
        # a second request gets the next id; the earlier one is kept
        _, req2 = request_for(rows, tmp)
        assert req2["id"] == "A-0002" and set(state()["requests"]) == {"A-0001", "A-0002"}
        ev = log()
        assert [e["event"] for e in ev] == ["request", "request"] and ev[0]["results_sha256"] == req["results_sha256"]


def test_moq_heavy_lines_get_no_suggested_quantity():
    # (unit price, MOQ text, the seller's stated sample qty, heavy?, suggested qty, line total)
    table = [
        (1.25, "80", None, False, 80, "100.00"),           # a total of exactly $100.00 is not over the limit
        (0.48, "500", None, True, None, None),             # $240.00: no longer suggested at full quantity
        (0.10, "1001", None, True, None, None),            # $100.10 is
        (11.00, "500", None, True, None, None),
        (0.55, "10,000 pieces", None, True, None, None),
        (2.00, "300", None, True, None, None),             # $600
        (11.00, "500", 25, False, 25, "275.00"),           # the seller's own sample quantity replaces the MOQ: not heavy
        (2.00, None, None, False, None, None),             # no MOQ stated: no qty either, but it is not "heavy"
        (1.00, "50", None, False, 50, "50.00"),
        (11.00, "500", 100, False, 100, "1100.00"),        # a big seller-stated quantity is flagged, not "heavy"
    ]
    for unit, moq, stated, heavy, qty, total in table:
        email = {key_for("X1", "Acme"): {"status": "pricing_provided", "sample_quantity": stated}} if stated else None
        with sandbox(email_state=email) as tmp:
            _, req = request_for([res_row("X1", "Acme", unit=unit, moq=moq)], tmp)
            line = req["lines"][0]
            assert (line["moq_heavy"], line["qty"], line["total"]) == (heavy, qty, total), (unit, moq, stated, line)
            assert ("needs_seller_quote_for_smaller_qty" in line["flags"]) == heavy
            assert ("large_total" in line["flags"]) == (not heavy and total is not None and D(total) > da.LARGE_TOTAL)
            if heavy:
                assert line["moq_total"] == str(da.money(D(str(unit)) * da.parse_moq(moq))) and line["qty_source"] == ""


def test_heavy_lines_have_their_own_section_and_stay_out_of_the_subtotal():
    rows = [res_row("PRICED1", "Acme", unit=1.25, moq="40"),
            res_row("HEAVY1", "Heavy Co", unit=0.55, moq="10,000 pieces"),            # $5,500.00 at the MOQ
            res_row("NOMOQ", "Plain Co", unit=2.00, moq=None),
            res_row("HEAVY2", "Heavier Co", unit=2.00, moq="300"),                    # $600.00 at the MOQ
            res_row("PRICED2", "Small Co", unit=0.50, moq="10")]
    quote = "needs a seller quote for a smaller quantity first"
    with sandbox() as tmp:
        path, req = request_for(rows, tmp, budget=D("200"))
        assert D(req["subtotal"]) == D("55.00") and req["left_out"] == ["HEAVY1", "HEAVY2"]
        assert req["left_out_moq_total"] == "6100.00" and req["lines_without_total"] == ["NOMOQ"]
        text = da.render_text(req)
        head, _, rest = text.partition("\n\n")
        assert "Products: 5 | Product subtotal at suggested quantities: $55.00 (covers 2 of 5 lines); no qty for NOMOQ (MOQ not stated)" in head
        assert ("Left out of that subtotal: 2 line(s) whose total at the MOQ is over $100.00 (HEAVY1, HEAVY2) - "
                "combined MOQ total $6,100.00. They need a seller quote for a smaller quantity first.") in head
        assert "subtotal (priced lines only) is within budget ($145.00 to spare)" in head
        section = "NEEDS A SELLER QUOTE FOR A SMALLER QUANTITY FIRST (total at the MOQ is over $100.00; no quantity is suggested"
        before, _, after = rest.partition(section)
        assert after and after.index("HEAVY1") < after.index("HEAVY2") < after.index("HOW TO REPLY")
        assert "HEAVY" not in before and "HEAVY" in after                   # heavy lines appear only inside their section
        assert [l.split(")")[0] for l in rest.splitlines() if re.match(r"\d+\) ", l)] == ["1", "2", "3", "4", "5"]   # numbering continues
        assert "4) HEAVY1" in after and "5) HEAVY2" in after
        h1 = after.split("4) HEAVY1")[1].split("\n\n")[0]
        assert "MOQ: 10,000 pieces | MOQ total: $5,500.00 | Suggested qty: none - " + quote in h1
        assert "Line total" not in h1 and "Flags: needs_seller_quote_for_smaller_qty" in h1
        assert quote not in before.split("NOMOQ")[0] and "Suggested qty: 40 (listing MOQ) | Line total: $50.00" in before
        assert "MOQ-driven" not in text
        assert "A line in the 'needs a seller quote' or 'does not clear the margin bar' section can still be approved" in text
        md = open(os.path.join(da.REPORTS_DIR, "approval_request_A-0001.md"), encoding="utf-8").read()
        assert "NEEDS A SELLER QUOTE FOR A SMALLER QUANTITY FIRST" in md and "Left out of that subtotal" in md
        # no heavy lines: no left-out sentence at all
        _, plain = request_for([rows[0]], tmp, "p.xlsx")
        assert "Left out of that subtotal" not in da.render_text(plain) and "NEEDS A SELLER QUOTE" not in da.render_text(plain)
        # a heavy line can still be approved with an explicit qty and max (but not without them)
        for bad in ("approve HEAVY1\ncap 50", "approve HEAVY1 qty 20\ncap 50"):
            try:
                da.parse_reply(bad)
                raise AssertionError("accepted " + bad)
            except da.ReplyError:
                pass
        _, stored, out = record(path, "approve HEAVY1 qty 20 max 0.60\ncap 50", request_id="A-0001")
        assert stored and "APPROVE HEAVY1 (Heavy Co): qty 20 x max $0.60 = $12.00 at most" in out
        d = state()["decisions"][0]
        assert (d["sku"], d["qty"], d["line_cap"]) == ("HEAVY1", 20, "12.00") and "needs_seller_quote_for_smaller_qty" in d["flags"]
        order = da.cmd_export(path)["orders"][0]
        assert order["sku"] == "HEAVY1" and order["qty"] == 20 and "needs_seller_quote_for_smaller_qty" in order["flags"]


def test_record_is_bound_to_the_request_it_answered():
    reply = "approve VIC00001 qty 100 max 1.30\ncap 200"
    with sandbox() as tmp:
        path, r1 = request_for(standard_rows(), tmp)
        Clock.t = T0 + timedelta(hours=1)
        r2 = da.cmd_request(path)                                       # same file, a second request, an hour later
        assert (r1["id"], r2["id"]) == ("A-0001", "A-0002")
        f = os.path.join(tmp, "reply.txt")
        open(f, "w", encoding="utf-8").write(reply)
        # 1. two requests are open: there is no default. Nothing is stored and no prompt is reached.
        for attempt in (lambda: record(path, reply), lambda: cli_fails(["record", "--results", path, "--file", f])):
            try:
                msg = attempt()
                assert isinstance(msg, str)
            except da.AgentError as e:
                msg = str(e)
            assert "2 requests are open for r.xlsx" in msg and "A-0001 created 2026-10-05T12:00:00Z" in msg
            assert "A-0002 created 2026-10-05T13:00:00Z" in msg and "--request A-000N" in msg
        assert state()["decisions"] == [] and [e["event"] for e in log()] == ["request", "request"]
        # 2. an explicit request: the echo-back names it, its date, the file and its hash; typing another id records nothing
        plan, stored, out = record(path, reply, request_id="A-0001", typed_id="A-0002")
        assert not stored and "Request not confirmed. Nothing recorded." in out
        for needle in ["RECORDING AGAINST REQUEST A-0001", "Created:      2026-10-05T12:00:00Z", "Results file: r.xlsx",
                       f"File hash:    sha256 {da.file_sha256(path)}", "(picked with --request)"]:
            assert needle in out, needle
        assert "picked by default" not in out and state()["decisions"] == []
        # 3. recording against A-0001 stores it on A-0001 only; A-0002 is still open and unanswered
        _, stored, _ = record(path, reply, request_id="a-1", typed_id="A-0001")          # "a-1" is read as A-0001
        s = state()
        assert stored and {d["request_id"] for d in s["decisions"]} == {"A-0001"}
        assert s["requests"]["A-0001"]["reply"]["text"] == reply and s["requests"]["A-0002"]["reply"] is None
        # 4. exactly one request is open now (A-0002), so it may be the default - and the echo says so, with its own date
        _, stored, out = record(path, "approve FOA-020C qty 10 max 0.30\ncap 50", typed_id="A-0002")
        assert stored and "RECORDING AGAINST REQUEST A-0002" in out and "Created:      2026-10-05T13:00:00Z" in out
        assert "(picked by default: it is the only open request for this results file)" in out
        assert {d["sku"]: d["request_id"] for d in state()["decisions"]} == {"VIC00001": "A-0001", "FOA-020C": "A-0002"}
        # 5. nothing is open now; unknown ids, ids for another file and answered requests are all refused
        other = os.path.join(tmp, "other.xlsx")
        make_results(other, standard_rows())
        r3 = da.cmd_request(other)
        for kwargs, why in (({}, "No open request"), ({"request_id": "A-0009"}, "No request A-0009. Requests so far: A-0001, A-0002, A-0003"),
                            ({"request_id": r3["id"]}, "was written for other.xlsx, not r.xlsx"),
                            ({"request_id": "A-0001"}, "already has a recorded reply")):
            try:
                record(path, reply, **kwargs)
                raise AssertionError(f"accepted {kwargs}")
            except da.AgentError as e:
                assert why in str(e), (kwargs, str(e))
        assert len(state()["decisions"]) == 2
        # 6. through the CLI: --request is how a reply is bound
        answers("A-0003", "Slack", "Neeraj", "yes")
        out = run_cli(["record", "--results", other, "--request", "A-0003", "--file", f])
        assert "Recorded. Nothing has been ordered." in out and "RECORDING AGAINST REQUEST A-0003" in out


# ---------- the seller's quote (read from the email agent's state; no LLM) ----------

def qentry(**fields):
    """An email_state.json entry as the email agent now writes it: every quote field present, null unless stated."""
    entry = {"status": "pricing_provided", "summary": "stub", "price_basis": None, "currency": None, "evidence": {},
             "quote_warnings": [], "sample_available": "unknown", "branding_possible": "unknown"}
    entry.update(fields)
    return entry


def quote_line(row_kwargs, **fields):
    """Request line for one seller whose email-state entry carries `fields`."""
    with sandbox(email_state={key_for("Q1", "Acme"): qentry(**fields)}) as tmp:
        return request_for([res_row("Q1", "Acme", **row_kwargs)], tmp)[1]["lines"][0]


def test_large_total_is_100_dollars():
    assert da.LARGE_TOTAL == D("100")
    with sandbox() as tmp:
        _, req = request_for([res_row("HDFF", "Farsince", unit=0.48, moq="500 Pieces")], tmp)      # $240 at the MOQ
        line = req["lines"][0]
        assert line["moq_heavy"] and line["qty"] is None and req["left_out"] == ["HDFF"] and req["subtotal"] == "0.00"


def test_the_quote_is_shown_on_the_line():
    full = dict(sample_available="yes", sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD",
                moq=100, bulk_tiers=[{"min_qty": 100, "unit_price": "0.25", "evidence": "x"}, {"min_qty": 500, "unit_price": "0.20", "evidence": "y"}],
                branding_possible="yes", branding_fee="30.00", branding_min_qty=50, lead_time_days=5, production_lead_time_days=20,
                shipping_cost="35.00", shipping_terms="FOB", quote_valid_until="2026-11-15",
                quote_warnings=["bulk_tiers[2].min_qty: evidence quote not found in the reply text"])
    with sandbox(email_state={key_for("Q1", "Acme"): qentry(**full)}) as tmp:
        _, req = request_for([res_row("Q1", "Acme", unit=0.25, lcom=10.0, moq="100")], tmp)
        text = da.render_text(req)
        for needle in ["Seller quote:", "Sample: samples available: yes, qty 10, $0.30 per unit",
                       "Bulk: >= 100 pcs $0.25; >= 500 pcs $0.20", "MOQ 100 | branding yes (fee $30.00, min qty 50) | lead time sample 5 days, production 20 days",
                       "shipping $35.00 FOB | valid until 2026-11-15 | currency USD",
                       "Quote warnings (fields the email agent could not verify were left blank): bulk_tiers[2].min_qty"]:
            assert needle in text, needle
        assert "no seller quote yet" not in text
    # branding "no", unclear basis and a total fee read correctly too
    line = quote_line({}, sample_quantity=5, sample_total_price="45.00", price_basis="total_for_quantity", currency="USD", branding_possible="no")
    assert "Sample: qty 5, $45.00 total" in "\n".join(da.quote_block(line)) and "branding no" in "\n".join(da.quote_block(line))
    assert "(price basis unclear)" in "\n".join(da.quote_block(quote_line({}, sample_quantity=5, sample_unit_price="15", price_basis="unclear", currency="USD")))


def test_suggested_quantity_and_price_come_from_the_quote():
    # (listing unit, listing MOQ, quote fields, expected qty, expected unit, expected total, expected qty_source start, flag)
    table = [
        (1.25, "40", dict(sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD"), 10, "0.30", "3.00", "seller's quoted sample", None),
        (1.25, "40", dict(sample_quantity=3, sample_unit_price="0.335", price_basis="per_unit", currency="USD"), 3, "0.335", "1.01", "seller's quoted sample", None),   # half-up
        (1.25, "500", dict(sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD"), 10, "0.30", "3.00", "seller's quoted sample", None),    # no longer MOQ-heavy
        (1.25, "40", dict(sample_quantity=10, sample_unit_price="0.30", price_basis="unclear", currency="USD"), 10, "1.25", "12.50", "seller's stated sample quantity", "quote_price_basis_unclear"),
        (1.25, "40", dict(sample_quantity=10, sample_unit_price="0.30", price_basis="total_for_quantity", currency="USD"), 10, "1.25", "12.50", "seller's stated sample quantity", "quote_price_basis_unclear"),
        (1.25, "40", dict(sample_unit_price="0.30", price_basis="per_unit", currency="USD"), 40, "1.25", "50.00", "listing MOQ", None),             # no quoted qty: the previous rules
        (1.25, "40", dict(sample_quantity=10), 10, "1.25", "12.50", "seller's stated sample quantity", None),                                          # no quoted price: the previous rules
        (1.25, "40", dict(sample_quantity=10, sample_total_price="45.00", price_basis="total_for_quantity", currency="USD"), 10, "1.25", "12.50", "seller's stated sample quantity", None),
        (1.25, "40", dict(moq=100, branding_possible="yes"), 40, "1.25", "50.00", "listing MOQ", None),                                              # a quote with no price or qty at all
    ]
    for unit, moq, fields, qty, price, total, source, flag in table:
        line = quote_line(dict(unit=unit, moq=moq), **fields)
        assert (line["qty"], line["suggested_unit_price"], line["total"]) == (qty, price, total), (fields, line["qty"], line["suggested_unit_price"], line["total"])
        assert line["qty_source"].startswith(source), (fields, line["qty_source"])
        assert (flag in line["flags"]) if flag else "quote_price_basis_unclear" not in line["flags"], fields
        assert line["unit_price"] == str(D(str(unit)))                                # the listed price is still recorded
    text = da.render_text(quote_request(dict(sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD")))
    assert "Suggested qty: 10 (seller's quoted sample quantity and unit price) at $0.30 | Line total: $3.00" in text


def quote_request(fields, row_kwargs=None, tmp=None):
    with sandbox(email_state={key_for("Q1", "Acme"): qentry(**fields)}) as t2:
        return request_for([res_row("Q1", "Acme", **(row_kwargs or {"unit": 1.25, "moq": "40"}))], t2)[1]


def test_margin_check_uses_the_bulk_price_never_the_sample_price():
    cheap_sample = dict(sample_quantity=5, sample_unit_price="0.50", price_basis="per_unit", currency="USD")
    pricey_sample = dict(sample_quantity=5, sample_unit_price="5.00", price_basis="per_unit", currency="USD")
    tier = lambda *ts: [{"min_qty": q, "unit_price": p, "evidence": "e"} for q, p in ts]
    # (quote fields, expected basis starts with, expected pct, clears?, expects the 'no longer clears' flag)   lcom = 10.00
    table = [
        (dict(currency="USD", bulk_tiers=tier((100, "1.80"), (500, "1.20"))), "quoted bulk price $1.80 at >= 100 pcs", "82.0", True, False),   # lowest tier
        (dict(currency="USD", bulk_tiers=tier((500, "1.20"), (100, "1.80"))), "quoted bulk price $1.80 at >= 100 pcs", "82.0", True, False),   # order-independent
        (dict(currency="USD", bulk_tiers=tier((100, "2.50"))), "quoted bulk price $2.50 at >= 100 pcs", "75.0", False, True),
        (dict(currency="USD", bulk_tiers=tier((100, "2.00"))), "quoted bulk price $2.00 at >= 100 pcs", "80.0", True, False),                 # exactly 80%
        (dict(currency="USD", bulk_tiers=tier((100, "2.01"))), "quoted bulk price $2.01 at >= 100 pcs", "79.9", False, True),
        (pricey_sample, "listing price, not confirmed by the seller", "87.5", True, False),   # a $5.00 sample would fail 80% - never checked
        (cheap_sample, "listing price, not confirmed by the seller", "87.5", True, False),
        (dict(currency="CNY", bulk_tiers=tier((100, "30"))), "listing price, not confirmed by the seller", "87.5", True, False),       # non-USD: ignored
        (dict(branding_possible="yes", moq=100), "listing price, not confirmed by the seller", "87.5", True, False),
    ]
    for fields, basis, pct, clears, flagged in table:
        line = quote_line(dict(unit=1.25, lcom=10.0, moq="40"), **fields)
        m = line["margin"]
        assert m["basis"].startswith(basis) and m["pct"] == pct and m["clears"] is clears, (fields, m)
        assert ("quote_no_longer_clears_margin_bar" in line["flags"]) is flagged, fields
        text = "\n".join(da.line_text(1, line))
        if line["quote_comparisons"] and not m["quoted"]:           # a quote well above the listing and no bulk price (see its own test)
            assert "Margin check: bulk price not quoted; the listing price may not hold" in text and "clears the 80% bar" not in text
        else:
            assert (f"Margin check: {pct}% below L-Com on the {m['basis']} - " + ("clears" if clears else "DOES NOT clear")) in text
        assert ("QUOTE NO LONGER CLEARS THE MARGIN BAR" in text) is flagged
    no_quote = quote_line(dict(unit=1.25, lcom=10.0, moq="40"), status="no_reply")
    assert no_quote["margin"]["basis"] == "listing price, not confirmed by the seller"
    no_lcom = quote_line(dict(unit=1.25, lcom=None, moq="40"), currency="USD", bulk_tiers=tier((100, "1.80")))
    assert no_lcom["margin"]["basis"] == "n/a (no L-Com price)" and no_lcom["margin"]["pct"] is None
    assert "Margin check: n/a (no L-Com price)" in "\n".join(da.line_text(1, no_lcom))
    assert da.MIN_MARGIN_PCT == D("80")


def test_a_quote_more_than_ten_percent_above_the_listing_is_flagged():
    assert da.QUOTE_PRICE_TOLERANCE == D("0.10")
    base = dict(price_basis="per_unit", currency="USD", sample_quantity=5)
    tier = lambda q, p: [{"min_qty": q, "unit_price": p, "evidence": "e"}]
    table = [   # (quote fields, listed unit price, expected comparison texts)
        (dict(base, sample_unit_price="1.10"), 1.00, []),                                               # exactly +10%: not flagged
        (dict(base, sample_unit_price="1.11"), 1.00, ["quoted sample unit price $1.11 vs listed $1.00 (+11.0%)"]),
        (dict(base, sample_unit_price="0.90"), 1.00, []),                                               # cheaper than listed
        (dict(base, sample_unit_price="1.00", bulk_tiers=tier(100, "1.20")), 1.00, ["quoted bulk price at >= 100 pcs $1.20 vs listed $1.00 (+20.0%)"]),
        (dict(base, sample_unit_price="1.50", bulk_tiers=tier(100, "1.05")), 1.00, ["quoted sample unit price $1.50 vs listed $1.00 (+50.0%)"]),
        (dict(base, currency="CNY", sample_unit_price="22.50"), 1.00, []),                              # non-USD: never compared
    ]
    for fields, listed, want in table:
        line = quote_line(dict(unit=listed, lcom=100.0, moq="40"), **fields)
        assert line["quote_comparisons"] == want, (fields, line["quote_comparisons"])
        assert ("quote_above_listed_price" in line["flags"]) == bool(want)
        text = "\n".join(da.line_text(1, line))
        for comparison in want:
            assert f"QUOTE ABOVE LISTED PRICE (more than 10%): {comparison}" in text


def test_non_usd_or_unknown_currency_is_left_out_of_all_arithmetic():
    rows = [res_row("USD1", "Acme", unit=1.0, moq="40"), res_row("RMB1", "Acme", unit=1.0, moq="40", lcom=10.0)]
    email = {key_for("RMB1", "Acme"): qentry(sample_quantity=10, sample_unit_price="22.50", price_basis="per_unit", currency="CNY",
                                             bulk_tiers=[{"min_qty": 100, "unit_price": "1.00", "evidence": "e"}], shipping_cost="80")}
    with sandbox(email_state=email) as tmp:
        path, req = request_for(rows, tmp)
        by = {l["sku"]: l for l in req["lines"]}
        l = by["RMB1"]
        assert (l["qty"], l["total"], l["no_total_reason"]) == (None, None, "quoted in CNY; needs conversion")
        assert "quote_not_usd" in l["flags"] and "quote_above_listed_price" not in l["flags"]
        assert l["margin"]["basis"] == "listing price, not confirmed by the seller" and l["quote_comparisons"] == []
        assert D(req["subtotal"]) == D("40.00")                                      # only USD1: the RMB line adds nothing
        assert req["lines_without_total"] == ["RMB1"] and req["no_total_reasons"] == {"RMB1": "quoted in CNY; needs conversion"}
        text = da.render_text(req)
        assert "no qty for RMB1 (quoted in CNY; needs conversion)" in text
        assert "Quoted in CNY; needs conversion - the quote is left out of every total and margin check" in text
        assert "Suggested qty: none - quoted in CNY; needs conversion (not in the subtotal)" in text
        assert "Sample: qty 10, 22.50 CNY per unit" in text and "shipping 80 CNY" in text         # shown, never as dollars
        _, stored, _ = record(path, "approve RMB1 qty 10 max 1.00\ncap 50")                       # it can still be approved explicitly
        assert stored
    for currency in ("EUR", "unknown", None):                                       # EUR, unknown and a missing currency on a priced quote
        line = quote_line(dict(unit=1.0, moq="40"), sample_quantity=10, sample_unit_price="5", price_basis="per_unit", currency=currency)
        assert line["total"] is None and "quote_not_usd" in line["flags"], currency
        assert line["no_total_reason"] == f"quoted in {currency or 'unknown'}; needs conversion"
    quantity_only = quote_line(dict(unit=1.0, moq="40"), sample_quantity=10)         # no prices at all: currency is irrelevant
    assert quantity_only["total"] == "10.00" and "quote_not_usd" not in quantity_only["flags"]


def test_lines_the_seller_is_waiting_on_us_for_get_their_own_section():
    rows = [res_row("OK1", "Fine Co", unit=1.0, moq="40"),
            res_row("WAIT1", "Waiting Co", unit=1.0, moq="40"),
            res_row("WAITHEAVY", "Both Co", unit=1.0, moq="500"),                    # MOQ-heavy AND waiting: waiting wins
            res_row("HEAVY1", "Heavy Co", unit=1.0, moq="500")]
    email = {key_for("WAIT1", "Waiting Co"): qentry(status="needs_info", needs_from_us=["the drawing", "your planned order quantity"]),
             key_for("WAITHEAVY", "Both Co"): qentry(status="needs_info", needs_from_us=["shielding material"])}
    with sandbox(email_state=email) as tmp:
        _, req = request_for(rows, tmp)
        assert req["waiting"] == ["WAIT1", "WAITHEAVY"] and req["left_out"] == ["HEAVY1"] and D(req["subtotal"]) == D("40.00")
        assert req["lines_without_total"] == []                                        # waiting lines are not "no qty" lines
        text = da.render_text(req)
        head, _, rest = text.partition("\n\n")
        assert "Waiting on us: 2 line(s) (WAIT1, WAITHEAVY) - the seller asked us for something first; not in the subtotal." in head
        assert "Left out of that subtotal: 1 line(s) whose total at the MOQ is over $100.00 (HEAVY1)" in head
        normal, _, after = rest.partition("NEEDS A SELLER QUOTE FOR A SMALLER QUANTITY FIRST")
        heavy, _, waiting = after.partition("WAITING ON US (the seller asked us for something before it can quote; these are not in the subtotal):")
        assert "OK1" in normal and "WAIT1" not in normal and "HEAVY1" not in normal         # three sections, in this order
        assert "HEAVY1" in heavy and "WAIT" not in heavy
        assert waiting.index("WAIT1") < waiting.index("WAITHEAVY") < waiting.index("HOW TO REPLY")
        assert "Seller asked us for: the drawing; your planned order quantity" in waiting and "Seller asked us for: shielding material" in waiting
        assert "Suggested qty: none - waiting on us (not in the subtotal)" in waiting and "waiting_on_us" in waiting
        assert "Seller quote:\n     no prices or quantities quoted yet" in waiting             # no empty "Seller quote:" header
        assert [l.split(")")[0] for l in rest.splitlines() if re.match(r"\d+\) ", l)] == ["1", "2", "3", "4"]


def test_no_reply_yet_says_so():
    rows = [res_row("A1", "Acme"), res_row("B2", "Beta")]
    for email in (None, {}, {key_for("B2", "Beta"): {"status": "needs_info", "summary": "asked something"}}):   # no state / no entry / status only
        with sandbox(email_state=email) as tmp:
            _, req = request_for(rows, tmp)
            for l in req["lines"]:
                assert l["quote"] is None
            text = da.render_text(req)
            assert text.count("Seller quote: no seller quote yet - listing price only") == 2, email
            assert "Seller quote:\n" not in text
    # entries whose only content is "unknown" / empty are not a quote either
    assert da.entry_quote({"status": "pricing_provided", "sample_available": "unknown", "branding_possible": "unknown",
                           "bulk_tiers": [], "needs_from_us": None}) is None
    assert da.entry_quote("not a dict") is None and da.entry_quote({"sample_quantity": 5})["sample_quantity"] == 5


def test_export_carries_the_quote_as_it_was_requested():
    fields = dict(sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD",
                  bulk_tiers=[{"min_qty": 100, "unit_price": "0.25", "evidence": "100 pcs at $0.25"}],
                  evidence={"sample_quantity": "10 pcs", "sample_unit_price": "$0.30 each"},
                  quote_warnings=["moq: evidence quote not found in the reply text"])
    email = {key_for("Q1", "Acme"): qentry(**fields)}
    with sandbox(email_state=email) as tmp:
        path, _ = request_for([res_row("Q1", "Acme", unit=0.25, lcom=10.0, moq="40"), res_row("N1", "Plain Co")], tmp)
        record(path, "approve Q1 qty 10 max 0.35\napprove N1 qty 5 max 2\ncap 50")
        open(da.EMAIL_STATE_PATH, "w", encoding="utf-8").write("{}")              # the state changes after the request: no effect
        r = da.cmd_export(path)
        o = {x["sku"]: x for x in r["orders"]}
        q = o["Q1"]
        assert (q["quoted_unit_price"], q["quoted_quantity"], q["quoted_currency"]) == ("0.30", 10, "USD")
        assert q["quoted_bulk_tiers"] == [{"min_qty": 100, "unit_price": "0.25", "evidence": "100 pcs at $0.25"}]
        assert q["quote_evidence"] == {"sample_quantity": "10 pcs", "sample_unit_price": "$0.30 each"}
        assert q["quote_warnings"] == ["moq: evidence quote not found in the reply text"]
        n = o["N1"]
        assert (n["quoted_unit_price"], n["quoted_quantity"], n["quoted_currency"], n["quote_evidence"], n["quote_warnings"]) == (None, None, None, {}, [])
        assert {"quoted_unit_price", "quoted_quantity", "quoted_currency", "quote_evidence", "quote_warnings"} <= set(da.EXPORT_FIELDS)
        data = json.load(open(da.EXPORT_JSON, encoding="utf-8"))
        assert data["orders"][0]["quote_evidence"]["sample_unit_price"] == "$0.30 each"
        rows = list(csv.DictReader(open(da.EXPORT_CSV, encoding="utf-8", newline="").read().splitlines()[1:]))
        csv_q = {x["sku"]: x for x in rows}["Q1"]
        assert csv_q["quoted_unit_price"] == "0.30" and csv_q["quoted_currency"] == "USD"
        assert json.loads(csv_q["quote_evidence"]) == q["quote_evidence"] and json.loads(csv_q["quote_warnings"]) == q["quote_warnings"]
        assert json.loads(csv_q["quoted_bulk_tiers"])[0]["min_qty"] == 100
        assert {x["sku"]: x for x in rows}["N1"]["quoted_unit_price"] == ""


def test_the_quote_is_read_only_and_needs_no_llm():
    email = {key_for("Q1", "Acme"): qentry(sample_quantity=10, sample_unit_price="0.30", price_basis="per_unit", currency="USD")}
    with sandbox(email_state=email) as tmp:
        before = open(da.EMAIL_STATE_PATH, "rb").read()
        path, _ = request_for([res_row("Q1", "Acme")], tmp)
        record(path, "approve Q1 qty 10 max 0.35\ncap 50")
        da.cmd_status(path)
        da.cmd_export(path)
        assert open(da.EMAIL_STATE_PATH, "rb").read() == before


def quote_scenario(tmp):
    """Results rows plus an email_state.json built by the REAL email agent (fixture replies, stub LLM, code verification),
    so a mismatch in field names between the two agents would show up here. Returns the results path."""
    import email_agent as ea
    import test_email_agent as te
    raws = {name: raw for name, _, raw, _, _ in te.QUOTE_CASES}
    scenario = [   # (sku, seller, listing price, L-Com price, listing MOQ, fixture, stub extraction)
        ("ECF504-SC6", "Quote Co", 3.20, 20.54, "100", "quote_bulk_tiers.txt", raws["bulk tiers"]),               # quote-informed, clears the bar
        ("VIC00001", "Pricey Co", 1.25, 28.19, "40", "quote_per_unit.txt", raws["a per-unit sample price"]),       # quoted above the listing
        ("HDFF", "Margin Co", 1.00, 5.00, "40", "quote_bulk_tiers.txt", raws["bulk tiers"]),                       # bulk price fails the 80% bar
        ("TDG1026KS-C6", "Yuan Co", 0.55, 16.39, "20", "quote_rmb.txt", raws["a quote in RMB"]),                   # RMB
        ("FOA-020C", "Waiting Co", 0.25, 72.99, "10", "quote_needs_info.txt", raws["a reply with questions for us"]),   # waiting on us
        ("ECF504-AA", "Silent Co", 0.80, 24.79, "20", None, None),                                                  # no reply yet
    ]
    state, rows = {}, []
    for sku, seller, listing, lcom, moq, fixture, raw in scenario:
        row_ = res_row(sku, seller, unit=listing, lcom=lcom, moq=moq)
        rows.append(row_)
        if fixture:
            draft = ea.make_draft({"sku": sku, "product": "p", "description": "d", "manufacturer": seller, "email": "",
                                   "url": row_["Recommended URL"], "unit_price": None, "moq": "", "title": "", "note": ""})
            assert draft["key"] == key_for(sku, seller)                      # both agents key the entry the same way
            ea.record_reply(state, draft, te.read_fixture(fixture), summarize=lambda text, r={"status": "pricing_provided", "summary": "stub", **(raw or {})}: r)
    open(da.EMAIL_STATE_PATH, "w", encoding="utf-8").write(json.dumps(state, ensure_ascii=False))
    path = os.path.join(tmp, "quotes.xlsx")
    make_results(path, rows)
    return path


def test_quotes_recorded_by_the_email_agent_flow_into_the_request():
    with sandbox(email_state="{}") as tmp:
        path = quote_scenario(tmp)
        req = da.cmd_request(path)
        by = {l["sku"]: l for l in req["lines"]}
        # quote-informed: sample qty 5 at $3.20, bulk tiers shown, margin on the lowest bulk tier ($1.80 at >= 100 pcs)
        q = by["ECF504-SC6"]
        assert (q["qty"], q["suggested_unit_price"], q["total"]) == (5, "3.20", "16.00")
        assert q["margin"]["basis"] == "quoted bulk price $1.80 at >= 100 pcs" and q["margin"]["pct"] == "91.2" and q["margin"]["clears"]
        assert q["quote_comparisons"] == [] and q["flags"] == [] and q["quote"]["moq"] == 100 and q["quote"]["currency"] == "USD"
        # above listed: the sample price is far above the listing; both numbers are shown
        a = by["VIC00001"]
        assert a["quote_comparisons"] == ["quoted sample unit price $3.20 vs listed $1.25 (+156.0%)"] and "quote_above_listed_price" in a["flags"]
        # the bulk price fails the bar: flagged
        m = by["HDFF"]                                    # the quoted bulk price fails the bar: not suggested, out of the subtotal
        assert m["margin"]["clears"] is False and "quote_no_longer_clears_margin_bar" in m["flags"]
        assert (m["qty"], m["total"], m["no_total_reason"]) == (None, None, "quote does not clear the margin bar")
        assert req["margin_failed"] == ["HDFF"] and da.line_section(m) == "margin"
        # RMB: left out of every total, shown in its own currency
        r = by["TDG1026KS-C6"]
        assert r["quote_currency"] == "CNY" and r["total"] is None and r["no_total_reason"] == "quoted in CNY; needs conversion"
        # waiting on us
        w = by["FOA-020C"]
        assert w["waiting_on_us"] == ["drawing", "planned order quantity", "preferred shielding material"] and w["total"] is None
        # no reply yet
        assert by["ECF504-AA"]["quote"] is None and by["ECF504-AA"]["total"] == "16.00"
        # SC6 (5 x $3.20), VIC (5 x $3.20) and AA (20 x $0.80): HDFF is out of the subtotal now, so it is three $16.00 lines
        assert D(req["subtotal"]) == D("16.00") * 3 == D("48.00")
        assert req["waiting"] == ["FOA-020C"] and req["lines_without_total"] == ["TDG1026KS-C6"]
        text = da.render_text(req)
        assert "Product subtotal at suggested quantities: $48.00 (covers 3 of 6 lines)" in text
        assert "WAITING ON US" in text and "Quoted in CNY; needs conversion" in text and "no seller quote yet - listing price only" in text
        assert "QUOTE DOES NOT CLEAR THE MARGIN BAR - NOT SUGGESTED" in text
        # the sample price is far above the listing and there is no bulk price: no "clears the bar" on the listing price
        assert "Margin check: bulk price not quoted; the listing price may not hold" in text.split("2) VIC00001")[1].split("3)")[0]


def test_a_quote_that_misses_the_margin_bar_is_not_suggested():
    tier = lambda p: [{"min_qty": 100, "unit_price": p, "evidence": "e"}]
    cheap_sample = dict(sample_quantity=5, sample_unit_price="0.90", price_basis="per_unit", currency="USD")   # would be suggested if the bulk price passed
    base = dict(unit=1.0, lcom=10.0)
    rows = [res_row("OK1", "Fine Co", moq="40", **base),                                  # no quote
            res_row("FAIL1", "Fail Co", moq="40", **base),                                # bulk $2.50 = 75.0% below L-Com: misses the bar
            res_row("EDGE1", "Edge Co", moq="40", **base),                                # bulk $2.00 = exactly 80.0%: passes
            res_row("EDGE2", "Edge2 Co", moq="40", **base),                               # bulk $2.01 = 79.9%: misses
            res_row("FAILHEAVY", "FH Co", moq="500", **base),                             # misses the bar AND is MOQ-heavy: the margin section wins
            res_row("FAILWAIT", "FW Co", moq="40", **base),                               # misses the bar AND the seller waits on us: waiting wins
            res_row("SAMPLEONLY", "SO Co", moq="40", **base),                             # a $5.00 SAMPLE price never triggers this
            res_row("RMB1", "R Co", moq="40", **base)]                                    # non-USD bulk prices are ignored
    email = {key_for("FAIL1", "Fail Co"): qentry(**dict(cheap_sample, bulk_tiers=tier("2.50"))),
             key_for("EDGE1", "Edge Co"): qentry(**dict(cheap_sample, bulk_tiers=tier("2.00"))),
             key_for("EDGE2", "Edge2 Co"): qentry(**dict(cheap_sample, bulk_tiers=tier("2.01"))),
             key_for("FAILHEAVY", "FH Co"): qentry(currency="USD", bulk_tiers=tier("3.00")),
             key_for("FAILWAIT", "FW Co"): qentry(currency="USD", bulk_tiers=tier("3.00"), needs_from_us=["the drawing"]),
             key_for("SAMPLEONLY", "SO Co"): qentry(sample_quantity=5, sample_unit_price="5.00", price_basis="per_unit", currency="USD"),
             key_for("RMB1", "R Co"): qentry(currency="CNY", bulk_tiers=tier("30"))}
    with sandbox(email_state=email) as tmp:
        path, req = request_for(rows, tmp, budget=D("100"))
        by = {l["sku"]: l for l in req["lines"]}
        sections = {sku: da.line_section(l) for sku, l in by.items()}
        assert sections == {"OK1": "normal", "FAIL1": "margin", "EDGE1": "normal", "EDGE2": "margin", "FAILHEAVY": "margin",
                            "FAILWAIT": "waiting", "SAMPLEONLY": "normal", "RMB1": "normal"}, sections
        for sku in ("FAIL1", "EDGE2", "FAILHEAVY"):                                      # no suggested quantity, nothing in the subtotal
            l = by[sku]
            assert (l["qty"], l["total"], l["no_total_reason"]) == (None, None, "quote does not clear the margin bar"), sku
            assert "quote_no_longer_clears_margin_bar" in l["flags"] and "moq_not_stated" not in l["flags"] and not l["moq_heavy"]
        assert (by["EDGE1"]["qty"], by["EDGE1"]["total"]) == (5, "4.50")                  # exactly 80%: suggested at the sample qty and price
        assert by["FAILWAIT"]["total"] is None and "quote_no_longer_clears_margin_bar" in by["FAILWAIT"]["flags"]
        assert (by["SAMPLEONLY"]["qty"], by["SAMPLEONLY"]["total"]) == (5, "25.00") and "quote_no_longer_clears_margin_bar" not in by["SAMPLEONLY"]["flags"]
        assert by["RMB1"]["no_total_reason"] == "quoted in CNY; needs conversion" and not by["RMB1"]["margin_fail"]
        assert req["margin_failed"] == ["FAIL1", "EDGE2", "FAILHEAVY"] and req["waiting"] == ["FAILWAIT"] and req["left_out"] == []
        assert req["lines_without_total"] == ["RMB1"]                                     # margin-failed lines are not "no qty" lines
        assert D(req["subtotal"]) == D("40.00") + D("4.50") + D("25.00")                  # OK1 + EDGE1 + SAMPLEONLY
        text = da.render_text(req)
        head, _, rest = text.partition("\n\n")
        assert "Product subtotal at suggested quantities: $69.50 (covers 3 of 8 lines); no qty for RMB1 (quoted in CNY; needs conversion)" in head
        assert ("Left out of that subtotal: 3 line(s) (FAIL1, EDGE2, FAILHEAVY) whose quoted bulk price does not clear the 80% margin bar "
                "- not suggested.") in head
        assert "subtotal (priced lines only) is within budget ($30.50 to spare)" in head
        title = "QUOTE DOES NOT CLEAR THE MARGIN BAR - NOT SUGGESTED (the quoted bulk price is not at least 80% below L-Com; no quantity is suggested and these are not in the subtotal):"
        normal, _, after = rest.partition(title)
        assert after and "FAIL1" not in normal and "EDGE2" not in normal and "FAILHEAVY" not in normal       # only inside their section
        section, _, waiting = after.partition("WAITING ON US")
        assert section.index("FAIL1") < section.index("EDGE2") < section.index("FAILHEAVY") and "FAILWAIT" in waiting
        assert "NEEDS A SELLER QUOTE" not in text                                         # FAILHEAVY is listed once, in the margin section
        one = section.split("FAIL1 - ")[1].split("\n\n")[0]
        assert "Suggested qty: none - the quote does not clear the margin bar (not in the subtotal)" in one
        assert "QUOTE NO LONGER CLEARS THE MARGIN BAR" in one and "Line total" not in one
        assert [l.split(")")[0] for l in rest.splitlines() if re.match(r"\d+\) ", l)] == [str(i) for i in range(1, 9)]
        md = open(os.path.join(da.REPORTS_DIR, "approval_request_A-0001.md"), encoding="utf-8").read()
        assert "QUOTE DOES NOT CLEAR THE MARGIN BAR - NOT SUGGESTED" in md
        assert "'does not clear the margin bar' section can still be approved" in text
        # it can still be approved with an explicit qty and max (and not without them), and it exports like any other line
        _, stored, out = record(path, "approve FAIL1 qty 10 max 2.50\ncap 30", request_id="A-0001")
        assert stored and "APPROVE FAIL1 (Fail Co): qty 10 x max $2.50 = $25.00 at most" in out
        order = da.cmd_export(path)["orders"][0]
        assert (order["sku"], order["qty"], order["max_unit_price"]) == ("FAIL1", 10, "2.50")
        assert "quote_no_longer_clears_margin_bar" in order["flags"] and order["quoted_unit_price"] == "0.90"
        for bad in ("approve FAIL1\ncap 30", "approve FAIL1 qty 10\ncap 30"):
            try:
                da.parse_reply(bad)
                raise AssertionError("accepted " + bad)
            except da.ReplyError:
                pass


def test_a_quote_above_the_listing_without_a_bulk_price_does_not_claim_the_bar_clears():
    hold = "Margin check: bulk price not quoted; the listing price may not hold"
    tier = lambda p: [{"min_qty": 100, "unit_price": p, "evidence": "e"}]
    usd = dict(price_basis="per_unit", currency="USD", sample_quantity=5)
    # (row kwargs, quote fields, expected listing_may_not_hold, expected quote_above_listed_price flag)   listing $1.00, L-Com $10.00
    table = [
        (dict(), dict(usd, sample_unit_price="1.11"), True, True),                          # +11%, no bulk price: the listing price may not hold
        (dict(), dict(usd, sample_unit_price="5.00"), True, True),
        (dict(), dict(usd, sample_unit_price="1.10"), False, False),                        # exactly +10%: neither
        (dict(), dict(usd, sample_unit_price="0.50"), False, False),                        # cheaper than listed
        (dict(), dict(usd, sample_unit_price="1.50", bulk_tiers=tier("1.30")), False, True),     # a bulk price came with it: it is the one checked
        (dict(), dict(currency="USD", bulk_tiers=tier("1.30")), False, True),               # bulk-only quote above the listing
        (dict(), dict(usd, currency="CNY", sample_unit_price="22.50"), False, False),       # non-USD: never compared
        (dict(lcom=None), dict(usd, sample_unit_price="1.50"), False, True),                # no L-Com price: there is no listing claim to retract
    ]
    for row_kwargs, fields, expect_hold, expect_flag in table:
        line = quote_line(dict(dict(unit=1.0, lcom=10.0, moq="40"), **row_kwargs), **fields)
        text = "\n".join(da.line_text(1, line))
        assert line["margin"]["listing_may_not_hold"] is expect_hold, (fields, line["margin"])
        assert ("quote_above_listed_price" in line["flags"]) is expect_flag, fields          # the existing flag is kept
        assert (hold in text) is expect_hold, (fields, text)
        if expect_hold:
            assert "clears the 80% bar" not in text and "% below L-Com on the listing price" not in text
            assert "QUOTE ABOVE LISTED PRICE (more than 10%)" in text                        # and both numbers are still shown
            assert line["qty"] == 5 and line["total"] is not None                           # the line itself is still suggested
        elif line["margin"]["pct"] is not None and not line["quote_comparisons"] and fields.get("currency") != "CNY":
            assert "clears the 80% bar" in text
    # no quote at all keeps the old wording
    plain = quote_line(dict(unit=1.0, lcom=10.0, moq="40"), status="no_reply")
    assert "Margin check: 90.0% below L-Com on the listing price, not confirmed by the seller - clears the 80% bar" in "\n".join(da.line_text(1, plain))


def test_product_name_falls_back_to_the_keyword():
    """Built-in products come with no Product name in the sheet; the line must still say what it is."""
    with sandbox() as tmp:
        blank = res_row("FOA-020C", "FiberMania", unit=0.25, moq="10")
        blank["Product"] = None
        _, req = request_for([blank], tmp)
        assert req["lines"][0]["product"] == "FOA-020C keyword" and "1) FOA-020C - FOA-020C keyword" in da.render_text(req)
        named = res_row("VIC00001", "Eternalstar")
        _, req = request_for([named], tmp, "n.xlsx")
        assert req["lines"][0]["product"] == "VIC00001 product"                       # a real name still wins


def test_request_skips_rows_and_says_why():
    rows = [res_row("OK1", "Fine Co"),
            res_row("ERR1", rec="Error researching this product: RuntimeError: boom - re-run with --sku", maker=None, url=""),
            res_row("NONE1", rec="No candidate qualifies - accuracy only", maker=None, url="", unit=None),
            res_row("NONE2", rec="No candidates found - nothing to source.", maker=None, url="", unit=None),
            res_row("EXCL1", "Bad Co", url="http://shop/bad"),
            res_row("EXCL2", "Shenzhen Xiangtianzhong Tech", url="http://shop/x"),
            res_row("EXCL1B", "Bad Co", url="http://shop/bad"),     # same URL, other SKU: not excluded
            res_row("DECL1", "Declined Co"),
            res_row("NOPRICE", "Free Co", unit=None),
            res_row("NOPRICE2", "Zero Co", unit=0)]
    email = {key_for("DECL1", "Declined Co"): {"status": "dead_end", "summary": "discontinued"},
             key_for("OK1", "Fine Co"): {"status": "auto_reply_or_spam"}}      # a spam auto-reply is not a decline
    excl = [("EXCL1", "http://shop/bad", "bad listing"), ("EXCL2", "xiangtianzhong", "short cable")]
    with sandbox(exclusions=excl, email_state=email) as tmp:
        before = (open(da.REVIEWER_EXCLUSIONS_CSV, "rb").read(), open(da.EMAIL_STATE_PATH, "rb").read())
        path, req = request_for(rows, tmp)
        assert [l["sku"] for l in req["lines"]] == ["OK1", "EXCL1B"]
        kinds = {s["sku"]: s["kind"] for s in req["skipped"]}
        assert kinds == {"ERR1": "error", "NONE1": "none", "NONE2": "none", "EXCL1": "reviewer", "EXCL2": "reviewer",
                         "DECL1": "declined", "NOPRICE": "noprice", "NOPRICE2": "noprice"}
        summary = da.skip_summary(req["skipped"])
        assert summary == ("Skipped 8 row(s): 2 no recommended seller (NONE1, NONE2); 1 search agent error (ERR1); "
                           "2 excluded by reviewer (EXCL1, EXCL2); 1 seller declined (DECL1); 2 no usable unit price (NOPRICE, NOPRICE2).")
        assert summary in da.render_text(req)
        assert (open(da.REVIEWER_EXCLUSIONS_CSV, "rb").read(), open(da.EMAIL_STATE_PATH, "rb").read()) == before   # read only


def test_email_state_formats():
    row = [res_row("OK1", "Fine Co")]
    cases = [(None, "not found"), ("{not json", "not valid JSON"), ("[1, 2]", "format not recognised"),
             ('{"A|x": "a string"}', "format not recognised"), ('{"_source": {"results": "x"}}', ""), ("{}", "")]
    for content, note in cases:
        with sandbox(email_state=content) as tmp:
            _, req = request_for(row, tmp)
            assert [l["sku"] for l in req["lines"]] == ["OK1"], content               # always carries on
            assert (note in req["email_note"]) if note else req["email_note"] == "", (content, req["email_note"])
            if content and note in ("not valid JSON", "format not recognised"):
                assert "ignored" in req["email_note"] and "Note: " + req["email_note"] in da.render_text(req)
            if note:
                assert req["lines"][0]["seller_reply"] == "unknown (no email state)"


# ---------- results file format ----------

def test_missing_header_stops_with_a_clear_message():
    with sandbox() as tmp:
        for header in da.REQUIRED_HEADERS:
            path = os.path.join(tmp, "m.xlsx")
            make_results(path, standard_rows(), drop=[header])
            msg = cli_fails(["request", path])
            assert header in msg and "missing expected column" in msg and "Nothing was read" in msg, (header, msg)
        assert not os.path.exists(da.STATE_PATH) and not log()                       # nothing stored
        both = os.path.join(tmp, "two.xlsx")
        make_results(both, standard_rows(), drop=["Recommended URL", "Ordering Note"])
        assert "Recommended URL, Ordering Note" in cli_fails(["status", "--results", both])
        wb = openpyxl.Workbook()
        wb.active.title = "Other"
        wb.save(os.path.join(tmp, "nosheet.xlsx"))
        assert "no 'Results' sheet" in cli_fails(["request", os.path.join(tmp, "nosheet.xlsx")])
        assert "not found" in cli_fails(["request", os.path.join(tmp, "absent.xlsx")])


def test_columns_are_read_by_name_not_position():
    with sandbox() as tmp:
        straight, reversed_ = os.path.join(tmp, "a.xlsx"), os.path.join(tmp, "b.xlsx")
        make_results(straight, standard_rows())
        make_results(reversed_, standard_rows(), reverse=True)
        a = da.build_request(straight)["lines"]
        b = da.build_request(reversed_)["lines"]
        assert a == b and len(a) == 4
        extra = os.path.join(tmp, "c.xlsx")
        wb = openpyxl.load_workbook(straight)
        wb.active.insert_cols(1)                                                  # a new first column shifts everything
        wb.active.cell(1, 1).value = "Brand New Column"
        wb.save(extra)
        assert da.build_request(extra)["lines"] == a


# ---------- recording a reply ----------

def test_record_needs_explicit_yes_and_stores_the_exact_text():
    reply = "approve VIC00001 qty 100 max 1.30\n  Approve foa-020c qty 50 max $0.30\nreject HDFF\ncap $200\n"
    with sandbox() as tmp:
        path, req = request_for(standard_rows(), tmp)
        for confirm in ("", "y", "YES PLEASE", "ye", "no", "yes!"):             # anything but the word yes records nothing
            plan, stored, out = record(path, reply, confirm=confirm)
            assert not stored and "Nothing recorded." in out, confirm
            assert state()["decisions"] == [] and state()["requests"]["A-0001"]["reply"] is None
        for channel, approver in (("", "Neeraj"), ("Slack", "  ")):             # channel and approver are both required
            _, stored, _ = record(path, reply, channel=channel, approver=approver)
            assert not stored
        for typed in ("", "yes", "A-0002", "A-001", "request A-0001"):          # the request id must be typed exactly
            _, stored, out = record(path, reply, typed_id=typed)
            assert not stored and "Request not confirmed. Nothing recorded." in out, typed
        assert [e["event"] for e in log()] == ["request"] and state()["decisions"] == []
        plan, stored, out = record(path, reply, channel="Slack DM", approver="Neeraj K.", confirm="Yes", typed_id="a-0001")
        assert stored
        for needle in ["RECORDING AGAINST REQUEST A-0001", "Created:      2026-10-05T12:00:00Z", "Results file: r.xlsx",
                       f"File hash:    sha256 {da.file_sha256(path)}", "(picked by default: it is the only open request",
                       "UNDERSTOOD for request A-0001", "APPROVE VIC00001 (Eternalstar): qty 100 x max $1.30 = $130.00 at most",
                       "APPROVE FOA-020C (FiberMania): qty 50 x max $0.30 = $15.00 at most", "REJECT  HDFF (Farsince)",
                       "Cap: $200.00 | Committed at max prices: $145.00 | Left for shipping, duties and taxes: $55.00",
                       "No decision given for: BIG1", "the sender cannot be verified", "Nothing is ordered"]:
            assert needle in out, needle
        s = state()
        r = s["requests"]["A-0001"]["reply"]
        assert r["text"] == reply and r["channel"] == "Slack DM" and r["approver"] == "Neeraj K." and r["sender_verified"] is False
        assert r["recorded_at"] == "2026-10-05T12:00:00Z" and r["cap"] == "200.00"
        by = {d["sku"]: d for d in s["decisions"]}
        v = by["VIC00001"]
        assert (v["decision"], v["qty"], v["max_unit_price"], v["line_cap"]) == ("approve", 100, "1.30", "130.00")
        assert v["expires_at"] == "2026-10-12T12:00:00Z" and v["id"] == "A-0001-1" and v["key"] == key_for("VIC00001", "Eternalstar")
        assert by["HDFF"]["decision"] == "reject" and "expires_at" not in by["HDFF"]
        assert [e["event"] for e in log()] == ["request", "reply_recorded", "approval", "approval", "rejection"]
        assert log()[1]["text"] == reply and log()[1]["sender_verified"] is False
        # a request answers once; the file must still be the one Neeraj saw
        try:
            record(path, reply, request_id="A-0001")
            raise AssertionError("second reply accepted")
        except da.AgentError as e:
            assert "already has a recorded reply" in str(e)
        try:
            record(path, reply)                                                     # by default: nothing is open any more
            raise AssertionError("second reply accepted by default")
        except da.AgentError as e:
            assert "No open request" in str(e) and "Run 'request' first" in str(e)
        path2, _ = request_for(standard_rows(), tmp, "r2.xlsx")
        make_results(path2, standard_rows()[:2])                                   # the file changes after the request
        try:
            record(path2, reply, typed_id="A-0002")
            raise AssertionError("accepted a changed file")
        except da.AgentError as e:
            assert "hash mismatch" in str(e)
        try:
            da.cmd_record(os.path.join(tmp, "never_requested.xlsx"), reply)
            raise AssertionError("no request")
        except da.AgentError as e:
            assert "Run 'request' first" in str(e)


def test_rejected_replies_store_nothing_through_the_cli():
    table = INVALID_REPLIES + [("approve NOPE qty 5 max 1\ncap 50", "is not in the request"),
                               ("approve VIC00001 qty 100 max 1.30\ncap 100", "exceeds the cap"),
                               ("approve VIC00001 qty 1 max 1\nreject vic00001\ncap 9", "appears twice")]
    with sandbox() as tmp:
        path, _ = request_for(standard_rows(), tmp)
        before = (open(da.STATE_PATH, "rb").read(), open(da.LOG_PATH, "rb").read())
        for text, why in table:
            f = os.path.join(tmp, "reply.txt")
            open(f, "w", encoding="utf-8").write(text)
            msg = cli_fails(["record", "--results", path, "--file", f])               # no prompt may be reached
            assert "Reply rejected - nothing stored" in msg and why in msg, (text, msg)
            assert (open(da.STATE_PATH, "rb").read(), open(da.LOG_PATH, "rb").read()) == before, text


def test_max_below_listed_price_warns_but_is_recorded():
    with sandbox() as tmp:
        path, _ = request_for(standard_rows(), tmp)
        _, stored, out = record(path, "approve VIC00001 qty 10 max 1.00\ncap 50")
        assert stored and "WARNING: max $1.00 is below the listed unit price $1.25 - this line cannot be ordered at the listed price." in out
        assert state()["decisions"][0]["max_unit_price"] == "1.00"


def test_utf8_stdin_and_export_roundtrip():
    sku = "ÉCF-ÜBER1"
    reply = f"approve {sku} qty 10 max 2.00\ncap 25\n"
    with sandbox() as tmp:
        path, _ = request_for([res_row(sku, "Shénzhèn 深圳 Co", unit=1.5, moq="10")], tmp)
        answers("A-0001", "Slack 群", "李雷", "yes")
        out = run_cli(["record", "--results", path], stdin_bytes=reply.encode("utf-8"))      # a cp1252 pipe, as on Windows
        assert "Recorded. Nothing has been ordered." in out and f"APPROVE {sku}" in out
        s = state()
        assert s["requests"]["A-0001"]["reply"]["text"] == reply and s["decisions"][0]["approver"] == "李雷"
        run_cli(["export", "--results", path])
        data = json.load(open(da.EXPORT_JSON, encoding="utf-8"))
        assert data["orders"][0]["sku"] == sku and data["orders"][0]["seller"] == "Shénzhèn 深圳 Co"
        rows = list(csv.reader(open(da.EXPORT_CSV, encoding="utf-8", newline="")))
        assert rows[2][0] == sku and "李雷" in rows[2] and "Slack 群" in rows[2]
        assert any(e.get("approver") == "李雷" for e in log())                   # the log is UTF-8 too
        raw = open(da.LOG_PATH, "rb").read().decode("utf-8")
        assert "李雷" in raw                                                     # stored as characters, not \\u escapes


# ---------- expiry, revoke, drift ----------

def test_expiry_and_revoke():
    reply = "approve VIC00001 qty 100 max 1.30\napprove FOA-020C qty 50 max 0.30\ncap 200"
    with sandbox() as tmp:
        path, _ = request_for(standard_rows(), tmp)
        record(path, reply)
        def statuses():
            return {r["sku"]: r["status"] for r in da.cmd_status(path)[0]}
        def exported():
            return sorted(o["sku"] for o in da.cmd_export(path)["orders"])
        assert statuses()["VIC00001"] == "approved" and exported() == ["FOA-020C", "VIC00001"]
        Clock.t = T0 + timedelta(days=da.APPROVAL_TTL_DAYS) - timedelta(seconds=1)
        assert statuses()["VIC00001"] == "approved" and exported() == ["FOA-020C", "VIC00001"]    # one second left
        Clock.t = T0 + timedelta(days=da.APPROVAL_TTL_DAYS)
        assert statuses()["VIC00001"] == "expired" and exported() == []                           # exactly 7 days: expired
        r = da.cmd_export(path)
        assert r["dropped"]["expired"] == 2
        Clock.t = T0 + timedelta(days=1)
        assert exported() == ["FOA-020C", "VIC00001"]
        assert da.APPROVAL_TTL_DAYS == 7
        revoked = da.cmd_revoke("vic00001", path)
        assert [d["id"] for d in revoked] == ["A-0001-1"] and statuses()["VIC00001"] == "revoked"
        assert exported() == ["FOA-020C"] and da.cmd_export(path)["dropped"]["revoked"] == 1
        for sku, why in (("VIC00001", "no active approval to revoke"), ("NOPE", "no active approval to revoke")):
            try:
                da.cmd_revoke(sku, path)
                raise AssertionError("revoked twice / unknown")
            except da.AgentError as e:
                assert why in str(e)
        assert [e["event"] for e in log()].count("revoke") == 1
        out = run_cli(["revoke", "FOA-020C", "--results", path])
        assert "Revoked A-0001-2: FOA-020C" in out and exported() == []


def test_a_changed_seller_orphans_the_old_approval():
    reply = "approve VIC00001 qty 100 max 1.30\napprove FOA-020C qty 50 max 0.30\ncap 200"
    with sandbox() as tmp:
        v1, _ = request_for(standard_rows(), tmp, "v1.xlsx")
        record(v1, reply)
        assert sorted(o["sku"] for o in da.cmd_export(v1)["orders"]) == ["FOA-020C", "VIC00001"]
        # A re-run recommends a different seller for VIC00001 (new name AND new URL); FOA-020C is unchanged.
        v2_rows = [res_row("VIC00001", "Another Seller", url="http://shop/another", unit=1.0),
                   res_row("FOA-020C", "FiberMania", unit=0.25, lcom=72.99, moq="10 Pieces")]
        v2 = os.path.join(tmp, "v2.xlsx")
        make_results(v2, v2_rows)
        rows, orphans, _ = da.cmd_status(v2)
        by = {r["sku"]: r["status"] for r in rows}
        assert by == {"VIC00001": "not requested", "FOA-020C": "approved"} and orphans == ["VIC00001 (Eternalstar)"]
        r = da.cmd_export(v2)
        assert [o["sku"] for o in r["orders"]] == ["FOA-020C"] and r["orphans"] == ["VIC00001 (Eternalstar)"]
        assert "1 orphaned" in run_cli(["export", "--results", v2])
        assert "1 earlier approval(s) no longer match the recommended seller and are orphaned: VIC00001 (Eternalstar)" in run_cli(["status", "--results", v2])
        req2 = da.cmd_request(v2)
        assert req2["orphans"] == ["VIC00001 (Eternalstar)"] and "1 earlier approval(s) no longer match" in da.render_text(req2)
        # the same SKU at a different URL from the same seller is a different listing too
        v3 = os.path.join(tmp, "v3.xlsx")
        make_results(v3, [res_row("VIC00001", "Eternalstar", url="http://shop/other-listing")])
        assert da.cmd_status(v3)[1] == ["VIC00001 (Eternalstar)", "FOA-020C (FiberMania)"]   # both earlier approvals orphaned
        assert da.cmd_export(v3)["orders"] == []
        # without a URL the seller's name is the key
        assert listing_key("A1", "", "Acme Co.") != listing_key("A1", "", "Other Co.")


def test_status_shows_all_five_states_and_requested():
    with sandbox() as tmp:
        rows = standard_rows() + [res_row("LATER", "Later Co", unit=2.0, moq="5")]
        path, _ = request_for(rows, tmp)
        assert {r["status"] for r in da.cmd_status(path)[0]} == {"requested"}
        record(path, "approve VIC00001 qty 100 max 1.30\napprove FOA-020C qty 50 max 0.30\napprove BIG1 qty 10 max 12\nreject HDFF\ncap 5000")
        da.cmd_revoke("BIG1", path)
        # FOA-020C and VIC00001 expire; a fresh approval for VIC00001 would be a new request
        Clock.t = T0 + timedelta(days=8)
        by = {r["sku"]: r["status"] for r in da.cmd_status(path)[0]}
        assert by == {"VIC00001": "expired", "FOA-020C": "expired", "BIG1": "revoked", "HDFF": "rejected", "LATER": "requested"}, by
        Clock.t = T0 + timedelta(days=1)
        by = {r["sku"]: r["status"] for r in da.cmd_status(path)[0]}
        assert by["VIC00001"] == "approved" and by["LATER"] == "requested"
        out = run_cli(["status", "--results", path])
        assert "VIC00001" in out and "approved" in out and "expires 2026-10-12T12:00:00Z" in out and "5 line(s)." in out
        assert "approved by Neeraj" in out


# ---------- export ----------

def test_export_contents_and_notice():
    with sandbox() as tmp:
        rows = standard_rows()
        path, _ = request_for(rows, tmp)
        record(path, "approve VIC00001 qty 100 max 1.30\napprove BIG1 qty 500 max 10.50\nreject HDFF\ncap 6000",
               channel="Email", approver="Neeraj K.")
        # BIG1 is an MOQ-heavy line approved with an explicit qty; its listed price ($11.00) is above the approved max ($10.50)
        r = run_cli(["export", "--results", path])
        assert r.startswith(da.NOT_AN_ORDER) and "Exported 2 valid approval(s)" in r
        raw = open(da.EXPORT_JSON, encoding="utf-8").read()
        assert raw.lstrip().startswith('{\n  "notice": "NOT AN ORDER. Nothing has been purchased."')
        data = json.loads(raw)
        assert data["notice"] == da.NOT_AN_ORDER and data["valid_approvals"] == 2 and data["source_results_file"] == "r.xlsx"
        o = {x["sku"]: x for x in data["orders"]}
        v = o["VIC00001"]
        assert v == {"sku": "VIC00001", "product": "VIC00001 product", "seller": "Eternalstar", "listing_url": "http://shop/VIC00001",
                     "qty": 100, "max_unit_price": "1.30", "line_cap": "130.00", "approval_id": "A-0001-1",
                     "approved_by": "Neeraj K.", "approved_at": "2026-10-05T12:00:00Z", "expires_at": "2026-10-12T12:00:00Z",
                     "source_results_file": "r.xlsx", "flags": [], "request_id": "A-0001", "channel": "Email", "sender_verified": False,
                     "quoted_unit_price": None, "quoted_quantity": None, "quoted_currency": None, "quoted_bulk_tiers": None,
                     "quote_evidence": {}, "quote_warnings": []}
        assert o["BIG1"]["flags"] == ["needs_seller_quote_for_smaller_qty", "listed_price_above_approved_max"]
        assert "HDFF" not in o and "FOA-020C" not in o                           # rejected and undecided lines are not exported
        lines = open(da.EXPORT_CSV, encoding="utf-8").read().splitlines()
        assert lines[0] == da.NOT_AN_ORDER and lines[1].split(",")[:4] == ["sku", "product", "seller", "listing_url"]
        rows_ = list(csv.DictReader(lines[1:]))
        assert [x["sku"] for x in rows_] == ["VIC00001", "BIG1"] and rows_[1]["flags"] == "needs_seller_quote_for_smaller_qty;listed_price_above_approved_max" and rows_[0]["flags"] == ""
        assert set(rows_[0]) == set(da.EXPORT_FIELDS) and rows_[0]["sender_verified"] == "False"
        # a price rise after approval is flagged for order generation to see
        up = os.path.join(tmp, "up.xlsx")
        make_results(up, [res_row("VIC00001", "Eternalstar", unit=1.40)])
        flagged = da.cmd_export(up)["orders"][0]
        assert "listed_price_above_approved_max" in flagged["flags"]
        # nothing valid still writes both files with the notice at the top
        Clock.t = T0 + timedelta(days=30)
        empty = da.cmd_export(path)
        assert empty["orders"] == []
        assert json.load(open(da.EXPORT_JSON, encoding="utf-8"))["notice"] == da.NOT_AN_ORDER and json.load(open(da.EXPORT_JSON))["orders"] == []
        assert open(da.EXPORT_CSV, encoding="utf-8").read().splitlines()[0] == da.NOT_AN_ORDER


def test_audit_log_is_append_only_and_complete():
    with sandbox() as tmp:
        path, _ = request_for(standard_rows(), tmp)
        sizes, prefix = [], b""
        def step():
            data = open(da.LOG_PATH, "rb").read()
            assert data.startswith(prefix) and len(data) > len(prefix)            # only ever grows; old lines unchanged
            sizes.append(len(data))
            return data
        prefix = step()
        record(path, "approve VIC00001 qty 100 max 1.30\nreject HDFF\ncap 200")
        prefix = step()
        da.cmd_revoke("VIC00001", path)
        prefix = step()
        da.cmd_export(path)
        step()
        events = [e["event"] for e in log()]
        assert events == ["request", "reply_recorded", "approval", "rejection", "revoke", "export"], events
        assert all("at" in e and e["at"].endswith("Z") for e in log())
        da.cmd_status(path)                                                       # read-only commands add nothing
        assert len(open(da.LOG_PATH, "rb").read()) == sizes[-1]


def test_read_only_inputs_and_status_writes_nothing():
    email = {key_for("VIC00001", "Eternalstar"): {"status": "pricing_provided", "summary": "ok"}}
    with sandbox(exclusions=[("X", "nobody", "r")], email_state=email) as tmp:
        before = (open(da.REVIEWER_EXCLUSIONS_CSV, "rb").read(), open(da.EMAIL_STATE_PATH, "rb").read())
        path, _ = request_for(standard_rows(), tmp)
        record(path, "approve VIC00001 qty 100 max 1.30\ncap 200")
        da.cmd_export(path)
        da.cmd_revoke("VIC00001", path)
        assert (open(da.REVIEWER_EXCLUSIONS_CSV, "rb").read(), open(da.EMAIL_STATE_PATH, "rb").read()) == before
        snapshot = (open(da.STATE_PATH, "rb").read(), open(da.LOG_PATH, "rb").read())
        da.cmd_status(path)
        run_cli(["status", "--results", path])
        assert (open(da.STATE_PATH, "rb").read(), open(da.LOG_PATH, "rb").read()) == snapshot


def test_no_network_llm_or_order_code():
    src = open(os.path.join(HERE, "decision_agent.py"), encoding="utf-8").read()
    imports = set(re.findall(r"^(?:import|from)\s+([\w.]+)", src, re.MULTILINE))
    assert imports == {"argparse", "csv", "hashlib", "json", "os", "re", "sys", "datetime", "decimal", "openpyxl", "email_agent"}, imports
    used = re.findall(r"(anthropic|httpx|requests|urllib|smtplib|socket|subprocess|slack)\w*\s*[.(]", src, re.IGNORECASE)
    assert used == [], used                                    # no network, LLM, mail or Slack calls anywhere
    assert "float(" not in src                                 # money is Decimal throughout
    assert "email_agent" in src and "sourcing_agent" not in src.replace("sourcing_agent.py", "").replace("sourcing_results", "")


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            fn()
    print("decision agent tests passed")
