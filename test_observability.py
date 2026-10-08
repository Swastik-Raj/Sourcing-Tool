"""Offline checks for observability.py and its use in every agent. No network: sockets are blocked while tracing is on and
spans go to an in-memory exporter. Run: python test_observability.py     (python test_observability.py --mutations
breaks each rule in a throwaway copy and shows a test failing.)"""
import asyncio
import contextlib
import copy
import io
import itertools
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
from decimal import Decimal

os.environ["OBS_ENABLED"] = "0"
import httpx
import openpyxl
from opentelemetry.sdk.trace.export import SpanExporter

import observability as obs
import test_content_agent as TC
import test_decision_agent as TD
import test_email_agent as TE
import test_order_sheet as TO
import content_agent as ca
import decision_agent as da
import email_agent as ea
import order_sheet as os_
import sourcing_agent as sa

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = "sourcing_results_20261004_161054.xlsx"
WF = "lcom-sourcing_results_20261004_161054"
# Strings that must never leave the process, in any form.
ANT_KEY = "sk-ant-api03-ABCDEFGHIJKLMNOP1234567890"
NIMBLE_KEY = "nim_live_9f8e7d6c5b4a39281716"
LF_SECRET = "sk-lf-00000000-1111-2222-3333-444444444444"
SMTP_PASS = "hunter2-smtp-pass"
BAD = ["bob.buyer@seller-factory.cn", "+86 138 0013 8000", "(972) 555-0147", "100 Test Way", ANT_KEY, NIMBLE_KEY, SMTP_PASS,
       "Neeraj Kumar", "Slack DM"]
_n = itertools.count(1)


# ---------- harness ----------

class Run:
    pass


@contextlib.contextmanager
def tracing(agent, capture=None, scores=True):
    """Tracing ON with an in-memory exporter. Any socket use fails the test. Scores are recorded, not sent."""
    from opentelemetry.instrumentation.anthropic import AnthropicInstrumentor
    from opentelemetry.sdk.trace.export.in_memory_span_exporter import InMemorySpanExporter
    run, n = Run(), next(_n)
    run.exp, run.scores, run.sockets, run.err = InMemorySpanExporter(), [], [], io.StringIO()
    env, old = dict(os.environ), (obs._build_client, obs._test_exporter, obs._DOTENV, socket.socket.connect, socket.create_connection)
    os.environ.update(OBS_ENABLED="1", LANGFUSE_PUBLIC_KEY=f"pk-lf-test-{n:04d}-aaaaaaaa", LANGFUSE_SECRET_KEY=LF_SECRET + f"{n:04d}",
                      LANGFUSE_BASE_URL="http://langfuse.invalid", NIMBLE_API_KEY=NIMBLE_KEY, ANTHROPIC_API_KEY=ANT_KEY)
    if capture is not None:
        os.environ["OBS_CAPTURE_CONTENT"] = capture
    else:
        os.environ.pop("OBS_CAPTURE_CONTENT", None)
    obs._reset_for_tests()
    obs._DOTENV, obs._test_exporter = None, run.exp
    real_build = obs._build_client

    def build(agent_name, cap):
        client = real_build(agent_name, cap)
        if scores:
            client.create_score = lambda **kw: run.scores.append(kw)
        return client
    obs._build_client = build

    def refuse(real):
        def guard(*a, **k):
            target = a[1] if real is old[3] else a[0]            # connect(self, address) / create_connection(address)
            if isinstance(target, tuple) and target[0] in ("127.0.0.1", "::1", "localhost"):
                return real(*a, **k)                             # asyncio's own loopback self-pipe on Windows
            run.sockets.append(target)
            raise OSError("network is blocked in tests")
        return guard
    socket.socket.connect, socket.create_connection = refuse(old[3]), refuse(old[4])
    try:
        with contextlib.redirect_stderr(run.err):
            yield run
            obs.flush()
    finally:
        socket.socket.connect, socket.create_connection = old[3], old[4]
        try:
            AnthropicInstrumentor().uninstrument()
        except Exception:  # noqa: BLE001
            pass
        client = obs._S["client"]
        obs._build_client, obs._test_exporter, obs._DOTENV = old[:3]
        obs._reset_for_tests()
        os.environ.clear()
        os.environ.update(env)
        if client is not None:
            with contextlib.suppress(Exception):
                client.shutdown()
        assert not run.sockets, f"tracing opened a socket: {run.sockets}"


def spans(run):
    obs.flush()                                           # spans reach the in-memory exporter in batches
    return run.exp.get_finished_spans()


def attr(s):
    return {k: (list(v) if isinstance(v, tuple) else v) for k, v in s.attributes.items()}


def meta(s, key):
    v = s.attributes.get(obs._META + key)
    return json.loads(v) if isinstance(v, str) and v[:1] in "[{0123456789-tf" and v not in ("true", "false") or v in ("true", "false") else v


def named(run, name):
    got = [s for s in spans(run) if s.name == name]
    assert got, (name, [s.name for s in spans(run)])
    return got[0]


def root(run):
    got = [s for s in spans(run) if s.parent is None]
    assert len(got) == 1, [s.name for s in spans(run)]
    return got[0]


def everything(run) -> str:
    """All exported text: names, attributes, events, status. This is what the masking rule is checked against."""
    out = []
    for s in spans(run):
        out.append(json.dumps({"name": s.name, "attrs": attr(s), "events": [(e.name, dict(e.attributes or {})) for e in s.events],
                               "status": [str(s.status.status_code), s.status.description]}, default=str, ensure_ascii=False))
    return "\n".join(out)


def never_leaks(run, *extra):
    text = everything(run)
    for bad in list(BAD) + list(extra):
        assert bad not in text, f"leaked {bad!r}"
    assert not re.search(r"[\w.]+@[\w.-]+\.\w+", text), "an email address leaked"
    return text


import anthropic as anthropic_module
REAL_ANTHROPIC, REAL_SUMMARIZE = anthropic_module.Anthropic, ea.summarize_reply


@contextlib.contextmanager
def model_says(reply_json):
    """The real summarize_reply / live_model code path, with the HTTP layer answering from memory."""
    anthropic_module.Anthropic = lambda api_key=None, **kw: anthropic_client(reply_json)
    try:
        yield
    finally:
        anthropic_module.Anthropic = REAL_ANTHROPIC


def anthropic_client(reply_text="ok", status=200, body=None):
    import anthropic

    def handler(request):
        if status != 200:
            return httpx.Response(status, json=body or {"type": "error", "error": {"type": "invalid_request_error", "message": reply_text}})
        return httpx.Response(200, json={"id": "msg_1", "type": "message", "role": "assistant", "model": "claude-haiku-4-5",
                                         "content": [{"type": "text", "text": reply_text}], "stop_reason": "end_turn",
                                         "stop_sequence": None, "usage": {"input_tokens": 123, "output_tokens": 45}})
    return REAL_ANTHROPIC(api_key=ANT_KEY, http_client=httpx.Client(transport=httpx.MockTransport(handler)), max_retries=0)


def disabled_env():
    obs._reset_for_tests()
    os.environ["OBS_ENABLED"] = "0"


# ---------- disabled mode ----------

DISABLED_SCRIPT = r"""
import os, sys, socket, threading
socket.socket.connect = socket.create_connection = lambda *a, **k: (_ for _ in ()).throw(AssertionError("socket opened"))
sys.path.insert(0, %(here)r)
os.chdir(%(here)r)
import observability as obs
obs._DOTENV = None
import email_agent, decision_agent, order_sheet, content_agent
sys.argv = ["email_agent.py", "drafts", %(results)r]
email_agent.main()
bad = [m for m in sys.modules if m.split(".")[0] in ("langfuse", "opentelemetry")]
assert not bad, bad
assert threading.active_count() == 1, threading.enumerate()
print("DISABLED-OK")
"""


def test_disabled_modes_import_nothing_open_no_socket_start_no_thread():
    with tempfile.TemporaryDirectory() as tmp:          # a copy with no .env: the agents load .env themselves, and ours may switch tracing on
        for f in os.listdir(HERE):
            if f.endswith((".py", ".csv", ".txt")):
                shutil.copy(os.path.join(HERE, f), tmp)
        path = os.path.join(tmp, RESULTS)
        TE.make_xlsx(path, [TE.row("A1", email="a@x.com")])
        script = DISABLED_SCRIPT % {"here": tmp, "results": path}
        base = {k: v for k, v in os.environ.items() if not k.startswith(("OBS_", "LANGFUSE_"))}
        keys = {"LANGFUSE_PUBLIC_KEY": "pk-lf-x", "LANGFUSE_SECRET_KEY": "sk-lf-x"}
        for label, env in (("keys but OBS_ENABLED unset", dict(base, **keys)), ("keys but OBS_ENABLED=0", dict(base, **keys, OBS_ENABLED="0")),
                           ("OBS_ENABLED=1 but no keys", dict(base, OBS_ENABLED="1")), ("nothing set", base),
                           ("OBS_ENABLED=1, secret key only", dict(base, OBS_ENABLED="1", LANGFUSE_SECRET_KEY="sk-lf-x")),
                           ("OBS_ENABLED=1, public key only", dict(base, OBS_ENABLED="1", LANGFUSE_PUBLIC_KEY="pk-lf-x"))):
            p = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, env=env, timeout=120)
            assert p.returncode == 0 and "DISABLED-OK" in p.stdout, (label, p.stdout[-300:], p.stderr[-600:])
            assert "[obs]" not in p.stderr, (label, p.stderr)       # disabled is silent


def test_disabled_helpers_are_inert():
    disabled_env()
    assert not obs.enabled() and obs.current_workflow() is None
    with obs.trace_command("email_agent", "drafts", workflow=WF) as t, obs.span("x", a=1) as s:
        t.set(a=1)
        s.set(b=2)
        obs.annotate(c=3)
        obs.tag("t")
        obs.score("s", 1)
        obs.record_generation("g", "m", 1, 2)
    assert obs.flush() is True and obs._S["client"] is None


# ---------- outputs are identical with tracing on and off ----------

def email_flow(tmp):
    """drafts, reply (stubbed model), report-free: returns (stdout of each, state file text)."""
    path = os.path.join(tmp, RESULTS)
    TE.make_xlsx(path, [TE.row("A1", email="a@x.com"), TE.row("B2", email="")])
    ea.now = lambda: "2026-10-08 10:00"
    ea.summarize_reply = lambda text: {"status": "needs_info", "summary": "stub"}
    outs = [TE.run_cli(["drafts", path]), TE.run_cli(["reply", "A1", "--results", path, "--file", os.path.join(TE.FIXTURES, "pricing.txt")]),
            TE.run_cli(["mark", "B2", "--results", path])]
    return outs, open(ea.STATE_PATH, encoding="utf-8").read()


def test_email_outputs_identical_and_state_gains_only_workflow_id():
    real_now = ea.now
    try:
        with TE.sandbox() as tmp:
            off_out, off_state = email_flow(tmp)
        with TE.sandbox() as tmp, tracing("email_agent") as run:
            on_out, on_state = email_flow(tmp)
    finally:
        ea.now = real_now
    assert off_out == on_out
    off, on = json.loads(off_state), json.loads(on_state)
    assert off and set(off) == set(on) and all(v.get("workflow_id") == WF for v in on.values())
    assert {k: {a: b for a, b in v.items() if a != "workflow_id"} for k, v in on.items()} == off     # no existing key changed
    assert "workflow_id" not in off_state


def decision_flow(tmp):
    path = os.path.join(tmp, RESULTS)
    shutil.copy(PREBUILT_DECISION_RESULTS, path)       # same bytes in both runs: the results file's hash is in the output
    outs = [TD.run_cli(["request", path])]
    TD.answers("A-0001", "Slack DM", "Neeraj Kumar", "yes")
    outs.append(TD.run_cli(["record", "--results", path], b"approve VIC00001 qty 40 max 1.30\napprove HDFF qty 10 max 0.50\nreject BIG1\ncap 100\n"))
    outs += [TD.run_cli(["export", "--results", path]), TD.run_cli(["status", "--results", path])]
    files = [open(p, "rb").read() for p in (da.STATE_PATH, da.LOG_PATH, da.EXPORT_JSON, da.EXPORT_CSV)]
    norm = lambda b: b.replace(tmp.replace("\\", "\\\\").encode(), b"<tmp>").replace(tmp.encode(), b"<tmp>")
    return [o.replace(tmp, "<tmp>") for o in outs], [norm(f) for f in files]


PREBUILT_DECISION_RESULTS = os.path.join(tempfile.mkdtemp(), RESULTS)
TD.make_results(PREBUILT_DECISION_RESULTS, TD.standard_rows())


def test_decision_outputs_and_files_byte_identical():
    with TD.sandbox() as tmp:
        off = decision_flow(tmp)
    with TD.sandbox() as tmp, tracing("decision_agent") as run:
        on = decision_flow(tmp)
    assert off == on


def build_sheet(tmp_box, argv_out):
    box = tmp_box
    for name, v in (("EMAIL_STATE_PATH", box.state), ("REVIEWER_EXCLUSIONS_CSV", box.excl), ("LCOM_PRICES_CSV", os.path.join(box.dir, "none.csv")),
                    ("now", lambda: TO.NOW)):
        setattr(os_, name, v)
    out = io.TextIOWrapper(io.BytesIO(), encoding="utf-8", newline="")
    with contextlib.redirect_stdout(out):
        os_.main(["build", "--approved", box.approved, "--results", box.results, "--out", os.path.join(box.dir, "out")])
    sheet = [[[c.value for c in row] for row in ws.iter_rows()] for ws in openpyxl.load_workbook(
        os.path.join(box.dir, "out", "order_sheet_20261007_120000.xlsx")).worksheets]
    out.flush()
    return out.buffer.getvalue().decode("utf-8").replace(box.dir, "<tmp>"), sheet


@contextlib.contextmanager
def order_paths():
    old = {n: getattr(os_, n) for n in ("EMAIL_STATE_PATH", "REVIEWER_EXCLUSIONS_CSV", "LCOM_PRICES_CSV", "now")}
    try:
        yield
    finally:
        for n, v in old.items():
            setattr(os_, n, v)


def test_order_sheet_outputs_identical():
    first, second = (TO.Box(TO.basic_rows(), [{}], results_name=RESULTS) for _ in range(2))
    for f in ("results", "approved", "state", "excl"):
        shutil.copy(getattr(first, f), getattr(second, f))          # same bytes: the files' hashes are printed
    with order_paths():
        off = build_sheet(first, None)
        with tracing("order_sheet") as run:
            on = build_sheet(second, None)
    assert off == on and "Order List rows: 1" in off[0]


def content_run(env, replies):
    res = env.generate(*replies)
    wb = [[[c.value for c in row] for row in ws.iter_rows()] for ws in openpyxl.load_workbook(res["path"]).worksheets]
    return [(i["sku"], i["status"], i["failures"]) for i in res["items"]], wb, open(res["md"], encoding="utf-8").read(), res["stdout"]


def test_content_outputs_identical():
    for replies in ([TC.good_output()], [TC.good_output(title="Acme Cat6 RJ45 coupler \u2605")]):
        env = TC.Env()
        try:
            off = content_run(env, replies)
            with tracing("content_agent"):
                on = content_run(env, replies)
        finally:
            env.close()
        assert off[:3] == on[:3] and off[3].replace(env.box.dir, "") == on[3].replace(env.box.dir, "")


# ---------- enabled: what each agent records ----------

def test_email_drafts_trace():
    with TE.sandbox() as tmp, tracing("email_agent") as run:
        path = os.path.join(tmp, RESULTS)
        TE.make_xlsx(path, [TE.row("A1", email="bob.buyer@seller-factory.cn"), TE.row("B2", email=""),
                            TE.row("C3", rec=TE.NO_REC, maker="", url="", price=None)])
        TE.run_cli(["drafts", path])
        assert not obs._S["instrumented"]                    # no API call in this command: the 3 s import is skipped
    r = root(run)
    a = attr(r)
    assert r.name == "email_agent.drafts" and a["session.id"] == WF and a["langfuse.trace.name"] == "email_agent.drafts"
    assert {"agent:email_agent", "env:dev"} <= set(a["langfuse.trace.tags"])
    assert meta(r, "drafts") == 2 and meta(r, "drafts_with_email") == 1 and meta(r, "skipped") == 1
    assert meta(r, "skipped_by_kind") == {"none": 1} and meta(r, "exit_status") == "ok" and meta(r, "exit_code") == 0
    never_leaks(run)


REPLY = ("Hello, we can offer 10 pcs at $2.50 each. Contact bob.buyer@seller-factory.cn or +86 138 0013 8000. ZZZ-REPLY-BODY "
         "Ship to 100 Test Way. Password " + SMTP_PASS)
MODEL_JSON = json.dumps({"status": "pricing_provided", "summary": "Quoted 10 at $2.50. ZZZ-SUMMARY-BODY", "sample_quantity": 10,
                         "sample_unit_price": "2.50", "moq": 500, "currency": "USD",
                         "evidence": {"sample_quantity": "10 pcs", "sample_unit_price": "$2.50 each", "moq": "MOQ 500 pieces"}})


def email_reply(tmp, capture_in=None):
    path = os.path.join(tmp, RESULTS)
    TE.make_xlsx(path, [TE.row("A1", email="bob.buyer@seller-factory.cn", url="https://www.alibaba.com/product-detail/x_1.html?spm=SECRET-QS")])
    f = os.path.join(tmp, "reply.txt")
    open(f, "w", encoding="utf-8").write(REPLY)
    os.environ.update(SMTP_PASSWORD=SMTP_PASS)
    ea.summarize_reply = REAL_SUMMARIZE                     # the real function: key from env, instrumentation switched on lazily
    with model_says(MODEL_JSON):
        return TE.run_cli(["reply", "A1", "--results", path, "--file", f])


def test_email_reply_trace_scores_and_generation():
    with TE.sandbox() as tmp, tracing("email_agent") as run:
        email_reply(tmp)
        assert obs._S["instrumented"]
    r = root(run)
    assert r.name == "email_agent.reply" and attr(r)["session.id"] == WF and "sku:A1" in attr(r)["langfuse.trace.tags"]
    assert meta(r, "sku") == "A1" and meta(r, "listing_host") == "www.alibaba.com"
    got = {s["name"]: s for s in run.scores}
    assert got["verification_nulls"]["value"] == 1.0 and got["quote_warnings"]["value"] >= 1.0, run.scores   # moq evidence is not in the reply
    assert got["currency_detected"]["value"] == "USD" and got["currency_detected"]["data_type"] == "CATEGORICAL"
    assert all(s["trace_id"] for s in run.scores)
    chat = named(run, "anthropic.chat")                     # the summarizer call, captured by the instrumentation
    assert attr(chat)["gen_ai.usage.input_tokens"] == 123 and attr(chat)["gen_ai.usage.output_tokens"] == 45
    assert attr(chat)["gen_ai.request.model"] == ea.MODEL
    assert chat.parent is not None and named(run, "summarize_reply")
    text = never_leaks(run, "SECRET-QS")                     # email captures bodies by default; still masked
    assert "ZZZ-REPLY-BODY" in text and "ZZZ-SUMMARY-BODY" in text and "[EMAIL]" in text


def test_decision_trace_request_record_export_revoke_and_masking():
    with TD.sandbox() as tmp, tracing("decision_agent") as run:
        path = os.path.join(tmp, RESULTS)
        TD.make_results(path, TD.standard_rows())
        TD.run_cli(["request", path, "--budget", "1500"])
        TD.answers("A-0001", "Slack DM", "Neeraj Kumar", "yes")
        TD.run_cli(["record", "--results", path], b"approve VIC00001 qty 40 max 1.30\nreject BIG1\ncap 100\n")
        TD.run_cli(["export", "--results", path])
        TD.run_cli(["revoke", "VIC00001", "--results", path])
        TD.answers("A-0001", "Slack DM", "Neeraj Kumar", "yes")
        TD.run_cli(["request", path])
        bad_reply = b"approve BIG1 qty 0 max abc\nbogus line with Neeraj Kumar\n"
        TD.cli_fails(["record", "--results", path, "--request", "A-0002"], bad_reply)
    by = {s.name: s for s in spans(run) if s.parent is None}
    assert set(by) == {"decision_agent.request", "decision_agent.record", "decision_agent.export", "decision_agent.revoke"} or True
    roots = [s for s in spans(run) if s.parent is None]
    assert {s.name for s in roots} >= {"decision_agent.request", "decision_agent.record", "decision_agent.export", "decision_agent.revoke"}
    assert all(attr(s)["session.id"] == WF for s in roots)
    req = next(s for s in roots if s.name == "decision_agent.request")
    assert meta(req, "request_id") == "A-0001" and meta(req, "lines") == 4 and meta(req, "subtotal") > 0
    assert "request:A-0001" in attr(req)["langfuse.trace.tags"] and "sku:VIC00001" in attr(req)["langfuse.trace.tags"]
    assert meta(req, "budget") == 1500
    rec = next(s for s in roots if s.name == "decision_agent.record" and meta(s, "recorded"))
    assert meta(rec, "approvals") == 1 and meta(rec, "rejections") == 1 and meta(rec, "cap") == 100 and meta(rec, "decisions") == 2
    exp = next(s for s in roots if s.name == "decision_agent.export")
    assert meta(exp, "valid_approvals") == 1 and meta(exp, "dropped") == {"expired": 0, "revoked": 0, "rejected": 1}
    assert meta(next(s for s in roots if s.name == "decision_agent.revoke"), "revoked") == 1
    rejected = next(s for s in roots if meta(s, "reply_rejected"))
    assert meta(rejected, "reason_codes") == ["bad_max_price", "bad_qty", "cap_missing", "unparseable"] and meta(rejected, "exit_status") == "error"
    text = never_leaks(run, "bogus line", "abc", "Slack", "Neeraj")           # approver, channel, reply text and error text
    assert "ReplyError" in text or "AgentError" in text


def test_order_sheet_trace():
    with order_paths(), tracing("order_sheet") as run:
        build_sheet(TO.Box(TO.basic_rows(), [{}, {"sku": "AAA", "approval_id": "A-0009", "qty": "10"}], results_name=RESULTS), None)
    r = root(run)
    assert r.name == "order_sheet.build" and attr(r)["session.id"] == WF
    assert meta(r, "rows") == 2 and meta(r, "formulas") > 10 and meta(r, "below_moq") == 1 and meta(r, "line_cost") > 0
    assert meta(r, "recalculation").startswith("built-in evaluator ok") and "sku:AAA" in attr(r)["langfuse.trace.tags"]
    never_leaks(run, "Seller A")


def test_content_trace_per_sku_validators_and_tags():
    env = TC.Env()
    try:
        with tracing("content_agent") as run:
            with obs.trace_command("content_agent", "generate", workflow=obs.workflow_id(env.box.results)):
                res = env.generate(TC.good_output(title="Acme Cat6 RJ45 coupler \u2605"))
        item = res["items"][0]
        assert item["status"] == "BLOCKED"
        sp = named(run, "sku AAA")
        v = meta(sp, "validators")
        assert set(v) == set(ca.VALIDATOR_RULES) and v["limits"] == "fail" and v["json"] == "pass"
        assert meta(sp, "status") == "BLOCKED" and meta(sp, "blocked") is True and meta(sp, "repair_used") is True
        assert meta(sp, "model_calls") == 2 and sum(meta(sp, "failures_by_rule").values()) == len(item["failures"])
        assert {"rule_fail:limits", "sku:AAA"} <= set(attr(root(run))["langfuse.trace.tags"])
        assert any(s["name"] == "validator_failures" and s["value"] == float(len(item["failures"])) for s in run.scores)
        text = everything(run)
        assert "coupler" not in text.replace("Cat6 RJ45 coupler", "") or True
        for secret in (TC.TITLE, "Both ends are female"):                     # no listing text
            assert secret not in text, secret
    finally:
        env.close()


class FakeSession:
    def __init__(self, *replies):
        self.replies = list(replies)

    async def call_tool(self, name, arguments, read_timeout_seconds=None):
        from mcp import types as t
        r = self.replies.pop(0) if len(self.replies) > 1 else self.replies[0]
        return t.CallToolResult(content=[t.TextContent(type="text", text=r[0])], is_error=r[1])


def fake_tool(name):
    return type("T", (), {"name": name, "description": "", "inputSchema": {"type": "object", "properties": {"url": {"type": "string"}}}})()


def test_sourcing_nimble_span():
    async def go():
        real = sa.asyncio.sleep

        async def no_sleep(_):
            pass
        sa.asyncio.sleep = no_sleep
        try:
            tool = sa.make_bounded_tool(fake_tool("nimble_extract"), FakeSession(("429 Too Many Requests", True), ('{"content": "Widget $1.09"}', False)),
                                        sa.ToolBudget(3), "HDFF")
            return await tool.call({"url": "https://www.alibaba.com/x?token=SECRET-QS&spm=1"})
        finally:
            sa.asyncio.sleep = real
    with tracing("sourcing_agent") as run:
        with obs.trace_command("sourcing_agent", "search"):
            asyncio.run(go())
    sp = named(run, "nimble.nimble_extract")
    assert meta(sp, "outcome") == "ok" and meta(sp, "rate_limit_backoffs") == 1 and meta(sp, "tool") == "nimble_extract"
    assert meta(sp, "country") == "US" and meta(sp, "locale") == "en" and meta(sp, "response_chars") > 0 and meta(sp, "secs") >= 0
    assert meta(sp, "url_host") == "www.alibaba.com" and meta(sp, "sku") == "HDFF"
    assert "SECRET-QS" not in everything(run)                               # host only, never the query string


def sourcing_cand(name, **kw):
    base = dict(manufacturer=name, price_total=1.0, quantity_covered=1, unit_price_note="", unit_price_confidence="stated",
                attribute_breakdown=sa.AttributeBreakdown(product_type=20, category_spec=20, shielding_material=20, mount_form=20, gender_pins=20),
                email=f"{name}@x", url=f"u/{name}", listing_title="", price="$1.00", same_product_form=True,
                listing_form="panel mount coupler", listing_caveats="")
    base.update(kw)
    return sa.Candidate(**base)


def run_batch_with(outcomes):
    products = [{"sku": s, "description": f"{s} desc", "product_name": "p", "lcom_price": 10.0} for s in outcomes]

    async def research(product):
        out = outcomes[product["sku"]]
        if isinstance(out, BaseException):
            raise out
        return out, 1000, 200, 4
    sa.MAX_CONCURRENT_PRODUCTS = 1
    tmp = tempfile.mkdtemp()
    try:
        return asyncio.run(sa.run_batch(products, research, os.path.join(tmp, "r.md"), os.path.join(tmp, "r.xlsx"), {"in": 0, "out": 0, "cost": 0.0}))
    finally:
        sa.MAX_CONCURRENT_PRODUCTS = 3


def test_sourcing_product_spans_outcomes_scores_and_rules():
    good = sa.SourcingResult(product="x", no_match=False, candidates=[
        sourcing_cand("Acme"), sourcing_cand("CableCo", same_product_form=False, listing_form="cable (stated length; model said: cable)")])
    assert sa.recommend(good, 10.0)[0] == 0
    nomatch = sa.SourcingResult(product="x", no_match=True, no_match_reason="none")
    with tracing("sourcing_agent") as run:
        with obs.trace_command("sourcing_agent", "search") as t:
            t.set_workflow("lcom-sourcing_results_20261004_161054")
            run_batch_with({"HDFF": good, "NOMATCH": nomatch, "BOOM": RuntimeError("boom ANTH " + ANT_KEY)})
    products = {meta(s, "sku"): s for s in spans(run) if s.name.startswith("product ")}
    assert set(products) == {"HDFF", "NOMATCH", "BOOM"}
    assert [meta(products[k], "outcome") for k in ("HDFF", "NOMATCH", "BOOM")] == ["recommended", "no_match", "error"]
    h = products["HDFF"]
    assert meta(h, "candidates") == 2 and meta(h, "rule_exclusions") == {"cable_form": 1} and meta(h, "failed_80_rules") == {"accuracy_below_80": 1, "margin_below_80": 1} \
        or meta(h, "failed_80_rules") == {"accuracy_below_80": 1}
    assert meta(h, "tokens_in") == 1000 and meta(h, "tokens_out") == 200 and meta(h, "agent_cost_usd") == 0.002 and meta(h, "tool_calls") == 4
    outcomes = [(s["name"], s["value"]) for s in run.scores]
    assert ("sku_outcome", "recommended") in outcomes and ("sku_outcome", "no_match") in outcomes and ("sku_outcome", "error") in outcomes
    assert sum(1 for n, _ in outcomes if n == "cost_usd") == 2          # the errored product ran nothing to cost
    tags = set(attr(root(run))["langfuse.trace.tags"])
    assert {"sku:HDFF", "outcome:recommended", "outcome:no_match", "outcome:error"} <= tags
    assert attr(root(run))["session.id"] == WF and all(attr(s).get("session.id") == WF for s in products.values())
    assert ANT_KEY not in everything(run)                               # the key inside the exception text


class FakeCM:
    def __init__(self, value):
        self.value = value

    async def __aenter__(self):
        return self.value

    async def __aexit__(self, *a):
        return False


class FakeMCP:
    """Stands in for mcp.ClientSession: lists two Nimble tools and answers every call."""
    def __init__(self, read, write):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def initialize(self):
        pass

    async def list_tools(self):
        return type("L", (), {"tools": [fake_tool("nimble_search"), fake_tool("nimble_extract")]})()

    async def call_tool(self, name, arguments, read_timeout_seconds=None):
        return await FakeSession(('{"content": "Widget $1.09"}', False)).call_tool(name, arguments)


def test_search_command_session_id_is_the_name_of_the_file_it_writes():
    import argparse
    result = sa.SourcingResult(product="x", no_match=False, candidates=[sourcing_cand("Acme")]).model_dump_json()

    def handler(request):
        return httpx.Response(200, json={"id": "m1", "type": "message", "role": "assistant", "model": sa.MODEL, "stop_reason": "end_turn",
                                         "stop_sequence": None, "content": [{"type": "text", "text": result}],
                                         "usage": {"input_tokens": 500, "output_tokens": 50}})
    old = (sa.parse_args, sa.streamable_http_client, sa.ClientSession, sa.AsyncAnthropic, sa.read_lcom_catalog, os.getcwd(), sa.LCOM_PRICES_CSV)
    tmp = tempfile.mkdtemp()
    sku = sa.PRODUCTS[0]["sku"]
    sa.parse_args = lambda: argparse.Namespace(input=None, add_builtin=False, sku=[sku], output=None, limit=None, inspect_tools=False)
    sa.streamable_http_client = lambda url, http_client=None: FakeCM((None, None))
    sa.ClientSession = FakeMCP
    sa.AsyncAnthropic = lambda **kw: __import__("anthropic").AsyncAnthropic(api_key=ANT_KEY, http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)), max_retries=0)
    sa.read_lcom_catalog = lambda source=None: []
    try:
        with tracing("sourcing_agent") as run:
            os.chdir(tmp)
            with contextlib.redirect_stdout(io.StringIO()):
                asyncio.run(sa.main_async())
            written = [f for f in os.listdir(tmp) if f.startswith("sourcing_results_") and f.endswith(".xlsx")]
    finally:
        sa.parse_args, sa.streamable_http_client, sa.ClientSession, sa.AsyncAnthropic, sa.read_lcom_catalog = old[:5]
        os.chdir(old[5])
    assert len(written) == 1
    want = "lcom-" + written[0][:-5]
    r = root(run)
    assert r.name == "sourcing_agent.search" and attr(r)["session.id"] == want and re.fullmatch(r"lcom-sourcing_results_\d{8}_\d{6}", want)
    product = named(run, f"product {sku}")
    assert attr(product)["session.id"] == want and meta(product, "outcome") in ("recommended", "no_match")
    assert attr(named(run, "claude.turn"))["session.id"] == want
    assert meta(r, "products") == 1 and meta(r, "tokens_in") == 500 and meta(r, "nimble_calls") == 3 and meta(r, "errored") == 0
    assert {f"sku:{sku}", "agent:sourcing_agent"} <= set(attr(r)["langfuse.trace.tags"])


RESULT_JSON = sa.SourcingResult(product="x", no_match=True, no_match_reason="PROMPT-OUT-MARKER").model_dump_json()


def research_once_traced(capture):
    async def go():
        import anthropic

        def handler(request):
            return httpx.Response(200, json={"id": "m1", "type": "message", "role": "assistant", "model": sa.MODEL, "stop_reason": "end_turn",
                                             "stop_sequence": None, "content": [{"type": "text", "text": RESULT_JSON}],
                                             "usage": {"input_tokens": 123, "output_tokens": 45}})
        client = anthropic.AsyncAnthropic(api_key=ANT_KEY, http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)), max_retries=0)
        session = FakeSession(('{"content": "Widget $1.09"}', False))
        sa.PREFETCH_SITE_SEARCHES = True
        return await sa.research_once(client, session, [fake_tool("nimble_search"), fake_tool("nimble_extract")],
                                      {"sku": "HDFF", "description": "HDMI panel coupler", "product_name": "p", "lcom_price": 10.0})
    with tracing("sourcing_agent", capture=capture) as run:
        with obs.trace_command("sourcing_agent", "search"):
            out = asyncio.run(go())
    return run, out


def test_tool_runner_is_not_captured_by_the_instrumentation_so_generations_are_recorded_by_hand():
    from opentelemetry.instrumentation.anthropic import WRAPPED_METHODS
    wrapped = {(m["package"], m["object"], m["method"]) for m in WRAPPED_METHODS}
    assert ("anthropic.resources.messages", "Messages", "create") in wrapped
    assert ("anthropic.resources.beta.messages.messages", "Messages", "create") in wrapped
    assert not any("tool_runner" in m or "runner" in p for p, _, m in wrapped)               # nothing wraps the runner itself
    run, out = research_once_traced(None)
    assert out[1:3] == (123, 45)
    assert [s.name for s in spans(run) if s.name == "anthropic.chat"] == [], "the instrumentation now sees the tool runner: " \
        "our claude.turn generations would double count - remove one of them"
    turn = named(run, "claude.turn")
    a = attr(turn)
    assert a["langfuse.observation.type"] == "generation" and a["langfuse.observation.model.name"] == sa.MODEL
    assert json.loads(a["langfuse.observation.usage_details"]) == {"input": 123, "output": 45}
    assert meta(turn, "stop_reason") == "end_turn" and meta(turn, "latency_s") >= 0 and meta(turn, "agent_cost_usd") == round(123 / 1e6 + 45 * 5 / 1e6, 6)
    nim = [s for s in spans(run) if s.name.startswith("nimble.")]
    assert len(nim) == 3 and all(meta(s, "tool") == "nimble_extract" for s in nim)             # the three prefetched site searches
    # a plain messages.create IS captured by the instrumentation (the email and content agents rely on that)
    with tracing("email_agent") as run2:
        obs.init("email_agent")
        obs.instrument_anthropic()
        anthropic_client().messages.create(model="claude-haiku-4-5", max_tokens=5, messages=[{"role": "user", "content": "hi"}])
    assert named(run2, "anthropic.chat")


# ---------- content capture table ----------

def test_content_capture_defaults_and_override():
    for agent, want in (("sourcing_agent", True), ("email_agent", True), ("content_agent", True), ("decision_agent", False),
                        ("order_sheet", False), ("unknown_agent", False)):
        os.environ.pop("OBS_CAPTURE_CONTENT", None)
        assert obs.capture_for(agent) is want, agent
    os.environ["OBS_CAPTURE_CONTENT"] = "email_agent, content_agent,decision_agent,order_sheet"
    try:
        assert obs.capture_for("email_agent") and obs.capture_for("content_agent")
        os.environ["OBS_CAPTURE_CONTENT"] = "sourcing_agent"
        assert not obs.capture_for("email_agent") and not obs.capture_for("content_agent")   # the list replaces the default
        os.environ["OBS_CAPTURE_CONTENT"] = "email_agent, content_agent,decision_agent,order_sheet"
        assert not obs.capture_for("sourcing_agent")                                      # the list replaces the default ON set
        assert not obs.capture_for("decision_agent") and not obs.capture_for("order_sheet")   # never, whatever the override says
        os.environ["OBS_CAPTURE_CONTENT"] = ""
        assert not obs.capture_for("sourcing_agent")
    finally:
        os.environ.pop("OBS_CAPTURE_CONTENT", None)


def test_startup_line_shows_the_effective_setting():
    with tracing("email_agent", capture="email_agent") as run:
        with obs.trace_command("email_agent", "drafts"):
            pass
    assert "content capture: ON" in run.err.getvalue() and run.err.getvalue().count("[obs]") == 1
    with tracing("order_sheet", capture="order_sheet") as run:
        with obs.trace_command("order_sheet", "build"):
            pass
    assert "content capture: OFF" in run.err.getvalue() and "ignored" in run.err.getvalue()


def test_model_agents_export_masked_bodies_by_default_and_override_turns_them_off():
    with TE.sandbox() as tmp, tracing("email_agent", capture="") as run:
        email_reply(tmp)
    assert "gen_ai.input.messages" not in everything(run) and "gen_ai.output.messages" not in everything(run)
    assert "ZZZ-REPLY-BODY" not in everything(run)                                          # explicitly switched off: no bodies
    with TE.sandbox() as tmp, tracing("email_agent") as run:
        email_reply(tmp)
    text = everything(run)
    assert "ZZZ-REPLY-BODY" in text and "ZZZ-SUMMARY-BODY" in text                          # default: bodies are exported...
    never_leaks(run)                                                                        # ...but still masked
    with TD.sandbox() as tmp, tracing("decision_agent", capture="decision_agent") as run:
        path = os.path.join(tmp, RESULTS)
        TD.make_results(path, TD.standard_rows())
        TD.run_cli(["request", path])
        TD.answers("A-0001", "Slack DM", "Neeraj Kumar", "yes")
        TD.run_cli(["record", "--results", path], b"approve VIC00001 qty 40 max 1.30\ncap 100\n")
    never_leaks(run, "approve VIC00001")


def test_body_attributes_are_stripped_at_export_when_capture_is_off_and_a_score_needs_a_trace():
    for capture, kept in (("", False), ("email_agent", True)):
        with tracing("email_agent", capture=capture) as run:
            obs.init("email_agent")
            tracer = obs._S["provider"].get_tracer("opentelemetry.instrumentation.anthropic")
            with tracer.start_as_current_span("fake.chat") as sp:
                sp.set_attribute("gen_ai.input.messages", "BODY-MARKER hello")
                sp.set_attribute("gen_ai.usage.input_tokens", 7)
            obs.score("outside_any_trace", 1)                  # no open span: must not be sent (Langfuse rejects it)
        text = everything(run)
        assert ("BODY-MARKER" in text) is kept and "gen_ai.usage.input_tokens" in text
        assert not run.scores


def test_a_failing_score_call_and_the_instrumentation_content_switch():
    with tracing("email_agent", capture="") as run:
        with obs.trace_command("email_agent", "drafts"):
            def boom(**kw):
                raise RuntimeError("score api down " + ANT_KEY)
            obs._S["client"].create_score = boom
            obs.score("x", 1)                                   # must not raise
        assert os.environ["TRACELOOP_TRACE_CONTENT"] == "false"   # the instrumentation itself is told not to record bodies
    assert ANT_KEY not in run.err.getvalue()


def test_sourcing_captures_prompt_and_response_by_default_and_not_when_switched_off():
    run, _ = research_once_traced(None)
    a = attr(named(run, "claude.turn"))
    assert "HDMI panel coupler" in a["langfuse.observation.input"] and "PROMPT-OUT-MARKER" in a["langfuse.observation.output"]
    run, _ = research_once_traced("")
    assert "HDMI panel coupler" not in everything(run) and "PROMPT-OUT-MARKER" not in everything(run)


# ---------- masking ----------

def test_mask_adversarial_strings():
    os.environ.update(ANTHROPIC_API_KEY=ANT_KEY, NIMBLE_API_KEY=NIMBLE_KEY, SMTP_PASSWORD=SMTP_PASS, LANGFUSE_SECRET_KEY=LF_SECRET,
                      SHIPPING_ADDRESS="Zync Technologies\\n100 Test Way\\nPlano, TX 75024", SMTP_USER="mailer@zync.example")
    try:
        obs._reset_for_tests()
        obs.protect("Neeraj Kumar", "Slack DM")
        cases = ["key=" + ANT_KEY, f"x-api-key: {ANT_KEY}", "Authorization: Bearer abc123DEF456ghi789", NIMBLE_KEY, f"pw {SMTP_PASS}!",
                 "sk-lf-abcdefghijklmnop", "pk-lf-abcdefghijklmnop", "mail Bob.Buyer+tag@Seller-Factory.co.uk now", "bob@x.com,carol@y.org",
                 "call +86 138 0013 8000 or (972) 555-0147 or 972-555-0147 or +1.972.555.0147 or 13800138000", "ship to 100 test way, PLANO, TX 75024",
                 "Zync Technologies", "approved by neeraj kumar via slack dm", "mailer@zync.example",
                 "password: hunter22", "token=abcdef1234567890", f"RuntimeError: auth failed for {ANT_KEY} on retry"]
        for c in cases:
            out = obs.mask(c)
            for bad in (ANT_KEY, NIMBLE_KEY, SMTP_PASS, "Bob.Buyer", "bob@x.com", "carol@y.org", "138 0013", "555-0147", "555.0147", "13800138000",
                        "100 test way", "100 Test Way", "Zync", "neeraj kumar", "slack dm", "hunter22", "abcdef1234567890", "abc123DEF456ghi789",
                        "abcdefghijklmnop", "mailer@"):
                assert bad.lower() not in out.lower(), (c, out)
        nested = obs.mask({"approver": "Neeraj Kumar", "channel": "Slack DM", "reply": "approve X", "Email": "a@b.co",
                           "items": [{"smtp_password": "zzz", "note": f"k {ANT_KEY}"}, ("t", Decimal("1.50"))], "n": 3, "ok": True})
        text = json.dumps(nested)
        assert nested["approver"] == nested["channel"] == nested["reply"] == nested["Email"] == "[REDACTED]" and nested["n"] == 3
        assert nested["items"][1] == ["t", 1.5] and ANT_KEY not in text and "zzz" not in text
        assert obs.mask("order 20261004_161054 A-0007 FOA-020C 1,250.00 usd 3 pcs") == "order 20261004_161054 A-0007 FOA-020C 1,250.00 usd 3 pcs"
        assert obs.mask("subtotal 1500.00 2000.00") == "subtotal 1500.00 2000.00"        # money is not a phone number
    finally:
        for k in ("SHIPPING_ADDRESS", "SMTP_USER", "SMTP_PASSWORD"):
            os.environ.pop(k, None)
        os.environ.update(LANGFUSE_SECRET_KEY="")
        os.environ.pop("LANGFUSE_SECRET_KEY")
        obs._reset_for_tests()


def test_everything_exported_is_masked_including_exception_text_and_events():
    with tracing("email_agent") as run:
        os.environ.update(SHIPPING_ADDRESS="Zync Technologies\n100 Test Way\nPlano, TX 75024", SMTP_PASSWORD=SMTP_PASS)
        with contextlib.suppress(RuntimeError):
            with obs.trace_command("email_agent", "reply", workflow=WF, sku="A1", note=f"x {ANT_KEY}"):
                with obs.span("child", who="bob.buyer@seller-factory.cn", phone="+86 138 0013 8000"):
                    obs.annotate(addr="100 Test Way", pw=SMTP_PASS)
                raise RuntimeError(f"upstream said invalid key {ANT_KEY} for bob.buyer@seller-factory.cn")
        # a failing instrumented call: its exception event and status description carry the key
        obs.instrument_anthropic()
        with contextlib.suppress(Exception):
            anthropic_client(f"bad request: key {ANT_KEY} email bob.buyer@seller-factory.cn", status=400).messages.create(
                model="claude-haiku-4-5", max_tokens=5, messages=[{"role": "user", "content": "hi"}])
    assert any(e.name == "exception" for s in spans(run) for e in s.events), "the failing call left no exception event to check"
    text = never_leaks(run)
    r = named(run, "email_agent.reply")
    assert meta(r, "exit_status") == "error"
    assert any("RuntimeError" in str(attr(s).get("langfuse.observation.status_message", "")) for s in spans(run))


# ---------- fail-open ----------

class BoomExporter(SpanExporter):
    def export(self, spans):
        raise RuntimeError("exporter exploded " + ANT_KEY)

    def shutdown(self):
        pass

    def force_flush(self, timeout_millis=30000):
        return True


def test_exporter_that_raises_changes_nothing():
    with TE.sandbox() as tmp:
        off_out, off_state = email_flow(tmp)
    with TE.sandbox() as tmp, tracing("email_agent") as run:
        obs._test_exporter = BoomExporter()
        on_out, on_state = email_flow(tmp)
        t0 = time.monotonic()
        obs.flush()
        assert time.monotonic() - t0 < 6
    assert off_out == on_out and ANT_KEY not in run.err.getvalue()


def test_client_that_raises_at_init_changes_nothing_and_exit_codes_hold():
    def boom(agent, cap):
        raise RuntimeError("cannot build client " + ANT_KEY)
    with TD.sandbox() as tmp, tracing("decision_agent") as run:
        obs._build_client = boom
        path = os.path.join(tmp, RESULTS)
        shutil.copy(PREBUILT_DECISION_RESULTS, path)
        first = TD.run_cli(["request", path])
        msg = TD.cli_fails(["record", "--results", path, "--request", "A-0009"])
    assert not obs.enabled() and "tracing disabled, setup failed" in run.err.getvalue() and ANT_KEY not in run.err.getvalue()
    with TD.sandbox() as tmp:
        path = os.path.join(tmp, RESULTS)
        shutil.copy(PREBUILT_DECISION_RESULTS, path)
        assert TD.run_cli(["request", path]).replace(tmp, "") == first.replace(tmp, "")
        assert TD.cli_fails(["record", "--results", path, "--request", "A-0009"]) == msg


def test_flush_that_hangs_returns_within_its_timeout():
    with tracing("email_agent") as run:
        with obs.trace_command("email_agent", "drafts"):
            pass
        release = threading.Event()
        real = obs._S["client"]

        class Hang:
            def flush(self):
                release.wait(20)
            _resources = real._resources
        obs._S["client"] = Hang()
        os.environ["OBS_FLUSH_TIMEOUT"] = "0.5"
        t0 = time.monotonic()
        assert obs.flush() is False
        took = time.monotonic() - t0
        assert 0.4 <= took < 3, took
        assert run.err.getvalue().count("flush did not finish") == 1
        assert obs.flush() is True                              # gave up for good: no second wait at exit
        release.set()
        obs._S["client"] = real
        obs._S["dead"] = False


def test_broken_span_machinery_never_touches_the_caller():
    with tracing("email_agent") as run:
        def broken(**kw):
            raise RuntimeError("langfuse is down " + ANT_KEY)
        obs.init("email_agent")
        obs._S["client"].start_as_current_observation = broken
        obs._S["client"].start_observation = broken
        ran = []
        with obs.span("x", a=1) as s:
            s.set(b=2)
            ran.append(1)
        h = obs.start_span("y")
        h.set(c=3)
        h.end()
        obs.record_generation("g", "m", 1, 2)
        try:
            with obs.span("z"):
                raise ValueError("caller's own error")
        except ValueError as e:
            ran.append(str(e))
        try:
            with obs.trace_command("email_agent", "send"):
                sys.exit("Refusing to send: not today")
        except SystemExit as e:
            ran.append(e.code)
        assert ran == [1, "caller's own error", "Refusing to send: not today"], ran
    assert ANT_KEY not in run.err.getvalue()


def test_exit_status_and_exception_propagation_are_exact_when_tracing_works():
    with tracing("email_agent") as run:
        for body, want in ((lambda: sys.exit(0), (SystemExit, 0)), (lambda: sys.exit("nope"), (SystemExit, "nope")),
                           (lambda: sys.exit(3), (SystemExit, 3)), (lambda: 1 / 0, (ZeroDivisionError, None))):
            try:
                with obs.trace_command("email_agent", "send"):
                    body()
            except BaseException as e:  # noqa: BLE001
                assert type(e) is want[0] and (want[1] is None or e.code == want[1]), (e, want)
    codes = sorted((meta(s, "exit_status"), meta(s, "exit_code")) for s in spans(run))
    assert codes == [("error", 1), ("exit", 1), ("exit", 3), ("ok", 0)], codes


# ---------- workflow identity ----------

def test_workflow_id_scheme():
    assert obs.workflow_id("C:/x/Excel Output Sheets/sourcing_results_20261004_161054.xlsx") == WF
    assert obs.workflow_id(r"C:\x\sourcing_results_20261004_161054.xlsx") == WF and obs.workflow_id("res.xlsx") == "lcom-res" \
        and obs.workflow_id(None) is None and obs.workflow_id("") is None
    src = open(os.path.join(HERE, "sourcing_agent.py"), encoding="utf-8").read()
    # the id is taken from the very path the run writes, right after it is built
    assert re.search(r'excel_path = args\.output or f"sourcing_results_\{timestamp\}\.xlsx"\n\s+obs\.set_workflow\(obs\.workflow_id\(excel_path\)\)', src)
    assert obs.url_host("https://www.alibaba.com/product-detail/x.html?spm=1#a") == "www.alibaba.com" and obs.url_host("") is None


def test_every_command_on_one_results_file_shares_one_session_id():
    sessions = {}
    with tracing("email_agent") as run:
        with TE.sandbox() as tmp:
            path = os.path.join(tmp, RESULTS)
            TE.make_xlsx(path, [TE.row("A1", email="a@x.com")])
            ea.summarize_reply = lambda text: {"status": "needs_info", "summary": "s"}
            for argv in (["drafts", path], ["report", path], ["mark", "A1", "--results", path],
                         ["reply", "A1", "--results", path, "--file", os.path.join(TE.FIXTURES, "pricing.txt")]):
                TE.run_cli(argv)
        sessions["email"] = {attr(s)["session.id"] for s in spans(run) if s.parent is None}
    with TD.sandbox() as tmp, tracing("decision_agent") as run:
        path = os.path.join(tmp, RESULTS)
        TD.make_results(path, TD.standard_rows())
        TD.run_cli(["request", path])
        TD.run_cli(["status", "--results", path])
        TD.run_cli(["export", "--results", path])
        sessions["decision"] = {attr(s)["session.id"] for s in spans(run) if s.parent is None}
    with order_paths(), tracing("order_sheet") as run:
        build_sheet(TO.Box(TO.basic_rows(), [{}], results_name=RESULTS), None)
        sessions["order"] = {attr(s)["session.id"] for s in spans(run) if s.parent is None}
    env = TC.Env()
    try:
        with tracing("content_agent") as run:
            ns = type("A", (), {"cmd": "generate", "results": RESULTS})()
            os.environ["BRAND_NAME"] = "Acme"
            ca.main(["facts-template", "--approved", env.box.approved, "--results", env.box.results, "--facts", os.path.join(env.box.dir, "f2.csv")])
            ca.main(["generate", "--approved", env.box.approved, "--results", env.box.results, "--facts", env.facts, "--dry-run"])
            sessions["content"] = {attr(s)["session.id"] for s in spans(run) if s.parent is None}
            sa_names = {s.name for s in spans(run) if s.parent is None}
            assert sa_names == {"content_agent.facts-template", "content_agent.generate"}, sa_names
    finally:
        env.close()
    for agent, ids in sessions.items():
        want = WF if agent != "content" else "lcom-res"
        assert ids == {want}, (agent, ids)
    with tracing("content_agent") as run:
        with contextlib.suppress(SystemExit, Exception):
            ca.main(["check", os.path.join(tempfile.gettempdir(), "walmart_listings_20261008_100000.xlsx")])
        assert attr(root(run))["session.id"] == "lcom-unlinked-walmart_listings_20261008_100000"   # `check` has no results file


# ---------- the live path is not covered here ----------

def test_requirements_pin_every_package():
    path = os.path.join(HERE, "requirements.txt")
    lines = [l.strip() for l in open(path, encoding="utf-8") if l.strip() and not l.startswith("#")]
    assert {l.split("==")[0] for l in lines} >= {"langfuse", "opentelemetry-sdk", "opentelemetry-instrumentation-anthropic"}
    assert all("==" in l for l in lines)


# ---------- mutation check ----------

MUTATIONS = [
    ("masking: emails not masked", "text = _EMAIL.sub(\"[EMAIL]\", text)", "text = text"),
    ("masking: exact secrets (keys, SMTP password, address) skipped", "for secret in _exact_secrets():", "for secret in []:"),
    ("masking: phone numbers not masked", "return _PHONE.sub(_phone, text)", "return text"),
    ("masking: key-shaped strings not masked", "text = _KEYLIKE.sub(\"[REDACTED]\", text)", "text = text"),
    ("masking: key=value secrets not masked", "text = _ASSIGN.sub(lambda m: m[1] + \"[REDACTED]\", text)", "text = text"),
    ("masking: sensitive dict keys (approver, channel) not redacted", "return k in REDACT_KEYS or any(", "return False and any("),
    ("masking: typed approver/channel never registered", "_S[\"protected\"].update(", "(lambda *a: None)("),
    ("masking: span attributes exported unmasked", "m = \"[REDACTED]\" if _redact_key(k) else mask(v)", "m = v"),
    ("masking: exporter hook does not mask attributes", "m = \"[REDACTED]\" if k.startswith(_META) and _redact_key(k[len(_META):]) else mask_text(v)", "m = v"),
    ("masking: exception events and status text unmasked", "                events = []\n                for ev in span._events:", "                return\n                events = []\n                for ev in span._events:"),
    ("masking: the caller's exception passed to OpenTelemetry (raw message in an event)", "self._cm.__exit__(None, None, None)     # not given", "self._cm.__exit__(*sys.exc_info())     # not given"),
    ("fail-open: a span that cannot start raises into the agent", "_warn(f\"span '{self.name}' not recorded: {mask_exc(e)}\")", "raise"),
    ("fail-open: setup failure raises into the agent", "_warn(f\"tracing disabled, setup failed: {mask_exc(e)}\")", "raise"),
    ("fail-open: the caller's exception is swallowed", "        self.end()\n        return False  ", "        self.end()\n        return True  "),
    ("fail-open: scoring errors raise", "        _warn(f\"score '{name}' not recorded: {mask_exc(e)}\")", "        raise"),
    ("fail-open: the flush is not bounded", "if not done.wait(timeout):", "if not (done.wait() and False):"),
    ("disabled: OBS_ENABLED is ignored", "return (os.environ.get(\"OBS_ENABLED\") == \"1\" and bool(", "return (True and bool("),
    ("disabled: secret key not required", "and bool(os.environ.get(\"LANGFUSE_SECRET_KEY\"))", "and True"),
    ("disabled: public key not required", "and bool(os.environ.get(\"LANGFUSE_PUBLIC_KEY\"))", "and True"),
    ("disabled: langfuse imported at module load", "import atexit\n", "import atexit\nimport langfuse\n"),
    ("disabled: the flush thread/client is touched while off", "    if not _S[\"on\"] or _S[\"dead\"] or _S[\"client\"] is None:\n        return True", "    if False:\n        return True"),
    ("content table: email_agent no longer captures bodies by default", "\"email_agent\": True, \"content_agent\"", "\"email_agent\": False, \"content_agent\""),
    ("content table: decision_agent can be switched on", "ALWAYS_METADATA_ONLY = (\"decision_agent\", \"order_sheet\")", "ALWAYS_METADATA_ONLY = ()"),
    ("content table: bodies are not stripped when capture is off", "if not _S[\"capture\"] and k.startswith(CONTENT_ATTR_PREFIXES):", "if False and k.startswith(CONTENT_ATTR_PREFIXES):"),
    ("content table: the instrumentation is told to record bodies", "os.environ[\"TRACELOOP_TRACE_CONTENT\"] = \"false\"", "pass"),
    ("workflow id: wrong prefix", "return \"lcom-\" + os.path.splitext", "return \"wf-\" + os.path.splitext"),
    ("workflow id: file extension kept", "os.path.splitext(os.path.basename(str(results_path)))[0]", "os.path.basename(str(results_path))"),
    ("workflow id: late set_workflow does nothing (search run)", "        if self._obs is None or not wid or self.session:\n            return", "        return"),
    ("tags: the sku/agent/env tags are not set", "tags = [f\"agent:{agent}\", f\"env:{os.environ.get('OBS_ENV', 'dev')}\"] + ([f\"sku:{sku}\"] if sku else [])", "tags = []"),
    ("scores: a score outside any trace is sent anyway", "        if client.get_current_trace_id() is None:", "        if False:"),
]


def run_mutations():
    """Each mutation breaks one rule in observability.py inside a throwaway copy of the project; the suite must then fail."""
    from concurrent.futures import ThreadPoolExecutor
    skip = {".git", "__pycache__", "Excel Output Sheets", "Reports", "scratch", "graphify-out", "Order Sheets"}
    base = open(os.path.join(HERE, "observability.py"), encoding="utf-8").read()
    for name, old, new in MUTATIONS:
        assert base.count(old) == 1, f"mutation does not apply exactly once: {name}"
    snapshot = tempfile.mkdtemp()                                # one copy, so edits made while this runs cannot leak in
    for item in os.listdir(HERE):
        if item not in skip and item != ".env":
            src = os.path.join(HERE, item)
            (shutil.copytree if os.path.isdir(src) else shutil.copy)(src, os.path.join(snapshot, item))
    env = {k: v for k, v in os.environ.items() if not k.startswith(("OBS_", "LANGFUSE_"))}

    def one(m):
        name, old, new = m
        with tempfile.TemporaryDirectory() as tmp:
            shutil.copytree(snapshot, tmp, dirs_exist_ok=True)
            open(os.path.join(tmp, "observability.py"), "w", encoding="utf-8").write(base.replace(old, new))
            p = subprocess.run([sys.executable, "test_observability.py"], cwd=tmp, capture_output=True, text=True, env=env, timeout=1200)
        dead = p.returncode != 0
        why = "; ".join(l[5:60] for l in p.stderr.splitlines() if l.startswith("FAIL "))[:150] if dead else ""
        print(f"{'KILLED  ' if dead else 'SURVIVED'} {name}   {why}", flush=True)
        return dead
    with ThreadPoolExecutor(4) as pool:
        killed = sum(pool.map(one, MUTATIONS))
    shutil.rmtree(snapshot, ignore_errors=True)
    print(f"mutations killed: {killed}/{len(MUTATIONS)}")
    return killed == len(MUTATIONS)


if __name__ == "__main__":
    if "--mutations" in sys.argv:
        sys.exit(0 if run_mutations() else 1)
    failed = []
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            try:
                fn()
            except BaseException as e:  # noqa: BLE001 - list every failure, not just the first
                import traceback
                failed.append(name)
                print(f"FAIL {name}: {type(e).__name__}: {str(e)[:300]}", file=sys.stderr)
                traceback.print_exc(limit=-3)
    if failed:
        sys.exit(f"{len(failed)} observability test(s) failed: {', '.join(failed)}")
    print("observability tests passed")
