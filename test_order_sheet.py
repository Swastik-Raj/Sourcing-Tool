"""Offline tests for order_sheet.py: small synthetic fixtures, plus the real approved_orders.csv if one exists.
Run: python test_order_sheet.py"""
import csv
import json
import os
import re
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import order_sheet as os_   # noqa: E402  (module is swapped by the mutation harness)
from decision_agent import EXPORT_FIELDS, NOT_AN_ORDER   # noqa: E402

NOW = datetime(2026, 10, 7, 12, 0, tzinfo=timezone.utc)
URL_A = "https://www.alibaba.com/product-detail/A-cable_111.html"
URL_B = "https://www.aliexpress.us/item/222.html"
URL_C = "https://www.made-in-china.com/product/ccc/Thing.html"
HEADS = ["SKU", "Product", "Keyword", "L-Com Unit Price", "L-Com Price Source", "L-Com Price Date", "Recommendation",
         "Recommended Manufacturer", "Recommended URL", "Ordering Note"]
for _n in range(1, 6):
    HEADS += [f"Manufacturer {_n} {s}" for s in ("Name", "Accuracy", "Listed Price", "Unit Price", "MOQ", "URL", "Match Tier", "Comment", "vs. L-Com Price")]


def cand(name, url, listed, unit, moq, tier="Exact", comment="", acc=95):
    return dict(name=name, url=url, listed=listed, unit=unit, moq=moq, tier=tier, comment=comment, acc=acc)


def results_file(path, rows):
    """rows: list of dict(sku, keyword, rec=(name,url) or None, note, cands=[cand])."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(HEADS)
    for r in rows:
        rec = r.get("rec") or ("", "")
        line = {"SKU": r["sku"], "Product": r["keyword"], "Keyword": r["keyword"], "L-Com Unit Price": 5,
                "L-Com Price Source": r.get("lsrc", "l-com.com"), "L-Com Price Date": "2026-10-01", "Recommendation": "x",
                "Recommended Manufacturer": rec[0], "Recommended URL": rec[1], "Ordering Note": r.get("note", "")}
        for n, c in enumerate(r["cands"], start=1):
            line.update({f"Manufacturer {n} Name": c["name"], f"Manufacturer {n} Accuracy": c["acc"],
                         f"Manufacturer {n} Listed Price": c["listed"], f"Manufacturer {n} Unit Price": c["unit"],
                         f"Manufacturer {n} MOQ": c["moq"], f"Manufacturer {n} URL": c["url"],
                         f"Manufacturer {n} Match Tier": c["tier"], f"Manufacturer {n} Comment": c["comment"],
                         f"Manufacturer {n} vs. L-Com Price": "-20%"})
        ws.append([line.get(h) for h in HEADS])
    wb.save(path)


def approved_file(path, lines, first=None):
    """lines: dict overrides per row. Writes the decision agent's export layout."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([first if first is not None else NOT_AN_ORDER])
        w.writerow(EXPORT_FIELDS)
        for i, d in enumerate(lines, start=1):
            base = {"sku": "AAA", "product": "p", "seller": "Seller A", "listing_url": URL_A, "qty": "200", "max_unit_price": "1.00",
                    "line_cap": "200.00", "approval_id": f"A-{i:04d}", "approved_by": "Neeraj", "approved_at": "2026-10-06T10:00:00Z",
                    "expires_at": "2026-10-13T10:00:00Z", "source_results_file": "res.xlsx", "flags": "", "channel": "slack"}
            base.update(d)
            w.writerow([base.get(k, "") for k in EXPORT_FIELDS])


class Box:
    """A throwaway folder holding the fixtures; build() reads and writes only inside it."""
    def __init__(self, rows, lines, state=None, excl=None, first=None, results_name="res.xlsx"):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name
        self.results = os.path.join(self.dir, results_name)
        self.approved = os.path.join(self.dir, "approved_orders.csv")
        self.state, self.excl = os.path.join(self.dir, "email_state.json"), os.path.join(self.dir, "reviewer_exclusions.csv")
        results_file(self.results, rows)
        approved_file(self.approved, lines, first)
        json.dump(state or {}, open(self.state, "w"))
        with open(self.excl, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["sku", "supplier_or_url", "reason"])
            for e in excl or []:
                w.writerow(e)

    def build(self, duty=Decimal("0.35"), at=NOW):
        return os_.build(self.approved, self.results, duty, os.path.join(self.dir, "out"), self.state, self.excl,
                         os.path.join(self.dir, "none.csv"), at)

    def sheet(self, b, formulas=True):
        return openpyxl.load_workbook(b["path"], data_only=not formulas)


def basic_rows():
    return [dict(sku="AAA", keyword="Test cable", rec=("Seller A", URL_A), note="Check the plug type.",
                 cands=[cand("Seller A", URL_A, "$0.90", 0.90, "100 pieces"), cand("Seller B", URL_B, "$1.20", 1.20, "50 pieces", "Close",
                                                                                   "Cable is longer."), cand("Seller C", URL_C, "$0.80", 0.80, "10 pieces")])]


def refuses(fn, *words):
    try:
        fn()
    except os_.OrderSheetError as e:
        assert all(w in str(e) for w in words), str(e)
        return
    raise AssertionError("expected a refusal")


def test_happy_path_builds_three_sheets_and_the_reference_layout():
    box = Box(basic_rows(), [{}])
    b = box.build()
    wb = box.sheet(b)
    assert wb.sheetnames == ["Order List", "Other options", "Notes"]
    ws = wb["Order List"]
    assert [c.value for c in ws[7]][:16] == os_.ORDER_HEADERS[:16]
    assert ws["F8"].value == "Open listing" and ws["F8"].hyperlink.target == URL_A
    assert ws["I8"].value == 200 and ws["I8"].fill.fgColor.rgb.endswith("FFFF00") and ws["I8"].font.color.rgb.endswith("0000FF")
    assert ws["E8"].value == "Alibaba" and ws["Q8"].value == "add-to-cart when logged in"
    assert ws["B9"].value == "Total" and ws.freeze_panes == "C8"
    assert "Check the plug type." in ws["P8"].value
    assert "listing price, not a seller quote" in ws["O8"].value
    assert b["totals"]["K"] == Decimal("180.00") and b["totals"]["L"] == Decimal("243.00")


def test_header_validation_and_not_an_order_check():
    box = Box(basic_rows(), [{}], first="approved orders")
    refuses(box.build, "does not start with", "NOT AN ORDER")
    box = Box(basic_rows(), [{}])
    lines = open(box.approved, encoding="utf-8").read().replace(",listing_url,", ",url,", 1)
    open(box.approved, "w", encoding="utf-8").write(lines)
    refuses(box.build, "missing required column", "listing_url")
    box = Box(basic_rows(), [])
    refuses(box.build, "zero valid")
    box = Box(basic_rows(), [{"qty": "abc"}])
    refuses(box.build, "zero valid", "qty")
    box = Box(basic_rows(), [{}])
    r = open(box.results, "rb").read()
    refuses(lambda: os_.read_results(os.path.join(box.dir, "nope.xlsx")), "not found")
    wb = openpyxl.load_workbook(box.results)
    wb["Results"].cell(1, HEADS.index("Ordering Note") + 1).value = "Renamed"
    wb.save(box.results)
    refuses(box.build, "missing column", "Ordering Note")


def test_columns_are_read_by_header_name_not_position():
    box = Box(basic_rows(), [{}])
    lines = open(box.approved, encoding="utf-8").read().splitlines()
    head = lines[1].split(",")
    swap = head[:]
    swap[0], swap[1] = swap[1], swap[0]                           # sku <-> product columns swapped in the header AND rows
    rows = [lines[0], ",".join(swap)] + [",".join([c.split(",")[1], c.split(",")[0]] + c.split(",")[2:]) for c in lines[2:]]
    open(box.approved, "w", encoding="utf-8").write("\n".join(rows) + "\n")
    assert box.build()["rows"][0]["sku"] == "AAA"


def test_expired_approvals_are_skipped_and_listed():
    box = Box(basic_rows(), [{"approved_at": "2026-09-20T10:00:00Z", "expires_at": "2026-09-27T10:00:00Z"}])
    refuses(box.build, "AAA", "expired")
    box = Box(basic_rows(), [{"approved_at": "2026-09-20T10:00:00Z", "expires_at": "2026-12-31T10:00:00Z"}])    # file's own date is wrong
    refuses(box.build, "expired")                                                                                   # 7-day rule re-checked
    box = Box(basic_rows() + [dict(sku="BBB", keyword="K2", rec=("S", URL_B), cands=[cand("S", URL_B, "$1", 1, "10")])],
              [{"approved_at": "2026-09-20T10:00:00Z", "expires_at": "2026-09-27T10:00:00Z"}, {"sku": "BBB", "listing_url": URL_B, "qty": "20"}])
    b = box.build()
    assert [r["sku"] for r in b["rows"]] == ["BBB"] and [ap["sku"] for ap, _ in b["expired"]] == ["AAA"]


def test_matching_is_by_sku_and_url_never_sku_alone():
    box = Box(basic_rows(), [{"listing_url": "https://www.alibaba.com/product-detail/Other_999.html"}])
    refuses(box.build, "AAA", "listing URL not found")
    box = Box(basic_rows(), [{"sku": "NOPE"}])
    refuses(box.build, "NOPE", "SKU not found")
    box = Box(basic_rows(), [{"listing_url": URL_B, "seller": "Seller B", "max_unit_price": "2"}])                  # a different listing of AAA
    b = box.build()
    assert b["rows"][0]["supplier"] == "Seller B" and b["rows"][0]["unit"] == Decimal("1.2")
    box = Box(basic_rows(), [{"listing_url": URL_A + "?spm=abc#x"}])                                                  # tracking params don't matter
    assert box.build()["rows"][0]["unit"] == Decimal("0.9")
    box = Box(basic_rows(), [{}, {"sku": "NOPE"}])
    b = box.build()
    assert len(b["rows"]) == 1 and len(b["could_not_build"]) == 1                                                     # one bad line doesn't sink the rest


def test_below_moq_flag_is_in_the_sheet_and_the_console_summary():
    box = Box(basic_rows(), [{"qty": "50"}])
    b = box.build()
    assert any(w.startswith("BELOW MOQ: AAA") for w in b["warnings"])
    ws = box.sheet(b)["Order List"]
    assert ws["J8"].value.startswith("=IF(") and "BELOW MOQ" in ws["J8"].value
    assert os_.recalc_builtin(b["path"])[0][("Order List", "J8")] == "BELOW MOQ"
    import io, contextlib
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        os_.print_summary(b, box.approved, box.results)
    assert "WARNING: BELOW MOQ: AAA" in out.getvalue()
    ok = Box(basic_rows(), [{"qty": "100"}]).build()
    assert not any("BELOW MOQ" in w for w in ok["warnings"])
    rows = basic_rows()
    rows[0]["cands"][0]["moq"] = ""
    nbox = Box(rows, [{}])
    nb = nbox.build()
    assert os_.recalc_builtin(nb["path"])[0][("Order List", "J8")] == "MOQ NOT STATED"


def test_large_moq_total_warns_with_sku_and_amount():
    rows = [dict(sku="BIG", keyword="Cheap", rec=("E-HONG", URL_B), cands=[cand("E-HONG", URL_B, "$0.55", 0.55, "10,000 pieces")])]
    box = Box(rows, [{"sku": "BIG", "listing_url": URL_B, "qty": "200", "max_unit_price": "1"}])
    b = box.build()
    w = [x for x in b["warnings"] if x.startswith("MOQ COST OVER LIMIT")]
    assert len(w) == 1 and "BIG" in w[0] and "$5,500.00" in w[0], b["warnings"]
    assert not [x for x in Box(basic_rows(), [{}]).build()["warnings"] if "OVER LIMIT" in x]                         # $90 MOQ cost: quiet


def test_non_usd_rows_are_flagged_and_left_out_of_totals():
    state = {os_.listing_key("AAA", URL_A, "Seller A"): {"status": "quoted", "currency": "CNY", "sample_unit_price": "6.5",
                                                         "evidence": {"sample_unit_price": "6.5 RMB"}}}
    rows = basic_rows() + [dict(sku="BBB", keyword="K2", rec=("S", URL_B), cands=[cand("S", URL_B, "$1", 1, "10")])]
    box = Box(rows, [{}, {"sku": "BBB", "listing_url": URL_B, "qty": "20"}], state)
    b = box.build()
    assert b["non_usd"][0][1] == "CNY" and b["totals"]["K"] == Decimal("20.00")
    ws = box.sheet(b)["Order List"]
    assert ws["K8"].value == "not in total" and ws["K9"].value == "=G9*I9" and ws["K10"].value == "=SUM(K8:K9)"
    assert "NOT in the totals" in ws["P8"].value


def test_reviewer_excluded_rows_only_in_other_options():
    box = Box(basic_rows(), [{}, {"listing_url": URL_B, "seller": "Seller B", "max_unit_price": "2"}], excl=[["AAA", "Seller B", "wrong cable length"]])
    b = box.build()
    assert [r["supplier"] for r in b["rows"]] == ["Seller A"] and len(b["excluded"]) == 1
    wb = box.sheet(b)
    assert "Seller B" not in [wb["Order List"].cell(r, 4).value for r in range(8, 12)]
    other = [(r[1].value, r[6].value) for r in wb["Other options"].iter_rows(min_row=2)]
    assert ("Seller B", "Excluded by reviewer_exclusions.csv: wrong cable length") in other
    assert "excluded by reviewer_exclusions.csv" in [r for r in wb["Notes"].iter_rows(values_only=True) if r[0] == "Left out of this sheet"][0][1]


def test_other_options_hold_the_other_candidates_with_the_files_own_reason():
    box = Box(basic_rows(), [{}])
    b = box.build()
    ws = box.sheet(b)["Other options"]
    assert [ws.cell(1, c).value for c in range(1, 9)] == os_.OTHER_HEADERS
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    assert set(rows) == {"Seller B", "Seller C"}
    assert rows["Seller B"][4].value == "=C2*D2" and "Cable is longer" in rows["Seller B"][6].value and "Close" in rows["Seller B"][6].value
    assert rows["Seller B"][7].value == URL_B


def test_a_quote_is_used_only_when_usd_and_evidence_verified():
    key = os_.listing_key("AAA", URL_A, "Seller A")
    tier = {"min_qty": 200, "unit_price": "0.70", "evidence": "200+ pcs $0.70"}
    ok = {"status": "quoted", "currency": "USD", "bulk_tiers": [tier], "evidence": {}}
    assert Box(basic_rows(), [{}], {key: ok}).build()["rows"][0]["unit"] == Decimal("0.70")
    assert "seller quote, bulk tier" in Box(basic_rows(), [{}], {key: ok}).build()["rows"][0]["how_verified"]
    bare = {**ok, "bulk_tiers": [{"min_qty": 200, "unit_price": "0.70"}]}                                           # no evidence
    r = Box(basic_rows(), [{}], {key: bare}).build()["rows"][0]
    assert r["unit"] == Decimal("0.9") and "listing price, not a seller quote" in r["how_verified"]
    small = Box(basic_rows(), [{"qty": "100"}], {key: ok}).build()["rows"][0]                                       # tier not met at 100
    assert small["unit"] == Decimal("0.9")
    sample = {"status": "quoted", "currency": "USD", "sample_unit_price": "1.10", "sample_quantity": 5, "price_basis": "per_unit",
              "evidence": {"sample_unit_price": "$1.10 each"}}
    assert Box(basic_rows(), [{"qty": "100"}], {key: sample}).build()["rows"][0]["unit"] == Decimal("0.9")            # qty beyond the sample
    unknown = {**ok, "currency": "unknown"}
    b = Box(basic_rows(), [{}], {key: unknown}).build()
    assert b["rows"][0]["unit"] == Decimal("0.9") and b["non_usd"]
    other_listing = {os_.listing_key("AAA", URL_B, "Seller B"): ok}                                                 # a quote for another listing
    assert Box(basic_rows(), [{}], other_listing).build()["rows"][0]["unit"] == Decimal("0.9")


def test_price_ranges_say_range_high_end():
    rows = [dict(sku="RNG", keyword="Range", rec=("R", URL_C), cands=[cand("R", URL_C, "US$3.82-5.22", 5.22, "10 pieces")])]
    box = Box(rows, [{"sku": "RNG", "listing_url": URL_C, "qty": "20", "max_unit_price": "6"}])
    r = box.build()["rows"][0]
    assert r["unit"] == Decimal("5.22") and "range, high end" in r["how_verified"] and "high end is used" in r["check"]
    assert r["how"] == "inquiry-based, no checkout seen"


    b = Box(basic_rows(), [{"max_unit_price": "0.50"}]).build()
    assert any(w.startswith("PRICE OVER APPROVED MAX: AAA") for w in b["warnings"])


def test_platform_table_and_unknown_platform():
    assert os_.platform_of("https://www.alibaba.com/x")[1] == "add-to-cart when logged in"
    assert os_.platform_of("https://www.aliexpress.us/x") == ("AliExpress", "add-to-cart when logged in")
    assert os_.platform_of("https://m.made-in-china.com/x")[1] == "inquiry-based, no checkout seen"
    assert os_.platform_of("https://example.org/x") == ("example.org", "unknown, check the listing")
    assert os_.platform_of("")[1] == "unknown, check the listing"


def test_formulas_present_and_nothing_computed_is_hardcoded():
    rows = basic_rows() + [dict(sku="BBB", keyword="K2", rec=("S", URL_B), cands=[cand("S", URL_B, "$1", 1, "10")])]
    box = Box(rows, [{}, {"sku": "BBB", "listing_url": URL_B, "qty": "20"}])
    b = box.build()
    ws = box.sheet(b)["Order List"]
    for x in (8, 9):
        assert ws[f"J{x}"].value == f'=IF(H{x}="","MOQ NOT STATED",IF(I{x}<H{x},"BELOW MOQ","ok"))'
        assert ws[f"K{x}"].value == f"=G{x}*I{x}" and ws[f"L{x}"].value == f"=K{x}*(1+$C$4)"
        assert ws[f"M{x}"].value == f"=G{x}*MAX(200,H{x})" and ws[f"N{x}"].value == f"=G{x}*MAX(500,H{x})"
    for col in "KLMN":
        assert ws[f"{col}10"].value == f"=SUM({col}8:{col}9)"
    assert all(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith("=") for r in (8, 9, 10) for c in range(11, 15))
    assert ws["G8"].value == 0.9 and ws["I8"].value == 200                                                          # inputs are plain numbers
    wb = box.sheet(b)
    assert wb.calculation.fullCalcOnLoad
    for sheet in wb.sheetnames:                                                                                     # no error values anywhere
        for row in wb[sheet].iter_rows(values_only=True):
            assert not any(v in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!") for v in row)


def test_total_equals_sum_of_lines_and_the_decimal_cross_check():
    rows = basic_rows() + [dict(sku="BBB", keyword="K2", rec=("S", URL_B), cands=[cand("S", URL_B, "$0.33", 0.33, "10")])]
    box = Box(rows, [{"qty": "300"}, {"sku": "BBB", "listing_url": URL_B, "qty": "7", "max_unit_price": "1"}])
    b = box.build()
    values, count = os_.recalc_builtin(b["path"])
    k8, k9, k10 = (values[("Order List", f"K{x}")] for x in (8, 9, 10))
    assert k10 == k8 + k9 == Decimal("270") + Decimal("2.31")
    assert b["totals"]["K"] == Decimal("272.31") and b["totals"]["L"] == Decimal("367.62")       # 272.31 * 1.35 = 367.6185 -> half up
    assert values[("Order List", "M10")] == Decimal("0.9") * 200 + Decimal("0.33") * 200 and b["totals"]["N"] == Decimal("615.00")


def test_verification_fails_loudly_on_a_wrong_total():
    box = Box(basic_rows(), [{}])
    b = box.build()
    wb = openpyxl.load_workbook(b["path"])
    wb["Order List"]["K9"] = "=SUM(K8:K8)+1"
    wb.save(b["path"])
    refuses(lambda: os_.verify_workbook(b["path"], b["rows"], Decimal("0.35")), "Verification failed", "K9")
    wb["Order List"]["K9"] = 180
    wb.save(b["path"])
    refuses(lambda: os_.verify_workbook(b["path"], b["rows"], Decimal("0.35")), "not a formula")
    wb["Order List"]["K9"] = "=G8/0"
    wb.save(b["path"])
    refuses(lambda: os_.verify_workbook(b["path"], b["rows"], Decimal("0.35")), "division by zero")


def test_duty_cell_drives_the_duty_column():
    box = Box(basic_rows(), [{}])
    b = box.build(duty=Decimal("0.10"))
    wb = box.sheet(b)
    ws = wb["Order List"]
    assert ws["C4"].value == 0.1 and ws["C4"].fill.fgColor.rgb.endswith("FFFF00") and "TariffTracker" in ws["D4"].value
    assert b["totals"]["L"] == Decimal("198.00")
    ws["C4"] = 0.5
    wb.save(b["path"])
    assert os_.recalc_builtin(b["path"])[0][("Order List", "L8")] == Decimal("270")                                 # change C4, L follows
    refuses(lambda: box.build(duty=Decimal("1.5")), "--duty")


def test_notes_sheet_names_files_and_hashes():
    box = Box(basic_rows(), [{}])
    b = box.build()
    notes = {r[0]: r[1] for r in box.sheet(b)["Notes"].iter_rows(values_only=True) if r[0]}
    assert {"Cart link", "Quantities", "Samples", "Differences vs. results file", "Price caveats", "Duty", "Sources"} <= set(notes)
    assert os_.sha256(box.results) in notes["Sources"] and os_.sha256(box.approved) in notes["Sources"] and "res.xlsx" in notes["Sources"]
    assert "No cart link" in notes["Cart link"] and "UNVERIFIED ASSUMPTION" in notes["Cart link"] and "TariffTracker" in notes["Duty"]


def test_the_script_only_reads_and_never_calls_out():
    box = Box(basic_rows(), [{}])
    before = {p: open(p, "rb").read() for p in (box.approved, box.results, box.state, box.excl)}
    box.build()
    assert before == {p: open(p, "rb").read() for p in before}
    src = open(os.path.join(HERE, "order_sheet.py"), encoding="utf-8").read()
    imports = set(re.findall(r"^(?:import|from)\s+([\w.]+)", src, re.MULTILINE))
    assert not imports & {"requests", "httpx", "urllib.request", "anthropic", "socket", "smtplib"}, imports
    for name in ("sourcing_agent", "nimble"):
        assert name not in src.lower().replace("sourcing_results", "")


def test_the_real_approved_orders_file_if_there_is_one():
    path = os.path.join(HERE, "approved_orders.csv")
    if not os.path.exists(path):
        print("  (no real approved_orders.csv here; covered by the sandbox dry run)")
        return
    approved, _ = os_.read_approved(path)
    assert all(a["qty"] > 0 for a in approved)


def test_money_rounds_half_up_in_the_cross_check():
    rows = [dict(sku="TIE", keyword="Tie", rec=("T", URL_B), cands=[cand("T", URL_B, "$0.10", 0.10, "1 piece")])]
    box = Box(rows, [{"sku": "TIE", "listing_url": URL_B, "qty": "1", "max_unit_price": "1"}])
    b = box.build(duty=Decimal("0.25"))                                    # 0.10 * 1.25 = 0.125 exactly
    assert b["totals"]["L"] == Decimal("0.13")                             # half-up (half-even would give 0.12)


def test_excluded_only_sku_and_csv_currency_and_unproven_sample_price():
    rows = basic_rows() + [dict(sku="CCC", keyword="K3", rec=("S", URL_B), cands=[cand("S", URL_B, "$1", 1, "10")])]
    box = Box(rows, [{}, {"sku": "CCC", "listing_url": URL_B, "seller": "S", "max_unit_price": "2"}], excl=[["CCC", "S", "wrong part"]])
    b = box.build()
    other = [(r[0].value, r[1].value, r[6].value) for r in box.sheet(b)["Other options"].iter_rows(min_row=2)]
    assert ("CCC", "S", "Excluded by reviewer_exclusions.csv: wrong part") in other and all(r["sku"] != "CCC" for r in b["rows"])
    key = os_.listing_key("AAA", URL_A, "Seller A")
    sample = {"status": "quoted", "currency": "USD", "sample_unit_price": "0.40", "sample_quantity": 500, "price_basis": "per_unit", "evidence": {}}
    assert Box(basic_rows(), [{}], {key: sample}).build()["rows"][0]["unit"] == Decimal("0.9")            # no evidence -> listing price
    sample["evidence"] = {"sample_unit_price": "$0.40 each"}
    assert Box(basic_rows(), [{}], {key: sample}).build()["rows"][0]["unit"] == Decimal("0.40")           # proven, within the sample qty
    cny = Box(basic_rows(), [{"quoted_currency": "CNY"}])
    assert cny.build()["non_usd"][0][1] == "CNY"                                                           # the export's own currency counts too


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            fn()
    print("order sheet tests passed")
