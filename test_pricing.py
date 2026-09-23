"""Offline check of unit-price / L-Com margin / recommendation logic. Run: python test_pricing.py"""
import os
import tempfile

import openpyxl

from sourcing_agent import (
    AttributeBreakdown, Candidate, SourcingResult, implausible_units, is_garbled, recommend, unit_price, write_excel_report,
)

FULL = AttributeBreakdown(product_type=20, category_spec=20, shielding_material=20, mount_form=20, gender_pins=20)
NINETY = AttributeBreakdown(product_type=20, category_spec=20, shielding_material=20, mount_form=10, gender_pins=20)


def cand(name, total, qty, conf, breakdown=FULL):
    return Candidate(manufacturer=name, price_total=total or 0, quantity_covered=qty or 0, unit_price_note="",
                     unit_price_confidence=conf, attribute_breakdown=breakdown, email=f"{name}@x", url=f"u/{name}")


def result(*cands):
    return SourcingResult(product="x", no_match=False, candidates=list(cands))


# Bundle price is divided down; ambiguous never yields a unit price.
assert unit_price(cand("A", 19.75, 10, "inferred")) == 1.975
assert unit_price(cand("A", 19.75, 10, "ambiguous")) is None
assert unit_price(cand("A", 19.75, None, "stated")) is None

# 100% match barely under L-Com loses to 90% match with a big gap.
idx, reason = recommend(result(cand("Close", 21.0, 1, "stated"), cand("Cheap", 4.0, 1, "stated", NINETY),
                               cand("Mystery", 1.0, 1, "ambiguous")), 22.59)
assert idx == 1, reason

# 30-79% cheaper no longer qualifies (bar is 80%).
idx, reason = recommend(result(cand("Half", 11.0, 1, "stated")), 22.59)
assert idx is None and ">= 80% cheaper" in reason and "do not source" in reason, reason

# HDFF regression: price/MOQ gave $0.0032 "inferred" next to two $1.09 "stated" -> implausible, not picked.
hdff = result(cand("A", 1.09, 1, "stated"), cand("B", 1.09, 1, "stated"), cand("FARSINCE", 1.58, 500, "inferred"))
assert implausible_units(hdff) == {2}
idx, reason = recommend(hdff, 23.79)
assert idx == 0, reason

# Implausible price can't make a product clear the bar on its own.
idx, reason = recommend(result(cand("Real", 20.0, 1, "stated"), cand("Bug", 1.58, 500, "inferred")), 23.79)
assert idx is None and "implausible" in reason, reason

# Stated beats a (plausible) inferred price at equal accuracy, even though inferred is cheaper.
idx, reason = recommend(result(cand("Inferred", 1.5, 1, "inferred"), cand("Stated", 2.0, 1, "stated")), 23.79)
assert idx == 1, reason

# Garbled model output (real ECF504-BAS values) is detected; clean output isn't.
bad = cand("Do nggu an", 2.10, 1, "stated")
bad.url, bad.distributor_site = "ht tps ://w ww .al ib ab a.c om/ pro duc t-d eta il/ US B-a", " al ib ab a.c om"
good = cand("Dongguan", 2.10, 1, "stated")
good.url, good.distributor_site = "https://www.alibaba.com/product-detail/USB-a-Male-to-USB-B_1601687478124.html", "alibaba.com"
assert is_garbled(result(good, bad)) and not is_garbled(result(good)) and not is_garbled(None)

# No L-Com price -> no recommendation.
assert recommend(result(cand("A", 1.0, 1, "stated")), None)[0] is None

# Excel: Recommended columns filled + green; manufacturer blocks not green; blank when nothing qualifies.
path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
write_excel_report([
    ({"sku": "S1", "description": "d", "lcom_price": 22.59},
     result(cand("Close", 15.0, 1, "stated"), cand("Cheap", 2.0, 1, "stated"))),
    ({"sku": "S2", "description": "d", "lcom_price": 22.59}, result(cand("Close", 21.0, 1, "stated"))),
], path)
wb = openpyxl.load_workbook(path)
ws = wb["Results"]
header = [c.value for c in ws[1]]
rec = {h: header.index(h) for h in ("Recommended Manufacturer", "Recommended Email",
                                    "Recommended URL", "Recommended Unit Price")}
row = ws[2]
assert [row[i].value for i in rec.values()] == ["Cheap", "Cheap@x", "u/Cheap", 2.0]
assert all(row[i].fill.fgColor.rgb.endswith("C6EFCE") for i in rec.values())
assert not any(c.fill.fgColor.rgb.endswith("C6EFCE") for c in row[max(rec.values()) + 1:])
assert all(ws[3][i].value is None for i in rec.values())
assert not any(c.fill.fgColor.rgb.endswith("C6EFCE") for c in ws[3])
assert wb["Comparison"].cell(3, 8).value == "YES"
print("ok")
