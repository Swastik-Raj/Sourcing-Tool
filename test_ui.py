"""Offline tests for the web UI: FastAPI TestClient, the stub agent only, no network, no real agent runs.
Run:  $env:OBS_ENABLED = "0"; python test_ui.py"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path
from types import SimpleNamespace

os.environ["OBS_ENABLED"] = "0"
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from fastapi.testclient import TestClient

import observability as obs
from ui import commands, jobs as jobs_mod, security, workflows
from ui.app import create_app
from ui.commands import JobRejected
from ui.config import Settings, load_settings
from ui.jobs import JobBusy, JobManager, write_json_atomic

ORIGIN = {"origin": "http://127.0.0.1:8000"}
FAKE_KEY = "sk-ant-FAKEKEYVALUE1234567890"
TESTS = []


def test(f):
    TESTS.append(f)
    return f


def make_env(stub=True):
    root = Path(tempfile.mkdtemp(prefix="ui_test_"))
    project = root / "project"
    (project / "Excel Output Sheets").mkdir(parents=True)
    shutil.copy(Path(__file__).with_name("email_agent.py"), project)    # the UI reads rows with the PROJECT's own email agent, never the one beside the UI
    for name in ("sourcing_results_20260101_100000.xlsx", "sourcing_results_20260202_100000.xlsx"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(["SKU", "Product", "Recommended URL"])
        ws.append(["AB-1", "Cable", "https://example.test/a"])
        ws.append(["CD-2", "Adapter", ""])
        wb.save(project / "Excel Output Sheets" / name)
        t = time.time() - (100000 if "0101" in name else 0)
        os.utime(project / "Excel Output Sheets" / name, (t, t))
    s = Settings(project=project, data=project / "ui_data", enable_stub=stub)
    app = create_app(s)
    client = TestClient(app, base_url="http://127.0.0.1:8000", raise_server_exceptions=False)
    return SimpleNamespace(root=root, project=project, settings=s, app=app, client=client, jobs=app.state.jobs)


def cleanup(env):
    for j in env.jobs.list():
        env.jobs.cancel(j["id"])
    time.sleep(0.3)
    shutil.rmtree(env.root, ignore_errors=True)


def token(env):
    html = env.client.get("/jobs").text
    return re.search(r'name="csrf_token" value="([0-9a-f]+)"', html)[1]


def raises(exc, fn, *a, **kw):
    try:
        fn(*a, **kw)
    except exc as e:
        return e
    raise AssertionError(f"expected {exc.__name__} from {getattr(fn, '__name__', fn)}{a}{kw}")


class Skip(Exception):
    """Raised by a test that cannot run on this system. The runner counts it and prints the reason; it is never silent."""


def native(*parts) -> str:
    """A path inside the project as the UI hands it to an agent: this system's own separator (a backslash on Windows)."""
    return str(Path(*parts))


def process_gone(pid: int) -> bool:
    """Whether a process has ended. Windows: tasklist no longer lists it. Linux: /proc has no entry, or only a finished one
    that its parent has not collected yet."""
    if os.name == "nt":
        return str(pid) not in subprocess.run(["tasklist", "/FI", f"PID eq {pid}"], capture_output=True, text=True).stdout
    try:
        return Path(f"/proc/{pid}/stat").read_text().rsplit(")", 1)[1].split()[0] == "Z"
    except OSError:
        return True


# ---------- allowlist ----------

@test
def allowlist_rejects_unknown_commands_bad_skus_paths_and_flags(env):
    j = env.jobs
    raises(JobRejected, j.start, "rm.everything")
    raises(JobRejected, j.start, "email_agent.send")            # not a command id: sending exists only as 'emails.send', behind its gated route
    raises(JobRejected, j.start, "emails.send")                 # ...and that needs a results file and the agent's --confirm-count
    raises(JobRejected, j.start, "emails.send", {"results": "sourcing_results_20260101_100000.xlsx"})
    raises(JobRejected, j.start, "emails.send", {"results": "sourcing_results_20260101_100000.xlsx", "confirm_count": "1", "yes": "1"})
    for cmd in ("emails.send", "approvals.record", "approvals.export"):     # the generic job page can never start one
        r = env.client.post("/jobs/run", data={"command": cmd, "csrf_token": token(env)}, headers=ORIGIN, follow_redirects=False)
        assert r.status_code == 400 and "not run from this page" in r.text, cmd
    assert j.list() == []
    for bad in ("x; calc", "-x", "a b", "../x", "A" * 41, "", "a\nb", "$(x)", "a|b"):
        raises(JobRejected, j.start, "stub.read", {"mode": "ok", "sku": bad})
    for bad in ("C:\\Windows\\win.ini", "..\\x.xlsx", "/etc/passwd", "Excel Output Sheets\\sourcing_results_20260101_100000.xlsx", "nope.xlsx"):
        raises(JobRejected, j.start, "stub.read", {"mode": "ok", "results": bad})
    raises(JobRejected, j.start, "stub.read", {"mode": "ok", "yes": "--yes"})         # extra flag/param
    raises(JobRejected, j.start, "stub.read", {"mode": "--help"})                       # not in the enum
    raises(JobRejected, j.start, "stub.read", {})                                       # missing mode
    argv = env.jobs.commands["stub.read"].argv({"mode": "ok", "sku": "C&P9M", "results": "sourcing_results_20260101_100000.xlsx"})
    assert argv == ["ok", "--sku", "C&P9M", "--results", native("Excel Output Sheets", "sourcing_results_20260101_100000.xlsx")], argv
    assert not any(c.startswith("stub") for c in JobManager(Settings(project=env.project, data=env.project / "d2")).commands), "stub must be off by default"


# ---------- state-changing lock ----------

@test
def lock_refuses_second_state_job_and_releases(env):
    j = env.jobs
    first = j.start("stub.state", {"mode": "slow"})
    raises(JobBusy, j.start, "stub.state", {"mode": "ok"})
    assert j.running_state_job()["id"] == first["id"]
    ro = j.start("stub.read", {"mode": "ok"})                                         # read-only runs alongside
    assert j.wait(ro["id"])["status"] == "succeeded"
    assert j.cancel(first["id"])
    assert j.wait(first["id"])["status"] == "cancelled"
    done = j.start("stub.state", {"mode": "ok"})                                      # released after cancel
    assert j.wait(done["id"])["status"] == "succeeded"
    again = j.start("stub.state", {"mode": "fail"})                                   # released after success
    assert j.wait(again["id"])["status"] == "failed"
    last = j.start("stub.state", {"mode": "ok"})                                      # released after failure
    assert j.wait(last["id"])["status"] == "succeeded"
    assert j.running_state_job() is None


@test
def runner_uses_argument_list_no_shell_closed_stdin(env):
    seen, real = {}, subprocess.Popen
    def spy(argv, **kw):
        seen.update(argv=argv, **kw)
        return real(argv, **kw)
    jobs_mod.subprocess.Popen = spy
    try:
        m = env.jobs.wait(env.jobs.start("stub.read", {"mode": "ok", "sku": "C&P9M"})["id"])
    finally:
        jobs_mod.subprocess.Popen = real
    assert isinstance(seen["argv"], list) and seen["argv"][:3] == [sys.executable, "-X", "utf8"], seen["argv"]
    assert seen["argv"][-3:] == ["ok", "--sku", "C&P9M"], seen["argv"]
    assert seen["shell"] is False and seen["stdin"] == subprocess.DEVNULL and Path(seen["cwd"]) == env.project
    assert seen["env"]["PYTHONIOENCODING"] == "utf-8" and seen["env"]["PYTHONUNBUFFERED"] == "1"
    # how the process is started, per system: on Windows exactly as before the Linux tree-kill was added
    common = {"argv", "cwd", "env", "shell", "stdin", "stdout", "stderr", "creationflags", "encoding", "errors"}
    if os.name == "nt":
        assert set(seen) == common and seen["creationflags"] == subprocess.CREATE_NO_WINDOW | subprocess.CREATE_NEW_PROCESS_GROUP, sorted(seen)
    else:
        assert set(seen) == common | {"start_new_session"} and seen["start_new_session"] is True and seen["creationflags"] == 0, sorted(seen)
    assert m["status"] == "succeeded" and "sku=C&P9M" in env.jobs.read_log(m["id"])[0]


@test
def failed_job_records_exit_code_meaning_and_tail(env):
    m = env.jobs.wait(env.jobs.start("stub.read", {"mode": "fail"})["id"])
    assert m["exit_code"] == 3 and "not one of the agent's documented codes" in m["exit_meaning"], m
    assert any("failing on purpose" in line for line in m["tail"]), m["tail"]


@test
def cancel_kills_the_whole_tree_and_warns(env):
    m = env.jobs.start("stub.state", {"mode": "slow"})
    child = None
    for _ in range(100):
        match = re.search(r"child pid (\d+)", env.jobs.read_log(m["id"])[0])
        if match:
            child = int(match[1])
            break
        time.sleep(0.1)
    assert child, "stub never reported its child"
    assert env.jobs.cancel(m["id"])
    done = env.jobs.wait(m["id"])
    assert done["status"] == "cancelled" and done["note"], done
    time.sleep(0.5)
    assert process_gone(child), "grandchild process survived the cancel"


@test
def interrupted_jobs_are_marked_on_startup(env):
    meta = {"id": "20260101-000000-abcdef", "command": "stub.read", "label": "x", "args": [], "state_changing": True,
            "status": "running", "started_at": "2026-01-01T00:00:00", "finished_at": None, "duration_s": None,
            "exit_code": None, "exit_meaning": "", "pid": 1, "note": "", "tail": []}
    write_json_atomic(env.settings.jobs / f"{meta['id']}.json", meta)
    fresh = JobManager(env.settings)
    got = fresh.get(meta["id"])
    assert got["status"] == "interrupted" and got["note"], got
    assert fresh.running_state_job() is None                                          # never resumed, never holds the lock


@test
def metadata_write_is_atomic(env):
    target = env.settings.jobs / "atomic.json"
    write_json_atomic(target, {"v": 1})
    calls, real = [], os.replace
    def spy(a, b):
        calls.append((str(a), str(b)))
        raise OSError("disk full")
    jobs_mod.os.replace = spy
    try:
        raises(OSError, write_json_atomic, target, {"v": 2})
    finally:
        jobs_mod.os.replace = real
    assert json.loads(target.read_text()) == {"v": 1}, "original must survive a failed write"
    assert calls and calls[0][0].endswith(".tmp") and calls[0][1] == str(target), calls
    write_json_atomic(target, {"v": 3})
    assert json.loads(target.read_text()) == {"v": 3}


# ---------- masking and escaping ----------

@test
def masking_removes_keys_emails_phones_from_logs(env):
    for raw in (f"key {FAKE_KEY}", "ANTHROPIC_API_KEY=abc123xyz789", "mail boss@example.com", "call +1 214 555 0100", "Bearer abcdefghijklmnop123"):
        out = jobs_mod.mask_line(raw)
        assert FAKE_KEY not in out and "boss@example.com" not in out and "555 0100" not in out and "abc123xyz789" not in out and "abcdefghijklmnop123" not in out, (raw, out)
    m = env.jobs.wait(env.jobs.start("stub.read", {"mode": "secret"})["id"])
    log = env.jobs.read_log(m["id"])[0]
    assert "sk-ant-abcdefgh12345678" not in log and "boss@example.com" not in log and "555 0100" not in log, log
    assert "[REDACTED]" in log and "[EMAIL]" in log, log


@test
def hostile_log_text_is_escaped(env):
    m = env.jobs.wait(env.jobs.start("stub.read", {"mode": "hostile"})["id"])
    html = env.client.get(f"/jobs/{m['id']}").text
    assert "<script>alert" not in html and "<img src=x" not in html, "log text must be escaped"
    assert "&lt;script&gt;alert" in html


# ---------- Host, Origin, CSRF ----------

@test
def host_header_is_checked(env):
    assert env.client.get("/health").status_code == 200
    assert env.client.get("/health", headers={"host": "evil.example"}).status_code == 400
    assert env.client.get("/health", headers={"host": "127.0.0.1:9999"}).status_code == 400
    assert env.client.get("/health", headers={"host": "localhost:8000"}).status_code == 200


@test
def origin_is_checked_on_post(env):
    t = token(env)
    body = {"command": "stub.read", "mode": "ok", "csrf_token": t}
    assert env.client.post("/jobs/run", data=body, headers={"origin": "http://evil.example"}, follow_redirects=False).status_code == 403
    assert env.client.post("/jobs/run", data=body, follow_redirects=False).status_code == 403            # no Origin, no Referer
    assert env.client.post("/jobs/run", data=body, headers={"origin": "null"}, follow_redirects=False).status_code == 403
    assert env.client.post("/jobs/run", data=body, headers=ORIGIN, follow_redirects=False).status_code == 303


@test
def csrf_token_is_required(env):
    t = token(env)
    body = {"command": "stub.read", "mode": "ok"}
    assert env.client.post("/jobs/run", data=body, headers=ORIGIN, follow_redirects=False).status_code == 403
    assert env.client.post("/jobs/run", data={**body, "csrf_token": "0" * 64}, headers=ORIGIN, follow_redirects=False).status_code == 403
    other = TestClient(env.app, base_url="http://127.0.0.1:8000", raise_server_exceptions=False)
    assert other.post("/jobs/run", data={**body, "csrf_token": t}, headers=ORIGIN, follow_redirects=False).status_code == 403   # token of another session
    r = env.client.post("/jobs/run", data={**body, "csrf_token": t}, headers=ORIGIN, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/jobs/")
    jid = r.headers["location"].rsplit("/", 1)[1]
    assert env.client.post(f"/jobs/{jid}/cancel", headers=ORIGIN, follow_redirects=False).status_code == 403
    assert env.client.post("/jobs/run", data={"command": "nope", "csrf_token": t}, headers=ORIGIN).status_code == 400


@test
def security_headers_are_sent(env):
    h = env.client.get("/").headers
    assert h["content-security-policy"] == "default-src 'self'"
    assert h["x-content-type-options"] == "nosniff" and h["x-frame-options"] == "DENY" and h["referrer-policy"] == "no-referrer"
    assert env.client.get("/health", headers={"host": "evil.example"}).headers["x-frame-options"] == "DENY"


# ---------- downloads ----------

def make_link(link: Path, target: Path) -> bool:
    try:
        os.symlink(target, link, target_is_directory=True)
        return True
    except OSError:
        r = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(target)], capture_output=True)   # junction: no admin needed
        return r.returncode == 0


@test
def downloads_stay_inside_output_folders(env):
    s = env.settings
    (s.project / "Reports").mkdir()
    (s.project / "Reports" / "ok.csv").write_text("a,b\n")
    (s.project / "Reports" / "state.json").write_text("{}")
    (s.project / "Reports" / "tool.py").write_text("print(1)")
    (s.project / "email_state.json").write_text("{}")
    (env.root / "outside").mkdir()
    (env.root / "outside" / "secret.csv").write_text("top secret")
    (env.root / "secret.csv").write_text("top secret")
    ok = security.safe_output_file
    assert ok(s, "Reports", "ok.csv") == (s.project / "Reports" / "ok.csv").resolve()
    for folder, rel in [("Reports", "../secret.csv"), ("Reports", "..\\secret.csv"), ("Reports", "..\\..\\project\\email_state.json"),
                        ("Reports", str(env.root / "secret.csv")), ("Reports", "C:\\Windows\\win.ini"), ("Reports", "/etc/passwd"),
                        ("Reports", "\\\\server\\share\\x.csv"), ("Reports", "ok.csv\x00.txt"), ("Reports", "state.json"), ("Reports", "tool.py"),
                        ("Reports", "missing.csv"), ("Reports", ""), ("Reports", "."), ("..", "project/email_state.json"),
                        ("ui_data", "errors.log"), ("Reports", "ok.csv:stream")]:
        assert ok(s, folder, rel) is None, (folder, rel)
    if make_link(s.project / "Reports" / "link", env.root / "outside"):
        assert ok(s, "Reports", "link/secret.csv") is None, "link/junction escape must be refused"
        assert env.client.get("/files/Reports/link/secret.csv").status_code == 404
    else:
        raise AssertionError("could not create a symlink or junction to test escapes")
    assert env.client.get("/files/Reports/ok.csv").status_code == 200
    for url in ("/files/Reports/state.json", "/files/Reports/..%2fsecret.csv", "/files/Reports/%2e%2e/secret.csv", "/files/Nope/ok.csv"):
        assert env.client.get(url).status_code == 404, url


# ---------- secrets, pages, workflows ----------

@test
def secrets_never_appear_in_any_page(env):
    saved = {k: os.environ.get(k) for k in ("ANTHROPIC_API_KEY", "SMTP_PASSWORD", "SHIPPING_ADDRESS", "NIMBLE_API_KEY")}
    os.environ.update(ANTHROPIC_API_KEY=FAKE_KEY, SMTP_PASSWORD="hunter2-fake-password", SHIPPING_ADDRESS="9 Fake Rd\\nNowhere, ZZ 00000", NIMBLE_API_KEY="nimble-fake-secret-value")
    try:
        job = env.jobs.wait(env.jobs.start("stub.read", {"mode": "secret"})["id"])
        pages = ["/", "/health", "/jobs", f"/jobs/{job['id']}", f"/jobs/{job['id']}/log", "/step/search", "/step/order-sheet", "/nope", "/jobs/20260101-000000-000000"]
        text = "".join(env.client.get(p).text for p in pages)
        for secret in (FAKE_KEY, "hunter2-fake-password", "9 Fake Rd", "nimble-fake-secret-value"):
            assert secret not in text, f"{secret!r} leaked"
        health = env.client.get("/health").text
        assert "ANTHROPIC_API_KEY" in health and "SMTP_HOST" in health and "Not set" in health and "Needed for" in health
    finally:
        for k, v in saved.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)


@test
def unhandled_errors_show_only_a_reference(env):
    def boom():
        raise RuntimeError(f"db password {FAKE_KEY}")
    real, env.jobs.list = env.jobs.list, boom
    r = env.client.get("/jobs")
    env.jobs.list = real
    assert r.status_code == 500 and FAKE_KEY not in r.text and "RuntimeError" not in r.text and "quote this reference" in r.text
    log = (env.settings.data / "errors.log").read_text()
    assert FAKE_KEY not in log and "RuntimeError" in log


@test
def workflow_id_matches_the_agents_and_lists_newest_first(env):
    wfs = workflows.list_workflows(env.settings)
    assert [w["id"] for w in wfs] == ["lcom-sourcing_results_20260202_100000", "lcom-sourcing_results_20260101_100000"], wfs
    for w in wfs:
        assert w["id"] == obs.workflow_id(str(Path("Excel Output Sheets") / w["file"])) == obs.workflow_id(env.settings.results_dir / w["file"])
    here = Path(__file__).parent
    for agent in ("email_agent", "decision_agent", "order_sheet", "content_agent"):
        assert "obs.workflow_id(" in (here / f"{agent}.py").read_text(encoding="utf-8"), agent
    assert "obs.workflow_id(excel_path)" in (here / "sourcing_agent.py").read_text(encoding="utf-8")


@test
def overview_steps_only_claim_what_files_show(env):
    page = env.client.get("/").text
    assert "lcom-" not in page and "2 products" in page and "1 can be drafted" in page and "1 have no recommended seller" in page
    steps = {s.name: s for s in workflows.steps_for(env.settings, "sourcing_results_20260202_100000.xlsx")}
    assert [s for s in steps] == list(workflows.STEPS)
    assert steps["Search"].status == "Files found" and steps["Emails"].status == "Not started"
    assert steps["Approvals"].blocked == "" and steps["Order sheet"].blocked and steps["Listing content"].blocked
    (env.project / "decision_state.json").write_text(json.dumps({"next_request": 2, "decisions": [], "requests": {
        "A-0001": {"id": "A-0001", "results_file": "sourcing_results_20260202_100000.xlsx", "reply": None, "lines": []},
        "A-0002": {"id": "A-0002", "results_file": "sourcing_results_20260101_100000.xlsx", "reply": None, "lines": []}}}))
    (env.project / "email_state.json").write_text("{not json")
    steps = {s.name: s for s in workflows.steps_for(env.settings, "sourcing_results_20260202_100000.xlsx")}
    assert steps["Approvals"].status == "Files found" and "1 request(s)" in steps["Approvals"].detail and "still waiting" in steps["Approvals"].detail
    assert steps["Emails"].status == "Blocked" and steps["Emails"].blocked
    gates = {"emails": "Nothing is sent yet. The next page lists every recipient and asks you to type the count.",
             "approvals": "These three are typed by you every time.", "order-sheet": "No cart link is possible",
             "listing-content": "Nothing is published"}
    for (slug, gate), name in zip(gates.items(), workflows.STEPS[1:]):                 # real screens now: each says what its gate is
        page = env.client.get(f"/step/{slug}").text
        assert "Not available yet" not in page and f"<h1>{name}</h1>" in page and "Pipeline status" in page, slug
        assert gate in page, slug
    assert "Which products" in env.client.get("/step/search").text                    # Stage 2a: the Search step is a real screen
    assert env.client.get("/step/bogus").status_code == 404


@test
def running_banner_and_job_pages(env):
    m = env.jobs.start("stub.state", {"mode": "slow"})
    page = env.client.get("/").text
    assert "A job is running" in page and "Cancel" in page
    assert m["id"] in env.client.get("/jobs").text
    log = env.client.get(f"/jobs/{m['id']}/log?offset=0").json()
    assert log["status"] == "running"
    assert env.client.get("/jobs/..%2f..%2fx").status_code in (404, 400)
    assert env.client.get("/jobs/not-an-id").status_code == 404
    env.jobs.cancel(m["id"])
    env.jobs.wait(m["id"])


@test
def loopback_only_startup(env):
    import run_ui
    saved = {k: os.environ.get(k) for k in ("UI_HOST", "UI_ALLOW_NON_LOOPBACK", "UI_DATA_DIR")}
    os.environ["UI_DATA_DIR"] = str(env.root / "data")         # create_app makes this folder: keep it out of the project
    try:
        calls = []                                                  # uvicorn is faked so a broken guard cannot start a server
        sys.modules["uvicorn"], real = SimpleNamespace(run=lambda *a, **k: calls.append(k)), sys.modules.get("uvicorn")
        os.environ.pop("UI_ALLOW_NON_LOOPBACK", None)
        for h in ("0.0.0.0", "192.168.1.5", "example.com", "::"):
            os.environ["UI_HOST"] = h
            raises(SystemExit, run_ui.main)
        assert not calls, "server must not start on a non-loopback host"
        os.environ["UI_HOST"] = "127.0.0.1"
        import contextlib
        import io
        said = io.StringIO()
        with contextlib.redirect_stdout(said):
            run_ui.main()
        assert calls and calls[0]["host"] == "127.0.0.1", calls
        assert said.getvalue() == "L-Com agents UI: http://127.0.0.1:8000/   (Ctrl+C to stop)\n", said.getvalue()     # the default start says what it always said
        # a container listens on every interface (the override is set on purpose there); the address to open and the
        # hosts the UI answers to are still this computer's own names, never 0.0.0.0
        os.environ.update(UI_HOST="0.0.0.0", UI_ALLOW_NON_LOOPBACK="1")
        said = io.StringIO()
        with contextlib.redirect_stdout(said):
            run_ui.main()
        assert calls[-1]["host"] == "0.0.0.0" and calls[-1]["port"] == 8000, calls[-1]
        assert "L-Com agents UI: http://127.0.0.1:8000/" in said.getvalue() and "http://0.0.0.0" not in said.getvalue(), said.getvalue()
        assert "WARNING: listening on 0.0.0.0 with NO LOGIN" in said.getvalue()
        assert security.allowed_hosts(load_settings()) == {"127.0.0.1:8000", "localhost:8000"}
    finally:
        if real is not None:
            sys.modules["uvicorn"] = real
        for k, v in saved.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)

# ================= Stage 1.5: wording, labels, contrast, Home counts =================

import colorsys
from datetime import datetime

from ui import wording

TERM_VARIANTS = ["L-Com", "LCom", "lcom", "l com", "L-COM", "L_Com", "lcom-"]
STEM_A, STEM_B = "sourcing_results_20260101_100000", "sourcing_results_20260202_100000"


def has_term(text: str) -> bool:
    return wording.contains_reference_name(text)


def add_results(env, name, rows, mtime=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["SKU", "Product", "Recommended Manufacturer", "Recommended URL"])
    for r in rows:
        ws.append(r)
    path = env.project / "Excel Output Sheets" / name
    wb.save(path)
    if mtime:
        os.utime(path, (mtime, mtime))
    return path


def urls_in(html: str) -> list:
    return re.findall(r'(?:href|src|action)="([^"]*)"', html)


@test
def wording_table(env):
    cases = {"L-Com price": "Reference price", "L-Com Prices": "Reference prices", "LCom unit price": "Reference unit price",
             "L-Com SKU": "SKU", "lcom_prices.csv": "reference_prices.csv", "lcom-sourcing_results_20261004_161054": "sourcing_results_20261004_161054",
             "l com": "reference", "L_COM": "reference", "the L-Com catalogue": "the reference catalogue",
             "Telecom": "Telecom", "Lcommerce": "Lcommerce", "plain text 12.00": "plain text 12.00"}
    for raw, want in cases.items():
        assert wording.display(raw) == want, (raw, wording.display(raw))
    assert wording.display(None) is None and wording.display(5) == 5
    assert wording.strip_prefix("lcom-abc") == "abc" and wording.strip_prefix("abc") == "abc"
    from markupsafe import Markup
    assert isinstance(wording.display(Markup("<b>L-Com</b>")), Markup)       # already-rendered markup is left alone


@test
def reference_name_never_reaches_a_page_title_or_url(env):
    p = env.project
    (p / "Order Sheets").mkdir()
    (p / "Order Sheets" / "order_sheet_LCom_1.xlsx").write_bytes(b"x")
    (p / "Walmart Listings").mkdir()
    (p / "Walmart Listings" / "walmart_listings_L-Com.xlsx").write_bytes(b"x")
    (p / "approved_orders.json").write_text(json.dumps({"source_results_file": STEM_B + ".xlsx", "valid_approvals": 1, "generated_at": "L-Com time"}))
    add_results(env, "sourcing_results_20260303_100000.xlsx", [["L-Com SKU 1", "LCom price cable", "L-Com Maker", "http://x/lcom"]])
    job = env.jobs.wait(env.jobs.start("stub.read", {"mode": "term"})["id"])
    assert "L-Com" in env.jobs.read_log(job["id"])[0], "the stored log must stay verbatim"
    pages = ["/", "/jobs", "/health", f"/jobs/{job['id']}", f"/jobs/{job['id']}/log?offset=0", "/jobs/not-an-id", "/nope", "/step/bogus",
             "/?wf=nonsense"] + [f"/step/{s}" for s in ("search", "emails", "approvals", "order-sheet", "listing-content")]
    for w in workflows.list_workflows(env.settings):
        pages += [f"/?wf={w['stem']}", f"/step/order-sheet?wf={w['stem']}", f"/step/listing-content?wf={w['stem']}"]
    for page in pages:
        r = env.client.get(page, follow_redirects=True)
        body = r.text
        title = re.search(r"<title>(.*?)</title>", body, re.S)
        assert not has_term(body), (page, [m for m in re.findall(r"(?i).{20}l[\s_-]?com.{20}", body)][:3])
        assert not title or not has_term(title[1])
        for u in urls_in(body):
            assert not has_term(u), (page, u)
    assert "Reference price 12.00" in env.client.get(f"/jobs/{job['id']}").text
    for variant in TERM_VARIANTS:
        assert not has_term(wording.display(f"x {variant} y")) or variant == "lcom-", variant


@test
def plain_stem_urls_and_old_prefixed_urls_redirect(env):
    html = env.client.get("/").text
    assert f"?wf={STEM_B}" in html and "lcom" not in html.lower()
    for old, new in [(f"/?wf=lcom-{STEM_A}", f"/?wf={STEM_A}"), (f"/step/search?wf=LCOM-{STEM_A}", f"/step/search?wf={STEM_A}")]:
        r = env.client.get(old, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == new, (old, r.status_code, r.headers.get("location"))
    page = env.client.get(f"/?wf={STEM_A}").text
    assert f'value="{STEM_A}" selected' in page
    assert workflows.internal_id(STEM_A) == "lcom-" + STEM_A == obs.workflow_id(STEM_A + ".xlsx")      # the internal id is unchanged


@test
def friendly_labels_recent_five_and_show_all(env):
    assert workflows.friendly_label(datetime(2026, 10, 4, 16, 10), 15) == "Oct 4, 4:10 PM, 15 products"
    assert workflows.friendly_label(datetime(2026, 10, 4, 0, 5), 1) == "Oct 4, 12:05 AM, 1 product"
    assert workflows.friendly_label(datetime(2026, 10, 4, 12, 0), None) == "Oct 4, 12:00 PM"
    base = datetime(2026, 3, 1, 9, 30).timestamp()
    for i in range(5):                                          # 2 existing + 5 = 7 workflows
        add_results(env, f"sourcing_results_2026030{i}_090000.xlsx", [["A", "p", "m", "u"]] * (i + 1), base + i * 86400 + 99999999)
    wfs = workflows.list_workflows(env.settings)
    assert len(wfs) == 7 and all("lcom" not in w["label"].lower() and "_" not in w["label"] for w in wfs)
    assert wfs[0]["label"].endswith("5 products") and len({w["label"] for w in wfs}) == 7
    html = env.client.get("/").text
    recent = html.split("Recent workflows")[1].split("<details>")[0]
    assert recent.count("<li>") == 5 and "Show all (7)" in html
    for f in list((env.project / "Excel Output Sheets").glob("sourcing_results_2026030*.xlsx")):
        f.unlink()
    for f in list((env.project / "Excel Output Sheets").glob("*.xlsx"))[2:]:
        f.unlink()
    assert "Show all" not in env.client.get("/").text                      # 2 workflows: no cut-off, no details


@test
def home_count_separates_carrying_a_recommendation_from_being_draftable(env):
    for f in (env.project / "Excel Output Sheets").glob("*.xlsx"):
        f.unlink()
    add_results(env, "sourcing_results_20260404_100000.xlsx",
                [["A1", "p", "Acme", "http://a/1"], ["B2", "p", "Beta", "http://b/2"], ["C3", "p", "", ""], ["D4", "p", "", ""]])
    (env.project / "reviewer_exclusions.csv").write_text("sku,supplier_or_url,reason\nB2,beta,not acceptable\n", encoding="utf-8")
    text = env.client.get("/").text
    assert "4 products" in text and "1 can be drafted" in text and "2 have no recommended seller" in text, text
    assert "1 are excluded by the reviewer" in text and "unverified" in text
    assert "7 with a recommended seller" not in text and "recommended sellers" not in text
    d = workflows.draft_counts(env.settings, env.project / "Excel Output Sheets" / "sourcing_results_20260404_100000.xlsx")
    assert d == {"draftable": 1, "no_seller": 2, "excluded": 1, "errors": 0}, d
    # when the agent's rules cannot be run, only what the file says is shown, labelled precisely
    real = workflows.draft_counts
    workflows.draft_counts = lambda s, p: None
    try:
        text = env.client.get("/").text
    finally:
        workflows.draft_counts = real
    assert "rows carry a recommendation; the Emails step shows how many can be drafted" in text and "can be drafted;" not in text


@test
def displayed_log_has_a_note_wording_and_whole_lines_only(env):
    job = env.jobs.wait(env.jobs.start("stub.read", {"mode": "term"})["id"])
    page = env.client.get(f"/jobs/{job['id']}").text
    assert page.count("Wording differs from the raw log file.") == 1 and not has_term(page)
    assert "Reference price 12.00" in page and "reference_prices.csv" in page
    assert "L-Com price 12.00" in env.jobs.log_path(job["id"]).read_text(encoding="utf-8"), "the stored log is never altered"
    log = env.client.get(f"/jobs/{job['id']}/log?offset=0").json()
    assert not has_term(log["text"])
    f = env.settings.jobs / "20260101-000000-aaaaaa.log"
    f.write_text("first line\nsecond half-writ", encoding="utf-8")
    assert env.jobs.read_log("20260101-000000-aaaaaa", 0, whole_lines_only=True) == ("first line\n", 11)
    assert env.jobs.read_log("20260101-000000-aaaaaa", 11, whole_lines_only=True) == ("", 11)
    assert env.jobs.read_log("20260101-000000-aaaaaa", 0)[0].endswith("half-writ")


@test
def app_name_is_configurable_and_neutral_by_default(env):
    html = env.client.get("/").text
    assert "<title>Workflow - Sourcing Tool</title>" in html and 'class="brand" href="/">Sourcing Tool<' in html
    other = Settings(project=env.project, data=env.project / "d3", app_name="Acme Sourcing")
    c = TestClient(create_app(other), base_url="http://127.0.0.1:8000")
    html = c.get("/health").text
    assert "<title>Health - Acme Sourcing</title>" in html and "Sourcing Tool" not in html
    saved = os.environ.get("UI_APP_NAME")
    os.environ["UI_APP_NAME"] = "From Env"
    try:
        assert load_settings().app_name == "From Env"
    finally:
        os.environ.pop("UI_APP_NAME", None) if saved is None else os.environ.__setitem__("UI_APP_NAME", saved)
    assert load_settings().app_name == "Sourcing Tool"
    assert "Sourcing Tool" not in "".join(p.read_text(encoding="utf-8") for p in (Path(__file__).parent / "ui" / "templates").glob("*.html"))


def parse_tokens():
    css = (Path(__file__).parent / "ui" / "static" / "app.css").read_text(encoding="utf-8")
    light_block = re.search(r":root\s*\{(.*?)\}", css, re.S)[1]
    dark_block = re.search(r"prefers-color-scheme:\s*dark\)\s*\{\s*:root\s*\{(.*?)\}", css, re.S)[1]
    grab = lambda block: dict(re.findall(r"--([\w-]+):\s*(#[0-9a-fA-F]{6})", block))
    light = grab(light_block)
    return light, {**light, **grab(dark_block)}


def luminance(hexcolor: str) -> float:
    rgb = [int(hexcolor[i:i + 2], 16) / 255 for i in (1, 3, 5)]
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in rgb]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def contrast(a: str, b: str) -> float:
    hi, lo = sorted((luminance(a), luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


TEXT_PAIRS = [("text", "bg"), ("text", "surface"), ("muted", "bg"), ("muted", "surface"), ("accent", "bg"), ("accent", "surface"),
              ("accent-contrast", "accent"), ("danger-contrast", "danger"), ("success", "surface"), ("warning", "surface"),
              ("danger", "surface"), ("info", "surface"), ("success", "bg"), ("warning", "bg"), ("danger", "bg"), ("info", "bg")]
UI_PAIRS = [("border", "surface"), ("border", "bg"), ("success", "surface"), ("warning", "surface"), ("danger", "surface"),
            ("info", "surface"), ("accent", "surface")]


@test
def colour_tokens_meet_wcag_contrast_in_both_themes(env):
    light, dark = parse_tokens()
    for theme, tokens in (("light", light), ("dark", dark)):
        for fg, bg in TEXT_PAIRS:
            assert contrast(tokens[fg], tokens[bg]) >= 4.5, (theme, fg, bg, round(contrast(tokens[fg], tokens[bg]), 2))
        for fg, bg in UI_PAIRS:
            assert contrast(tokens[fg], tokens[bg]) >= 3.0, (theme, fg, bg, round(contrast(tokens[fg], tokens[bg]), 2))
    css = (Path(__file__).parent / "ui" / "static" / "app.css").read_text(encoding="utf-8")
    body = css.split("* { box-sizing")[1]                          # everything after the token block uses var(), no literal colours
    assert not re.search(r"#[0-9a-fA-F]{3,6}\b", body) and not re.search(r"\brgba?\(", body), "components must use tokens only"
    assert "http" not in css and "@import" not in css and "url(" not in css


@test
def every_status_has_a_distinct_badge_with_text(env):
    css = (Path(__file__).parent / "ui" / "static" / "app.css").read_text(encoding="utf-8")
    kinds = ["not-started", "files-found", "blocked", "running", "done", "failed", "cancelled", "interrupted"]
    markers = []
    for k in kinds:
        m = re.search(rf"\.badge-{k}::before[^{{]*\{{\s*content:\s*\"([^\"]+)\"", css)
        assert m, k
        markers.append(m[1])
    assert len(set(markers)) == len(kinds), markers
    from ui.app import JOB_BADGES
    assert set(JOB_BADGES.values()) <= {"Running", "Done", "Failed", "Cancelled", "Interrupted"}
    for j in ("ok", "fail"):
        env.jobs.wait(env.jobs.start("stub.read", {"mode": j})["id"])
    html = env.client.get("/jobs").text
    assert "badge-done" in html and ">Done<" in html and ">Failed<" in html and "No jobs yet" not in html


# ================= Stage 2a: the Search screen (stub agent only; no real search is ever run) =================

from contextlib import contextmanager

from ui import search as search_mod

PRESET = ["HDFF", "FOA-020C", "LCSP1050"]
NIMBLE_FAKE = "nimble-fake-secret-value"
HOSTILE = "<script>alert(1)</script>"


def fake_agent(env, skus=("HDFF", "FOA-020C", "LCSP1050", "C&P9M", "ECF504-AA")):
    """A sourcing_agent.py that holds only the built-in list. It is parsed by the UI and never run (the search command
    runs the stub in these tests)."""
    desc = {"HDFF": f"L-Com HDMI coupler {HOSTILE}", "FOA-020C": f"Fiber adapter {HOSTILE}", "LCSP1050": "LCom surge protector"}
    products = [{"sku": k, "description": desc.get(k, f"Product {k}"), "lcom_price": 10.0} for k in skus]
    (env.project / "sourcing_agent.py").write_text("PRODUCTS = " + repr(products) + "\n", encoding="utf-8")


@contextmanager
def env_vars(**kv):
    saved = {k: os.environ.get(k) for k in kv}
    try:
        for k, v in kv.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)
        yield
    finally:
        for k, v in saved.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)


def keys(mode="ok"):
    return env_vars(ANTHROPIC_API_KEY=FAKE_KEY, NIMBLE_API_KEY=NIMBLE_FAKE, UI_STUB_SEARCH=mode)


def post(env, url, **data):
    return env.client.post(url, data={"csrf_token": token(env), **data}, headers=ORIGIN, follow_redirects=False)


def started(r) -> str:
    assert r.status_code == 303 and "/step/search?job=" in r.headers["location"], (r.status_code, r.text[-600:])
    return r.headers["location"].split("job=")[1]


def wait_log(env, jid, needle, timeout=25):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if needle in env.jobs.read_log(jid)[0]:
            return
        time.sleep(0.1)
    raise AssertionError(f"{needle!r} never appeared in the log:\n{env.jobs.read_log(jid)[0]}")


def clean(html: str, where: str = "") -> str:
    """No reference name, no secret, no unescaped hostile text."""
    assert not has_term(html), (where, re.findall(r"(?i).{30}l[\s_-]?com.{20}", html)[:3])
    for secret in (FAKE_KEY, NIMBLE_FAKE):
        assert secret not in html, (where, "secret leaked")
    assert "<script>alert" not in html, (where, "hostile text was not escaped")
    return html


@test
def search_allowlist_validates_skus_flags_paths_and_limits(env):
    fake_agent(env)
    (env.project / "input.xlsx").write_bytes((env.project / "Excel Output Sheets" / f"{STEM_A}.xlsx").read_bytes())
    j, out = env.jobs, "sourcing_results_20260505_101010.xlsx"
    for bad in (["NOPE"], ["-x"], ["HDFF; calc"], ["HDFF", "../x"], ["--inspect-tools"], "HDFF", [], [""]):
        raises(JobRejected, j.start, "search", {"sku": bad, "output": out})
    for extra in ({"inspect-tools": "1"}, {"inspect_tools": "1"}, {"add_builtin": "1"}, {"mode": "ok"}, {"results": f"{STEM_A}.xlsx"}):
        raises(JobRejected, j.start, "search", {"output": out, **extra})
    for bad in ("C:\\x.xlsx", f"..\\{out}", f"Excel Output Sheets\\{out}", "sourcing_results_1.xlsx", "x.xlsx", "", out + " --inspect-tools"):
        raises(JobRejected, j.start, "search", {"output": bad})
    raises(JobRejected, j.start, "search", {})                                          # the output is required and server-built
    for bad in ("0", "-1", "201", "99999", "1e3", "5; x", "", " 5", "٣"):
        raises(JobRejected, j.start, "search", {"input": "input.xlsx", "limit": bad, "output": out})
    raises(JobRejected, j.start, "search", {"limit": "5", "output": out})               # a limit needs an input file
    raises(JobRejected, j.start, "search", {"sku": ["HDFF"], "input": "input.xlsx", "output": out})
    for bad in ("C:\\Windows\\win.ini", "..\\input.xlsx", "nope.xlsx", f"Excel Output Sheets\\{STEM_A}.xlsx", str(env.project / "input.xlsx")):
        raises(JobRejected, j.start, "search", {"input": bad, "output": out})
    cmd = j.commands["search"]
    assert cmd.state_changing and cmd.argv({"sku": ["hdff", "C&P9M"], "output": out}) == \
        ["search", "--sku", "HDFF", "C&P9M", "--output", native("Excel Output Sheets", out)]
    assert cmd.argv({"input": "input.xlsx", "limit": "007", "output": out})[1:] == \
        ["--input", "input.xlsx", "--limit", "7", "--output", native("Excel Output Sheets", out)]
    assert not j.list(), "a rejected request must start nothing"
    # the production allowlist: the real agent, never the stub, whatever the stub's own switch says
    with env_vars(UI_STUB_SEARCH="slow"):
        prod = JobManager(Settings(project=env.project, data=env.project / "d4")).commands
    assert prod["search"].script == "sourcing_agent.py" and prod["search"].fixed == ()
    assert not any("stub" in c.script or "stub" in c.id for c in prod.values()), "the stub must never be in the production allowlist"
    assert sorted(p.flag for p in prod["search"].params.values()) == ["--add-builtin", "--input", "--limit", "--output", "--sku"]
    # --add-builtin: a switch, only with an input file, never with chosen SKUs, and only ever 'on'
    assert cmd.argv({"input": "input.xlsx", "limit": "2", "add_builtin": "1", "output": out})[1:] == \
        ["--input", "input.xlsx", "--limit", "2", "--add-builtin", "--output", native("Excel Output Sheets", out)]
    for bad in ({"add_builtin": "1"}, {"sku": ["HDFF"], "add_builtin": "1"}, {"input": "input.xlsx", "add_builtin": "0"},
                {"input": "input.xlsx", "add_builtin": "true"}, {"input": "input.xlsx", "add_builtin": "--inspect-tools"},
                {"input": "input.xlsx", "add_builtin": ""}, {"input": "input.xlsx", "add_builtin": ["1"]}):
        raises(JobRejected, j.start, "search", {**bad, "output": out})
    assert post(env, "/jobs/run", command="search").status_code == 400                # the diagnostics form cannot start a search
    assert post(env, "/jobs/run", command="search", mode="ok").status_code == 400
    assert not j.list()


@test
def preset_refuses_when_a_sku_is_missing_from_the_built_in_list(env):
    fake_agent(env, ("FOA-020C", "LCSP1050", "C&P9M"))
    try:
        search_mod.resolve_selection(env.settings, "preset")
        raise AssertionError("the preset must refuse")
    except JobRejected as e:
        assert "HDFF" in str(e)
    page = env.client.get("/step/search").text
    assert "Not available: HDFF is not in the built-in product list" in page and 'value="preset" data-count="3" disabled' in page
    with keys():
        for url, extra in (("/search/review", {}), ("/search/start", {"confirm_count": "3"})):
            r = post(env, url, choice="preset", **extra)
            assert r.status_code == 400 and "HDFF is not in the built-in product list" in r.text, (url, r.status_code)
    assert not env.jobs.list()
    fake_agent(env)
    assert [p[0] for p in search_mod.resolve_selection(env.settings, "preset")["products"]] == PRESET
    real = search_mod.builtin_products(Settings(project=Path(__file__).parent, data=env.root / "d5"))   # the real list, parsed not run
    assert {"FOA-020C", "LCSP1050", "C&P9M"} <= {sku for sku, _ in real} and all(desc for _, desc in real), real


@test
def typed_count_is_checked_against_the_server_side_count(env):
    fake_agent(env)
    with keys():
        review = post(env, "/search/review", choice="preset")
        assert review.status_code == 200 and "About 3 products, roughly $0.24 and 3 minutes." in review.text
        assert "This is an estimate from one 15-product batch (about $0.08 and about 56 seconds per product), not a measurement of this run." in review.text
        assert "spends money and calls external services" in review.text and "Type the number of products to confirm" in review.text
        for typed in ("2", "", "three", "3.0", "-3", "33"):
            r = post(env, "/search/start", choice="preset", confirm_count=typed)
            assert r.status_code == 400 and "not the number of products selected" in r.text, typed
        # a tampered browser: says 3 everywhere, but the selection it posts is 5 products
        r = post(env, "/search/start", choice="choose", sku=["HDFF", "FOA-020C", "LCSP1050", "C&P9M", "ECF504-AA"], confirm_count="3", count="3", expected="3")
        assert r.status_code == 400 and "About 5 products" in r.text
        assert post(env, "/search/start", choice="choose", sku=["HDFF", "NOPE"], confirm_count="2").status_code == 400
        assert post(env, "/search/start", choice="choose", confirm_count="0").status_code == 400      # nothing ticked
        assert post(env, "/search/start", choice="bogus", confirm_count="1").status_code == 400
        assert not env.jobs.list(), "a wrong count must start nothing"
        jid = started(post(env, "/search/start", choice="choose", sku=["HDFF", "hdff", "C&P9M"], confirm_count=" 2 "))
        meta = env.jobs.wait(jid)
    assert meta["status"] == "succeeded" and meta["search"]["expected"] == 2 and meta["args"][:4] == ["search", "--sku", "HDFF", "C&P9M"], meta


@test
def missing_keys_disable_start_and_the_server_refuses_a_forced_post(env):
    fake_agent(env)
    with env_vars(ANTHROPIC_API_KEY=None, NIMBLE_API_KEY=None):
        page = env.client.get("/step/search").text
        assert "A search cannot be started." in page and "ANTHROPIC_API_KEY and NIMBLE_API_KEY are not set" in page
        assert page.count("Not set") == 2 and re.search(r'<button type="submit" class="btn-primary" disabled>Continue', page)
        assert re.search(r'class="btn-primary" disabled>Start search', post(env, "/search/review", choice="preset").text)
        r = post(env, "/search/start", choice="preset", confirm_count="3")
        assert r.status_code == 400 and "Nothing was started: ANTHROPIC_API_KEY, NIMBLE_API_KEY is not set" in r.text
        assert not env.jobs.list()
        with env_vars(ANTHROPIC_API_KEY=FAKE_KEY):                                    # one of the two is still missing
            r = post(env, "/search/start", choice="preset", confirm_count="3")
            assert r.status_code == 400 and "NIMBLE_API_KEY is not set" in r.text and not env.jobs.list()
        # the agents read the project's .env themselves, so a key there counts as set; its value is never shown
        (env.project / ".env").write_text("ANTHROPIC_API_KEY=sk-ant-FILEKEY0000000000\nNIMBLE_API_KEY=nimble-file-secret\n", encoding="utf-8")
        text = env.client.get("/step/search").text + env.client.get("/health").text
        assert "A search cannot be started." not in text and "sk-ant-FILEKEY" not in text and "nimble-file-secret" not in text
        assert not re.search(r'disabled>Continue', text)


@test
def a_search_takes_the_exclusive_lock(env):
    fake_agent(env)
    with keys("slow"):
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        assert env.jobs.get(jid)["state_changing"] and env.jobs.running_state_job()["id"] == jid
        r = post(env, "/search/start", choice="preset", confirm_count="3")
        assert r.status_code == 409 and "Nothing was started" in r.text and "Another job that changes data is running" in r.text
        raises(JobBusy, env.jobs.start, "stub.state", {"mode": "ok"})                 # no state-changing job alongside a search
        assert env.jobs.wait(env.jobs.start("stub.read", {"mode": "ok"})["id"])["status"] == "succeeded"   # read-only still runs
        assert len([m for m in env.jobs.list() if m["command"] == "search"]) == 1
        env.jobs.cancel(jid)
        env.jobs.wait(jid)
        other = env.jobs.start("stub.state", {"mode": "slow"})                        # and the other way round
        page = env.client.get("/step/search").text
        assert "A search cannot be started while another job is running" in page and "stub (changes state)" in page
        assert re.search(r'disabled>Continue', page)
        assert post(env, "/search/start", choice="preset", confirm_count="3").status_code == 409
        env.jobs.cancel(other["id"])
        env.jobs.wait(other["id"])
        assert len([m for m in env.jobs.list() if m["command"] == "search"]) == 1


@test
def output_path_is_server_built_contained_and_matches_homes_workflow_id(env):
    fake_agent(env)
    name = search_mod.new_output_name(datetime(2026, 5, 6, 7, 8, 9))
    assert name == "sourcing_results_20260506_070809.xlsx" and search_mod.OUTPUT_RE.fullmatch(search_mod.new_output_name())
    rule = search_mod.output_rule(env.settings)
    assert rule(name) == native("Excel Output Sheets", name)
    assert (env.project / rule(name)).resolve().parent == env.settings.results_dir.resolve()
    (env.settings.results_dir / name).write_bytes(b"x")
    raises(JobRejected, rule, name)                                                     # collision: never overwrite
    real = search_mod.new_output_name
    with keys():
        search_mod.new_output_name = lambda now=None: name
        try:
            r = post(env, "/search/start", choice="preset", confirm_count="3")
        finally:
            search_mod.new_output_name = real
        assert r.status_code == 400 and "already exists" in r.text and not env.jobs.list()
        assert (env.settings.results_dir / name).read_bytes() == b"x"
        (env.settings.results_dir / name).unlink()
        shutil.rmtree(env.settings.results_dir)                                         # the UI makes the folder; the agent would not
        meta = env.jobs.wait(started(post(env, "/search/start", choice="preset", confirm_count="3")))
    made = meta["search"]["output"]
    assert meta["status"] == "succeeded" and meta["args"][-2:] == ["--output", native("Excel Output Sheets", made)]
    wf = workflows.list_workflows(env.settings)[0]
    assert wf["file"] == made and wf["id"] == obs.workflow_id(meta["args"][-1]) == workflows.internal_id(Path(made).stem)


@test
def progress_parser_reads_the_agents_real_lines_and_survives_garbage(env):
    src = (Path(__file__).parent / "sourcing_agent.py").read_text(encoding="utf-8")      # the formats the parser is written against
    for real_line in ('log(f"Product {idx + 1} of {len(products)} ...")', 'log(f"  Tokens:     in={tin} out={tout}  (~${cost:.4f})")',
                      'log(f"  Error researching this product: {reason}")', 'log(f"!!! BILLING / CREDIT / AUTH ERROR - stopping the batch: {reason}")',
                      'print(f"Missing required environment variable(s): {\', \'.join(missing)}")', "Estimated total cost: ~${totals['cost']:.4f}",
                      'print(f"Report written to: {os.path.abspath(report_path)}")', 'print(f"[{sku}] {message}" if sku else message)',
                      'print(f"Saving results after every product to {report_path} / {excel_path}")',
                      "print(f\"Model: {MODEL}  |  Products: {len(products)}  |  Sites: {', '.join(SOURCING_SITES)}\")"):
        assert real_line in src, f"the agent no longer prints: {real_line}"
    log = "\n".join([
        "Model: claude-haiku-4-5  |  Products: 3  |  Sites: aliexpress.com, alibaba.com, made-in-china.com",
        "Saving results after every product to sourcing_report_20260101_101010.md / Excel Output Sheets\\x.xlsx",
        "[A-1] Product 1 of 3 ...", "[C&P9M] Product 2 of 3 ...", "[Z] Product 3 of 3 ...",
        "[A-1]   Tokens:     in=12000 out=900  (~$0.0165)", "[A-1]   Time:       41s  (Nimble 30s, Anthropic+other 11s, 5 tool calls)",
        "[C&P9M]   Error researching this product: TimeoutError"])
    p = search_mod.parse_progress(log)
    assert (p["total"], p["started"], p["done"], p["errors"], p["current"], p["cost"]) == (3, 3, 2, 1, ["Z"], 0.0165), p
    assert p["report"] == "sourcing_report_20260101_101010.md" and p["final_cost"] is None and not p["billing"] and not p["missing"]
    p = search_mod.parse_progress(log + "\n[Z] !!! BILLING / CREDIT / AUTH ERROR - stopping the batch: x\n"
                                  "Total tokens: in=12000 out=900  |  Estimated total cost: ~$0.0165\n"
                                  "Missing required environment variable(s): ANTHROPIC_API_KEY, NIMBLE_API_KEY")
    assert p["billing"] and p["final_cost"] == 0.0165 and p["missing"] == ["ANTHROPIC_API_KEY", "NIMBLE_API_KEY"]
    hostile = ["[x] Product 99999999999999999999 of 3 ...", "[<script>] Product 1 of 2 ...", "[a]   Tokens:     in=abc out=1  (~$1.0)",
               "[a]   Tokens:     in=1 out=1  (~$nan)", "[a]   Tokens:     in=1 out=1  (~$1e999)", "\x00\x01", "[" * 5000, "] Product",
               "Missing required environment variable(s): " + FAKE_KEY, "Total tokens: in=1 out=1  |  Estimated total cost: ~$-5",
               "Model: x  |  Products: -1  |  Sites: y", "[a] Product 5 of 2 ...", "Report written to: C:\\Windows\\win.ini", "\ufffd" * 300]
    p = search_mod.parse_progress("\n".join(hostile))
    assert p["total"] == 2 and p["done"] == 0 and p["cost"] == 0.0 and p["missing"] == [] and p["report"] is None and p["started"] == 2, p
    assert search_mod.parse_progress("") == search_mod.parse_progress(None) and search_mod.parse_progress("")["total"] is None
    # through the real runner: the stored log is masked, and the lines must still be readable afterwards
    fake_agent(env)
    with keys():
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        env.jobs.wait(jid)
    body = env.client.get(f"/jobs/{jid}/log?offset=0").json()
    g = body["progress"]
    assert (g["total"], g["done"], g["started"], g["current"], g["errors"]) == (3, 3, 3, [], 0), g
    assert g["cost"] == g["final_cost"] == round(sum((12000 + i) / 1e6 + 900 * 5 / 1e6 for i in range(3)), 4), g
    assert g["report"] and (env.project / g["report"]).is_file()
    assert not has_term(json.dumps(body)) and FAKE_KEY not in json.dumps(body)       # JSON is put on the page with textContent, not as HTML
    assert "progress" not in env.client.get(f"/jobs/{env.jobs.wait(env.jobs.start('stub.read', {'mode': 'ok'})['id'])['id']}/log").json()


@test
def cancel_keeps_the_partial_results_file_and_the_run_is_shown_as_partial(env):
    fake_agent(env)
    with keys("stubborn"):                                                              # ignores polite stop requests
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        wait_log(env, jid, "Product 2 of 3")
        out = env.settings.results_dir / env.jobs.get(jid)["search"]["output"]
        assert out.is_file()
        page = clean(env.client.get(f"/step/search?job={jid}").text, "running")
        assert "Products finished so far are kept in the results file. The product in progress is lost." in page
        assert 'id="p-bar" max="3" value="1"' in page and "the agent's own figure" in page and "badge-running" in page
        assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page                         # the hostile description, escaped in the log
        prog = env.client.get(f"/jobs/{jid}/log?offset=0").json()["progress"]
        assert (prog["done"], prog["total"], prog["current"]) == (1, 3, ["FOA-020C"]) and prog["cost"] == 0.0165, prog
        home = env.client.get("/").text
        assert "A job is running" in home and "View progress or cancel" in home and home.count("badge-running") >= 2   # banner, rail and table
        assert f"/jobs/{jid}/cancel" not in home, "a search is only cancelled from the page that shows the warning"
        assert env.client.get("/step/search", follow_redirects=False).headers["location"] == f"/step/search?job={jid}"
        child = int(re.search(r"child pid (\d+)", env.jobs.read_log(jid)[0])[1])
        assert env.client.post(f"/jobs/{jid}/cancel", headers=ORIGIN, follow_redirects=False).status_code == 403   # CSRF
        r = post(env, f"/jobs/{jid}/cancel")
        assert r.status_code == 303 and r.headers["location"] == f"/step/search?job={jid}"
        done = env.jobs.wait(jid)
    assert done["status"] == "cancelled" and "saved after each product" in done["note"], done
    time.sleep(0.5)
    assert process_gone(child)
    assert out.is_file() and workflows.products_in(out) == 1, "the finished product must be kept"
    page = clean(env.client.get(f"/step/search?job={jid}").text, "cancelled")
    assert "Partial results." in page and "holds 1 of the 3 products asked for" in page and "Cancelled" in page
    assert "The search job has ended" not in page and ">Done<" not in page
    named = search_mod.parse_progress(env.jobs.read_log(jid)[0])["report"]             # a stopped agent never prints where the report went
    assert "so no report was copied." in page and f"The report file is <code>{named}</code> in the project folder." in page
    assert "/files/Reports/" not in page and not (env.project / "Reports").exists()
    steps ={st.name: st for st in workflows.steps_for(env.settings, out.name)}
    assert steps["Search"].status == "Files found" and "badge-running" not in env.client.get("/").text


@test
def exit_1_and_a_billing_stop_are_read_in_plain_language(env):
    fake_agent(env)
    with keys("nokey"):
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        meta = env.jobs.wait(jid)
    page = clean(env.client.get(f"/step/search?job={jid}").text, "failed")
    assert meta["status"] == "failed" and meta["exit_code"] == 1
    assert "Exit code 1" in page and "a required setting is missing: ANTHROPIC_API_KEY, NIMBLE_API_KEY" in page
    assert "Last lines" in page and "No results file was produced." in page and "Next: seller emails" not in page and "sk-ant-..." not in page
    with keys("billing"):
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        meta = env.jobs.wait(jid)
    page = clean(env.client.get(f"/step/search?job={jid}").text, "billing")
    assert meta["exit_code"] == 0, "the real agent ends normally after a billing stop, so the page must read the log"
    assert "billing, credit or sign-in problem" in page and "Partial results." in page and "2 of them recorded as errors" in page
    assert page.count(">Error") >= 2 and "1 can be drafted" in page and "2 could not be searched" in page and ">Done<" not in page


@test
def results_table_is_honest_about_recommended_draftable_and_excluded(env):
    fake_agent(env)
    (env.project / "reviewer_exclusions.csv").write_text("sku,supplier_or_url,reason\nLCSP1050,stub maker 2,not acceptable\n", encoding="utf-8")
    form = clean(env.client.get("/step/search").text, "form")
    assert "reference HDMI coupler" in form and "&lt;script&gt;alert(1)&lt;/script&gt;" in form and form.count('name="sku"') == 5
    with keys():
        review = clean(post(env, "/search/review", choice="choose", sku=PRESET).text, "review")
        assert "reference HDMI coupler" in review and "&lt;script&gt;" in review and review.count('type="hidden" name="sku"') == 3
        jid = started(post(env, "/search/start", choice="preset", confirm_count="3"))
        meta = env.jobs.wait(jid)
    name = meta["search"]["output"]
    page = clean(env.client.get(f"/step/search?job={jid}").text, "finished")
    assert "The search job has ended." in page and ">Done<" not in page and "Partial results" not in page
    assert page.count("badge-unverified") == 1, "only the draftable recommendation is marked; an excluded one is not a recommendation"
    assert "Excluded by reviewer list" in page and "No recommendation" in page and "Check the pack size" in page
    assert "3 products in the results file: 1 can be drafted; 1 have no recommended seller; 1 are excluded by the reviewer" in page
    assert "Stub Maker 0" in page and "alibaba.com" in page and "$1.25" in page and "90%" in page
    assert "Reference unit price" in page and "$10.00" in page and "reference list, 2026-01-01" in page          # carried by the file, relabelled
    assert "Downloaded files keep their original wording" in page and "Next: seller emails" in page
    assert f'/step/emails?wf={Path(name).stem}"' in page and f'value="{Path(name).stem}" selected' in page       # the rail follows the run
    assert env.client.get(f"/files/Excel Output Sheets/{name}").status_code == 200
    report = search_mod.parse_progress(env.jobs.read_log(jid)[0])["report"]
    original = env.project / report
    assert f'href="/files/Reports/{report}"' in page and "the original stays in the project folder" in page
    assert original.is_file() and (env.project / "Reports" / report).read_bytes() == original.read_bytes()      # copied, not moved
    assert env.client.get(f"/files/Reports/{report}").status_code == 200
    for u in urls_in(page):
        assert not has_term(u), u
    assert not has_term(re.search(r"<title>(.*?)</title>", page, re.S)[1])
    assert "Open the search view" in env.client.get(f"/jobs/{jid}").text
    home = env.client.get("/").text
    assert "Files found" in home and ">Done<" not in home
    # without the email agent's rules the page says the exclusion list was not applied, and never invents a count
    real, workflows._cache = (workflows.draft_counts, workflows.draft_rows), {}
    workflows.draft_counts = workflows.draft_rows = lambda s, p: None        # the project's email agent cannot be loaded
    try:
        page = env.client.get(f"/step/search?job={jid}").text
    finally:
        workflows.draft_counts, workflows.draft_rows = real
        workflows._cache = {}
    assert "exclusion list could not be applied on this page" in page and "can be drafted;" not in page and page.count("badge-unverified") == 2


@test
def search_routes_check_csrf_host_and_origin(env):
    fake_agent(env)
    with keys():
        t = token(env)
        for url in ("/search/review", "/search/start"):
            body = {"choice": "preset", "confirm_count": "3"}
            assert env.client.post(url, data=body, headers=ORIGIN, follow_redirects=False).status_code == 403                      # no token
            assert env.client.post(url, data={**body, "csrf_token": "0" * 64}, headers=ORIGIN, follow_redirects=False).status_code == 403
            assert env.client.post(url, data={**body, "csrf_token": t}, headers={"origin": "http://evil.example"}, follow_redirects=False).status_code == 403
            assert env.client.post(url, data={**body, "csrf_token": t}, follow_redirects=False).status_code == 403                 # no Origin
            assert env.client.post(url, data={**body, "csrf_token": t}, headers={**ORIGIN, "host": "evil.example"}, follow_redirects=False).status_code == 400
            other = TestClient(env.app, base_url="http://127.0.0.1:8000", raise_server_exceptions=False)
            assert other.post(url, data={**body, "csrf_token": t}, headers=ORIGIN, follow_redirects=False).status_code == 403      # another session's token
        assert env.client.get("/step/search", headers={"host": "evil.example"}).status_code == 400
        assert not env.jobs.list()


@test
def input_file_choice_is_a_listed_file_with_a_bounded_limit(env):
    fake_agent(env)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, None, None])
    ws.append(["SKU", "Product", "Keyword"])
    for row in (["IN-1", "Cable", "L-Com cable"], [None, None, None], ["IN-2", "Adapter", ""], ["IN-3", "Plug", "Plug kw"]):
        ws.append(row)
    wb.save(env.project / "products.xlsx")
    (env.project / "notes.txt").write_text("x")
    assert search_mod.input_files(env.settings) == ["products.xlsx"]                      # only .xlsx, only the project folder itself
    assert search_mod.input_products(env.settings, "products.xlsx") == [("IN-1", "L-Com cable"), ("IN-2", "Adapter"), ("IN-3", "Plug kw")]
    page = env.client.get("/step/search").text
    assert '<option value="products.xlsx">' in page and "sourcing_results_2026" not in page.split("input_file")[1].split("</select>")[0]
    with keys():
        r = clean(post(env, "/search/review", choice="file", input_file="products.xlsx", limit="2").text)
        assert "About 2 products, roughly $0.16 and 2 minutes." in r and "first 2 rows" in r
        for bad in ({"input_file": "..\\products.xlsx"}, {"input_file": str(env.project / "products.xlsx")}, {"input_file": "notes.txt"},
                    {"input_file": "products.xlsx", "limit": "0"}, {"input_file": "products.xlsx", "limit": "999"}, {"input_file": ""}):
            assert post(env, "/search/start", choice="file", confirm_count="3", **bad).status_code == 400, bad
        assert post(env, "/search/start", choice="file", input_file="products.xlsx", limit="2", confirm_count="3").status_code == 400
        assert not env.jobs.list()
        meta = env.jobs.wait(started(post(env, "/search/start", choice="file", input_file="products.xlsx", limit="2", confirm_count="2")))
    assert meta["status"] == "succeeded" and meta["args"][1:5] == ["--input", "products.xlsx", "--limit", "2"], meta
    assert "--add-builtin" not in meta["args"] and workflows.products_in(env.settings.results_dir / meta["search"]["output"]) == 2
    # "Also include the built-in products": unchecked by default, counted on the server the way the agent adds them (after the limit)
    assert re.search(r'<input type="checkbox" name="add_builtin" value="1" >', page) and "Also include the built-in products" in page
    sel = search_mod.resolve_selection(env.settings, "file", (), "products.xlsx", "2", "1")
    assert [p[0] for p in sel["products"]] == ["IN-1", "IN-2", "HDFF", "FOA-020C", "LCSP1050", "C&P9M", "ECF504-AA"] and sel["params"]["add_builtin"] == "1"
    assert "add_builtin" not in search_mod.resolve_selection(env.settings, "preset", (), "", "", "1")["params"]   # only the file choice has it
    raises(JobRejected, search_mod.resolve_selection, env.settings, "file", (), "products.xlsx", "2", "yes")
    time.sleep(1.1)                                                                       # results file names are per second
    with keys():
        r = clean(post(env, "/search/review", choice="file", input_file="products.xlsx", limit="2", add_builtin="1").text)
        assert "About 7 products, roughly $0.56 and 7 minutes." in r and "first 2 rows, plus the 5 built-in products" in r
        assert 'name="add_builtin" value="1"' in r and r.count("<li><strong>") == 7
        assert 'name="add_builtin" value=""' in post(env, "/search/review", choice="file", input_file="products.xlsx", limit="2").text
        before = len(env.jobs.list())
        for typed in ("2", "5", "6"):                                                     # the file rows alone, or the built-ins alone, is the wrong count
            assert post(env, "/search/start", choice="file", input_file="products.xlsx", limit="2", add_builtin="1", confirm_count=typed).status_code == 400
        assert post(env, "/search/start", choice="file", input_file="products.xlsx", limit="2", confirm_count="7").status_code == 400   # box not ticked: 2
        assert len(env.jobs.list()) == before
        meta = env.jobs.wait(started(post(env, "/search/start", choice="file", input_file="products.xlsx", limit="2", add_builtin="1", confirm_count="7")))
    assert meta["status"] == "succeeded" and meta["search"]["expected"] == 7, meta
    assert meta["args"][1:6] == ["--input", "products.xlsx", "--limit", "2", "--add-builtin"], meta
    assert workflows.products_in(env.settings.results_dir / meta["search"]["output"]) == 7
    assert "plus the 5 built-in products" in env.client.get(f"/step/search?job={meta['id']}").text
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Price"])
    wb.save(env.project / "nosku.xlsx")
    raises(JobRejected, search_mod.input_products, env.settings, "nosku.xlsx")


@test
def built_in_list_is_the_one_the_agents_sku_lookup_uses(env):
    src = (Path(__file__).parent / "sourcing_agent.py").read_text(encoding="utf-8")      # what the UI's reading is written against
    assert "builtin = read_lcom_catalog() or PRODUCTS" in src
    assert "products, missing = pick_skus(products + (builtin if args.input else []), args.sku)" in src
    assert "if args.input and args.add_builtin:\n            products = products + builtin" in src
    import ast
    real_tree = ast.parse(src)
    assert search_mod._catalog_is_still_a_stub(real_tree), "the agent's catalog reader now does something: the UI list must be taught its source"
    literal = next(ast.literal_eval(n.value) for n in real_tree.body if isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") == "PRODUCTS")
    real = search_mod.builtin_products(Settings(project=Path(__file__).parent, data=env.root / "d6"))
    assert real == [(p["sku"], p["description"]) for p in literal] and len(real) >= 5      # the same list, in the same order
    stub_fn = 'def read_lcom_catalog(source=None) -> list:\n    """Future input path."""\n    return []\n'
    fake_agent(env)
    agent = env.project / "sourcing_agent.py"
    base = agent.read_text(encoding="utf-8")
    agent.write_text(base + stub_fn, encoding="utf-8")
    assert len(search_mod.builtin_products(env.settings)) == 5
    # a catalog reader that does real work: the UI must not go on offering PRODUCTS, which the agent would no longer use
    for body in ("    return [{'sku': 'X', 'description': 'y'}]\n", "    rows = fetch()\n    return rows\n", "    return []\n    # unreachable\n    return 1\n", "    return list()\n"):
        time.sleep(0.02)
        agent.write_text(base + "def read_lcom_catalog(source=None) -> list:\n" + body, encoding="utf-8")
        os.utime(agent, (time.time() + 5, time.time() + 5 + len(body)))                   # a new modified time, so the cache is not used
        assert search_mod.builtin_products(env.settings) == [], body
    page = env.client.get("/step/search").text
    assert "The built-in product list could not be read." in page and 'name="sku"' not in page and 'value="preset" data-count="3" disabled' in page
    with keys():
        assert post(env, "/search/start", choice="preset", confirm_count="3").status_code == 400
        assert post(env, "/search/start", choice="all", confirm_count="5").status_code == 400
    raises(JobRejected, env.jobs.start, "search", {"sku": ["HDFF"], "output": "sourcing_results_20260505_101010.xlsx"})
    assert not env.jobs.list()


@test
def report_is_copied_into_reports_only_after_strict_checks(env):
    s, name = env.settings, "sourcing_report_20260101_101010.md"
    line = lambda p: f"noise\nReport written to: {p}\nExcel results written to: x\n"
    ok = lambda r: r["copied"] and not r["problem"]
    assert search_mod.copy_report(s, "no such line here") == {"name": None, "copied": False, "problem": "The agent's log has no 'Report written to' line, so no report was copied."}
    assert search_mod.copy_report(s, None)["copied"] is False
    assert search_mod.copy_report(s, line(env.project / name))["copied"] is False and not (env.project / "Reports").exists()   # no such file
    (env.project / name).write_text("the report", encoding="utf-8")
    outside = env.root / name
    outside.write_text("outside the project", encoding="utf-8")
    (env.project / "Excel Output Sheets" / name).write_text("in a sub-folder", encoding="utf-8")
    (env.project / "sourcing_report_20260202_101010.md").mkdir()
    for fname in ("notes.md", "sourcing_report_x.md", "email_state.json", "sourcing_report_20260101_101010.md.txt", ".env"):
        (env.project / fname).write_text("secret or unrelated", encoding="utf-8")
    refused = [outside, env.project / ".." / name, env.project / "Excel Output Sheets" / name, env.project / "Excel Output Sheets" / ".." / ".." / name,
               name, f".\\{name}", f"..\\{name}", env.project / "sourcing_report_20260202_101010.md", env.project / "sourcing_report_20260303_101010.md",
               env.project / "notes.md", env.project / "sourcing_report_x.md", env.project / "email_state.json", env.project / ".env",
               env.project / "sourcing_report_20260101_101010.md.txt", f"{env.project / name}\x00.md", f"\\\\server\\share\\{name}", "C:\\Windows\\win.ini", "", "  "]
    cwd = os.getcwd()
    os.chdir(env.project)                                                                # so a relative path would find the file if it were allowed
    try:
        for p in refused:
            r = search_mod.copy_report(s, line(p))
            assert r["copied"] is False and "did not pass the checks" in r["problem"] or "no 'Report written to' line" in r["problem"], (str(p), r)
            assert r["name"] in (None, name, "sourcing_report_20260202_101010.md", "sourcing_report_20260303_101010.md"), r
    finally:
        os.chdir(cwd)
    assert not (env.project / "Reports").exists() or not list((env.project / "Reports").iterdir()), "a refused path must copy nothing"
    link = env.project / "sourcing_report_20260404_101010.md"
    try:
        os.symlink(outside, link)                                                         # needs a privilege on Windows; skipped when unavailable
    except OSError:
        link = None
    if link:
        assert search_mod.copy_report(s, line(link))["copied"] is False
    # the good path: last line wins, a path through '..' that really is in the project folder is fine, original untouched
    good = search_mod.copy_report(s, line(outside) + line(env.project / "Excel Output Sheets" / ".." / name))
    assert ok(good) and good["name"] == name
    assert (env.project / "Reports" / name).read_text(encoding="utf-8") == "the report" == (env.project / name).read_text(encoding="utf-8")
    (env.project / name).write_text("changed later", encoding="utf-8")
    assert ok(search_mod.copy_report(s, line(env.project / name)))
    assert (env.project / "Reports" / name).read_text(encoding="utf-8") == "the report", "an existing copy is never overwritten"
    assert (env.project / name).read_text(encoding="utf-8") == "changed later"
    shutil.rmtree(env.project / "Reports")                                                # Reports that is a link to somewhere else: refused
    (env.root / "elsewhere").mkdir()
    assert make_link(env.project / "Reports", env.root / "elsewhere")
    assert search_mod.copy_report(s, line(env.project / name))["copied"] is False and not list((env.root / "elsewhere").iterdir())


@test
def job_retention_prunes_only_old_finished_jobs_inside_the_jobs_folder(env):
    jdir = env.settings.jobs
    def job(jid, status, finished):
        write_json_atomic(jdir / f"{jid}.json", {"id": jid, "command": "stub.read", "label": "x", "args": [], "state_changing": False,
                                                 "status": status, "started_at": "2020-01-01T00:00:00", "finished_at": finished,
                                                 "duration_s": 1, "exit_code": 0, "exit_meaning": "", "pid": 1, "note": "", "tail": []})
        (jdir / f"{jid}.log").write_text("log", encoding="utf-8")
    recent = datetime.now().isoformat(timespec="seconds")
    job("20200101-000000-aaaaaa", "succeeded", "2020-01-01T00:01:00")                  # old and finished: goes
    job("20200101-000000-bbbbbb", "failed", recent)                                      # finished recently: stays
    job("20200101-000000-cccccc", "running", None)                                       # still running: stays however old
    job("20200101-000000-dddddd", "cancelled", "2020-01-01T00:01:00")                  # old and finished: goes
    job("20200101-000000-abcabc", "running", "2020-01-01T00:01:00")                    # marked running: stays whatever its dates say
    keep = [jdir / "notes.txt", jdir / "atomic.json", jdir / "20190101-000000-eeeeee.log", jdir / "x.json.tmp",
            env.project / "approvals_log.jsonl", env.project / "email_state.json", env.settings.data / "errors.log",
            env.project / "20200101-000000-aaaaaa.json", env.settings.results_dir / f"{STEM_A}.xlsx"]
    for f in keep:
        if not f.exists():
            f.write_text('{"id": "20200101-000000-aaaaaa", "status": "succeeded", "finished_at": "2020-01-01T00:00:00"}', encoding="utf-8")
    assert env.jobs.prune() == 2
    left = sorted(p.name for p in jdir.iterdir())
    assert not [n for n in left if "aaaaaa" in n or "dddddd" in n], left
    assert all((jdir / f"20200101-000000-{x}.{ext}").exists() for x in ("bbbbbb", "cccccc", "abcabc") for ext in ("json", "log")), left
    assert all(f.exists() for f in keep), [str(f) for f in keep if not f.exists()]
    job("20200101-000000-ffffff", "succeeded", "2020-01-01T00:01:00")
    off = JobManager(Settings(project=env.project, data=env.settings.data, enable_stub=True, retention_days=0))   # 0 disables, also at startup
    assert off.prune() == 0 and (jdir / "20200101-000000-ffffff.json").exists()
    assert "20200101-000000-ffffff" not in env.client.get("/jobs").text and not (jdir / "20200101-000000-ffffff.log").exists()   # the Jobs page prunes
    job("20200101-000000-999999", "succeeded", "2020-01-01T00:01:00")
    JobManager(env.settings)                                                             # and so does startup
    assert not (jdir / "20200101-000000-999999.json").exists()
    for raw, want in (("7", 7), ("abc", 30), ("-5", 0), ("0", 0), (None, 30)):
        with env_vars(UI_LOG_RETENTION_DAYS=raw):
            assert load_settings().retention_days == want, raw
    assert "# UI_LOG_RETENTION_DAYS=" in (Path(__file__).parent / ".env.example").read_text(encoding="utf-8")


# ================= running in a container: the time zone, reference files, start-up =================

@test
def health_page_names_the_time_zone(env):
    from datetime import timedelta, timezone
    from ui import pipeline
    for hours, minutes, name, want in ((5, 30, None, "UTC+05:30"), (-8, 0, None, "UTC-08:00"), (0, 0, None, "UTC+00:00"), (-3, -30, None, "UTC-03:30"),
                                       (13, 45, None, "UTC+13:45"), (5, 30, "IST", "UTC+05:30 (IST)"), (-6, 0, "Central Standard Time", "UTC-06:00 (Central Standard Time)")):
        zone = timezone(timedelta(hours=hours, minutes=minutes), name) if name else timezone(timedelta(hours=hours, minutes=minutes))
        assert pipeline.zone_label(datetime(2026, 1, 15, 12, 0, tzinfo=zone)) == want, (hours, minutes, name)
    off = datetime.now().astimezone().utcoffset().total_seconds()              # this computer's own zone, worked out separately
    here = f"UTC{'+' if off >= 0 else '-'}{int(abs(off)) // 3600:02d}:{int(abs(off)) % 3600 // 60:02d}"
    assert pipeline.zone_label().startswith(here), (pipeline.zone_label(), here)
    page = env.client.get("/health").text
    assert f"Times on these pages are shown in <strong>{here}" in page, here
    # it is the zone the pages convert to: an agent's UTC time is shown shifted by exactly that offset
    shown = datetime.strptime(pipeline._shown("2026-01-15T12:00:00+00:00"), "%Y-%m-%d %H:%M")
    local = datetime(2026, 1, 15, 12, 0, tzinfo=timezone.utc).astimezone()
    assert shown == local.replace(tzinfo=None) and pipeline.zone_label(local).startswith("UTC" + local.strftime("%z")[:3])


@test
def health_page_lists_changed_reference_files_by_name_only(env):
    from ui.config import REFERENCE_FILES, changed_reference_files
    shipped = env.root / "shipped"
    shipped.mkdir()
    texts = {n: f"SHIPPED-CONTENT-{i}\nline two\n" for i, n in enumerate(REFERENCE_FILES)}
    for n, t in texts.items():
        (shipped / n).write_bytes(t.encode())
    s = Settings(project=env.project, data=env.project / "d5", shipped=shipped)
    c = TestClient(create_app(s), base_url="http://127.0.0.1:8000", raise_server_exceptions=False)
    note = "differ from the ones that came with this version"
    assert changed_reference_files(s) == [] and note not in c.get("/health").text          # the data folder has none of them yet
    for n, t in texts.items():
        (env.project / n).write_bytes(t.encode())
    assert changed_reference_files(s) == [] and note not in c.get("/health").text          # identical copies
    (env.project / "near_misses.csv").write_bytes(texts["near_misses.csv"].replace("\n", "\r\n").encode())
    assert changed_reference_files(s) == [], "only the line endings differ: that is not a change"
    (env.project / "reviewer_exclusions.csv").write_bytes(b"PERSONAL-EDIT-ZZTOP,secret seller note\n")
    (env.project / "email_template.txt").write_bytes(texts["email_template.txt"].encode() + b"one more line\n")
    assert changed_reference_files(s) == ["email_template.txt", "reviewer_exclusions.csv"]
    page = c.get("/health").text
    block = page.split(note)[1].split("</div>")[0]
    assert "email_template.txt, reviewer_exclusions.csv" in block and "near_misses" not in block and "The copies in your data folder are the ones in use" in block
    for content in ("PERSONAL-EDIT-ZZTOP", "secret seller note", "SHIPPED-CONTENT", "one more line"):
        assert content not in page, "the note names files; it never shows what is in them"
    (env.project / "lcom_prices.csv").write_bytes(b"changed\n")
    page = c.get("/health").text
    assert "reference_prices.csv" in page.split(note)[1].split("</div>")[0] and not has_term(page)      # shown under the UI's neutral name
    (shipped / "near_misses.csv").unlink()                                                                # nothing came with this version to compare with
    (env.project / "near_misses.csv").write_bytes(b"anything\n")
    assert "near_misses.csv" not in changed_reference_files(s)
    # outside a container nothing is compared at all: no shipped folder, no note, whatever the files hold
    assert env.settings.shipped is None and changed_reference_files(env.settings) == [] and note not in env.client.get("/health").text
    with env_vars(UI_SHIPPED_DIR=str(shipped)):
        assert load_settings().shipped == shipped.resolve()
    with env_vars(UI_SHIPPED_DIR=None):
        assert load_settings().shipped is None


def entrypoint_module():
    """docker\\entrypoint.py, loaded by path the way the container runs it."""
    import importlib.util
    spec = importlib.util.spec_from_file_location("lcom_entrypoint", Path(__file__).with_name("docker") / "entrypoint.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def fake_image(root: Path) -> Path:
    """A stand-in for the image's /app folder: every program and reference file, each with known content."""
    from ui.config import PROGRAM_FILES, REFERENCE_FILES
    app = root / "image_app"
    app.mkdir()
    for n in PROGRAM_FILES + REFERENCE_FILES:
        (app / n).write_bytes(f"# this version's {n}\r\nsecond line\r\n".encode())
    return app


def fake_zones(root: Path) -> Path:
    zones = root / "zones"
    (zones / "Asia").mkdir(parents=True)
    (zones / "Asia" / "Kolkata").write_bytes(b"TZif")
    (zones / "UTC").write_bytes(b"TZif")
    return zones


def tree(folder: Path) -> dict:
    """{relative path: sha256} of every file under a folder."""
    import hashlib
    return {p.relative_to(folder).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(folder.rglob("*")) if p.is_file()}


MOUNTED, NOT_MOUNTED = (lambda p: True), (lambda p: False)


@test
def container_start_refuses_a_missing_mount_and_the_source_folder(env):
    ep = entrypoint_module()
    app, zones, data = fake_image(env.root), fake_zones(env.root), env.root / "data_folder"
    data.mkdir()
    started, real_app = [], ep.APP
    ep.APP = app
    try:
        assert "no data folder is mounted" in str(raises(ep.Refused, ep.check_data, data, NOT_MOUNTED))
        raises(ep.Refused, ep.check_data, env.root / "no_such_folder", MOUNTED)
        ep.check_data(data, MOUNTED)                                                   # an empty mounted folder is fine
        assert list(data.iterdir()) == [], "the write check leaves nothing behind"
        with env_vars(TZ="Asia/Kolkata", PROJECT_DIR=None, UI_DATA_DIR=None, UI_SHIPPED_DIR=None, UI_HOST=None, UI_ALLOW_NON_LOOPBACK=None, UI_ENABLE_STUB=None):
            e = raises(SystemExit, ep.main, data, NOT_MOUNTED, zones, lambda: started.append("ui"))
            assert str(e).startswith("NOT STARTED: no data folder is mounted") and list(data.iterdir()) == []
            # the source folder mounted as the data folder: refused, and the working copy in it is not replaced
            planted = b"MY OWN WORKING COPY - DO NOT REPLACE\n"
            for mark in ep.SOURCE_MARKS:
                src = env.root / f"source_{mark.strip('.')}"
                src.mkdir()
                (src / mark).mkdir() if mark in (".git", "ui") else (src / mark).write_text("x")
                (src / "email_agent.py").write_bytes(planted)
                before = tree(src)
                e = raises(SystemExit, ep.main, src, MOUNTED, zones, lambda: started.append("ui"))
                assert "looks like the source folder" in str(e) and mark in str(e), (mark, str(e))
                assert tree(src) == before and (src / "email_agent.py").read_bytes() == planted, mark
            assert ep.SOURCE_MARKS == (".git", "ui", "run_ui.py", "Dockerfile")
        assert started == [], "a refused start must not start the UI"
    finally:
        ep.APP = real_app


@test
def container_start_needs_a_time_zone_it_knows(env):
    """The compose file passes TZ through as it is (empty when .env has none); the entrypoint is the one that refuses."""
    ep = entrypoint_module()
    app, zones, data = fake_image(env.root), fake_zones(env.root), env.root / "data_folder"
    data.mkdir()
    started, real_app = [], ep.APP
    ep.APP = app
    rest = dict(PROJECT_DIR=None, UI_DATA_DIR=None, UI_SHIPPED_DIR=None, UI_HOST=None, UI_ALLOW_NON_LOOPBACK=None, UI_ENABLE_STUB=None)
    try:
        for empty in (None, "", "   "):                                                # not in .env, empty, or only spaces
            with env_vars(TZ=empty, **rest):
                e = str(raises(SystemExit, ep.main, data, MOUNTED, zones, lambda: started.append("ui")))
            assert e.startswith("NOT STARTED: TZ is not set.") and "Add your time zone to .env, for example TZ=Asia/Kolkata" in e, (empty, e)
            assert "docker compose up -d" in e and list(data.iterdir()) == [] and started == [], (empty, e)
        for unknown in ("Asia/Kolkatta", "Mars/Olympus", "Asia", "../zones/UTC", "/etc/passwd", "Asia/Kolkata; x", " UTC", "UTC "):
            with env_vars(TZ=unknown, **rest):
                e = str(raises(SystemExit, ep.main, data, MOUNTED, zones, lambda: started.append("ui")))
            assert e.startswith("NOT STARTED: TZ in .env is not a time zone name this system knows.") and "such as Asia/Kolkata" in e, (unknown, e)
            assert unknown.strip() not in e.replace("Asia/Kolkata or", ""), "a refusal never repeats the value that was set"
            assert list(data.iterdir()) == [] and started == [], unknown
        import contextlib
        import io
        for known in ("Asia/Kolkata", "UTC"):                                            # accepted: the UI starts, once each
            with env_vars(TZ=known, **rest), contextlib.redirect_stdout(io.StringIO()) as said:
                ep.main(data, MOUNTED, zones, lambda: started.append(known))
            assert said.getvalue().endswith("data folder ready, starting the UI\n"), said.getvalue()
        assert started == ["Asia/Kolkata", "UTC"]
    finally:
        ep.APP = real_app


@test
def compose_never_fails_a_command_because_a_setting_is_missing(env):
    """`${NAME:?message}` would make stop, down, ps, logs and config fail too. The compose file has none, and passes TZ through."""
    compose = (Path(__file__).parent / "docker-compose.yml").read_text(encoding="utf-8")
    assert ":?" not in compose and "${TZ:-}" in compose, "TZ is passed through as it is; the container refuses, not compose"
    assert [line.strip() for line in compose.splitlines() if line.strip().startswith("TZ:")] == ['TZ: "${TZ:-}"'], "one TZ line, passed through"

@test
def container_start_writes_only_the_program_files_and_leaves_everything_else_alone(env):
    from ui.config import PROGRAM_FILES, REFERENCE_FILES
    ep = entrypoint_module()
    app, data = fake_image(env.root), env.root / "data_folder"
    data.mkdir()
    yours = {"email_state.json": b'{"k": 1}', "decision_state.json": b"{}", "approvals_log.jsonl": b'{"a": 1}\n', "approved_orders.csv": b"NOT AN ORDER\n",
             "approved_orders.json": b"{}", "product_facts.csv": b"sku\n", ".env": b'SMTP_PASSWORD="FAKE-not-a-password"\n', "email_state.json.lock": b"{}",
             "notes.txt": b"mine", "Excel Output Sheets/sourcing_results_20260101_100000.xlsx": b"PK-not-really",
             "Reports/approval_request_A-0001.txt": b"request", "Uploads/input_20260101_000000_abcdef.xlsx": b"PK",
             "ui_data/jobs/20260101-000000-abcdef.log": b"log", "email_agent.py": b"# an older version's email agent\n",
             "reviewer_exclusions.csv": b"sku,supplier_or_url,reason\nMINE,x,my own edit\n"}
    for name, content in yours.items():
        (data / name).parent.mkdir(parents=True, exist_ok=True)
        (data / name).write_bytes(content)
    kept = [n for n in yours if n != "email_agent.py"]
    added = [n for n in REFERENCE_FILES if n != "reviewer_exclusions.csv"]
    note = "NOTE: your reviewer_exclusions.csv differs from the one in this version. Yours is kept and used. The Health page lists these files."
    said = []
    ep.install(app, data, said.append)
    after = tree(data)
    assert all((data / n).read_bytes() == (app / n).read_bytes() for n in PROGRAM_FILES), "every program file is this version's"
    assert all((data / n).read_bytes() == yours[n] for n in kept), "state, outputs, settings and a person's own reference file are untouched"
    assert all((data / n).read_bytes() == (app / n).read_bytes() for n in added), "a missing reference file is added"
    assert set(after) == set(kept) | set(PROGRAM_FILES) | set(added), "nothing else appears: no temp file, no folder of its own"
    assert sorted(said) == sorted([f"program file written: {n}" for n in PROGRAM_FILES] + [f"reference file added: {n}" for n in added] + [note])
    assert not any(v in " ".join(said) for v in ("FAKE-not-a-password", "my own edit")), "start-up lines name files, never what is in them"
    # a second start writes nothing at all
    stamps = {p: p.stat().st_mtime_ns for p in data.rglob("*") if p.is_file()}
    said.clear()
    ep.install(app, data, said.append)
    assert tree(data) == after and {p: p.stat().st_mtime_ns for p in data.rglob("*") if p.is_file()} == stamps and said == [note]
    # a new version of the image: its program files replace the old ones; a person's reference copy and all state stay
    (app / "decision_agent.py").write_bytes(b"# the next version\n")
    (app / "lcom_prices.csv").write_bytes(b"sku,price\nNEW,1\n")
    seeded = (data / "lcom_prices.csv").read_bytes()
    said.clear()
    ep.install(app, data, said.append)
    assert (data / "decision_agent.py").read_bytes() == b"# the next version\n" and "program file written: decision_agent.py" in said
    assert (data / "lcom_prices.csv").read_bytes() == seeded and any(line.startswith("NOTE: your lcom_prices.csv differs") for line in said)
    assert all((data / n).read_bytes() == yours[n] for n in kept) and len(said) == 3, said
    # a program file's name taken by a folder: refused before anything is written, even a file that would have changed
    (app / "sourcing_agent.py").write_bytes(b"# changed again\n")
    (data / "statefile.py").unlink()
    (data / "statefile.py").mkdir()
    before = tree(data)
    assert "statefile.py in the data folder is not an ordinary file" in str(raises(ep.Refused, ep.install, app, data, said.append))
    assert tree(data) == before and (data / "sourcing_agent.py").read_bytes() != b"# changed again\n"
    (data / "statefile.py").rmdir()
    (app / "order_sheet.py").unlink()                                                  # an image that lacks a program file
    assert "order_sheet.py is missing" in str(raises(ep.Refused, ep.install, app, data, said.append)) and tree(data) == before
    assert ep.DATA == Path("/data") and set(PROGRAM_FILES) & set(REFERENCE_FILES) == set() and len(PROGRAM_FILES) == 7


@test
def container_start_hands_the_data_folder_to_the_ui(env):
    ep = entrypoint_module()
    app, zones, data = fake_image(env.root), fake_zones(env.root), env.root / "data_folder"
    data.mkdir()
    names = ("PROJECT_DIR", "UI_DATA_DIR", "UI_SHIPPED_DIR", "UI_HOST", "UI_ALLOW_NON_LOOPBACK", "UI_ENABLE_STUB", "TZ")
    seen, real_app = {}, ep.APP

    def start():
        seen.update({k: os.environ.get(k) for k in names})
        seen["settings"] = load_settings()

    import contextlib
    import io
    ep.APP = app
    try:        # a settings file that tries to move the project, change the address or switch the test stub on does not get its way
        with env_vars(TZ="Asia/Kolkata", PROJECT_DIR=str(env.root / "somewhere_else"), UI_HOST="10.1.2.3", UI_ENABLE_STUB="1", UI_DATA_DIR=None,
                      UI_SHIPPED_DIR=None, UI_ALLOW_NON_LOOPBACK=None, SMTP_PASSWORD="FAKE-not-a-password"), contextlib.redirect_stdout(io.StringIO()) as said:
            ep.main(data, MOUNTED, zones, start)
    finally:
        ep.APP = real_app
    lines = said.getvalue().splitlines()
    assert lines[-1] == "data folder ready, starting the UI" and len(lines) == 12 and all(line.split(": ")[0] in ("program file written", "reference file added") for line in lines[:-1]), lines
    assert not any(v in said.getvalue() for v in ("FAKE-not-a-password", "Asia/Kolkata", "10.1.2.3")), "start-up never prints the value of a setting"
    s = seen.pop("settings")
    assert seen == {"PROJECT_DIR": str(data), "UI_DATA_DIR": str(data / "ui_data"), "UI_SHIPPED_DIR": str(app), "UI_HOST": "0.0.0.0",
                    "UI_ALLOW_NON_LOOPBACK": "1", "UI_ENABLE_STUB": None, "TZ": "Asia/Kolkata"}, seen
    assert s.project == data.resolve() and s.data == (data / "ui_data").resolve() and s.shipped == app.resolve() and s.host == "0.0.0.0" and not s.enable_stub
    assert security.allowed_hosts(s) == {"127.0.0.1:8000", "localhost:8000"}, "listening on every interface does not widen the hosts it answers to"
    assert (data / "email_agent.py").read_bytes() == (app / "email_agent.py").read_bytes()         # the agents the UI will run are in the data folder
    assert os.environ.get("PROJECT_DIR") != str(data), "the test's own environment is restored"


def main():
    failed, skipped = [], []
    for f in TESTS:
        env = make_env()
        try:
            f(env)
            print(f"ok   {f.__name__}")
        except Skip as e:
            skipped.append(f.__name__)
            print(f"skip {f.__name__}: {e}")
        except Exception:  # noqa: BLE001
            failed.append(f.__name__)
            print(f"FAIL {f.__name__}\n{traceback.format_exc()}")
        finally:
            cleanup(env)
    print(f"\n{len(TESTS) - len(failed) - len(skipped)}/{len(TESTS)} passed, {len(skipped)} skipped" + (f": {', '.join(skipped)}" if skipped else ""))
    if failed:
        sys.exit(1)
    print("ui tests passed")


if __name__ == "__main__":
    main()
