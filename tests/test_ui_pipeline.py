"""Offline tests for the pipeline UI: uploads, every tab's gates, the allowlist, and (at the end) the whole pipeline with
the real agents behind a test-only shim. No network, no real state: every test works in a throwaway project folder, and a
tripwire fails the run if anything opens a file of the real project.
Run:  $env:OBS_ENABLED = "0"; python tests\\test_ui_pipeline.py"""
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import zipfile
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace

os.environ["OBS_ENABLED"] = "0"
REAL = Path(__file__).resolve().parent.parent          # the real project folder: code is imported from it, data never read
sys.path.insert(0, str(REAL))

# ---------- the tripwire: the real project's data is never opened ----------
VIOLATIONS = []
_REAL = os.path.normcase(str(REAL)) + os.sep
_CODE_DIRS = tuple(os.path.normcase(str(REAL / d)) + os.sep for d in ("ui", "tests", "__pycache__", "scratch"))
_WATCHED = {"open": 0, "os.listdir": 0, "os.scandir": 0, "os.mkdir": 0, "os.remove": 0, "os.rename": 0, "os.rmdir": 0}


_SNAPSHOT = False                       # True only inside snapshot_real(), which reads the real state to prove it did not change
_CLEANING = None                        # the test's own temp folder, only while remove_temp() deletes it


def _tripwire(event, args):
    if _SNAPSHOT or event not in _WATCHED or not args or not isinstance(args[0], (str, bytes, os.PathLike)):
        return
    raw = os.fsdecode(args[0])
    if _CLEANING is not None and not os.path.isabs(raw):
        return                          # Linux's rmtree opens each entry by its bare name inside an open folder of the temp tree
    path = os.path.normcase(os.path.abspath(raw))
    if not (path + os.sep).startswith(_REAL) or (path + os.sep).startswith(_CODE_DIRS):
        return
    if event in ("os.listdir", "os.scandir") and path + os.sep == _REAL:     # Python's import system lists the folder to find modules
        return
    reading_code = event == "open" and os.path.dirname(path) + os.sep == _REAL and path.endswith((".py", "email_template.txt")) \
        and not (isinstance(args[1], str) and set(args[1]) & set("wax+"))
    if not reading_code:                    # agent sources are read to copy them into a test project; nothing else is touched
        VIOLATIONS.append(f"{event} {path}")
        raise AssertionError(f"the real project folder was used: {event} {path}")


sys.addaudithook(_tripwire)

import openpyxl                                         # noqa: E402
from fastapi.testclient import TestClient               # noqa: E402

from ui import commands, config, search, security, uploads, workflows   # noqa: E402
from ui.app import create_app                           # noqa: E402
from ui.commands import JobRejected                     # noqa: E402
from ui.config import MAX_PRODUCT_ROWS, MAX_UPLOAD_BYTES, Settings   # noqa: E402
from ui.jobs import JobBusy, JobManager                 # noqa: E402

ORIGIN = {"origin": "http://127.0.0.1:8000"}
FAKE = {"ANTHROPIC_API_KEY": "sk-ant-FAKE-KEY-pipeline-000111222", "NIMBLE_API_KEY": "nimble-FAKE-secret-333444",
        "SMTP_HOST": "smtp.example.invalid", "SMTP_USER": "sender@example.invalid", "SMTP_PASSWORD": "FAKE-smtp-password-555",
        "SHIPPING_ADDRESS": "Test Receiving\\n1 Test Street\\nTestville, TX 00000", "BRAND_NAME": "TestBrand",
        "SENDER_NAME": "Test Sender"}
SECRETS = (FAKE["ANTHROPIC_API_KEY"], FAKE["NIMBLE_API_KEY"], FAKE["SMTP_PASSWORD"], "1 Test Street", "Testville", "+1 555 010 0199")
BUILTIN = [("HDFF", "HDMI coupler"), ("FOA-020C", "Fiber adapter"), ("LCSP1050", "Surge protector")]
RESULT_HEAD = ["SKU", "Product", "Keyword", "Recommendation", "Recommended Manufacturer", "Recommended Email", "Recommended URL",
               "Recommended Unit Price", "Ordering Note"]
TESTS = []


def test(f):
    TESTS.append(f)
    return f


def raises(exc, fn, *a, **kw):
    try:
        fn(*a, **kw)
    except exc as e:
        return e
    raise AssertionError(f"expected {exc.__name__} from {getattr(fn, '__name__', fn)}{a}{kw}")


class Skip(Exception):
    """Raised by a test that cannot run on this system. The runner counts it and prints the reason; it is never silent."""


def native(*parts) -> str:
    """A path inside the project as the UI hands it to an agent: this system's own separator (a backslash on Windows)."""
    return str(Path(*parts))


def link_folder(link: Path, target: Path) -> bool:
    """A folder that is really another folder: a junction on Windows (needs no special right), a symlink elsewhere."""
    if os.name == "nt":
        return subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(target)], capture_output=True).returncode == 0
    os.symlink(target, link, target_is_directory=True)
    return True


def unlink_folder(link: Path) -> None:
    """Removes the link itself and leaves its target alone."""
    os.rmdir(link) if os.name == "nt" else os.unlink(link)


@contextmanager
def env_vars(**kv):
    saved = {k: os.environ.get(k) for k in kv}
    try:
        for k, v in kv.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)
        yield
    finally:
        for k, v in saved.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)


def make_env(stub=True, popen=None, agents=("email_agent.py",)):
    root = Path(tempfile.mkdtemp(prefix="ui_pipe_"))
    project = root / "project"
    (project / "Excel Output Sheets").mkdir(parents=True)
    for name in agents:
        shutil.copy(REAL / name, project)
    if "sourcing_agent.py" not in agents:                # only its built-in list is read (parsed, never run)
        (project / "sourcing_agent.py").write_text("PRODUCTS = " + repr([{"sku": k, "description": d} for k, d in BUILTIN]) + "\n", encoding="utf-8")
    s = Settings(project=project, data=project / "ui_data", enable_stub=stub)
    app = create_app(s, popen=popen)
    client = TestClient(app, base_url="http://127.0.0.1:8000", raise_server_exceptions=False)
    return SimpleNamespace(root=root, project=project, settings=s, app=app, client=client, jobs=app.state.jobs)


def cleanup(env):
    for j in env.jobs.list():
        env.jobs.cancel(j["id"])
    time.sleep(0.2)
    remove_temp(env.root)


def remove_temp(root: Path) -> None:
    """Deletes a test's own temp folder. Only while this runs does the tripwire let relative names through, and only for a
    folder under the system temp folder that is neither the real project nor inside it; absolute paths are checked as always."""
    global _CLEANING
    root = Path(root).resolve()
    temp = Path(tempfile.gettempdir()).resolve()
    assert root.is_relative_to(temp) and root != temp and not root.is_relative_to(REAL) and not REAL.is_relative_to(root), root
    _CLEANING = root
    try:
        shutil.rmtree(root, ignore_errors=True)
    finally:
        _CLEANING = None


def token(env):
    return re.search(r'name="csrf_token" value="([0-9a-f]+)"', env.client.get("/step/search").text)[1]


def post(env, url, files=None, **data):
    return env.client.post(url, data={"csrf_token": token(env), **data}, files=files, headers=ORIGIN, follow_redirects=False)


def workbook(rows, sheet="Sheet1", macro=False) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = sheet
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    if not macro:
        return buf.getvalue()
    out = io.BytesIO()                                   # the same workbook carrying a macro project, as an .xlsm does
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"not a real macro")
    return out.getvalue()


SHEET = [["SKU", "Product", "Keyword", "lcom sale price"], ["IN-1", "Cable", "Cat6 patch cable", 4.5], ["IN-2", "Adapter", "USB adapter", 9],
         ["C&P9M", "Connector", "DB9 connector", 2]]


def upload(env, data=None, name="My products.xlsx"):
    return post(env, "/search/upload", files={"sheet": (name, data if data is not None else workbook(SHEET),
                                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def uploaded_value(r) -> str:
    assert r.status_code == 303 and "/step/search?file=Uploads%2Finput_" in r.headers["location"], (r.status_code, r.text[-500:])
    return "Uploads/" + r.headers["location"].split("Uploads%2F")[1]


def results_file(env, name, rows, mtime=None):
    """rows: (sku, seller, email, url). A results workbook with the columns the email agent reads."""
    wb = openpyxl.Workbook()
    wb.active.title = "Results"
    wb.active.append(RESULT_HEAD)
    for sku, seller, email, url in rows:
        wb.active.append([sku, f"Product {sku}", f"Keyword {sku}", "Buy from Candidate 1" if seller else "No candidate qualifies",
                          seller, email, url, 1.25 if seller else None, ""])
    path = env.project / "Excel Output Sheets" / name
    wb.save(path)
    if mtime:
        os.utime(path, (mtime, mtime))
    return path


class FakePopen:
    """Stands in for subprocess.Popen: records how it was called, 'prints' two lines, exits 0."""
    calls = []

    def __init__(self, argv, **kw):
        FakePopen.calls.append((argv, kw))
        self.pid, self.stdout = 4242, iter(["fake line one\n", "fake line two\n"])

    def wait(self):
        return 0

    def poll(self):
        return 0


# ================= Phase 1: the allowlist and the search command =================

@test
def argv_for_every_allowlisted_command(env):
    cmds = env.jobs.commands
    up = uploaded_value(upload(env))
    wb = openpyxl.Workbook()
    wb.active.append(["SKU"])
    wb.active.append(["ROOT-1"])
    wb.save(env.project / "products.xlsx")
    out = search.new_output_name()
    tail = ["--output", native("Excel Output Sheets", out)]
    path = native("Uploads", up[8:])
    argv = cmds["search"].argv
    assert argv({"output": out}) == ["search", *tail]
    assert argv({"sku": ["hdff", "FOA-020C"], "output": out}) == ["search", "--sku", "HDFF", "FOA-020C", *tail]      # the list's own spelling
    assert argv({"input": up, "output": out}) == ["search", "--input", path, *tail]
    assert argv({"input": "products.xlsx", "limit": "2", "add_builtin": "1", "output": out}) == ["search", "--input", "products.xlsx", "--limit", "2", "--add-builtin", *tail]
    assert argv({"input": up, "sku": ["in-2", "C&P9M"], "output": out}) == ["search", "--sku", "IN-2", "C&P9M", "--input", path, *tail]
    for agent in commands.AGENTS:
        assert cmds[f"help.{agent}"].argv({}) == ["--help"] and not cmds[f"help.{agent}"].state_changing
    assert cmds["search"].state_changing
    real = JobManager(Settings(project=env.project, data=env.project / "d2")).commands["search"]     # without the stub: the agent itself
    assert real.script == "sourcing_agent.py" and real.argv({"input": up, "output": out}) == ["--input", path, *tail]
    assert "inspect" not in " ".join(real.params) and all(p.flag != "--inspect-tools" for p in real.params.values())


@test
def allowlist_refuses_unknown_commands_flags_and_values(env):
    j = env.jobs
    up = uploaded_value(upload(env))
    out = search.new_output_name()
    for cmd in ("rm", "search; calc", "sourcing_agent.py", "help.rm", "../search", "", "email_agent.py send", "python"):
        raises(JobRejected, j.start, cmd, {})
    for bad in ({"inspect_tools": "1"}, {"yes": "1"}, {"output": "..\\x.xlsx"}, {"output": str(env.project / "x.xlsx")}, {},
                {"input": "Uploads/../sourcing_agent.py"}, {"input": "Uploads\\" + up[8:]}, {"input": str(env.project / up)},
                {"input": "Uploads/input_20260101_000000_abcdef.xlsx"}, {"input": up, "limit": "0"}, {"input": up, "limit": str(MAX_PRODUCT_ROWS + 1)},
                {"input": up, "sku": ["HDFF"]}, {"input": up, "sku": ["IN-1"], "limit": "1"}, {"input": up, "sku": ["IN-1"], "add_builtin": "1"},
                {"sku": ["IN-1"]}, {"sku": ["x; calc"]}, {"sku": ["--inspect-tools"]}, {"sku": "IN-1", "input": up}, {"limit": "2"},
                {"input": up, "add_builtin": "yes"}):
        raises(JobRejected, j.start, "search", {"output": out, **bad} if bad else {})
    assert j.list() == [], "a refused command must not leave a job behind"


@test
def jobs_run_through_a_replaceable_popen(env):
    FakePopen.calls = []
    fake = make_env(popen=FakePopen)
    try:
        out = search.new_output_name()
        meta = fake.jobs.wait(fake.jobs.start("search", {"sku": ["HDFF"], "output": out})["id"])
        (argv, kw), = FakePopen.calls
        assert argv[:3] == [sys.executable, "-X", "utf8"] and argv[3].endswith("stub_agent.py")
        assert argv[4:] == ["search", "--sku", "HDFF", "--output", native("Excel Output Sheets", out)]
        assert kw["shell"] is False and kw["stdin"] == subprocess.DEVNULL and Path(kw["cwd"]) == fake.project
        if os.name == "nt":                  # how the process is started, per system: on Windows exactly as before
            assert kw["creationflags"] == subprocess.CREATE_NO_WINDOW | subprocess.CREATE_NEW_PROCESS_GROUP and "start_new_session" not in kw
        else:                                # elsewhere: its own process group, so a cancel can stop the agent's children too
            assert kw["start_new_session"] is True and kw["creationflags"] == 0
        assert meta["status"] == "succeeded" and fake.jobs.read_log(meta["id"])[0] == "fake line one\nfake line two\n"
        assert meta["cwd"] == str(fake.project) and meta["script"] == argv[3]
    finally:
        cleanup(fake)


# ================= Phase 1: upload =================

@test
def upload_rejections(env):
    good = workbook(SHEET)
    cases = {
        "macro workbook saved as .xlsm": ("book.xlsm", workbook(SHEET, macro=True)),
        "macro workbook renamed to .xlsx": ("book.xlsx", workbook(SHEET, macro=True)),
        "old .xls": ("book.xls", good),
        "csv": ("book.csv", b"SKU,Product\nA,B\n"),
        "csv renamed to .xlsx": ("book.xlsx", b"SKU,Product\nA,B\n"),
        "a zip that is not a workbook": ("book.xlsx", _zip({"hello.txt": b"hi"})),
        "no extension": ("book", good),
        "double extension": ("book.xlsx.exe", good),
        "no SKU heading": ("book.xlsx", workbook([["Part", "Name"], ["A1", "Cable"]])),
        "SKU heading below row 5": ("book.xlsx", workbook([["x"]] * 5 + [["SKU"], ["A1"]])),
        "heading but no rows": ("book.xlsx", workbook([["SKU", "Product"]])),
        "too many rows": ("book.xlsx", workbook([["SKU"]] + [[f"S{i}"] for i in range(MAX_PRODUCT_ROWS + 1)])),
        "unpacks too large": ("book.xlsx", _zip({"[Content_Types].xml": b"x", "xl/workbook.xml": b"0" * (uploads.MAX_UNZIPPED_BYTES + 1)})),
        "empty file": ("book.xlsx", b""),
    }
    for why, (name, data) in cases.items():
        r = upload(env, data, name)
        assert r.status_code == 400 and "Nothing was uploaded" in r.text, (why, r.status_code)
        raises(JobRejected, uploads.save_workbook, env.settings, name, data)
    assert not env.settings.uploads_dir.exists() or not list(env.settings.uploads_dir.iterdir()), "a refused upload must leave no file"
    big = upload(env, good + b"0" * (MAX_UPLOAD_BYTES + security.BODY_OVERHEAD), "book.xlsx")                # refused before the body is read
    assert big.status_code == 413
    just_over = good + b"0" * (MAX_UPLOAD_BYTES - len(good) + 1)
    r = upload(env, just_over, "book.xlsx")                                                                      # one byte over, read and refused
    assert r.status_code == 400 and "larger than 5 MB" in r.text
    raises(JobRejected, uploads.check_workbook, "book.xlsx", just_over)
    assert not env.settings.uploads_dir.exists() or not list(env.settings.uploads_dir.iterdir())
    assert post(env, "/search/upload").status_code == 400                                                        # no file at all
    exactly = workbook([["SKU"]] + [[f"S{i}"] for i in range(MAX_PRODUCT_ROWS)])                              # the cap itself is allowed
    assert uploads.check_workbook("book.XLSX", exactly)["count"] == MAX_PRODUCT_ROWS
    assert search.MAX_LIMIT == MAX_PRODUCT_ROWS == config.MAX_PRODUCT_ROWS == 200, "one constant caps rows and --limit"
    # an Uploads folder that is a link to somewhere else is refused: nothing is written through it
    (env.root / "elsewhere").mkdir()
    shutil.rmtree(env.settings.uploads_dir, ignore_errors=True)
    assert link_folder(env.settings.uploads_dir, env.root / "elsewhere")
    r = upload(env, good, "book.xlsx")
    assert r.status_code == 400 and "link to another place" in r.text and not list((env.root / "elsewhere").iterdir())
    unlink_folder(env.settings.uploads_dir)


def _zip(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


@test
def upload_gets_a_server_name_and_is_never_edited(env):
    data = workbook(SHEET)
    hostile = ["..\\..\\evil.xlsx", "../../evil.xlsx", "C:\\Windows\\evil.xlsx", "/etc/evil.xlsx", "evil%00.txt.xlsx", "<script>alert(1)</script>.xlsx",
               "CON.xlsx", "a" * 300 + ".xlsx"]
    values = [uploaded_value(upload(env, data, name)) for name in hostile]
    saved = sorted(p.name for p in env.settings.uploads_dir.iterdir())
    assert len(saved) == len(hostile) and all(search.UPLOAD_RE.fullmatch(n) for n in saved), saved        # only server-built names
    assert sorted(v[8:] for v in values) == saved and search.upload_files(env.settings) == sorted(values, reverse=True)
    everything = [p for p in env.root.rglob("*") if p.is_file()]
    assert not [p for p in everything if "evil" in p.name.lower() or "script" in p.name.lower()], "the client's file name reached the disk"
    assert not (env.root / "evil.xlsx").exists() and not (env.project / "evil.xlsx").exists()
    page = env.client.get(f"/step/search?file={values[0]}").text
    assert "What was read from" in page and "<script>alert" not in page and "evil" not in page
    for label in ("Products (rows with a SKU)", "Keyword column", "Product column", "Reference sale price column"):
        assert label in page, label
    assert page.count(">Found<") == 3 and ">3<" in page and "Cat6 patch cable" in page and "IN-1" in page
    bare = uploaded_value(upload(env, workbook([["SKU"], ["ONLY-1"]])))
    page = env.client.get(f"/step/search?file={bare}").text
    assert page.count(">Not found<") == 3 and "only the SKU to go on" in page
    # a search on the upload leaves its bytes exactly as they were
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        r = post(env, "/search/start", choice="file", input_file=values[0], rows="all", confirm_count="3")
        assert r.status_code == 303, r.text[-400:]
        assert env.jobs.wait(r.headers["location"].split("job=")[1])["status"] == "succeeded"
    assert (env.project / values[0]).read_bytes() == data
    assert sorted(p.name for p in env.settings.uploads_dir.iterdir()) == sorted(saved + [bare[8:]]), "no copy or temp file is left in Uploads"
    for bad in ("Uploads/../sourcing_agent.py", "Uploads/nope.xlsx", "..\\x.xlsx", str(env.project / values[0]), "Uploads\\" + values[0][8:]):
        assert env.client.get("/step/search", params={"file": bad}).status_code == 400, bad


@test
def sku_selection_uses_input_and_sku(env):
    up = uploaded_value(upload(env))
    sel = search.resolve_selection(env.settings, "file", ["in-2", "C&P9M", "IN-2"], up)
    assert sel["params"] == {"input": up, "sku": ["IN-2", "C&P9M"]} and [p[0] for p in sel["products"]] == ["IN-2", "C&P9M"]
    for bad in (["HDFF"], ["IN-1", "NOPE"], ["IN-1; calc"], [""]):                              # a built-in SKU is not in this file
        raises(JobRejected, search.resolve_selection, env.settings, "file", bad, up)
    raises(JobRejected, search.resolve_selection, env.settings, "file", ["IN-1"], up, "1")      # not with a limit
    raises(JobRejected, search.resolve_selection, env.settings, "file", ["IN-1"], up, "", "1")  # not with the built-in list
    review = post(env, "/search/review", choice="file", input_file=up, rows="pick", sku=["IN-2", "C&P9M"])
    assert review.status_code == 200 and "About 2 products, roughly $0.16 and 2 minutes." in review.text and "chosen products" in review.text
    assert 'name="sku" value="IN-2"' in review.text and 'name="sku" value="C&amp;P9M"' in review.text
    assert post(env, "/search/review", choice="file", input_file=up, rows="pick").status_code == 400          # nothing ticked
    assert "About 3 products" in post(env, "/search/review", choice="file", input_file=up, rows="all", sku=["IN-2"], limit="1").text
    assert "About 1 product, roughly $0.08 and 1 minute." in post(env, "/search/review", choice="file", input_file=up, rows="first", sku=["IN-2"], limit="1").text
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        r = post(env, "/search/start", choice="file", input_file=up, sku=["IN-2", "C&P9M"], confirm_count="2")
        meta = env.jobs.wait(r.headers["location"].split("job=")[1])
    assert meta["args"][:6] == ["search", "--sku", "IN-2", "C&P9M", "--input", native("Uploads", up[8:])] and meta["status"] == "succeeded", meta
    out = openpyxl.load_workbook(env.project / "Excel Output Sheets" / meta["search"]["output"])["Results"]
    assert [r[0].value for r in out.iter_rows(min_row=2)] == ["IN-2", "C&P9M"]


@test
def search_needs_the_cost_confirmation(env):
    up = uploaded_value(upload(env))
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        review = post(env, "/search/review", choice="file", input_file=up, rows="all")
        assert "About 3 products, roughly $0.24 and 3 minutes." in review.text
        assert "This is an estimate from one 15-product batch (about $0.08 and about 56 seconds per product), not a measurement of this run." in review.text
        assert 'name="confirm_count"' in review.text and 'name="confirm_count" value' not in review.text       # never pre-filled
        for typed in ("", " ", "2", "4", "three", "3.0", "-3", "03x"):
            r = post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count=typed)
            assert r.status_code == 400 and "Nothing was started" in r.text, typed
        assert post(env, "/search/start", choice="file", input_file=up, rows="all").status_code == 400         # the field left out
        assert env.jobs.list() == [] and not list((env.project / "Excel Output Sheets").glob("*.xlsx"))
        ok = post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count="3")
        assert ok.status_code == 303
        env.jobs.wait(ok.headers["location"].split("job=")[1])
    assert search.estimate(15) == {"count": 15, "cost": "1.20", "minutes": 14, "seconds_each": 56, "cost_each": "0.08"}


# ================= security on every route =================

def all_routes(env):
    """[(method, path with its parameters filled in)] for every route the app serves."""
    out = []
    for r in env.app.routes:
        for m in sorted(getattr(r, "methods", None) or []):
            if m in ("GET", "POST") and not r.path.startswith("/static"):
                out.append((m, r.path.replace("{slug}", "emails").replace("{jid}", "20260101-000000-abcdef").replace("{folder}", "Reports")
                            .replace("{rel:path}", "x.csv").replace("{name}", "approved_orders.csv")))
    return out


@test
def every_post_route_checks_csrf(env):
    posts = [p for m, p in all_routes(env) if m == "POST"]
    assert {"/search/upload", "/search/review", "/search/start", "/jobs/run"} <= set(posts), posts
    good = token(env)
    for path in posts:
        for tok in (None, "", "0" * 64, good[:-1] + ("0" if good[-1] != "0" else "1")):
            r = env.client.post(path, data={} if tok is None else {"csrf_token": tok}, headers=ORIGIN, follow_redirects=False)
            assert r.status_code == 403, (path, tok, r.status_code)
        fresh = TestClient(env.app, base_url="http://127.0.0.1:8000")                    # another session's token is not this one's
        assert fresh.post(path, data={"csrf_token": good}, headers=ORIGIN, follow_redirects=False).status_code == 403, path
    assert env.jobs.list() == []


@test
def host_and_origin_are_checked_on_every_route(env):
    tok = token(env)
    for method, path in all_routes(env):
        for host in ("evil.example", "127.0.0.1:9999", "localhost", "127.0.0.1.evil.example:8000", "evil.example:8000") + WRONG_HOSTS:
            r = env.client.request(method, path, headers={"host": host.replace("PORT", "8000"), **ORIGIN}, data={"csrf_token": tok} if method == "POST" else None)
            assert r.status_code == 400 and r.text == "Invalid Host header.", (method, path, host, r.status_code)
        if method == "POST":
            for headers in ({}, {"origin": "http://evil.example"}, {"origin": "null"}, {"referer": "http://evil.example/x"},
                            {"origin": "https://127.0.0.1:8000"}, {"origin": "http://127.0.0.1:8001"}):
                r = env.client.post(path, data={"csrf_token": tok}, headers=headers, follow_redirects=False)
                assert r.status_code == 403 and "Cross-site" in r.text, (path, headers, r.status_code)
            ok = env.client.post(path, data={"csrf_token": tok}, headers={"referer": "http://localhost:8000/step/search"}, follow_redirects=False)
            assert ok.status_code != 403 or "Cross-site" not in ok.text, path          # our own page as Referer passes this check
    assert env.jobs.list() == []
    assert security.allowed_hosts(env.settings) == {"127.0.0.1:8000", "localhost:8000"}
    # as in a container: listening on every interface, reached through one published port on this computer. The hosts it
    # answers to are still only this computer's two names with that port; nothing is accepted because of where it listens.
    s = Settings(project=env.project, data=env.project / "d3", host="0.0.0.0", port=8010)
    contained = create_app(s)
    c = TestClient(contained, base_url="http://127.0.0.1:8010", raise_server_exceptions=False)
    assert security.allowed_hosts(s) == {"127.0.0.1:8010", "localhost:8010"}
    assert c.get("/health").status_code == 200 and c.get("/health", headers={"host": "localhost:8010"}).status_code == 200
    page_token = re.search(r'name="csrf_token" value="([0-9a-f]+)"', c.get("/step/search").text)[1]
    for host in WRONG_HOSTS + ("127.0.0.1:8000", "localhost:8000", "lcom-ui:PORT", "[::1]:PORT", "evil.example:PORT"):
        host = host.replace("PORT", "8010")
        r = c.get("/health", headers={"host": host})
        assert r.status_code == 400 and r.text == "Invalid Host header.", (host, r.status_code)
        r = c.post("/jobs/run", data={"csrf_token": page_token, "command": "help.email_agent"}, headers={"host": host, "origin": "http://127.0.0.1:8010"})
        assert r.status_code == 400 and r.text == "Invalid Host header.", (host, r.status_code)
    for origin in ("http://0.0.0.0:8010", "http://host.docker.internal:8010", "http://127.0.0.1:8000", "http://LOCALHOST:8010", "http://localhost.:8010"):
        r = c.post("/jobs/run", data={"csrf_token": page_token, "command": "help.email_agent"}, headers={"origin": origin}, follow_redirects=False)
        assert r.status_code == 403 and "Cross-site" in r.text, (origin, r.status_code)
    assert contained.state.jobs.list() == [], "a refused request started nothing"
    ok = c.post("/jobs/run", data={"csrf_token": page_token, "command": "help.email_agent"}, headers={"origin": "http://127.0.0.1:8010"}, follow_redirects=False)
    assert ok.status_code == 303, "the right host and origin are accepted through the same port"
    contained.state.jobs.wait(ok.headers["location"].rsplit("/", 1)[1])


# Host headers that must be refused wherever the UI runs. PORT is replaced by the port in use: the address it listens on,
# Docker's name for the host, the right name with a wrong port or no port, an upper-case name, and names with a trailing dot.
WRONG_HOSTS = ("0.0.0.0:PORT", "host.docker.internal:PORT", "127.0.0.1:1", "127.0.0.1", "LOCALHOST:PORT", "Localhost:PORT", "localhost.:PORT",
               "127.0.0.1.:PORT")


# ================= the strip, the results file, secrets =================

def state_key(sku: str, url: str) -> str:
    """The email agent's key, written out by hand: SKU upper-cased, then the listing's host and path."""
    host_path = url.split("://", 1)[1].split("?")[0].rstrip("/")
    return f"{sku.upper()}|{host_path.split('/', 1)[0].lower()}{'/' + host_path.split('/', 1)[1] if '/' in host_path else ''}"


def strip_of(env, stem, slug="emails") -> str:
    page = env.client.get(f"/step/{slug}?wf={stem}").text
    return page.split('<section class="strip"')[1].split("</section>")[0]


@test
def status_strip_follows_the_files(env):
    old, new = "sourcing_results_20260101_100000", "sourcing_results_20260202_100000"
    url_a, url_b, url_c = "https://www.Alibaba.com/product/a1/?spm=1", "https://maker-b.example/item/b2", "https://maker-c.example/item/c3"
    results_file(env, old + ".xlsx", [("a1", "Maker A", "a@maker-a.example", url_a)], mtime=time.time() - 9000)
    results_file(env, new + ".xlsx", [("B2", "Maker B", "b@maker-b.example", url_b), ("C3", "Maker C", "", url_c), ("D4", "", "", "")])
    strip = strip_of(env, new)
    for text in ("Results file exists", "3 products", "2 can be drafted", "0 marked sent", "0 replies recorded", "No request yet",
                 "No order sheet for this file", "No facts file yet", "Not generated since the last export"):
        assert text in strip, (text, strip)
    assert 'Next step: <a href="/step/emails?wf=' + new in strip
    assert "<form" not in strip and "<button" not in strip and "<input" not in strip, "the strip only links; it starts nothing"
    assert env.jobs.list() == []
    # email_state.json is one file for every results file: an entry counts for a results file only when its key
    # (SKU + the recommended listing's host and path) is the key of one of that file's draftable rows
    assert state_key("a1", url_a) == "A1|www.alibaba.com/product/a1" and state_key("B2", url_b) == "B2|maker-b.example/item/b2"
    (env.project / "email_state.json").write_text(json.dumps({
        state_key("a1", url_a): {"sku": "a1", "seller": "Maker A", "sent": "2026-01-01 10:00"},
        state_key("B2", url_b): {"sku": "B2", "seller": "Maker B", "sent": "2026-02-02 10:00", "reply": "hello", "status": "pricing_provided"},
        "B2|other-seller.example/item/old": {"sku": "B2", "seller": "A seller no longer recommended", "sent": "2026-01-15 10:00", "reply": "x"},
        "Z9|nowhere.example/z": {"sku": "Z9", "seller": "Unrelated", "sent": "2026-01-15 10:00"}}), encoding="utf-8")
    strip = strip_of(env, new)
    assert "1 marked sent" in strip and "1 reply recorded" in strip and "2 can be drafted" in strip, strip      # not 3 or 4: only this file's keys
    strip = strip_of(env, old)
    assert "1 marked sent" in strip and "0 replies recorded" in strip and "1 can be drafted" in strip, strip
    steps = {st.name: st for st in workflows.steps_for(env.settings, new + ".xlsx")}
    assert "1 seller(s) for this results file" in steps["Emails"].detail and "covers every results file" not in steps["Emails"].detail
    assert workflows.next_step(list(steps.values())).name == "Approvals"
    # approvals: request, reply, approval with its expiry, export
    soon = time.strftime("%Y-%m-%dT%H:%M:%S+00:00", time.gmtime(time.time() + 3 * 86400))
    past = time.strftime("%Y-%m-%dT%H:%M:%S+00:00", time.gmtime(time.time() - 86400))
    state = {"next_request": 3, "requests": {"A-0001": {"id": "A-0001", "results_file": new + ".xlsx", "reply": None, "created_at": past, "lines": []},
                                              "A-0002": {"id": "A-0002", "results_file": old + ".xlsx", "reply": None, "created_at": past, "lines": []}},
             "decisions": []}
    (env.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
    strip = strip_of(env, new)
    assert "1 request written" in strip and "0 replies recorded" in strip and "0 approvals valid" in strip and "Not exported" in strip, strip
    state["requests"]["A-0001"]["reply"] = {"approver": "Someone"}
    state["decisions"] = [{"id": "D-1", "request_id": "A-0001", "decision": "approve", "sku": "B2", "expires_at": soon},
                          {"id": "D-2", "request_id": "A-0001", "decision": "approve", "sku": "C3", "expires_at": past},          # expired
                          {"id": "D-3", "request_id": "A-0001", "decision": "approve", "sku": "D4", "expires_at": soon, "revoked_at": past},
                          {"id": "D-4", "request_id": "A-0002", "decision": "approve", "sku": "a1", "expires_at": soon}]            # the other file's
    (env.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
    strip = strip_of(env, new)
    assert "1 reply recorded" in strip and "1 approval valid, first expires " in strip and "Not exported" in strip, strip
    from datetime import datetime
    local = f"{datetime.fromisoformat(soon).astimezone():%Y-%m-%d %H:%M}"                    # this computer's time, the same as the Approvals table shows
    assert f"first expires {local}" in strip and f'data-label="Expires">{local}</td>' in env.client.get(f"/step/approvals?wf={new}").text, (local, strip)
    (env.project / "approved_orders.json").write_text(json.dumps({"source_results_file": old + ".xlsx", "valid_approvals": 1, "generated_at": "x", "orders": []}))
    assert "Not exported" in strip_of(env, new), "an export made for another results file is not this file's"
    (env.project / "approved_orders.json").write_text(json.dumps({"source_results_file": new + ".xlsx", "valid_approvals": 1, "generated_at": "x", "orders": []}))
    strip = strip_of(env, new)
    assert ">Exported<" in strip and "Not built since the last export" in strip and 'Next step: <a href="/step/order-sheet?wf=' in strip, strip
    # order sheet and listings count only when they are newer than the export
    for folder, name in (("Order Sheets", "order_sheet_20260101_000000.xlsx"), ("Walmart Listings", "walmart_listings_20260101_000000.xlsx")):
        (env.project / folder).mkdir()
        (env.project / folder / name).write_bytes(b"x")
        os.utime(env.project / folder / name, (time.time() - 5000, time.time() - 5000))
    strip = strip_of(env, new)
    assert "Not built since the last export" in strip and "Not generated since the last export" in strip, strip
    for folder, name in (("Order Sheets", "order_sheet_20260303_000000.xlsx"), ("Walmart Listings", "walmart_listings_20260303_000000.xlsx")):
        (env.project / folder / name).write_bytes(b"x")
        os.utime(env.project / folder / name, (time.time() + 5, time.time() + 5))
    (env.project / "Uploads").mkdir()
    (env.project / "Uploads" / "facts_20260303_000000_abcdef.csv").write_text("sku,fact_id,field,value,status,source\n")
    strip = strip_of(env, new, "listing-content")
    for text in ("Built: order_sheet_20260303_000000.xlsx", "Facts file: facts_20260303_000000_abcdef.csv", "Generated: walmart_listings_20260303_000000.xlsx",
                 "The files show every step has produced its output."):
        assert text in strip, (text, strip)
    assert "Next step:" not in strip and 'aria-current="step"' in strip


@test
def results_file_choice_is_validated_server_side(env):
    results_file(env, "sourcing_results_20260101_100000.xlsx", [("A1", "Maker A", "", "https://a.example/1")], mtime=time.time() - 9000)
    results_file(env, "sourcing_results_20260202_100000.xlsx", [("B2", "Maker B", "", "https://b.example/2"), ("C3", "", "", "")])
    (env.project / "elsewhere.xlsx").write_bytes((env.project / "Excel Output Sheets" / "sourcing_results_20260101_100000.xlsx").read_bytes())
    for slug in ("emails", "approvals", "order-sheet", "listing-content"):
        picked = env.client.get(f"/step/{slug}?wf=sourcing_results_20260101_100000").text
        assert "Status for <code>sourcing_results_20260101_100000.xlsx</code>" in picked and "1 product<" in picked
        for bad in ("..\\elsewhere", "../elsewhere", "elsewhere", str(env.project / "elsewhere"), "sourcing_results_20260101_100000.xlsx", "nope", ""):
            page = env.client.get(f"/step/{slug}", params={"wf": bad}).text                 # anything not in the folder: the newest file, never a path
            assert "Status for <code>sourcing_results_20260202_100000.xlsx</code>" in page and "elsewhere" not in page, (slug, bad)
        assert 'name="step" value="' in picked and '<label for="wf">Results file</label>' in picked     # the selector is on every tab
    r = env.client.get("/", params={"wf": "sourcing_results_20260101_100000", "step": "Approvals"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/step/approvals?wf=sourcing_results_20260101_100000"
    assert env.client.get("/", params={"wf": "x", "step": "https://evil.example"}, follow_redirects=False).status_code == 200
    rule = commands.results_file_rule(env.settings)
    assert rule("sourcing_results_20260101_100000.xlsx") == native("Excel Output Sheets", "sourcing_results_20260101_100000.xlsx")
    for bad in ("elsewhere.xlsx", "..\\elsewhere.xlsx", "Excel Output Sheets\\sourcing_results_20260101_100000.xlsx", str(env.project / "elsewhere.xlsx"), ""):
        raises(JobRejected, rule, bad)
    nav = env.client.get("/step/emails").text.split('<nav class="nav"')[1].split("</nav>")[0]
    assert re.findall(r">([^<>]+)</a>", nav) == ["Home", "Search", "Emails", "Approvals", "Order sheet", "Listing content", "Jobs", "Health"], nav


@test
def planted_secrets_never_reach_a_page_or_a_log(env):
    results_file(env, "sourcing_results_20260202_100000.xlsx", [("B2", "Maker B", "b@maker-b.example", "https://b.example/2")])
    (env.project / ".env").write_text("".join(f'{k}="{v}"\n' for k, v in FAKE.items()) + 'PHONE="+1 555 010 0199"\n', encoding="utf-8")
    up = uploaded_value(upload(env))
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        job = env.jobs.wait(env.jobs.start("stub.read", {"mode": "secret"})["id"])
        r = post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count="3")
        sid = r.headers["location"].split("job=")[1]
        env.jobs.wait(sid)
        pages = [p for m, p in all_routes(env) if m == "GET"] + [f"/jobs/{job['id']}", f"/jobs/{job['id']}/log", f"/step/search?job={sid}", f"/jobs/{sid}/log",
                                                                   f"/step/search?file={up}", "/nope", "/step/search", "/step/approvals", "/step/order-sheet",
                                                                   "/step/listing-content"]
        seen = "".join(env.client.get(p).text + json.dumps(dict(env.client.get(p).headers)) for p in pages)
        seen += post(env, "/search/review", choice="file", input_file=up, rows="all").text + upload(env, b"junk", "x.xlsx").text
        seen += "".join(p.read_text(encoding="utf-8", errors="replace") for p in env.settings.jobs.glob("*"))        # every stored log and job record
    for secret in SECRETS + ("sk-ant-abcdefgh12345678",):
        assert secret not in seen, f"{secret!r} leaked"
    health = env.client.get("/health").text
    assert "SHIPPING_ADDRESS" in health and "ANTHROPIC_API_KEY" in health and ">Set<" in health        # names and yes/no only


# ================= the project's own email agent, read-only =================

@test
def importing_the_email_agent_has_no_side_effects(env):
    """The UI loads the project's email_agent.py to read rows with the agent's own rules. Importing it must run nothing:
    no connection, no output, no .env read, no file written, and it must be the PROJECT's copy."""
    (env.project / ".env").write_text("".join(f'{k}="{v}"\n' for k, v in FAKE.items()), encoding="utf-8")
    probe = r'''
import io, json, os, socket, sys
from pathlib import Path
opened, net = [], []
def hook(event, args):
    if event == "open" and isinstance(args[0], (str, bytes, os.PathLike)):
        opened.append((os.path.abspath(os.fsdecode(args[0])), str(args[1])))
    if event.startswith("socket.") and event not in ("socket.__new__", "socket.gethostname"):
        net.append(event)
    if event in ("subprocess.Popen", "os.system", "os.exec", "os.spawn"):
        net.append(event)
sys.addaudithook(hook)
env_before = dict(os.environ)
out, err = io.StringIO(), io.StringIO()
real_out, real_err = sys.stdout, sys.stderr
sys.stdout, sys.stderr = out, err
sys.path.insert(0, sys.argv[2])
from ui import workflows
from ui.config import Settings
project = Path(sys.argv[1])
mod = workflows.agent_module(Settings(project=project, data=project / "ui_data"))
sys.stdout, sys.stderr = real_out, real_err
written = [p for p, mode in opened if set(mode) & set("wax+") and "__pycache__" not in p]
print(json.dumps({"file": mod.__file__, "stdout": out.getvalue(), "stderr": err.getvalue(), "net": net, "written": written,
                  "dotenv": [p for p, _ in opened if os.path.basename(p) == ".env"], "env_same": dict(os.environ) == env_before,
                  "model_client": any(m == "anthropic" or m.startswith("anthropic.") for m in sys.modules),
                  "state_path": mod.STATE_PATH, "has": [n for n in ("load_results", "load_exclusions", "listing_key", "make_draft") if hasattr(mod, n)]}))
'''
    p = subprocess.run([sys.executable, "-B", "-c", probe, str(env.project), str(REAL)], capture_output=True, text=True, timeout=120,
                       cwd=env.project, env={**os.environ, "OBS_ENABLED": "0"})
    assert p.returncode == 0, p.stderr[-800:]
    r = json.loads(p.stdout.strip().splitlines()[-1])
    assert Path(r["file"]) == env.project / "email_agent.py" and Path(r["state_path"]) == env.project / "email_state.json", r
    assert r["stdout"] == "" and r["stderr"] == "", r
    assert r["net"] == [] and r["written"] == [] and r["dotenv"] == [] and r["env_same"] is True, r
    assert r["model_client"] is False, "importing the email agent must not load the model client"
    assert r["has"] == ["load_results", "load_exclusions", "listing_key", "make_draft"]
    assert p.stdout.count("\n") == 1 and not any(s in p.stdout + p.stderr for s in SECRETS)
    assert not (env.project / "email_state.json").exists()
    # and in this process: the module the pages use is the project's, not the one beside the UI's code
    results_file(env, "sourcing_results_20260202_100000.xlsx", [("B2", "Maker B", "b@maker-b.example", "https://b.example/2")])
    assert "1 can be drafted" in env.client.get("/").text
    assert Path(workflows.agent_module(env.settings).__file__) == env.project / "email_agent.py"
    assert "email_agent" not in sys.modules, "the UI must not import the email agent from its own folder"


@test
def the_real_project_folder_is_never_used(env):
    assert REAL not in (env.project, *env.project.parents) and env.settings.project == env.project and env.project.is_relative_to(tempfile.gettempdir())
    for target in (REAL / "README.md", REAL / ".env", REAL / "email_state.json", REAL / "Excel Output Sheets"):       # the tripwire works
        e = raises(AssertionError, (os.listdir if target.suffix == "" and target.name != ".env" else open), target)
        assert "the real project folder was used" in str(e)
    raises(AssertionError, open, REAL / "no_such_folder_for_this_test" / "x.log", "a")                                # (a path that could not be created anyway)
    with open(REAL / "statefile.py", encoding="utf-8") as f:                                                          # code may be read: it is copied into test projects
        assert f.read(3)
    assert len(VIOLATIONS) == 5, VIOLATIONS
    VIOLATIONS.clear()
    results_file(env, "sourcing_results_20260202_100000.xlsx", [("B2", "Maker B", "b@maker-b.example", "https://b.example/2")])
    up = uploaded_value(upload(env))
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        for method, path in all_routes(env):
            if method == "GET":
                env.client.get(path)
        r = post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count="3")
        meta = env.jobs.wait(r.headers["location"].split("job=")[1])
    assert VIOLATIONS == [], VIOLATIONS
    assert Path(meta["cwd"]) == env.project and Path(meta["args"][meta["args"].index("--input") + 1]).parts[0] == "Uploads"
    made = {p.relative_to(env.project).parts[0] for p in env.project.rglob("*") if p.is_file()}
    assert {"Uploads", "Excel Output Sheets", "ui_data"} <= made, made


@test
def tripwire_stays_strong_outside_its_own_cleanup(env):
    """The one allowance (bare names while a test deletes its own temp folder) changes nothing else: it still stops an
    absolute path into the real project, even during that cleanup; a bare name of a real file when the working folder is
    the real project; and rmtree of the real project folder. Deleting is switched off while that last one is tried, so
    even a broken tripwire could not remove anything."""
    global _CLEANING
    e = raises(AssertionError, open, REAL / "README.md")                                    # an absolute path into the real project
    assert "the real project folder was used" in str(e)
    _CLEANING = env.root.resolve()                                                          # ...also in the middle of a cleanup
    try:
        raises(AssertionError, open, REAL / "README.md")
        raises(AssertionError, os.listdir, REAL / "Excel Output Sheets")
    finally:
        _CLEANING = None
    here = os.getcwd()
    os.chdir(REAL)                                                                          # a bare name, outside any cleanup
    try:
        raises(AssertionError, open, "README.md")
        raises(AssertionError, open, os.path.join("no_such_folder_for_this_test", "x.log"), "a")   # (a path that could not be created anyway)
    finally:
        os.chdir(here)
    real_delete = (os.unlink, os.rmdir, os.remove)

    def refuse(*a, **k):
        raise PermissionError("deleting is switched off in this test")
    os.unlink = os.rmdir = os.remove = refuse
    try:
        e = raises(AssertionError, shutil.rmtree, REAL)                                     # stopped at its first look, before any delete
        assert "the real project folder was used" in str(e), e
        assert len(VIOLATIONS) >= 6, VIOLATIONS
        VIOLATIONS.clear()
        e = raises(AssertionError, remove_temp, REAL)                                       # the allowance is only for a test's temp folder
        assert "the real project folder was used" not in str(e) and VIOLATIONS == [], "refused by remove_temp itself, before any look"
    finally:
        os.unlink, os.rmdir, os.remove = real_delete
    assert REAL.is_dir() and (REAL / "tests" / "test_ui_pipeline.py").is_file()


# ================= Phase 2: Emails and Approvals =================

RES = "sourcing_results_20260202_100000"
SELLERS = [("B2", "Maker B", "b@maker-b.example", "https://maker-b.example/item/b2"),
           ("C3", "Maker C", "c@maker-c.example", "https://maker-c.example/item/c3"),
           ("D4", "Maker D", "", "https://maker-d.example/item/d4"),          # no address: inquiry form only
           ("E5", "", "", "")]                                                 # no recommendation


def fake_env(agents=("email_agent.py",)):
    """A project whose commands are recorded instead of run (the offline end-to-end test runs them for real)."""
    FakePopen.calls = []
    env = make_env(popen=FakePopen, agents=agents)
    results_file(env, RES + ".xlsx", SELLERS)
    return env


def last_argv(env):
    """The agent's arguments of the newest recorded command, after `python -X utf8`, with the script as its file name."""
    argv = FakePopen.calls[-1][0]
    assert argv[:3] == [sys.executable, "-X", "utf8"] and Path(argv[3]).parent == env.project, argv
    return [Path(argv[3]).name, *argv[4:]]


def started_job(env, r):
    assert r.status_code == 303 and "job=" in r.headers["location"], (r.status_code, r.text[-700:])
    return env.jobs.wait(r.headers["location"].split("job=")[1])


def refused(r, text="Nothing was", status=400):
    assert r.status_code == status and text in r.text, (r.status_code, text, r.text[-600:])
    return r.text


def open_request(env, rid="A-0001", results=RES + ".xlsx", **more):
    path = env.project / "decision_state.json"
    state = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"next_request": 1, "requests": {}, "decisions": []}
    state["requests"][rid] = {"id": rid, "results_file": results, "reply": None, "created_at": "2026-02-03T12:00:00+00:00", "lines": [{}, {}], **more}
    path.write_text(json.dumps(state), encoding="utf-8")
    (env.project / "Reports").mkdir(exist_ok=True)
    (env.project / "Reports" / f"approval_request_{rid}.txt").write_text(f"REQUEST {rid}\nHOW TO REPLY\n  approve <SKU> qty <whole number> max <max unit price>\n", encoding="utf-8")
    return state


@test
def argv_for_every_pipeline_command(env):
    results_file(env, RES + ".xlsx", SELLERS)
    c, file, rel = env.jobs.commands, RES + ".xlsx", native("Excel Output Sheets", RES + ".xlsx")
    name = env.jobs.save_input("pasted text")
    full = str((env.settings.data / "inputs" / name).resolve())
    assert c["emails.drafts"].argv({"results": file}) == ["drafts", rel]
    assert c["emails.drafts"].argv({"results": file, "near_miss": "1"}) == ["drafts", rel, "--near-miss"]
    assert c["emails.send"].argv({"results": file, "confirm_count": "2"}) == ["send", rel, "--confirm-count", "2"]
    assert c["emails.send"].argv({"results": file, "confirm_count": "3", "test_recipient": "me@example.invalid"}) == \
        ["send", rel, "--confirm-count", "3", "--test-recipient", "me@example.invalid"]
    assert c["emails.mark"].argv({"sku": "C&P9M", "results": file}) == ["mark", "C&P9M", "--results", rel]
    assert c["emails.reply"].argv({"sku": "B2", "results": file, "file": name}) == ["reply", "B2", "--results", rel, "--file", full]
    assert c["approvals.request"].argv({"results": file}) == ["request", rel]
    assert c["approvals.request"].argv({"results": file, "budget": "1500.50"}) == ["request", rel, "--budget", "1500.50"]
    record = {"results": file, "request": "A-0001", "file": name, "confirm_request": "A-0001", "channel": "email", "approver": "Zed Quillfeather"}
    assert c["approvals.record"].argv(record) == ["record", "--results", rel, "--request", "A-0001", "--file", full, "--confirm-request", "A-0001",
                                                  "--channel", "email", "--approver", "Zed Quillfeather"]
    assert c["approvals.revoke"].argv({"sku": "B2", "results": file}) == ["revoke", "B2", "--results", rel]
    assert c["approvals.export"].argv({"results": file}) == ["export", "--results", rel]
    scripts = {k: (v.script, v.fixed[0]) for k, v in c.items() if "." in k and not k.startswith(("help.", "stub."))}
    assert scripts == {"emails.drafts": ("email_agent.py", "drafts"), "emails.send": ("email_agent.py", "send"), "emails.mark": ("email_agent.py", "mark"),
                       "emails.reply": ("email_agent.py", "reply"), "approvals.request": ("decision_agent.py", "request"),
                       "approvals.record": ("decision_agent.py", "record"), "approvals.revoke": ("decision_agent.py", "revoke"),
                       "approvals.export": ("decision_agent.py", "export"), **PHASE3_SCRIPTS}, scripts
    assert all(v.state_changing for k, v in c.items() if k in scripts), "every pipeline command takes the one-job lock"
    flags = {p.flag for k in scripts for p in c[k].params.values()}
    assert not flags & {"--yes", "--allow-listing-facts", "--dry-run", "--out", "--duty", "--inspect-tools"}, flags
    # values a person could type are held to their rules
    for cmd, bad in (("emails.send", {"results": file, "confirm_count": "2", "yes": "1"}), ("emails.send", {"results": file, "confirm_count": "-1"}),
                     ("emails.send", {"results": file, "confirm_count": "2", "test_recipient": "--yes"}), ("emails.send", {"results": file}),
                     ("emails.send", {"results": "..\\x.xlsx", "confirm_count": "2"}), ("emails.mark", {"sku": "--results", "results": file}),
                     ("emails.reply", {"sku": "B2", "results": file, "file": "..\\..\\email_state.json"}),
                     ("emails.reply", {"sku": "B2", "results": file, "file": str(env.project / "sourcing_agent.py")}),
                     ("approvals.request", {"results": file, "budget": "1e9"}), ("approvals.request", {"results": file, "budget": "--x"}),
                     ("approvals.record", {**record, "approver": ""}), ("approvals.record", {**record, "approver": "   "}),
                     ("approvals.record", {**record, "approver": "--channel"}), ("approvals.record", {**record, "approver": "a\nb"}),
                     ("approvals.record", {**record, "channel": ""}), ("approvals.record", {**record, "confirm_request": "A-0002"}),
                     ("approvals.record", {**record, "confirm_request": ""}), ("approvals.record", {k: v for k, v in record.items() if k != "approver"}),
                     ("approvals.record", {**record, "request": "A-1"}), ("approvals.record", {**record, "approver": "x" * 81}),
                     ("approvals.revoke", {"sku": "B2; calc", "results": file}), ("approvals.export", {"results": file, "out": "x"})):
        raises(JobRejected, c[cmd].argv, bad)
    assert c["approvals.record"].recorded(c["approvals.record"].argv(record)).count("(typed by a person, not stored)") == 3


PHASE3_SCRIPTS = {}


@test
def send_needs_the_typed_count(env):
    fake = fake_env()
    try:
        with env_vars(**FAKE):
            page = fake.client.get(f"/step/emails?wf={RES}").text
            assert "Review 2 emails to sellers" in page and "b@maker-b.example" in page and "No address: use the seller" in page
            review = post(fake, "/emails/send/review", wf=RES)
            assert review.status_code == 200 and "2 emails would be sent" in review.text and "This sends real emails to sellers" in review.text
            for who in ("b@maker-b.example", "c@maker-c.example", "Maker B", "Maker C"):                  # the recipients, in full
                assert who in review.text, who
            assert "Maker D" not in review.text.split("<table")[1] and 'name="confirm_count" value' not in review.text
            for typed in ("", " ", "1", "3", "two", "2.0", "-2", "02x"):
                refused(post(fake, "/emails/send", wf=RES, confirm_count=typed), "Nothing was sent")
            refused(post(fake, "/emails/send", wf=RES), "Nothing was sent")                               # the field left out
            assert FakePopen.calls == [] and fake.jobs.list() == [], "a wrong count must start nothing"
            meta = started_job(fake, post(fake, "/emails/send", wf=RES, confirm_count="2"))
            assert last_argv(fake) == ["email_agent.py", "send", native("Excel Output Sheets", RES + ".xlsx"), "--confirm-count", "2"]
            assert meta["command"] == "emails.send" and meta["state_changing"] and meta["step"] == "emails"
            # one seller already marked sent: the count the server works out changes, and yesterday's number is refused
            (fake.project / "email_state.json").write_text(json.dumps({state_key("B2", SELLERS[0][3]): {"sku": "B2", "sent": "2026-02-03 10:00"}}))
            refused(post(fake, "/emails/send", wf=RES, confirm_count="2"), "Nothing was sent")
            started_job(fake, post(fake, "/emails/send", wf=RES, confirm_count="1"))
            assert last_argv(fake)[-2:] == ["--confirm-count", "1"] and len(FakePopen.calls) == 2
            # everything already sent: nothing to send, whatever is typed
            (fake.project / "email_state.json").write_text(json.dumps({state_key(k, u): {"sku": k, "sent": "x"} for k, _, _, u in SELLERS[:2]}))
            for typed in ("0", "2", ""):
                refused(post(fake, "/emails/send", wf=RES, confirm_count=typed), "nothing to send")
            (fake.project / "email_state.json").unlink()
            # another results file, or none: refused
            for wf in ("nope", "", "..\\x", RES + ".xlsx"):
                refused(post(fake, "/emails/send", wf=wf, confirm_count="2"), "Choose a results file")
        # a setting the send needs is missing: refused here, and the agent would refuse too
        for missing in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SHIPPING_ADDRESS"):
            with env_vars(**{**FAKE, missing: None}):
                text = refused(post(fake, "/emails/send", wf=RES, confirm_count="2"), f"{missing} is not set")
                assert "disabled" in fake.client.get(f"/step/emails?wf={RES}").text.split("Review 2 emails")[0].rsplit("<button", 1)[1], missing
                assert not any(s in text for s in SECRETS)
        assert len(FakePopen.calls) == 2
    finally:
        cleanup(fake)


@test
def send_passes_only_the_agents_confirm_flag(env):
    fake = fake_env()
    try:
        with env_vars(**FAKE):
            started_job(fake, post(fake, "/emails/send", wf=RES, confirm_count="2", yes="1", near_miss="1", force="1"))
            argv = FakePopen.calls[-1][0]
            assert argv[4:] == ["send", native("Excel Output Sheets", RES + ".xlsx"), "--confirm-count", "2"], argv        # nothing a form adds reaches the agent
            assert "--yes" not in argv and FakePopen.calls[-1][1]["stdin"] == subprocess.DEVNULL                 # no terminal, no blanket yes
            kw = FakePopen.calls[-1][1]
            assert kw["shell"] is False and Path(kw["cwd"]) == fake.project
            passed = {k for k in kw["env"] if k.startswith(("SMTP_", "SHIPPING"))}
            assert passed == {"SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SHIPPING_ADDRESS"}                      # the agent's own gate reads these itself
        src = (REAL / "ui" / "commands.py").read_text(encoding="utf-8") + (REAL / "ui" / "app.py").read_text(encoding="utf-8")
        assert '"--yes"' not in src and "'--yes'" not in src.replace("--yes, --allow", "")
    finally:
        cleanup(fake)


@test
def test_send_is_marked_and_still_gated(env):
    fake = fake_env()
    try:
        with env_vars(**FAKE):
            review = post(fake, "/emails/send/review", wf=RES, test_recipient="me@example.invalid")
            assert review.status_code == 200 and "This is a TEST" in review.text and "3 emails would be sent" in review.text
            assert review.text.count("<td data-label=\"Goes to\">me@example.invalid</td>") == 3 and "nobody: no address" in review.text
            assert "Send the TEST emails" in review.text and "Send to sellers" not in review.text
            for typed in ("", "2", "4"):
                refused(post(fake, "/emails/send", wf=RES, test_recipient="me@example.invalid", confirm_count=typed), "Nothing was sent")
            for bad in ("b@maker-b.example", "B@Maker-B.example", "-x@example.invalid", "me @example.invalid", "me", "me@x", "a@b.c --yes", "me@example.invalid\n--yes"):
                refused(post(fake, "/emails/send/review", wf=RES, test_recipient=bad))
                refused(post(fake, "/emails/send", wf=RES, test_recipient=bad, confirm_count="3"))
            assert FakePopen.calls == []
            started_job(fake, post(fake, "/emails/send", wf=RES, test_recipient="me@example.invalid", confirm_count="3"))
            assert last_argv(fake) == ["email_agent.py", "send", native("Excel Output Sheets", RES + ".xlsx"), "--confirm-count", "3", "--test-recipient", "me@example.invalid"]
    finally:
        cleanup(fake)


@test
def pages_mask_user_info_in_addresses_but_not_email_addresses(env):
    """A results file's text reaches pages as data (seller names in tables and drop-downs). The user name and password
    inside an address are masked wherever a page prints a value, attributes and the JSON endpoint included; a plain email
    address is shown in full; and forms built from such rows still send the right SKU and count."""
    fake = fake_env()
    planted = "http://pageuser:FAKE-page-pass@seller.example/item"
    try:
        results_file(fake, RES + ".xlsx", [("B2", f"Maker {planted}", "b@maker-b.example", planted + "/b2"),
                                           ("C3", "Maker C", "c@maker-c.example", "https://maker-c.example/item/c3")])
        with env_vars(**FAKE):
            pages = {p: fake.client.get(p).text for p in (f"/step/emails?wf={RES}", f"/step/approvals?wf={RES}", f"/?wf={RES}", "/jobs")}
            review = post(fake, "/emails/send/review", wf=RES).text
            emails = pages[f"/step/emails?wf={RES}"]
            assert '<td data-label="Seller">Maker http://[REDACTED]@seller.example/item</td>' in emails, "the sellers table"
            assert 'B2 (Maker http://[REDACTED]@seller.example/item)</option>' in emails, "a drop-down's text"
            for path, html in list(pages.items()) + [("send review", review)]:
                assert "pageuser" not in html and "FAKE-page-pass" not in html, path
            # attribute values go through the same filter as text: rendered with the app's own template environment
            out = fake.app.state.templates.env.from_string('<input value="{{ v }}" title="{{ v }}"><td data-label="{{ v }}">{{ v }}</td>').render(v=f"Maker {planted}")
            assert out.count("Maker http://[REDACTED]@seller.example/item") == 4 and "pageuser" not in out and "FAKE-page-pass" not in out, out
            # plain email addresses are shown in full where the page lists recipients
            assert "b@maker-b.example" in review and "c@maker-c.example" in review and "2 emails would be sent" in review
            # forms built from such rows send the right SKU and the typed count, unchanged
            assert 'value="B2"' in emails
            started_job(fake, post(fake, "/emails/mark", wf=RES, sku="B2", confirm="1"))
            assert last_argv(fake) == ["email_agent.py", "mark", "B2", "--results", native("Excel Output Sheets", RES + ".xlsx")]
            meta = started_job(fake, post(fake, "/emails/send", wf=RES, confirm_count="2"))
            assert last_argv(fake)[-2:] == ["--confirm-count", "2"]
        # the JSON endpoint and the job page mask too, even a line that reached a stored log without the log masking
        with open(fake.jobs.log_path(meta["id"]), "a", encoding="utf-8") as log:
            log.write(f"an older line: {planted}/raw, contact b@maker-b.example\n")
        body = fake.client.get(f"/jobs/{meta['id']}/log?offset=0").json()
        page = fake.client.get(f"/jobs/{meta['id']}").text
        for text in (json.dumps(body), page):
            assert "pageuser" not in text and "FAKE-page-pass" not in text and "http://[REDACTED]@seller.example/item/raw" in text, text[-400:]
        assert "b@maker-b.example" in body["text"], "an email address in the log text is left as it is by the page filter"
        # downloads are the person's own files, given back exactly as they are
        raw = fake.client.get(f"/files/Excel Output Sheets/{RES}.xlsx")
        assert raw.status_code == 200 and raw.content == (fake.project / "Excel Output Sheets" / f"{RES}.xlsx").read_bytes()
    finally:
        cleanup(fake)


@test
def drafts_mark_and_reply_run_the_agents_own_commands(env):
    fake = fake_env()
    (fake.project / "near_misses.csv").write_text("sku,supplier_or_url,concern,question\n", encoding="utf-8")
    try:
        rel = native("Excel Output Sheets", RES + ".xlsx")
        meta = started_job(fake, post(fake, "/emails/drafts", wf=RES))
        assert last_argv(fake) == ["email_agent.py", "drafts", rel] and meta["state_changing"] is True                # drafts takes the lock too
        page = fake.client.get(f"/step/emails?wf={RES}&job={meta['id']}").text
        assert "fake line one" in page and "Write email drafts" in page and "drafts only" in page
        started_job(fake, post(fake, "/emails/drafts", wf=RES, near_miss="1"))
        assert last_argv(fake) == ["email_agent.py", "drafts", rel, "--near-miss"]
        refused(post(fake, "/emails/drafts", wf=RES, near_miss="yes"))
        # mark: a product from this file's drafts, and a tick
        n = len(FakePopen.calls)
        refused(post(fake, "/emails/mark", wf=RES, sku="D4"), "Tick the box")
        for sku in ("E5", "NOPE", "", "--results", "D4; calc"):
            refused(post(fake, "/emails/mark", wf=RES, sku=sku, confirm="1"), "Choose a product")
        assert len(FakePopen.calls) == n
        started_job(fake, post(fake, "/emails/mark", wf=RES, sku="D4", confirm="1"))
        assert last_argv(fake) == ["email_agent.py", "mark", "D4", "--results", rel]
        # reply: the pasted text is saved by the server, passed with --file, and removed when the command has ended
        text = "Hello, 10 pcs at $2.50 each.\nRegards, <script>alert(1)</script> 测试"
        refused(post(fake, "/emails/reply", wf=RES, sku="B2", reply=text), "ANTHROPIC_API_KEY is not set")
        with env_vars(**FAKE):
            refused(post(fake, "/emails/reply", wf=RES, sku="B2", reply="  \n "), "Paste the seller")
            refused(post(fake, "/emails/reply", wf=RES, sku="NOPE", reply=text), "Choose a product")
            assert len(FakePopen.calls) == n + 1 and not list((fake.settings.data / "inputs").glob("*")) if (fake.settings.data / "inputs").exists() else True
            seen = {}
            real_init = FakePopen.__init__

            def spy(self, argv, **kw):
                seen["text"] = Path(argv[argv.index("--file") + 1]).read_text(encoding="utf-8")
                real_init(self, argv, **kw)
            FakePopen.__init__ = spy
            try:
                started_job(fake, post(fake, "/emails/reply", wf=RES, sku="b2", reply=text))
            finally:
                FakePopen.__init__ = real_init
            argv = last_argv(fake)
            assert argv[:5] == ["email_agent.py", "reply", "b2", "--results", rel] and argv[5] == "--file"
            saved = Path(argv[6])
            assert saved.parent == (fake.settings.data / "inputs").resolve() and re.fullmatch(r"[0-9a-f]{16}\.txt", saved.name)
            assert seen["text"].replace("\r\n", "\n") == text, "the agent reads exactly what was pasted"
            assert not saved.exists() and not list((fake.settings.data / "inputs").glob("*")), "pasted text is removed once the command has ended"
    finally:
        cleanup(fake)


@test
def approvals_needs_a_tick_without_an_email_record(env):
    fake = fake_env(agents=("email_agent.py", "decision_agent.py"))
    try:
        rel = native("Excel Output Sheets", RES + ".xlsx")
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        assert "seller reply status is unavailable" in page and 'name="ack" value="1" required' in page
        text = refused(post(fake, "/approvals/request", wf=RES), "Nothing was written")
        assert "seller reply status is unavailable" in text and FakePopen.calls == []
        refused(post(fake, "/approvals/request", wf=RES, ack="yes"), "Nothing was written")
        started_job(fake, post(fake, "/approvals/request", wf=RES, ack="1"))
        assert last_argv(fake) == ["decision_agent.py", "request", rel]                    # the tick is the UI's warning; the agent's rules are untouched
        started_job(fake, post(fake, "/approvals/request", wf=RES, ack="1", budget=" 1500 "))
        assert last_argv(fake) == ["decision_agent.py", "request", rel, "--budget", "1500"]
        for bad in ("abc", "-5", "1,500", "1e3", "--budget"):
            refused(post(fake, "/approvals/request", wf=RES, ack="1", budget=bad))
        (fake.project / "email_state.json").write_text("{}")
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        assert "seller reply status is unavailable" not in page and 'name="ack"' not in page
        started_job(fake, post(fake, "/approvals/request", wf=RES))                        # with an email record no tick is asked for
        assert len(FakePopen.calls) == 3
    finally:
        cleanup(fake)


TYPED = {"confirm_request": "A-0001", "channel": "ChannelXyzzy", "approver": "Zed Quillfeather"}
REPLY = "approve B2 qty 10 max 1.30\nreject C3\ncap 50"


@test
def approval_fields_are_never_prefilled(env):
    fake = fake_env(agents=("email_agent.py", "decision_agent.py"))
    try:
        open_request(fake)
        (fake.project / "email_state.json").write_text("{}")

        def fields(html):
            form = html.split('action="/approvals/record"')[1].split("</form>")[0]
            return form, {m[1]: m[0] for m in re.findall(r'(<input type="text"[^>]*name="(confirm_request|channel|approver)"[^>]*>)', form)}
        form, inputs = fields(fake.client.get(f"/step/approvals?wf={RES}").text)
        assert set(inputs) == {"confirm_request", "channel", "approver"}
        for name, tag in inputs.items():
            assert "value=" not in tag and "placeholder=" not in tag and 'autocomplete="off"' in tag and "required" in tag, tag
        assert "checked" not in form and "<datalist" not in form and "<select" not in form and "A-0001</textarea>" not in form
        # each typed field missing or blank: refused, nothing started, the pasted reply not kept
        for name in TYPED:
            for value in (None, "", "   "):
                data = {**TYPED, name: value} if value is not None else {k: v for k, v in TYPED.items() if k != name}
                refused(post(fake, "/approvals/record", wf=RES, request_id="A-0001", reply=REPLY, **data))
        refused(post(fake, "/approvals/record", wf=RES, request_id="A-0001", reply="", **TYPED), "Paste the manager")
        assert FakePopen.calls == [] and fake.jobs.list() == []
        assert not (fake.settings.data / "inputs").exists() or not list((fake.settings.data / "inputs").iterdir())
        meta = started_job(fake, post(fake, "/approvals/record", wf=RES, request_id="A-0001", reply=REPLY, **TYPED))
        argv = last_argv(fake)
        assert argv[:5] == ["decision_agent.py", "record", "--results", native("Excel Output Sheets", RES + ".xlsx"), "--request"] and argv[5] == "A-0001"
        assert argv[8:] == ["--confirm-request", "A-0001", "--channel", "ChannelXyzzy", "--approver", "Zed Quillfeather"], argv
        # the UI keeps no copy of what was typed: not in its job record, not in a cookie, not in any later form
        stored = "".join(p.read_text(encoding="utf-8") for p in fake.settings.data.rglob("*") if p.is_file())
        assert "Quillfeather" not in stored and "ChannelXyzzy" not in stored and meta["args"].count("(typed by a person, not stored)") == 3
        assert set(fake.client.cookies.keys()) == {"ui_sid"} and "Quillfeather" not in str(dict(fake.client.cookies))
        # the agent has recorded it (written here as the agent would): the page shows it as a fact, and the form is empty again
        state = open_request(fake, "A-0002")
        state["requests"]["A-0001"]["reply"] = {"text": REPLY, "approver": "Zed Quillfeather", "channel": "ChannelXyzzy"}
        state["decisions"] = [{"id": "A-0001-1", "request_id": "A-0001", "decision": "approve", "sku": "B2", "seller": "Maker B", "qty": 10,
                               "approver": "Zed Quillfeather", "channel": "ChannelXyzzy", "recorded_at": "2026-02-03T10:00:00+00:00",
                               "expires_at": "2099-01-01T00:00:00+00:00", "revoked_at": None}]
        (fake.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        form, inputs = fields(page)
        assert "Zed Quillfeather" in page.split("Approvals recorded")[1].split("</table>")[0]           # shown where the agent's record is listed
        assert "Quillfeather" not in form and "ChannelXyzzy" not in form and "A-0001" not in form       # A-0001 is answered; A-0002 is the open one
        assert all("value=" not in tag for tag in inputs.values()) and "checked" not in form
        for other in ("/step/emails", "/step/search", "/", "/jobs", f"/jobs/{meta['id']}"):
            assert "Quillfeather" not in fake.client.get(other).text and "ChannelXyzzy" not in fake.client.get(other).text, other
    finally:
        cleanup(fake)


@test
def record_confirmation_must_match_the_picked_request(env):
    fake = fake_env(agents=("email_agent.py", "decision_agent.py"))
    try:
        open_request(fake, "A-0001")
        open_request(fake, "A-0002")
        open_request(fake, "A-0003", results="sourcing_results_20250101_000000.xlsx")           # another results file's request
        state = open_request(fake, "A-0004", reply={"approver": "Earlier"})                      # already answered
        results_file(fake, "sourcing_results_20250101_000000.xlsx", SELLERS[:1], mtime=time.time() - 99999)
        for picked, typed in (("A-0001", "A-0002"), ("A-0002", "A-0001"), ("A-0001", "A-1"), ("A-0001", "A-00011"), ("A-0001", "0001"),
                              ("A-0001", "A-0001 A-0002"), ("A-0001", "")):
            text = refused(post(fake, "/approvals/record", wf=RES, request_id=picked, reply=REPLY, **{**TYPED, "confirm_request": typed}))
            assert "Nothing was recorded" in text or "is not a request id" in text, (picked, typed)
        for picked in ("A-0003", "A-0004", "A-9999", "", "A-0001' OR 1=1"):                      # not an open request of THIS results file
            refused(post(fake, "/approvals/record", wf=RES, request_id=picked, reply=REPLY, **{**TYPED, "confirm_request": picked}), "Pick the request")
        assert FakePopen.calls == [] and not list((fake.settings.data / "inputs").glob("*")) if (fake.settings.data / "inputs").exists() else FakePopen.calls == []
        started_job(fake, post(fake, "/approvals/record", wf=RES, request_id="A-0002", reply=REPLY, **{**TYPED, "confirm_request": " a-0002 "}))
        argv = last_argv(fake)
        assert argv[argv.index("--request") + 1] == "A-0002" == argv[argv.index("--confirm-request") + 1]
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        assert "A-0001: written 2026-02-03" in page and "A-0003" not in page and "REQUEST A-0001" in page       # this file's requests, with their text
        assert "A request does not expire; each approval recorded from its reply does, 7 days after" in page
        assert state
    finally:
        cleanup(fake)


@test
def approvals_show_expiry_and_revoke_and_export_are_gated(env):
    fake = fake_env(agents=("email_agent.py", "decision_agent.py"))
    try:
        rel = native("Excel Output Sheets", RES + ".xlsx")
        state = open_request(fake, reply={"approver": "Someone"})
        soon = time.strftime("%Y-%m-%dT%H:%M:%S+00:00", time.gmtime(time.time() + 3 * 86400))
        base = {"request_id": "A-0001", "decision": "approve", "seller": "Maker", "qty": 5, "approver": "Someone", "channel": "email",
                "recorded_at": "2026-02-03T10:00:00+00:00", "revoked_at": None}
        state["decisions"] = [{**base, "id": "A-0001-1", "sku": "B2", "expires_at": soon},
                              {**base, "id": "A-0001-2", "sku": "C3", "expires_at": "2026-02-10T10:00:00+00:00"},
                              {**base, "id": "A-0001-3", "sku": "D4", "expires_at": soon, "revoked_at": "2026-02-04T10:00:00+00:00"},
                              {**base, "id": "A-0001-4", "sku": "E5", "decision": "reject", "expires_at": soon}]
        (fake.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
        table = fake.client.get(f"/step/approvals?wf={RES}").text.split("Approvals recorded")[1].split("</table>")[0]
        rows = dict(re.findall(r'<th scope="row">(\w+)</th><td data-label="Status">(\w+)</td>', table))
        assert rows == {"B2": "Valid", "C3": "Expired", "D4": "Revoked", "E5": "Rejected"}, rows
        assert table.count('data-label="Expires">20') == 4 and "2026-02-10" in table                         # every approval shows its own expiry
        for sku in ("C3", "D4", "E5", "NOPE", "", "b2", "B2 --results x"):                                    # only a currently valid approval, typed exactly
            refused(post(fake, "/approvals/revoke", wf=RES, sku=sku), "Nothing was revoked")
        assert FakePopen.calls == []
        started_job(fake, post(fake, "/approvals/revoke", wf=RES, sku="B2"))
        assert last_argv(fake) == ["decision_agent.py", "revoke", "B2", "--results", rel]
        started_job(fake, post(fake, "/approvals/export", wf=RES))
        assert last_argv(fake) == ["decision_agent.py", "export", "--results", rel]
        # the export is offered for download only by its exact name, and only when it was made from this results file
        assert fake.client.get("/files/project/approved_orders.csv").status_code == 404
        (fake.project / "approved_orders.csv").write_text("NOT AN ORDER\nsku\nB2\n", encoding="utf-8")
        (fake.project / "approved_orders.json").write_text(json.dumps({"source_results_file": "sourcing_results_20250101_000000.xlsx", "valid_approvals": 1,
                                                                        "generated_at": "2026-02-03T10:00:00+00:00", "orders": []}))
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        assert "made from a different results file" in page and "sourcing_results_20250101_000000.xlsx" in page and "Download approved_orders.csv" not in page
        (fake.project / "approved_orders.json").write_text(json.dumps({"source_results_file": RES + ".xlsx", "valid_approvals": 1,
                                                                        "generated_at": "2026-02-03T10:00:00+00:00", "orders": []}))
        page = fake.client.get(f"/step/approvals?wf={RES}").text
        assert 'href="/files/project/approved_orders.csv"' in page and "holds 1 valid approval." in page
        r = fake.client.get("/files/project/approved_orders.csv")
        assert r.status_code == 200 and r.text.startswith("NOT AN ORDER")
        (fake.project / "secret.csv").write_text("secret")
        for name in ("approved_orders.json", "decision_state.json", "email_state.json", ".env", "secret.csv", "sourcing_agent.py", "APPROVED_ORDERS.CSV",
                     "approved_orders.csv%00", "..%2fapproved_orders.csv", "approved_orders.csv/", "Reports", "approved_orders.csv:stream"):
            assert fake.client.get(f"/files/project/{name}").status_code == 404, name
    finally:
        cleanup(fake)


@test
def one_job_lock_covers_every_state_changing_command(env):
    results_file(env, RES + ".xlsx", SELLERS)
    for agent in PHASE3_AGENTS:
        shutil.copy(REAL / agent, env.project)
    export(env)
    facts = facts_name(upload_facts(env))
    state = open_request(env)
    state["decisions"] = [{"id": "A-0001-1", "request_id": "A-0001", "decision": "approve", "sku": "B2", "expires_at": "2099-01-01T00:00:00+00:00"}]
    (env.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
    (env.project / "email_state.json").write_text("{}")
    up = uploaded_value(upload(env))
    with env_vars(**FAKE, UI_STUB_SEARCH="ok"):
        blocker = env.jobs.start("stub.state", {"mode": "slow"})
        attempts = {
            "/search/start": dict(choice="file", input_file=up, rows="all", confirm_count="3"),
            "/emails/drafts": dict(wf=RES), "/emails/send": dict(wf=RES, confirm_count="2"),
            "/emails/mark": dict(wf=RES, sku="D4", confirm="1"), "/emails/reply": dict(wf=RES, sku="B2", reply="hello"),
            "/approvals/request": dict(wf=RES), "/approvals/record": dict(wf=RES, request_id="A-0001", reply=REPLY, **TYPED),
            "/approvals/revoke": dict(wf=RES, sku="B2"), "/approvals/export": dict(wf=RES),
            "/order/build": dict(wf=RES), "/content/template": dict(wf=RES), "/content/generate": dict(wf=RES, facts=facts, confirm_count="2")}
        for path, data in attempts.items():
            r = post(env, path, **data)
            assert r.status_code == 409 and "Nothing was started" in r.text and "Another job that changes data is running" in r.text, (path, r.status_code, r.text[-500:])
        assert [j["id"] for j in env.jobs.list()] == [blocker["id"]], "a refused start must not leave a job behind"
        assert not list((env.settings.data / "inputs").glob("*")), "a refused start must not keep pasted text"
        assert (env.project / "email_state.json").read_text() == "{}" and not (env.project / "Order Sheets").exists()
        posts = {p for m, p in all_routes(env) if m == "POST"} - {"/search/upload", "/search/review", "/emails/send/review", "/jobs/run",
                                                                     "/jobs/20260101-000000-abcdef/cancel", *PHASE3_NOT_JOBS}
        assert posts == set(attempts), f"a route that starts a command is not covered by this test: {posts ^ set(attempts)}"
        env.jobs.cancel(blocker["id"])
        env.jobs.wait(blocker["id"])
    # the lock is in the runner itself, not only in the pages
    hold = env.jobs.start("stub.state", {"mode": "slow"})
    for cmd in [k for k, v in env.jobs.commands.items() if v.state_changing]:
        e = raises((JobBusy, JobRejected), env.jobs.start, cmd, {"mode": "ok"} if cmd.startswith("stub.") else {})
        assert isinstance(e, JobBusy) or cmd != "stub.state"
    raises(JobBusy, env.jobs.start, "emails.drafts", {"results": RES + ".xlsx"})
    raises(JobBusy, env.jobs.start, "approvals.export", {"results": RES + ".xlsx"})
    env.jobs.cancel(hold["id"])


PHASE3_NOT_JOBS = set()


@test
def new_pages_use_neutral_wording_and_escape_what_files_hold(env):
    hostile = "<script>alert(1)</script>"
    fake = fake_env(agents=("email_agent.py", "decision_agent.py"))
    try:
        results_file(fake, RES + ".xlsx", [("B2", f"L-Com {hostile}", "b@maker-b.example", "https://maker-b.example/item/b2"),
                                           ("C3", "LCom reseller", "", "https://maker-c.example/item/c3")])
        state = open_request(fake)
        (fake.project / "Reports" / "approval_request_A-0001.txt").write_text(f"L-Com unit price {hostile}\nlcom_prices.csv\n", encoding="utf-8")
        state["requests"]["A-0001"]["reply"] = {"approver": f"L-Com {hostile}"}
        state["decisions"] = [{"id": "A-0001-1", "request_id": "A-0001", "decision": "approve", "sku": "B2", "seller": f"L-Com {hostile}", "qty": 1,
                               "approver": f"lcom {hostile}", "channel": hostile, "expires_at": "2099-01-01T00:00:00+00:00"}]
        (fake.project / "decision_state.json").write_text(json.dumps(state), encoding="utf-8")
        with env_vars(**FAKE):
            pages = [fake.client.get(f"/step/{slug}?wf={RES}").text for slug in ("emails", "approvals", "order-sheet", "listing-content")]
            pages += [post(fake, "/emails/send/review", wf=RES).text, post(fake, "/emails/send/review", wf=RES, test_recipient="me@example.invalid").text,
                      post(fake, "/emails/send", wf=RES, confirm_count="9").text, post(fake, "/approvals/revoke", wf=RES, sku=hostile).text]
            pages += PHASE3_PAGES(fake) if callable(PHASE3_PAGES) else []
        from ui import wording
        for i, page in enumerate(pages):
            assert hostile not in page, f"page {i}: text from a file was not escaped"
            assert not wording.contains_reference_name(page), (i, re.findall(r"(?i).{40}l[\s_-]?com.{20}", page)[:2])
            assert "&lt;script&gt;" in page or i >= 2
    finally:
        cleanup(fake)


# ================= Phase 3: Order sheet and Listing content =================

PHASE3_SCRIPTS.update({"order.build": ("order_sheet.py", "build"), "content.template": ("content_agent.py", "facts-template"),
                       "content.generate": ("content_agent.py", "generate")})
FACT_HEAD = "sku,fact_id,field,value,status,source\n"
FACTS = (FACT_HEAD + "B2,F1,product_type,Cat6 coupler,sample_inspected,measured on the sample\nB2,F2,color,Black,page_read,seller page\n"
         "B2,F3,material,,,\nC3,F1,product_type,USB adapter,unverified,search keyword\nC3,F2,spec_line,Gold plated,listing_title,listing title\n").encode()
PHASE3_AGENTS = ("email_agent.py", "decision_agent.py", "order_sheet.py", "content_agent.py")


def export(env, source=RES + ".xlsx", skus=("B2", "C3"), expires="2099-06-15T12:00:00Z"):
    orders = [{"sku": k, "seller": f"Maker {k[0]}", "qty": 10, "expires_at": expires if i == 0 else "2020-01-01T00:00:00Z"} for i, k in enumerate(skus)]
    (env.project / "approved_orders.json").write_text(json.dumps({"notice": "NOT AN ORDER", "source_results_file": source, "valid_approvals": len(skus),
                                                                   "generated_at": "2026-02-04T10:00:00+00:00", "orders": orders}), encoding="utf-8")
    (env.project / "approved_orders.csv").write_text("NOT AN ORDER\nsku\n" + "".join(k + "\n" for k in skus), encoding="utf-8")


def upload_facts(env, data=FACTS, name="My facts.csv", wf=RES):
    return post(env, "/content/facts", wf=wf, files={"facts_file": (name, data, "text/csv")})


def facts_name(r) -> str:
    assert r.status_code == 303 and "facts=facts_" in r.headers["location"], (r.status_code, r.text[-500:])
    return r.headers["location"].split("facts=")[1]


def phase3_pages(fake) -> list:
    shutil.copy(REAL / "content_agent.py", fake.project)
    export(fake)
    name = facts_name(upload_facts(fake, FACT_HEAD.encode() + "B2,F1,product_type,L-Com <script>alert(1)</script>,lcom <script>alert(1)</script>,x\n".encode()))
    return [fake.client.get(f"/step/order-sheet?wf={RES}").text, fake.client.get(f"/step/listing-content?wf={RES}&facts={name}").text,
            post(fake, "/content/review", wf=RES, facts=name).text, post(fake, "/content/generate", wf=RES, facts=name, confirm_count="9").text]


PHASE3_PAGES = phase3_pages
PHASE3_NOT_JOBS.update({"/content/facts", "/content/review"})


@test
def argv_for_order_sheet_and_listing_commands(env):
    results_file(env, RES + ".xlsx", SELLERS)
    c, rel = env.jobs.commands, native("Excel Output Sheets", RES + ".xlsx")
    name = uploads.save_facts(env.settings, "facts.csv", FACTS, FACT_HEAD.strip().split(","))
    assert c["order.build"].argv({"results": RES + ".xlsx"}) == ["build", "--approved", "approved_orders.csv", "--results", rel]
    new = "facts_template_20260204_101500_abc123.csv"
    assert c["content.template"].argv({"results": RES + ".xlsx", "facts": new}) == ["facts-template", "--approved", "approved_orders.csv", "--results", rel,
                                                                                    "--facts", native("Uploads", new)]
    assert c["content.generate"].argv({"results": RES + ".xlsx", "facts": name}) == ["generate", "--approved", "approved_orders.csv", "--results", rel,
                                                                                    "--facts", native("Uploads", name)]
    (env.settings.uploads_dir / new).write_text("x")
    for cmd, bad in (("order.build", {"results": RES + ".xlsx", "approved": "other.csv"}), ("order.build", {"results": RES + ".xlsx", "duty": "0"}),
                     ("order.build", {"results": RES + ".xlsx", "out": "C:\\"}), ("order.build", {}),
                     ("content.template", {"results": RES + ".xlsx", "facts": new}),                              # never over an existing file
                     ("content.template", {"results": RES + ".xlsx", "facts": "product_facts.csv"}), ("content.template", {"results": RES + ".xlsx", "facts": "..\\x.csv"}),
                     ("content.template", {"results": RES + ".xlsx"}),
                     ("content.generate", {"results": RES + ".xlsx"}), ("content.generate", {"results": RES + ".xlsx", "facts": "product_facts.csv"}),
                     ("content.generate", {"results": RES + ".xlsx", "facts": new}), ("content.generate", {"results": RES + ".xlsx", "facts": f"Uploads\\{name}"}),
                     ("content.generate", {"results": RES + ".xlsx", "facts": "facts_20260101_000000_ffffff.csv"}),
                     ("content.generate", {"results": RES + ".xlsx", "facts": name, "allow_listing_facts": "1"}),
                     ("content.generate", {"results": RES + ".xlsx", "facts": name, "dry_run": "1"}),
                     ("content.generate", {"results": RES + ".xlsx", "facts": name, "sku": ["B2"]})):
        raises(JobRejected, c[cmd].argv, bad)
    from ui import pipeline
    agent = (REAL / "content_agent.py").read_text(encoding="utf-8")
    assert f"MAX_CALLS_PER_SKU = {pipeline.MAX_MODEL_CALLS_PER_PRODUCT} " in agent, "the page's 'up to N model calls per product' is the content agent's own limit"
    shutil.copy(REAL / "content_agent.py", env.project)
    assert pipeline.facts_columns(env.settings) == FACT_HEAD.strip().split(","), "the facts columns are read from the content agent's own list"


@test
def results_file_match_gate(env):
    """T24: an order sheet or listings are only built from an export made from the SELECTED results file."""
    fake = fake_env(agents=PHASE3_AGENTS)
    try:
        other = "sourcing_results_20250101_000000"
        results_file(fake, other + ".xlsx", SELLERS[:1], mtime=time.time() - 99999)
        with env_vars(**FAKE):
            name = facts_name(upload_facts(fake))
            attempts = {"/order/build": {}, "/content/template": {}, "/content/review": {"facts": name}, "/content/generate": {"facts": name, "confirm_count": "2"}}
            for path, data in attempts.items():                                             # no export at all
                refused(post(fake, path, wf=RES, **data), "There is no export yet")
            export(fake, source=other + ".xlsx")                                           # an export made from ANOTHER results file
            for path, data in attempts.items():
                text = refused(post(fake, path, wf=RES, **data), "Nothing was started")
                assert f"&#39;{other}.xlsx&#39;" in text and f"&#39;{RES}.xlsx&#39;" in text, (path, "the refusal must name both files")
            for slug in ("order-sheet", "listing-content"):
                page = fake.client.get(f"/step/{slug}?wf={RES}").text
                assert "belongs to a different results file" in page and f"<code>{other}.xlsx</code>" in page and f"<code>{RES}.xlsx</code>" in page
                assert "disabled" in page.split("Build the order sheet" if slug == "order-sheet" else "Write a facts template")[0].rsplit("<button", 1)[1]
            for bad in ("", None, 5, RES, RES + ".XLSX", " " + RES + ".xlsx", f"Excel Output Sheets\\{RES}.xlsx", str(fake.project / "Excel Output Sheets" / (RES + ".xlsx"))):
                export(fake, source=bad)                                                    # anything but the exact file name is not a match
                refused(post(fake, "/order/build", wf=RES), "Nothing was started")
            export(fake)
            (fake.project / "approved_orders.csv").unlink()                                 # the .json alone is not an export to build from
            refused(post(fake, "/order/build", wf=RES), "There is no export yet")
            assert FakePopen.calls == [] and fake.jobs.list() == []
            export(fake)                                                                    # made from the selected file: allowed
            started_job(fake, post(fake, "/order/build", wf=RES))
            assert last_argv(fake) == ["order_sheet.py", "build", "--approved", "approved_orders.csv", "--results", native("Excel Output Sheets", RES + ".xlsx")]
            started_job(fake, post(fake, "/content/template", wf=RES))
            assert last_argv(fake)[:6] == ["content_agent.py", "facts-template", "--approved", "approved_orders.csv", "--results", native("Excel Output Sheets", RES + ".xlsx")]
            started_job(fake, post(fake, "/content/generate", wf=RES, facts=name, confirm_count="2"))
            refused(post(fake, "/order/build", wf=other), "Nothing was started")           # and the same export is refused for the other file
            assert len(FakePopen.calls) == 3
            # expiry is shown on both tabs, and never used to hide the agent's own refusal: an expired approval is still passed to the agent
            for slug in ("order-sheet", "listing-content"):
                page = fake.client.get(f"/step/{slug}?wf={RES}").text
                rows = dict(re.findall(r'<th scope="row">(\w+)</th>.*?data-label="Now">([\w ]+)</td>', page, re.S))
                assert rows == {"B2": "Still valid", "C3": "Expired"} and "2099-06-15" in page and "Some approvals have expired since the export" in page, (slug, rows)
            export(fake, expires="2020-01-01T00:00:00Z")                                    # every approval expired: the agent decides and says so
            started_job(fake, post(fake, "/order/build", wf=RES))
            assert len(FakePopen.calls) == 4
    finally:
        cleanup(fake)


@test
def facts_upload_is_validated_and_saved_byte_for_byte(env):
    fake = fake_env(agents=PHASE3_AGENTS)
    try:
        export(fake)
        import hashlib
        cases = {
            "an Excel file": ("facts.xlsx", workbook([FACT_HEAD.strip().split(",")])),
            "an Excel file named .csv": ("facts.csv", workbook([FACT_HEAD.strip().split(",")])),
            "wrong columns": ("facts.csv", b"sku,fact,value\nB2,F1,x\n"), "a column missing": ("facts.csv", b"sku,fact_id,field,value,status\nB2,F1,a,b,c\n"),
            "an extra column": ("facts.csv", FACT_HEAD.strip().encode() + b",verified_by\nB2,F1,a,b,sample_inspected,x,ui\n"),
            "columns in another order": ("facts.csv", b"fact_id,sku,field,value,status,source\n"), "empty": ("facts.csv", b""),
            "not UTF-8": ("facts.csv", FACT_HEAD.encode() + "B2,F1,a,caf\xe9,page_read,x\n".encode("latin-1")),
            "binary": ("facts.csv", FACT_HEAD.encode() + b"B2,F1,a,\x00\x01,page_read,x\n"), "a text file": ("facts.txt", FACTS),
            "too large": ("facts.csv", FACT_HEAD.encode() + b"B2,F1,a,b,page_read,x\n" * 60000),
        }
        for why, (name, data) in cases.items():
            r = upload_facts(fake, data, name)
            assert r.status_code == 400 and "Nothing was uploaded" in r.text, (why, r.status_code)
        assert post(fake, "/content/facts", wf=RES).status_code == 400
        assert not fake.settings.uploads_dir.exists() or not list(fake.settings.uploads_dir.iterdir()), "a refused facts file must leave nothing behind"
        # accepted: saved under a server-built name, exactly as uploaded (BOM, CRLF, odd statuses and all)
        odd = "\ufeff".encode("utf-8") + FACTS.replace(b"\n", b"\r\n") + b"C3,F9,color,Blue,SAMPLE_INSPECTED ,typed oddly\r\nC3,F10,upc,,sample_inspected,\r\n"
        saved = facts_name(upload_facts(fake, odd, "..\\..\\evil facts.csv"))
        path = fake.settings.uploads_dir / saved
        assert commands.FACTS_RE.fullmatch(saved) and path.read_bytes() == odd and not list(fake.root.rglob("*evil*"))
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        page = fake.client.get(f"/step/listing-content?wf={RES}&facts={saved}").text
        counts = dict(re.findall(r'<th scope="row">([^<]+)</th><td data-label="Facts" class="num">(\d+)</td>', page))
        assert counts == {"sample_inspected": "1", "page_read": "1", "unverified": "1", "listing_title": "1", "SAMPLE_INSPECTED": "1"}, counts
        assert "5 filled-in facts for 2 products" in page and "this tool never sets or changes it" in page
        with env_vars(**FAKE):
            review = post(fake, "/content/review", wf=RES, facts=saved)
            assert review.status_code == 200 and "up to 3 model calls per product, so up to 6 calls for this run" in review.text
            started_job(fake, post(fake, "/content/generate", wf=RES, facts=saved, confirm_count="2"))
            assert last_argv(fake)[-2:] == ["--facts", native("Uploads", saved)]
            started_job(fake, post(fake, "/content/template", wf=RES))
        # nothing the UI did changed a byte of it: the UI never sets or upgrades a status
        assert hashlib.sha256(path.read_bytes()).hexdigest() == before and path.read_bytes() == odd
        assert sorted(p.name for p in fake.settings.uploads_dir.iterdir()) == [saved], "the UI wrote no other facts file (the template is the agent's to write)"
        src = "".join((REAL / "ui" / f).read_text(encoding="utf-8") for f in ("app.py", "pipeline.py", "uploads.py", "workflows.py", "commands.py"))
        for status in ("sample_inspected", "page_read"):
            assert f'"{status}"' not in src and f"'{status}'" not in src, f"the UI's code must not hold a fact status to write: {status}"
        assert not (fake.project / "product_facts.csv").exists(), "the project's own product_facts.csv is never created or touched"
    finally:
        cleanup(fake)


@test
def generate_needs_a_facts_file_and_the_typed_product_count(env):
    fake = fake_env(agents=PHASE3_AGENTS)
    try:
        export(fake)
        with env_vars(**FAKE):
            page = fake.client.get(f"/step/listing-content?wf={RES}").text
            assert "No filled facts file has been uploaded yet." in page and "/content/review" not in page                # no facts file: no way to generate
            for facts in ("", "product_facts.csv", "facts_20260101_000000_ffffff.csv", "..\\x.csv"):
                refused(post(fake, "/content/review", wf=RES, facts=facts), "Choose a facts file you uploaded")
                refused(post(fake, "/content/generate", wf=RES, facts=facts, confirm_count="2"), "Choose a facts file you uploaded")
            (fake.project / "product_facts.csv").write_bytes(FACTS)                                                         # a file beside the agent is not an upload
            refused(post(fake, "/content/generate", wf=RES, facts="product_facts.csv", confirm_count="2"), "Choose a facts file you uploaded")
            name = facts_name(upload_facts(fake))
            review = post(fake, "/content/review", wf=RES, facts=name)
            assert "<strong>2 products</strong> in the export: B2, C3." in review.text and 'name="confirm_count" value' not in review.text
            assert "This calls the AI service and spends money." in review.text
            for typed in ("", " ", "1", "3", "two", "2.0", "-2"):
                refused(post(fake, "/content/generate", wf=RES, facts=name, confirm_count=typed), "Nothing was started")
            refused(post(fake, "/content/generate", wf=RES, facts=name), "Nothing was started")
            assert FakePopen.calls == [] and fake.jobs.list() == []
            meta = started_job(fake, post(fake, "/content/generate", wf=RES, facts=name, confirm_count="2"))
            assert last_argv(fake) == ["content_agent.py", "generate", "--approved", "approved_orders.csv", "--results", native("Excel Output Sheets", RES + ".xlsx"),
                                       "--facts", native("Uploads", name)] and meta["command"] == "content.generate"
            export(fake, skus=("B2", "C3", "D4"))                                                                           # the export changed: so does the number
            refused(post(fake, "/content/generate", wf=RES, facts=name, confirm_count="2"), "Nothing was started")
            export(fake, skus=())
            refused(post(fake, "/content/generate", wf=RES, facts=name, confirm_count="0"), "holds no approved product")
            export(fake)
        for missing in ("ANTHROPIC_API_KEY", "BRAND_NAME"):
            with env_vars(**{**FAKE, missing: None}):
                refused(post(fake, "/content/generate", wf=RES, facts=name, confirm_count="2"), f"{missing} is not set")
        assert len(FakePopen.calls) == 1
    finally:
        cleanup(fake)


@test
def downloads_outside_the_folders_are_refused(env):
    results_file(env, RES + ".xlsx", SELLERS)
    s, ok = env.settings, security.safe_output_file
    for folder, name in (("Order Sheets", "order_sheet_20260204_101500.xlsx"), ("Walmart Listings", "walmart_listings_20260204_101500.xlsx"),
                         ("Walmart Listings", "walmart_listings_20260204_101500.md"), ("Uploads", "facts_template_20260204_101500_abc123.csv"),
                         ("Reports", "approval_request_A-0001.txt")):
        (s.project / folder).mkdir(exist_ok=True)
        (s.project / folder / name).write_bytes(b"content")
        assert env.client.get(f"/files/{folder}/{name}").content == b"content", (folder, name)
    (env.root / "outside").mkdir()
    (env.root / "outside" / "secret.csv").write_text("top secret")
    (s.project / "Uploads" / "notes.json").write_text("{}")
    (s.project / "Uploads" / "tool.py").write_text("print(1)")
    (s.project / ".env").write_text('SMTP_PASSWORD="FAKE-smtp-password-555"\n')
    (s.project / "email_state.json").write_text("{}")
    for folder in ("Uploads", "Order Sheets", "Walmart Listings"):
        for rel in ("../email_state.json", "..\\email_state.json", "../.env", "..\\..\\outside\\secret.csv", str(env.root / "outside" / "secret.csv"),
                    "C:\\Windows\\win.ini", "D:\\x.csv", "/etc/passwd", "\\\\server\\share\\x.csv", "notes.json", "tool.py", "missing.csv", "", ".",
                    "x.csv:stream", "facts.csv\x00.txt"):
            assert ok(s, folder, rel) is None, (folder, rel)
    link = s.project / "Uploads" / "link"
    assert link_folder(link, env.root / "outside"), "could not create a link to test an escape through it"
    assert ok(s, "Uploads", "link/secret.csv") is None and env.client.get("/files/Uploads/link/secret.csv").status_code == 404
    for url in ("/files/Uploads/..%2femail_state.json", "/files/Uploads/%2e%2e/.env", "/files/Uploads/notes.json", "/files/project/email_state.json",
                "/files/project/.env", "/files/ui_data/errors.log", "/files/Nope/x.csv", "/files/Uploads/..%5c..%5coutside%5csecret.csv"):
        r = env.client.get(url)
        assert r.status_code == 404 and "top secret" not in r.text and "FAKE-smtp" not in r.text, url
    unlink_folder(link)                                             # removing the link leaves its target alone
    assert (env.root / "outside" / "secret.csv").read_text() == "top secret"


# ================= Phase 4: the whole pipeline, offline, with the real agents =================

SHIM, FIXTURES = REAL / "tests" / "shim", REAL / "tests" / "fixtures"
E2E_FILES = ("sourcing_agent.py", "email_agent.py", "decision_agent.py", "order_sheet.py", "content_agent.py", "observability.py", "statefile.py",
             "email_template.txt")
REAL_STATE = ("email_state.json", "decision_state.json", "approvals_log.jsonl", "approved_orders.csv", "approved_orders.json", "product_facts.csv", ".env")
REAL_FOLDERS = ("Excel Output Sheets", "Reports", "Order Sheets", "Walmart Listings", "Uploads", "ui_data")
PRODUCTS = [["SKU", "Product", "Keyword", "lcom sale price"], ["TST-100", "Shielded Cat6 coupler", "Cat6 RJ45 shielded inline coupler", 10.0],
            ["TST-200", "USB panel coupler", "USB 2.0 panel mount coupler", 12.0], ["TST-300", "Mystery part", "Unfindable widget", 5.0]]
SELLER_ADDRESS = "sales@zyxtek.example"


def shim_vars(out: Path) -> dict:
    """The environment the agents run in: the shim switched on, fake settings only, tracing off."""
    return {**FAKE, "BRAND_NAME": "Acme", "OBS_ENABLED": "0", "LCOM_TEST_SHIM": "1", "LCOM_TEST_SHIM_DIR": str(out), "LCOM_TEST_FIXTURES": str(FIXTURES),
            "LCOM_TEST_FORBID_ROOT": str(REAL), "PYTHONPATH": str(SHIM), "ORDER_SHEET_RECALC_PY": None, "LANGFUSE_SECRET_KEY": None,
            "LANGFUSE_PUBLIC_KEY": None}


def run_like_the_ui(script: Path, cwd: Path, env: dict):
    """A subprocess started exactly as ui\\jobs.py starts an agent."""
    full = {**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUNBUFFERED": "1"}
    for k, v in env.items():
        full.pop(k, None) if v is None else full.__setitem__(k, v)
    how = {"creationflags": subprocess.CREATE_NO_WINDOW | subprocess.CREATE_NEW_PROCESS_GROUP} if os.name == "nt" else {"creationflags": 0, "start_new_session": True}
    p = subprocess.Popen([sys.executable, "-X", "utf8", str(script)], cwd=cwd, env=full, shell=False, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                         stderr=subprocess.STDOUT, encoding="utf-8", errors="replace", **how)
    text, _ = p.communicate(timeout=120)
    return p.returncode, text, p.pid


@test
def shim_self_test(env):
    out = env.root / "shim"
    out.mkdir()
    code, text, pid = run_like_the_ui(SHIM / "selftest.py", env.project, shim_vars(out))
    assert code == 0, text[-1500:]
    r = json.loads(text.strip().splitlines()[-1])
    assert r["sitecustomize"] is True and r["pid"] == pid and (out / f"shim_{pid}.ok").exists(), "the shim was not loaded by a process started like an agent"
    assert r["leaked"] == [], r["leaked"]
    assert r["blocked"] == ["socket.connect", "socket.connect_ex", "socket.sendto", "socket.create_connection", "socket.getaddrinfo",
                            "asyncio.open_connection ip", "asyncio.open_connection name", "loop.sock_connect", "httpx.AsyncClient"], r["blocked"]
    # the event loop differs by system (Windows connects through its own call, which the shim guards separately); the list
    # of blocked routes above is the same on both
    assert r["loop"] == ("ProactorEventLoop" if os.name == "nt" else "_UnixSelectorEventLoop") and r["loopback_ok"] is True
    assert r["tool_runner"] == {"turns": 1, "final": "Cat6 RJ45 shielded inline coupler"} and r["sync_client"] == "pricing_provided"
    assert r["nimble"][:2] == ["nimble_search", "nimble_extract"] and "Test listing page" in r["nimble"][2]
    assert json.loads((out / "smtp_sink.jsonl").read_text().splitlines()[0])["to"] == "test@example.invalid"
    assert (out / "net.log").read_text().count("\n") >= 9 and r["real_root"] == ["refused"] * 4 and (out / "realroot.log").read_text().count("\n") == 4
    # switched off, the same folder on PYTHONPATH changes nothing
    off = env.root / "off"
    off.mkdir()
    check = env.root / "check.py"
    check.write_text("import smtplib, socket, anthropic\nprint('untouched' if anthropic.AsyncAnthropic.__module__.startswith('anthropic') and "
                     "smtplib.SMTP.__module__ == 'smtplib' and socket.create_connection.__module__ == 'socket' else 'patched')\n")
    code, text, _ = run_like_the_ui(check, env.project, {**shim_vars(off), "LCOM_TEST_SHIM": None})
    assert code == 0 and text.strip() == "untouched" and not list(off.iterdir()), text[-300:]
    code, text, _ = run_like_the_ui(check, env.project, shim_vars(off))
    assert code == 0 and text.strip() == "patched"
    # fails closed: a shim that cannot install ends the process before the agent's first line
    code, text, _ = run_like_the_ui(check, env.project, {**shim_vars(off), "LCOM_TEST_SHIM_DIR": str(env.root / "missing" / "folder")})
    assert code == 97 and "untouched" not in text and "patched" not in text and "could not be installed" in text, (code, text[-300:])
    code, text, _ = run_like_the_ui(check, env.project, {**shim_vars(off), "LCOM_TEST_SHIM_DIR": None})
    assert code == 97 and "patched" not in text, (code, text[-300:])


def snapshot_real() -> dict:
    """Hashes of the real project's state files and a listing of its output folders (read-only), to compare before and after."""
    import hashlib
    global _SNAPSHOT
    _SNAPSHOT = True
    try:
        snap = {name: (hashlib.sha256((REAL / name).read_bytes()).hexdigest() if (REAL / name).is_file() else None) for name in REAL_STATE}
        for folder in REAL_FOLDERS:
            root = REAL / folder
            snap[folder] = sorted((str(p.relative_to(root)), p.stat().st_size, p.stat().st_mtime_ns) for p in root.rglob("*") if p.is_file()) if root.is_dir() else None
        snap["top"] = sorted(p.name for p in REAL.iterdir())
        return snap
    finally:
        _SNAPSHOT = False


@test
def end_to_end_offline(env):
    """Upload -> search -> drafts -> send -> reply -> request -> record -> export -> order sheet -> listings, through the
    UI's own routes, with the real agents as subprocesses behind the shim, in a throwaway project folder."""
    before = snapshot_real()
    cleanup(env)
    env = make_env(stub=False, agents=E2E_FILES)
    out = env.root / "shim"
    out.mkdir()
    jobs_run = []
    try:
        assert env.settings.project == env.project != REAL and not env.project.is_relative_to(REAL), "the UI under test must not use the real project folder"
        assert not any((d / ".env").exists() for d in (env.project, *env.project.parents)), "a .env above the test project could be read by an agent"
        assert all(Path(c.script).name == c.script for c in env.jobs.commands.values()), "commands run the agents of the project folder they are given"

        def done(r, expect="succeeded"):
            assert r.status_code == 303 and "job=" in r.headers["location"], (r.status_code, r.text[-900:])
            meta = env.jobs.wait(r.headers["location"].split("job=")[1], timeout=180)
            log = env.jobs.read_log(meta["id"])[0]
            assert meta["status"] == expect, (meta["command"], meta["status"], meta["exit_code"], log[-1500:])
            assert (out / f"shim_{meta['pid']}.ok").exists(), f"{meta['command']} ran without the test shim"
            assert Path(meta["cwd"]) == env.project and Path(meta["script"]).parent == env.project, meta
            jobs_run.append(meta)
            return meta, log

        with env_vars(**shim_vars(out)):
            # 1. upload and search
            up = uploaded_value(upload(env, workbook(PRODUCTS)))
            review = post(env, "/search/review", choice="file", input_file=up, rows="all")
            assert "About 3 products, roughly $0.24 and 3 minutes." in review.text
            refused(post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count="2"), "Nothing was started")
            assert env.jobs.list() == []
            (out / "hold").write_text("x")                                               # keeps the search running until removed
            r = post(env, "/search/start", choice="file", input_file=up, rows="all", confirm_count="3")
            assert r.status_code == 303, r.text[-600:]
            sid = r.headers["location"].split("job=")[1]
            try:
                wait_for(lambda: (out / "nimble.log").exists(), "the search to reach its first page fetch")
                assert env.jobs.get(sid)["status"] == "running"
                second = post(env, "/search/start", choice="file", input_file=up, rows="first", limit="1", confirm_count="1")
                assert second.status_code == 409 and "Another job that changes data is running" in second.text      # a second job mid-run is refused
                raises(JobBusy, env.jobs.start, "search", {"input": up, "output": search.new_output_name()})          # by the runner itself too
                assert len(env.jobs.list()) == 1
            finally:
                (out / "hold").unlink()
            search_meta, log = done(r)
            results = search_meta["search"]["output"]
            wf = results[:-5]
            path = env.project / "Excel Output Sheets" / results
            assert search_meta["args"][:2] == ["--input", native("Uploads", up[8:])] and path.is_file() and "Report written to:" in log
            sheet = openpyxl.load_workbook(path)["Results"]
            head = [c.value for c in sheet[1]]
            rows = {r[0]: dict(zip(head, r)) for r in sheet.iter_rows(min_row=2, values_only=True)}
            assert list(rows) == ["TST-100", "TST-200", "TST-300"], list(rows)
            assert rows["TST-100"]["Recommended Manufacturer"] == "Zyxtek Limited" and rows["TST-100"]["Recommended Email"] == SELLER_ADDRESS
            assert rows["TST-200"]["Recommended Manufacturer"] == "Qorvane Hardware" and not rows["TST-300"]["Recommended Manufacturer"]
            rel = native("Excel Output Sheets", results)

            # 2. emails: drafts, a refused and an accepted send, a test send, a form mark, a reply
            meta, log = done(post(env, "/emails/drafts", wf=wf))
            assert meta["args"] == ["drafts", rel] and "2 draft(s); 1 with an email address. Nothing sent." in log
            assert not (out / "smtp_sink.jsonl").exists() and not (env.project / "email_state.json").exists()
            page = env.client.get(f"/step/emails?wf={wf}").text
            assert "Review 1 email to sellers" in page and SELLER_ADDRESS in page
            for typed in ("2", "0", ""):
                refused(post(env, "/emails/send", wf=wf, confirm_count=typed), "Nothing was sent")
            assert not (out / "smtp_sink.jsonl").exists(), "a refused send reached the mail server"
            meta, log = done(post(env, "/emails/send", wf=wf, confirm_count="1"))
            assert meta["args"] == ["send", rel, "--confirm-count", "1"] and "Sent 1 email(s)." in log
            sink = [json.loads(line) for line in (out / "smtp_sink.jsonl").read_text(encoding="utf-8").splitlines()]
            assert [m["to"] for m in sink] == [SELLER_ADDRESS] and sink[0]["host"] == "smtp.example.invalid" and "1 Test Street" in sink[0]["body"]
            state = json.loads((env.project / "email_state.json").read_text(encoding="utf-8"))
            key = "TST-100|www.alibaba.com/product-detail/zyxtek-coupler_111.html"
            assert list(state) == [key] and "sent" in state[key]
            refused(post(env, "/emails/send", wf=wf, confirm_count="1"), "nothing to send")              # already sent: not sent twice
            meta, log = done(post(env, "/emails/send", wf=wf, test_recipient="me@example.invalid", confirm_count="2"))
            sink = [json.loads(line) for line in (out / "smtp_sink.jsonl").read_text(encoding="utf-8").splitlines()]
            assert [m["to"] for m in sink] == [SELLER_ADDRESS, "me@example.invalid", "me@example.invalid"] and all(m["subject"].startswith("[TEST] ") for m in sink[1:])
            assert json.loads((env.project / "email_state.json").read_text(encoding="utf-8")) == state, "a test send must not be recorded as sent"
            meta, log = done(post(env, "/emails/mark", wf=wf, sku="TST-200", confirm="1"))
            assert "marked as submitted via the inquiry form" in log
            meta, log = done(post(env, "/emails/reply", wf=wf, sku="TST-100", reply=(FIXTURES / "seller_reply.txt").read_text(encoding="utf-8")))
            state = json.loads((env.project / "email_state.json").read_text(encoding="utf-8"))
            assert state[key]["status"] == "pricing_provided" and "10 pcs" in state[key]["reply"] and len(state) == 2, state
            assert "pricing_provided" in log and not list((env.settings.data / "inputs").glob("*"))
            strip = strip_of(env, wf)
            assert "2 can be drafted" in strip and "2 marked sent" in strip and "1 reply recorded" in strip, strip

            # 3. approvals: request, a refused and an accepted record, export
            assert "seller reply status is unavailable" not in env.client.get(f"/step/approvals?wf={wf}").text
            meta, log = done(post(env, "/approvals/request", wf=wf, budget="100"))
            assert meta["args"] == ["request", rel, "--budget", "100"]
            request_text = (env.project / "Reports" / "approval_request_A-0001.txt").read_text(encoding="utf-8")
            assert "TST-100" in request_text and "A-0001" in request_text
            dstate = json.loads((env.project / "decision_state.json").read_text(encoding="utf-8"))
            assert dstate["requests"]["A-0001"]["results_file"] == results and dstate["requests"]["A-0001"]["reply"] is None
            page = env.client.get(f"/step/approvals?wf={wf}").text
            assert "A-0001: written " in page and "waiting for a reply" in page and "HOW TO REPLY" in page
            reply = (FIXTURES / "manager_reply.txt").read_text(encoding="utf-8")
            frozen = (env.project / "decision_state.json").read_bytes()
            typed = {"confirm_request": "A-0001", "channel": "test channel", "approver": "Test Approver"}
            for missing in ("approver", "channel", "confirm_request"):
                refused(post(env, "/approvals/record", wf=wf, request_id="A-0001", reply=reply, **{k: v for k, v in typed.items() if k != missing}))
            refused(post(env, "/approvals/record", wf=wf, request_id="A-0001", reply=reply, **{**typed, "confirm_request": "A-0002"}), "Nothing was recorded")
            assert (env.project / "decision_state.json").read_bytes() == frozen, "a refused record changed the approvals"
            meta, log = done(post(env, "/approvals/record", wf=wf, request_id="A-0001", reply=reply, **typed))
            assert "Recorded. Nothing has been ordered." in log and "Test Approver" not in json.dumps(meta)
            dstate = json.loads((env.project / "decision_state.json").read_text(encoding="utf-8"))
            decisions = {d["sku"]: d for d in dstate["decisions"]}
            assert decisions["TST-100"]["decision"] == "approve" and decisions["TST-100"]["approver"] == "Test Approver" and decisions["TST-100"]["qty"] == 10
            assert decisions["TST-200"]["decision"] == "reject" and decisions["TST-100"]["channel"] == "test channel"
            assert "1 approval valid, first expires " in strip_of(env, wf, "approvals")
            meta, log = done(post(env, "/approvals/export", wf=wf))
            assert meta["args"] == ["export", "--results", rel] and "Exported 1 valid approval(s)" in log
            exported = json.loads((env.project / "approved_orders.json").read_text(encoding="utf-8"))
            assert exported["source_results_file"] == results and [o["sku"] for o in exported["orders"]] == ["TST-100"]
            assert exported["orders"][0]["listing_url"] == rows["TST-100"]["Recommended URL"] and exported["orders"][0]["approved_by"] == "Test Approver"
            csv_bytes = env.client.get("/files/project/approved_orders.csv").content
            assert b"TST-100" in csv_bytes and b"NOT AN ORDER" in csv_bytes and csv_bytes == (env.project / "approved_orders.csv").read_bytes()

            # 4. order sheet, from that export and that results file
            meta, log = done(post(env, "/order/build", wf=wf))
            assert meta["args"] == ["build", "--approved", "approved_orders.csv", "--results", rel]
            sheets = list((env.project / "Order Sheets").glob("order_sheet_*.xlsx"))
            assert len(sheets) == 1 and sheets[0].stat().st_mtime >= (env.project / "approved_orders.csv").stat().st_mtime
            assert "TST-100" in str([[c.value for c in row] for ws in openpyxl.load_workbook(sheets[0]).worksheets for row in ws.iter_rows()])
            assert env.client.get(f"/files/Order Sheets/{sheets[0].name}").status_code == 200

            # 5. listing content: only with a facts file a person filled in
            meta, log = done(post(env, "/content/template", wf=wf))
            template, = list(env.settings.uploads_dir.glob("facts_template_*.csv"))
            assert meta["args"][-2:] == ["--facts", native("Uploads", template.name)] and template.read_text(encoding="utf-8-sig").startswith("sku,fact_id,field,value,status,source")
            assert "TST-100" in template.read_text(encoding="utf-8-sig") and env.client.get(f"/files/Uploads/{template.name}").status_code == 200
            refused(post(env, "/content/generate", wf=wf, facts="", confirm_count="1"), "Choose a facts file you uploaded")
            refused(post(env, "/content/generate", wf=wf, facts=template.name, confirm_count="1"), "Choose a facts file you uploaded")   # the unfilled template is not one
            assert not (env.project / "Walmart Listings").exists()
            filled = (FIXTURES / "facts_filled.csv").read_bytes()
            facts = facts_name(upload_facts(env, filled, "filled.csv", wf=wf))
            refused(post(env, "/content/generate", wf=wf, facts=facts, confirm_count="2"), "Nothing was started")
            meta, log = done(post(env, "/content/generate", wf=wf, facts=facts, confirm_count="1"))
            assert meta["args"] == ["generate", "--approved", "approved_orders.csv", "--results", rel, "--facts", native("Uploads", facts)]
            listings = sorted(p.name for p in (env.project / "Walmart Listings").iterdir())
            assert len(listings) == 2 and listings[0].endswith(".md") and listings[1].endswith(".xlsx") and listings[0][:-3] == listings[1][:-5], listings
            assert "TST-100: READY FOR REVIEW" in log, log[-1200:]
            assert (env.settings.uploads_dir / facts).read_bytes() == filled, "the facts file was changed"
            for name in listings:
                assert env.client.get(f"/files/Walmart Listings/{name}").status_code == 200
            strip = strip_of(env, wf, "listing-content")
            assert "The files show every step has produced its output." in strip and "Generated: walmart_listings_" in strip and "Built: order_sheet_" in strip, strip

            # every step's input was the previous step's real output, nothing left this computer, and nothing was started twice
            assert [m["command"] for m in jobs_run] == ["search", "emails.drafts", "emails.send", "emails.send", "emails.mark", "emails.reply", "approvals.request",
                                                        "approvals.record", "approvals.export", "order.build", "content.template", "content.generate"]
            assert all(rel in m["args"] for m in jobs_run[1:]) and len(env.jobs.list()) == len(jobs_run)
            assert not (out / "net.log").exists(), (out / "net.log").read_text()
            assert not (out / "realroot.log").exists(), (out / "realroot.log").read_text()
            model = (out / "model.log").read_text().split()
            assert model.count("search") == 3 and model.count("reply") == 1 and model.count("listing") == 1 and "UNRECOGNISED" not in model, model
            seen = "".join(p.read_text(encoding="utf-8", errors="replace") for p in env.settings.data.rglob("*") if p.is_file())
            seen += "".join(env.client.get(f"/step/{slug}?wf={wf}").text for slug in ("emails", "approvals", "order-sheet", "listing-content")) + env.client.get("/jobs").text
            for secret in SECRETS:
                assert secret not in seen, f"{secret!r} reached a page or a stored log"
            assert SELLER_ADDRESS not in "".join(p.read_text(encoding="utf-8", errors="replace") for p in env.settings.jobs.glob("*.log")), "addresses are masked in logs"
        assert VIOLATIONS == [], VIOLATIONS
        assert snapshot_real() == before, "the real project's state or output folders changed during the test"
    finally:
        (out / "hold").unlink(missing_ok=True)
        cleanup(env)


def wait_for(cond, what, timeout=60):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if cond():
            return
        time.sleep(0.1)
    raise AssertionError(f"timed out waiting for {what}")


def main():
    failed, skipped, wanted = [], [], sys.argv[1:]            # optional: names (or parts of names) of the tests to run
    chosen = [f for f in TESTS if not wanted or any(w in f.__name__ for w in wanted)]
    for f in chosen:
        env = make_env()
        try:
            f(env)
            assert VIOLATIONS == [], VIOLATIONS
            print(f"ok   {f.__name__}")
        except Skip as e:
            skipped.append(f.__name__)
            print(f"skip {f.__name__}: {e}")
        except Exception:  # noqa: BLE001
            failed.append(f.__name__)
            print(f"FAIL {f.__name__}\n{traceback.format_exc()}")
            VIOLATIONS.clear()
        finally:
            cleanup(env)
    print(f"\n{len(chosen) - len(failed) - len(skipped)}/{len(chosen)} passed, {len(skipped)} skipped" + (f": {', '.join(skipped)}" if skipped else "")
          + (f" (of {len(TESTS)} tests; others not run)" if wanted else ""))
    if failed:
        sys.exit(1)
    print("ui pipeline tests passed")


if __name__ == "__main__":
    main()
