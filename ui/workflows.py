"""Workflow overview, read-only. A workflow is one results file in 'Excel Output Sheets'. URLs use the plain file stem; the
internal id (observability.workflow_id, the agents' own function) is only mapped back on the server. Step statuses only say
what the files show."""
import importlib.util
import json
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import observability as obs

from .config import UPLOADS_FOLDER, Settings

STEPS = ("Search", "Emails", "Approvals", "Order sheet", "Listing content")
STEP_TEXT = {
    "Search": "Finds alternative suppliers for each product and compares their prices with the reference price.",
    "Emails": "Writes sample-order emails to the suppliers the search recommended, for a person to check and send.",
    "Approvals": "Turns the recommendations into a request for a manager and records the manager's reply.",
    "Order sheet": "Builds the Excel sheet a person uses to place the approved orders by hand.",
    "Listing content": "Drafts the product listing text for a person to review and upload by hand.",
}
# Badge vocabulary. "Files found" means files exist, nothing more: it does not mean they are correct or approved.
NOT_STARTED, FOUND, BLOCKED = "Not started", "Files found", "Blocked"


@dataclass
class Step:
    name: str
    status: str
    detail: str = ""
    blocked: str = ""
    items: list = field(default_factory=list)   # the status strip: [(text, True / False / None)]; None = a plain fact
    done: bool = False                           # the files show this step's main output exists (decides the "Next step" link)

    @property
    def slug(self) -> str:
        return self.name.lower().replace(" ", "-")


_agents: dict = {}


def agent_module(s: Settings, name: str = "email_agent"):
    """The project's OWN agent module, loaded from the project folder by path (never from wherever the UI's code lives),
    so the UI reads rows with the agent's rules instead of a copy of them. Used read-only: only functions that read files
    are called. Importing an agent runs no command, opens no connection and prints nothing (a test holds this)."""
    path = s.project / f"{name}.py"
    key = (str(path), path.stat().st_mtime)
    if key not in _agents:
        spec = importlib.util.spec_from_file_location(f"_ui_agent_{name}_{len(_agents)}", path)
        mod = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = mod
        spec.loader.exec_module(mod)
        _agents[key] = mod
    return _agents[key]


def internal_id(stem: str) -> str:
    """The agents' workflow id for a results file with this stem."""
    return obs.workflow_id(stem + ".xlsx")


_cache: dict = {}      # (path, mtime) -> value, so the pages stay fast


def _cached(key, make):
    if key not in _cache:
        _cache[key] = make()
    return _cache[key]


def _count_rows(path: Path) -> tuple:
    """(products, rows carrying a recommended URL), read-only."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Results"] if "Results" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        head = [str(c or "") for c in next(rows)]
        sku, url = head.index("SKU"), head.index("Recommended URL")
        n = rec = 0
        for r in rows:
            if r[sku]:
                n += 1
                rec += bool(r[url])
        return n, rec
    finally:
        wb.close()


def products_in(path: Path):
    try:
        return _cached(("n", str(path), path.stat().st_mtime), lambda: _count_rows(path)[0])
    except Exception:  # noqa: BLE001
        return None


def friendly_label(modified: datetime, products) -> str:
    hour = modified.hour % 12 or 12
    when = f"{modified:%b} {modified.day}, {hour}:{modified:%M} {'AM' if modified.hour < 12 else 'PM'}"
    return when + (f", {products} product{'s' if products != 1 else ''}" if products is not None else "")


def list_workflows(s: Settings) -> list:
    """Newest first: [{stem, id (internal), file, modified, label}]."""
    if not s.results_dir.is_dir():
        return []
    files = sorted((p for p in s.results_dir.glob("*.xlsx") if not p.name.startswith("~$")),
                   key=lambda p: p.stat().st_mtime, reverse=True)
    out, seen = [], {}
    for p in files:
        when = datetime.fromtimestamp(p.stat().st_mtime)
        label = friendly_label(when, products_in(p))
        seen[label] = seen.get(label, 0) + 1
        out.append({"stem": p.stem, "id": internal_id(p.stem), "file": p.name, "modified": when.strftime("%Y-%m-%d %H:%M"),
                    "label": label if seen[label] == 1 else f"{label} ({seen[label]})"})
    return out


def _json(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8")), ""
    except FileNotFoundError:
        return None, ""
    except (OSError, ValueError):
        return None, f"{path.name} could not be read"


def draft_counts(s: Settings, path: Path):
    """What the Emails step can draft for this file, from the email agent's OWN row rules (its load_results, read-only), not
    a copy of them. None if that cannot be run: the page then shows only what the file says."""
    rows = draft_rows(s, path)
    if rows is None:
        return None
    kinds = {k: sum(x["kind"] == k for x in rows["skipped"]) for k in ("none", "reviewer", "error")}
    return {"draftable": len(rows["drafts"]), "no_seller": kinds["none"], "excluded": kinds["reviewer"], "errors": kinds["error"]}


def draft_rows(s: Settings, path: Path):
    """{'drafts': the rows the email agent would draft (each with its state 'key'), 'skipped': the rest with a reason},
    from the email agent's OWN load_results, read-only. None if that cannot be run."""
    def make():
        try:
            ea = agent_module(s)
            drafts, skipped = ea.load_results(str(path), ea.load_exclusions(str(s.project / "reviewer_exclusions.csv")))
            return {"drafts": drafts, "skipped": skipped}
        except Exception:  # noqa: BLE001
            return None
    return _cached(("d", str(path), path.stat().st_mtime, _mtime(s.project / "reviewer_exclusions.csv"), _mtime(s.project / "email_agent.py")), make)


def _mtime(p: Path):
    try:
        return p.stat().st_mtime
    except OSError:
        return 0


def search_detail(s: Settings, path: Path) -> tuple:
    """(products, text). Counts come from the file and, for what can be drafted, from the email agent's own rules."""
    n, rec = _cached(("c", str(path), path.stat().st_mtime), lambda: _count_rows(path))
    d = draft_counts(s, path)
    plural = f"{n} product{'s' if n != 1 else ''}"
    if d is None:
        return n, (f"{plural} in the results file. {rec} row{'s' if rec != 1 else ''} carry a recommendation; the Emails step shows how "
                   "many can be drafted. Recommendations are the search's reading of seller pages and are unverified until a person checks them.")
    parts = [f"{d['draftable']} can be drafted"]
    if d["no_seller"]:
        parts.append(f"{d['no_seller']} have no recommended seller")
    if d["excluded"]:
        parts.append(f"{d['excluded']} are excluded by the reviewer's exclusion list")
    if d["errors"]:
        parts.append(f"{d['errors']} could not be searched")
    return n, (f"{plural} in the results file: " + "; ".join(parts) + ". Recommendations are unverified until a person checks them.")


def _money(v) -> str:
    if not isinstance(v, (int, float)):
        return str(v or "")
    return f"${v:,.2f}" if v >= 1 else f"${v:.4f}".rstrip("0").rstrip(".")


def _read_result_rows(s: Settings, path: Path) -> dict:
    import openpyxl
    from urllib.parse import urlsplit
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Results"] if "Results" in wb.sheetnames else wb[wb.sheetnames[0]]
        raw = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    head = [str(c or "") for c in raw[0]]
    found = draft_rows(s, path)             # the email agent's OWN row rules decide the outcome, as on Home
    # ponytail: keyed by SKU; a file listing one SKU twice shows one outcome for both
    kinds = {x["sku"]: x["kind"] for x in found["skipped"]} if found else None
    rows = []
    for values in raw[1:]:
        r = dict(zip(head, values))
        if not r.get("SKU"):
            continue
        sku, url, note = str(r["SKU"]), str(r.get("Recommended URL") or ""), str(r.get("Recommendation") or "")
        seller = str(r.get("Recommended Manufacturer") or "")
        if kinds is not None:
            outcome = {"error": "error", "none": "none", "reviewer": "excluded"}.get(kinds.get(sku), "recommended")
        else:                               # the file alone: the reviewer's exclusion list cannot be applied
            outcome = "error" if note.startswith("Error researching") else "recommended" if (seller or url) else "none"
        match = next((r.get(f"Manufacturer {n} Accuracy") for n in range(1, 10) if url and r.get(f"Manufacturer {n} URL") == url), None)
        host = urlsplit(url).netloc.lower()
        rows.append({"sku": sku, "description": str(r.get("Keyword") or r.get("Product") or ""), "outcome": outcome,
                     "supplier": seller, "platform": host[4:] if host.startswith("www.") else host,
                     "unit_price": _money(r.get("Recommended Unit Price")), "match": f"{match}%" if isinstance(match, (int, float)) else "",
                     "caveat": str(r.get("Ordering Note") or ""), "why": note if outcome in ("none", "error") else "",
                     "ref_price": _money(r.get("L-Com Unit Price")), "ref_source": str(r.get("L-Com Price Source") or ""),
                     "ref_date": str(r.get("L-Com Price Date") or "")})
    has_ref = "L-Com Unit Price" in head and any(x["ref_price"] for x in rows)
    return {"rows": rows, "has_ref": has_ref, "ref_label": "L-Com Unit Price", "exclusions_applied": kinds is not None}


def result_rows(s: Settings, path: Path):
    """The read-only results table for one results file: {'rows', 'has_ref', 'ref_label', 'exclusions_applied'}, or
    None when the file cannot be read. Cached by modified time (and the exclusion list's)."""
    try:
        return _cached(("r", str(path), path.stat().st_mtime, _mtime(s.project / "reviewer_exclusions.csv")),
                       lambda: _read_result_rows(s, path))
    except Exception:  # noqa: BLE001
        return None


def steps_for(s: Settings, results_name: str) -> list:
    path = s.results_dir / results_name
    out = []
    try:
        n, text = search_detail(s, path)
        d = draft_counts(s, path)
        search = Step("Search", FOUND, text, items=[("Results file exists", True), (f"{n} product{'s' if n != 1 else ''}", None)], done=True)
        can_draft = d["draftable"] if d else None
    except Exception:  # noqa: BLE001
        n = 0
        can_draft = None
        search = Step("Search", BLOCKED, "", "The results file could not be read. It may be open in Excel or damaged.")
    out.append(search)
    no_results = "There is no readable results file for this workflow." if search.status == BLOCKED else ""

    # Emails. email_state.json is one file for every results file; an entry belongs to THIS results file when its key
    # (the email agent's listing_key: SKU + the recommended listing's host and path) is the key of one of this file's
    # draftable rows. Without the agent's rules the keys are unknown, so the whole record is counted and said to be.
    dec, dec_err = _json(s.project / "decision_state.json")
    reqs = [r for r in (dec or {}).get("requests", {}).values() if isinstance(r, dict) and r.get("results_file") == results_name]
    email, err = _json(s.project / "email_state.json")
    rows = draft_rows(s, path) if search.status == FOUND else None
    record = {k: v for k, v in (email or {}).items() if not k.startswith("_") and isinstance(v, dict)}
    entries = [record[k] for k in {d["key"] for d in rows["drafts"]} if k in record] if rows else list(record.values())
    sent, replies = sum("sent" in e for e in entries), sum("reply" in e for e in entries)
    scope = "" if rows else " That record covers every results file, not only this one."
    strip = [(f"{can_draft} can be drafted" if can_draft is not None else "What can be drafted is unknown", None),
             (f"{sent} marked sent", sent > 0), (f"{replies} repl{'ies' if replies != 1 else 'y'} recorded", replies > 0)]
    if err:
        emails = Step("Emails", BLOCKED, "", err)
    elif not entries:
        why = no_results or ("Nothing in this file can be drafted." if can_draft == 0 else "")
        emails = Step("Emails", BLOCKED if why else NOT_STARTED, "No emails recorded yet.", why, strip, done=bool(reqs))
    else:
        emails = Step("Emails", FOUND, f"The email record has {len(entries)} seller(s){' for this results file' if rows else ''}: "
                      f"{sent} marked sent, {replies} with a reply.{scope}", items=strip, done=True)
    out.append(emails)

    exp, exp_err = _json(s.project / "approved_orders.json")
    mine = exp if isinstance(exp, dict) and exp.get("source_results_file") == results_name else None
    exported_at = _mtime(s.project / "approved_orders.json") if mine else 0
    if dec_err:
        approvals = Step("Approvals", BLOCKED, "", dec_err)
    elif not reqs:
        approvals = Step("Approvals", BLOCKED if no_results else NOT_STARTED, "No approval request has been written for this file.", no_results,
                         [("No request yet", False)])
    else:
        ids = {r["id"] for r in reqs}
        decs = [d for d in dec.get("decisions", []) if d.get("request_id") in ids]
        now = datetime.now().astimezone()
        valid = [d for d in decs if d.get("decision") == "approve" and not d.get("revoked_at") and (_when(d.get("expires_at")) or now) > now]
        live, waiting = len(valid), sum(not r.get("reply") for r in reqs)
        soonest = min((_when(d.get("expires_at")) for d in valid), default=None)
        approvals = Step("Approvals", FOUND, f"{len(reqs)} request(s); {waiting} still waiting for a recorded reply; "
                         f"{live} approval(s) currently valid (approvals expire).",
                         items=[(f"{len(reqs)} request{'s' if len(reqs) != 1 else ''} written", True),
                                (f"{len(reqs) - waiting} repl{'ies' if len(reqs) - waiting != 1 else 'y'} recorded", len(reqs) > waiting),
                                (f"{live} approval{'s' if live != 1 else ''} valid" + (f", first expires {soonest.astimezone():%Y-%m-%d %H:%M}" if soonest else ""), live > 0),   # local time, as on the tabs
                                ("Exported" if mine else "Not exported", bool(mine))])
    approvals.done = bool(mine)
    out.append(approvals)

    sheet = _newest(s.project / "Order Sheets", "order_sheet_*.xlsx")
    built = bool(mine and sheet and sheet.stat().st_mtime >= exported_at)       # a sheet older than the export was built from an earlier one
    if exp_err:
        order = Step("Order sheet", BLOCKED, "", exp_err)
    elif mine is None:
        why = "Needs a valid approval, exported from the Approvals step." if approvals.status != FOUND else ""
        order = Step("Order sheet", BLOCKED if why else NOT_STARTED, "No approved orders have been exported for this file.", why,
                     [("No order sheet for this file's export", False)])
    else:
        extra = f" The newest order sheet is {sheet.name} (sheets are not linked to a results file)." if sheet else " No order sheet file yet."
        order = Step("Order sheet", FOUND, f"{mine.get('valid_approvals', 0)} valid approval(s) exported at {mine.get('generated_at', '?')}.{extra}",
                     items=[(f"Built: {sheet.name}" if built else "Not built since the last export", built)], done=built)
    out.append(order)

    facts = newest_facts(s)
    listing = _newest(s.project / "Walmart Listings", "walmart_listings_*.xlsx")
    generated = bool(mine and listing and listing.stat().st_mtime >= exported_at)
    strip = [(f"Facts file: {facts.name}" if facts else "No facts file yet", bool(facts)),
             (f"Generated: {listing.name}" if generated else "Not generated since the last export", generated)]
    if listing:
        content = Step("Listing content", FOUND, f"Newest listing file: {listing.name} (not linked to a results file). Drafts need a person's review before upload.",
                       items=strip, done=generated)
    else:
        why = "Needs an exported approval first." if mine is None else ""
        content = Step("Listing content", BLOCKED if why else NOT_STARTED,
                       "The product facts file exists." if facts else "The product facts file has not been created yet.", why, strip)
    out.append(content)
    return out


def next_step(steps: list):
    """The first step whose main output the files do not show yet. Only ever used as a link; nothing is started from it."""
    return next((st for st in steps if not st.done), None)


def newest_facts(s: Settings):
    """The newest facts file a person uploaded (Uploads\\facts_*.csv), else a product_facts.csv in the project folder."""
    return _newest(s.project / UPLOADS_FOLDER, "facts_2*.csv") or (s.project / "product_facts.csv" if (s.project / "product_facts.csv").is_file() else None)


def _when(text):
    """An agent's ISO time as an aware datetime, or None."""
    try:
        t = datetime.fromisoformat(str(text).replace("Z", "+00:00"))
        return t if t.tzinfo else t.astimezone()
    except ValueError:
        return None


def _newest(folder: Path, pattern: str):
    found = sorted(folder.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True) if folder.is_dir() else []
    return found[0] if found else None
